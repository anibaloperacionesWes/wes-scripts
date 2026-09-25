"""
Una hoja por colegio CORMUP: lecturas de la cuenta vs WES (día completo y medio día).

Columnas:
  1 Mes
  2 Lectura inicial (fecha)
  3 Fecha lectura final y lectura (m³ del medidor)
  4 Diferencia entre lecturas (m³ cuenta / turbina)
  5 Consumo app WES entre las fechas de lectura (días completos)
  6 Consumo app WES medio día: día inicio 12:00–23:59 + intermedios + día final 00:00–12:00
  7 Diferencia de consumo (WES medio día − m³ cuenta)
    Azul = la app marca más que la cuenta; rojo = la cuenta marca más que la app.
    Verde = el mes se cobró a promedio.

Uso:
  python generar_hojas_lecturas_medio_dia_cormup.py
  python generar_hojas_lecturas_medio_dia_cormup.py --skip-download
"""

from __future__ import annotations

import argparse
import re
import shutil
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from facturacion_aguas_andinas_pdf import extraer_texto_pdf
from generar_comparativo_facturaciones_cormup_penalolen import (
    MESES_CORTOS,
    OUT_DIR,
    PDF_CACHE,
    Sitio,
    _cargar_sitios,
    _medidas_wes,
    descargar_facturaciones,
)
from generar_reporte_word import format_number_chilean
from wes_estilo_graficos_app import horas_api_chile

AZUL_APP = PatternFill("solid", fgColor="2E75B6")
ROJO_CTA = PatternFill("solid", fgColor="C00000")
VERDE_PROM = PatternFill("solid", fgColor="C6EFCE")
AZUL_HDR = PatternFill("solid", fgColor="003366")
GRIS = PatternFill("solid", fgColor="D9E1F2")
BLANCO = Font(color="FFFFFF", bold=True, size=9)
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def _parse_m3_cl(s: str) -> Optional[float]:
    s = (s or "").strip().replace(" ", "")
    if not s or s == "-":
        return None
    if "," in s:
        return float(s.replace(".", "").replace(",", "."))
    parts = s.split(".")
    if len(parts) > 1 and len(parts[-1]) == 3:
        return float("".join(parts))
    try:
        return float(s)
    except ValueError:
        return None


def _lecturas_medidor(txt: str) -> Tuple[Optional[float], Optional[float], Optional[float]]:
    """Lectura anterior, lectura actual y diferencia de lecturas (m³)."""
    ant = act = dif = None
    m = re.search(
        r"LECTURA\s+ANTERIOR\s+[0-9]{2}-[A-ZÁ]{3}-[0-9]{4}\s+([0-9\.\-]+)\s*m3",
        txt,
        re.I,
    )
    if m:
        ant = _parse_m3_cl(m.group(1))
    m = re.search(
        r"LECTURA\s+ACTUAL\s+[0-9]{2}-[A-ZÁ]{3}-[0-9]{4}\s+([0-9\.\-]+)\s*m3",
        txt,
        re.I,
    )
    if m:
        act = _parse_m3_cl(m.group(1))
    m = re.search(r"DIFERENCIA\s+DE\s+LECTURAS\s+([0-9\.\-]+)\s*m3", txt, re.I)
    if m:
        dif = _parse_m3_cl(m.group(1))
    return ant, act, dif


def _wes_dias_completos(por_dia: Dict[date, float], d0: date, d1: date) -> float:
    tot = 0.0
    d = d0
    while d <= d1:
        tot += float(por_dia.get(d, 0.0))
        d += timedelta(days=1)
    return tot


def _wes_medio_dia(
    node_id: str,
    d0: date,
    d1: date,
    por_dia: Dict[date, float],
    cache_h: Dict[Tuple[str, date], Dict[int, float]],
) -> float:
    def horas(dia: date) -> Dict[int, float]:
        key = (node_id, dia)
        if key not in cache_h:
            cache_h[key] = horas_api_chile(node_id, datetime.combine(dia, datetime.min.time()))
        return cache_h[key]

    if d1 < d0:
        return 0.0
    if d0 == d1:
        h = horas(d0)
        return sum(h.get(i, 0.0) for i in range(12, 24)) + sum(h.get(i, 0.0) for i in range(0, 12))

    h0 = horas(d0)
    h1 = horas(d1)
    ini = sum(h0.get(i, 0.0) for i in range(12, 24))
    fin = sum(h1.get(i, 0.0) for i in range(0, 12))
    medio = 0.0
    d = d0 + timedelta(days=1)
    while d <= d1 - timedelta(days=1):
        medio += float(por_dia.get(d, 0.0))
        d += timedelta(days=1)
    return ini + medio + fin


def _sheet_name(sitio: Sitio) -> str:
    raw = f"{sitio.node_id[-2:]} {sitio.node_name}"
    bad = r'[]:*?/\\'
    for ch in bad:
        raw = raw.replace(ch, " ")
    return raw[:31].strip()


def _filas_colegio(sitio: Sitio, cache_h) -> List[dict]:
    if not sitio.filas:
        return []
    d0 = min(f.lectura_anterior.date() for f in sitio.filas)
    d1 = max(f.lectura_actual.date() for f in sitio.filas)
    print(f"[WES] {sitio.node_id} {sitio.node_name} {d0} → {d1}", flush=True)
    try:
        meas = _medidas_wes(sitio.node_id, d0, d1)
    except Exception as e:
        print(f"  [WARN] API {sitio.node_id}: {e}", flush=True)
        meas = []
    por_dia = {m.date.date(): float(m.total_m3) for m in meas}
    out = []
    for f in sorted(sitio.filas, key=lambda x: x.emision):
        a, b = f.lectura_anterior.date(), f.lectura_actual.date()
        pdf = PDF_CACHE / f.carpeta / f.pdf_name
        lect_ini = lect_fin = dif_lect = None
        if pdf.is_file():
            try:
                lect_ini, lect_fin, dif_lect = _lecturas_medidor(extraer_texto_pdf(pdf))
            except Exception:
                pass
        dif_cta = float(dif_lect) if dif_lect is not None else float(f.m3_cuenta)
        wes_full = _wes_dias_completos(por_dia, a, b)
        wes_mid = _wes_medio_dia(sitio.node_id, a, b, por_dia, cache_h)
        out.append(
            {
                "mes": f"{MESES_CORTOS[f.emision.month]}-{f.emision.year}",
                "lect_ini_fecha": a,
                "lect_ini_m3": lect_ini,
                "lect_fin_fecha": b,
                "lect_fin_m3": lect_fin,
                "dif_lect": dif_cta,
                "wes_full": wes_full,
                "wes_mid": wes_mid,
                "dif_cons": wes_mid - dif_cta,
                "estimado": f.estimado,
                "boleta": f.boleta,
            }
        )
    return out


def _write_sheet(ws, sitio: Sitio, filas: List[dict]) -> None:
    ws.append(
        [
            "Mes",
            "Lectura inicial",
            "Fecha lectura final y lectura",
            "Diferencia entre lecturas (m³)",
            "Consumo app WES (fechas completas)",
            "Consumo app WES (medio día)",
            "Diferencia consumo (WES ½ día − cuenta)",
        ]
    )
    for col in range(1, 8):
        c = ws.cell(1, col)
        c.fill = AZUL_HDR
        c.font = BLANCO
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    ws.append(
        [
            sitio.node_name,
            sitio.node_id,
            "",
            "",
            "",
            "",
            "",
        ]
    )
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=7)
    ws.cell(2, 1).font = Font(bold=True, size=11, color="003366")

    tot_cta = tot_full = tot_mid = 0.0
    for r in filas:
        ini = r["lect_ini_fecha"].strftime("%d-%m-%Y")
        if r["lect_ini_m3"] is not None:
            ini = f"{ini}  ({format_number_chilean(r['lect_ini_m3'], 0)} m³)"
        fin = r["lect_fin_fecha"].strftime("%d-%m-%Y")
        if r["lect_fin_m3"] is not None:
            fin = f"{fin}  ({format_number_chilean(r['lect_fin_m3'], 0)} m³)"
        ws.append(
            [
                r["mes"],
                ini,
                fin,
                round(r["dif_lect"], 1),
                round(r["wes_full"], 1),
                round(r["wes_mid"], 1),
                round(r["dif_cons"], 1),
            ]
        )
        row = ws.max_row
        tot_cta += r["dif_lect"]
        tot_full += r["wes_full"]
        tot_mid += r["wes_mid"]
        dif = r["dif_cons"]
        fill = None
        font = Font(bold=True, color="FFFFFF", size=10)
        if dif > 0.05:
            fill = AZUL_APP
        elif dif < -0.05:
            fill = ROJO_CTA
        if fill is not None:
            ws.cell(row, 7).fill = fill
            ws.cell(row, 7).font = font
            ws.cell(row, 6).fill = fill
            ws.cell(row, 6).font = font
        if r["estimado"]:
            for col in (1, 2, 3, 4):
                ws.cell(row, col).fill = VERDE_PROM
            ws.cell(row, 1).font = Font(bold=True, size=10, color="006100")
        for col in range(1, 8):
            ws.cell(row, col).border = THIN
            ws.cell(row, col).alignment = Alignment(horizontal="center")
        for col in (4, 5, 6, 7):
            ws.cell(row, col).number_format = "#,##0.0"

    ws.append(
        [
            "TOTAL",
            "",
            "",
            round(tot_cta, 1),
            round(tot_full, 1),
            round(tot_mid, 1),
            round(tot_mid - tot_cta, 1),
        ]
    )
    last = ws.max_row
    tdif = tot_mid - tot_cta
    for col in range(1, 8):
        cell = ws.cell(last, col)
        cell.font = Font(bold=True)
        cell.border = THIN
        cell.fill = GRIS
        cell.alignment = Alignment(horizontal="center")
    for col in (4, 5, 6, 7):
        ws.cell(last, col).number_format = "#,##0.0"
    if tdif > 0.05:
        ws.cell(last, 7).fill = AZUL_APP
        ws.cell(last, 7).font = Font(bold=True, color="FFFFFF")
    elif tdif < -0.05:
        ws.cell(last, 7).fill = ROJO_CTA
        ws.cell(last, 7).font = Font(bold=True, color="FFFFFF")

    ws.freeze_panes = "A3"
    for i, w in enumerate([12, 28, 28, 22, 28, 26, 28], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    nota = last + 2
    ws.cell(
        nota,
        1,
        "Verde = ese mes se cobró a promedio/estimado. "
        "Azul = la app WES marca más que la cuenta. Rojo = la cuenta marca más que la app.",
    )
    ws.merge_cells(start_row=nota, start_column=1, end_row=nota, end_column=7)
    ws.cell(nota, 1).font = Font(size=9, italic=True, color="006100")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:2"


def _write_resumen(ws, bloques: List[Tuple[Sitio, List[dict]]]) -> None:
    ws.append(
        [
            "Colegio",
            "Nodo",
            "Meses",
            "m³ cuenta (Σ dif. lecturas)",
            "m³ WES días completos",
            "m³ WES medio día",
            "Dif. (WES ½ día − cuenta)",
        ]
    )
    for col in range(1, 8):
        c = ws.cell(1, col)
        c.fill = AZUL_HDR
        c.font = BLANCO
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    tot = [0.0, 0.0, 0.0]
    for sitio, filas in bloques:
        cta = sum(r["dif_lect"] for r in filas)
        full = sum(r["wes_full"] for r in filas)
        mid = sum(r["wes_mid"] for r in filas)
        dif = mid - cta
        ws.append([sitio.node_name, sitio.node_id, len(filas), cta, round(full, 1), round(mid, 1), round(dif, 1)])
        tot[0] += cta
        tot[1] += full
        tot[2] += mid
        row = ws.max_row
        if dif > 0.05:
            ws.cell(row, 7).fill = AZUL_APP
            ws.cell(row, 7).font = Font(bold=True, color="FFFFFF")
        elif dif < -0.05:
            ws.cell(row, 7).fill = ROJO_CTA
            ws.cell(row, 7).font = Font(bold=True, color="FFFFFF")
        for col in range(1, 8):
            ws.cell(row, col).border = THIN
            ws.cell(row, col).alignment = Alignment(horizontal="center")
        ws.cell(row, 1).alignment = Alignment(horizontal="left")
        for col in (4, 5, 6, 7):
            ws.cell(row, col).number_format = "#,##0.0"
    ws.append(["TOTAL CORMUP", "000008", "", tot[0], round(tot[1], 1), round(tot[2], 1), round(tot[2] - tot[0], 1)])
    for col in range(1, 8):
        cell = ws.cell(ws.max_row, col)
        cell.font = Font(bold=True)
        cell.fill = GRIS
        cell.border = THIN
    for col in (4, 5, 6, 7):
        ws.cell(ws.max_row, col).number_format = "#,##0.0"
    ws.freeze_panes = "A2"
    for i, w in enumerate([28, 12, 10, 24, 22, 20, 26], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    nota = ws.max_row + 2
    ws.cell(nota, 1, "Azul = la app WES marca más que la cuenta. Rojo = la cuenta marca más que la app.")
    ws.merge_cells(start_row=nota, start_column=1, end_row=nota, end_column=7)
    ws.cell(
        nota + 1,
        1,
        "WES medio día: día de lectura inicial 12:00–23:59; días intermedios completos; "
        "día de lectura final 00:00–12:00. Verde = ese mes se cobró a promedio/estimado.",
    )
    ws.merge_cells(start_row=nota + 1, start_column=1, end_row=nota + 1, end_column=7)


def generar(skip_download: bool = False) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    generado = datetime.now(timezone.utc).astimezone()
    ts = generado.strftime("%Y%m%d_%H%M")

    if skip_download and any(PDF_CACHE.glob("*/*.pdf")):
        locales = {p.name: p for p in PDF_CACHE.iterdir() if p.is_dir()}
        print(f"[INFO] Reuso cache {PDF_CACHE}", flush=True)
    else:
        if PDF_CACHE.exists():
            print("[INFO] Facturaciones actualizadas: redescargo PDFs", flush=True)
            shutil.rmtree(PDF_CACHE)
        locales = descargar_facturaciones(PDF_CACHE)

    sitios = [s for s in _cargar_sitios(locales) if s.filas]
    sitios.sort(key=lambda s: s.node_id)
    cache_h: Dict[Tuple[str, date], Dict[int, float]] = {}
    bloques = []
    for s in sitios:
        bloques.append((s, _filas_colegio(s, cache_h)))

    wb = Workbook()
    wr = wb.active
    wr.title = "Resumen"
    _write_resumen(wr, bloques)
    for sitio, filas in bloques:
        ws = wb.create_sheet(_sheet_name(sitio))
        _write_sheet(ws, sitio, filas)

    stem = f"Lecturas_vs_WES_medio_dia_CORMUP_{ts}"
    out_xlsx = OUT_DIR / f"{stem}.xlsx"
    wb.save(out_xlsx)
    print(f"[OK] Excel: {out_xlsx}", flush=True)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        import subprocess
        import tempfile

        with tempfile.TemporaryDirectory() as td:
            subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--nologo",
                    "--nofirststartwizard",
                    "--convert-to",
                    "pdf:calc_pdf_Export",
                    "--outdir",
                    td,
                    str(out_xlsx),
                ],
                check=False,
                capture_output=True,
                text=True,
            )
            produced = Path(td) / (out_xlsx.stem + ".pdf")
            if produced.is_file():
                out_pdf = out_xlsx.with_suffix(".pdf")
                shutil.copy2(produced, out_pdf)
                print(f"[OK] PDF:   {out_pdf}", flush=True)
    return out_xlsx


def main() -> int:
    if sys.platform == "win32":
        for s in (sys.stdout, sys.stderr):
            try:
                s.reconfigure(encoding="utf-8", line_buffering=True)
            except Exception:
                pass
    ap = argparse.ArgumentParser(description="Hojas por colegio CORMUP: lecturas vs WES medio día")
    ap.add_argument("--skip-download", action="store_true")
    args = ap.parse_args()
    generar(skip_download=args.skip_download)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
