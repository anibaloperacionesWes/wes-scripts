"""
Auditoría WES estilo ICCO — CORMUP Peñalolén, vacaciones Fiestas Patrias 2026.

Misma estructura que el informe de auditoría ICCO:
  Portada → Índice → Metodología → Registros de consumos → Resultados y Conclusiones

Periodos (7+7 días homólogos):
  Con WES (control vacaciones): 14–20/09/2026
  Sin WES (línea base):         07–13/09/2026

Cohorte: 10 colegios con control (excluye Tobalaba por pulso y los 3 sin válvula).

Uso:
  python generar_auditoria_cormup_vacaciones_sept2026.py
"""

from __future__ import annotations

import json
import sys
from datetime import date, datetime, timezone
from pathlib import Path
from typing import List, Optional

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_BREAK, WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

import auditoria_cpa_icco_renca_grafico as graf
import generar_informe_auditoria_icco_renca_word as icco
from auditoria_cpa_icco_renca_grafico import Periodo, generar_png_barras_con_sin
from generar_comparativo_cormup_vacaciones_sept2026 import (
    COLEGIOS,
    CON_CONTROL_FIN,
    CON_CONTROL_INI,
    HORARIOS_ESPECIALES,
    SIN_CONTROL,
    SIN_CONTROL_FIN,
    SIN_CONTROL_INI,
    TOBALABA,
    evaluar_colegios,
    evaluar_fuera_comparativo,
    grafico_ahorro_clp,
    grafico_barras_colegios,
    grafico_totales_diarios,
    precio_referencia_clp,
)
from generar_reporte_word import (
    add_picture_with_pagination,
    convertir_word_a_pdf,
    format_currency_chilean,
    format_number_chilean,
)
from wes_estilo_graficos_app import COLOR_BARRA_WES, COLOR_NOCHE

ROOT = Path(__file__).resolve().parent
OUT_BASE = ROOT / "reports" / "CORMUP" / "AUDITORIA" / "VACACIONES_SEP2026"

PERIODO_REF = Periodo(
    "Con control vacaciones (14-09 a 20-09-2026)",
    tuple(
        date(2026, 9, d)
        for d in range(CON_CONTROL_INI.day, CON_CONTROL_FIN.day + 1)
    ),
)
PERIODO_AUD = Periodo(
    "Sin control (07-09 a 13-09-2026)",
    tuple(
        date(2026, 9, d)
        for d in range(SIN_CONTROL_INI.day, SIN_CONTROL_FIN.day + 1)
    ),
)


def _png_barras_totales(tot_con: float, tot_sin: float, out_png: Path) -> Path:
    """Barras Con WES vs Sin WES (agregado cohorte), estilo ICCO."""
    out_png.parent.mkdir(parents=True, exist_ok=True)
    # Reutiliza el helper del módulo gráfico (etiquetas de periodo).
    tmp = generar_png_barras_con_sin(
        tot_con,
        tot_sin,
        out_png.parent,
        "cormup_vac_agg",
        ref=PERIODO_REF,
        aud=PERIODO_AUD,
    )
    if tmp and tmp.is_file() and tmp.resolve() != out_png.resolve():
        out_png.write_bytes(tmp.read_bytes())
        try:
            tmp.unlink()
        except OSError:
            pass
        return out_png
    # Fallback local
    fig, ax = plt.subplots(figsize=(7.2, 4.4))
    ax.bar(
        ["Con WES\n14–20/09", "Sin WES\n07–13/09"],
        [tot_con, tot_sin],
        color=[COLOR_BARRA_WES, COLOR_NOCHE],
        width=0.55,
    )
    ax.set_ylabel("Consumo (m³)")
    ax.set_title("CORMUP Peñalolén — totales cohorte con control (7 días)")
    ax.grid(axis="y", linestyle=":", alpha=0.4)
    fig.tight_layout()
    fig.savefig(out_png, dpi=150)
    plt.close(fig)
    return out_png


def _tabla_colegios(doc: Document, filas) -> None:
    headers = (
        "Nodo",
        "Establecimiento",
        "Sin WES (m³)",
        "Con WES (m³)",
        "Ahorro (m³)",
        "Ahorro %",
    )
    tbl = doc.add_table(rows=1 + len(filas), cols=len(headers))
    icco._apply_metricas_clave_table_style(tbl, doc)
    for j, h in enumerate(headers):
        cell = tbl.rows[0].cells[j]
        cell.text = h
        icco._set_cell_shading_hex(cell, icco._TABLA_CUANTITATIVOS_FILL_HEADER)
        for par in cell.paragraphs:
            par.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            for run in par.runs:
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = icco._HEADING_COLOR
                run.font.name = icco._BODY_FONT

    for i, f in enumerate(filas, start=1):
        pct = f.ahorro_pct
        vals = (
            f.node_id,
            f.nombre,
            format_number_chilean(f.m3_sin, 1),
            format_number_chilean(f.m3_con, 1),
            format_number_chilean(max(0.0, f.ahorro_m3), 1),
            f"{format_number_chilean(pct, 1)} %" if pct is not None else "—",
        )
        for j, v in enumerate(vals):
            cell = tbl.rows[i].cells[j]
            cell.text = v
            for par in cell.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8)
                    run.font.name = icco._BODY_FONT
    gap = doc.add_paragraph("")
    gap.paragraph_format.space_after = Pt(6)


def _metodologia_cormup(doc: Document) -> None:
    icco._section_heading_icco(doc, "Metodología")
    nd = len(PERIODO_REF.dias)
    rc = icco._rango_fechas_prosa_cl(PERIODO_REF.dias[0], PERIODO_REF.dias[-1])
    rs = icco._rango_fechas_prosa_cl(PERIODO_AUD.dias[0], PERIODO_AUD.dias[-1])
    icco._p_justify(
        doc,
        "Se generó una auditoría hídrica agregada al estilo ICCO para comprobar de forma empírica "
        "el efecto de ahorro asociado al servicio WES «Control Inteligente de Agua Potable» en la "
        f"cohorte de {len(COLEGIOS)} colegios CORMUP de Peñalolén con control hidráulico activo. "
        f"Se definieron dos periodos de {nd} días cada uno: con control especial por vacaciones "
        f"entre {rc} y con control desactivado (línea base) entre {rs}, de modo que ambas series "
        "son comparables y permiten dimensionar la diferencia de consumo atribuible al servicio WES.",
    )
    icco._p_justify(
        doc,
        "En la elaboración de los indicadores se consideran, para cada nodo, solo registros con fecha "
        "válida; el promedio diario de un periodo es la suma de los consumos diarios dividida por la "
        "cantidad de días del periodo. El ahorro estimado en volumen (m³) se expresa como "
        "max(0, Σ_sin − Σ_con) y el ahorro porcentual como (ahorro_m³ / Σ_sin) × 100. La valorización "
        "económica orientativa usa el precio de referencia del servicio (CLP/m³).",
    )
    icco._p_justify(
        doc,
        "Quedan fuera del cuadro comparativo: Tobalaba (000008-04) por medidor de pulso (lectura "
        "anómala / no comparable) y los tres establecimientos solo monitoreo sin válvula de control "
        "(Eduardo de la Barra, Alicura y Likankura). Sus consumos se informan aparte, sin atribuir "
        "ahorro al control WES. En la semana con control aplicaron horarios especiales de vacaciones "
        "(cortes programados, habilitaciones por obras y patinaje), documentados en Registros.",
    )


def _registros_cormup(
    doc: Document,
    *,
    tot_con: float,
    tot_sin: float,
    filas,
    png_barras: Path,
    png_colegios: Path,
    png_diario: Path,
    png_clp: Path,
) -> None:
    icco._section_heading_icco(doc, "Registros de consumos")
    icco._p_justify(
        doc,
        f"Totales de la cohorte ({len(filas)} colegios): Con WES "
        f"{format_number_chilean(tot_con, 1)} m³ "
        f"({icco._formato_fechas_periodo_compacto(PERIODO_REF)}) frente a Sin WES "
        f"{format_number_chilean(tot_sin, 1)} m³ "
        f"({icco._formato_fechas_periodo_compacto(PERIODO_AUD)}).",
    )

    if png_barras.is_file():
        add_picture_with_pagination(doc, str(png_barras), width=Inches(5.8))
    if png_colegios.is_file():
        add_picture_with_pagination(doc, str(png_colegios), width=Inches(6.2))

    icco._section_heading_icco(doc, "Detalle por establecimiento", space_after_pt=6)
    _tabla_colegios(doc, filas)

    if png_diario.is_file():
        add_picture_with_pagination(doc, str(png_diario), width=Inches(6.0))
    if png_clp.is_file():
        add_picture_with_pagination(doc, str(png_clp), width=Inches(6.0))

    icco._section_heading_icco(doc, "Horarios especiales (semana Con WES)", space_after_pt=6)
    for linea in HORARIOS_ESPECIALES:
        icco._p_justify(doc, f"• {linea}")


def _anexo_fuera(doc: Document, fuera) -> None:
    icco._section_heading_icco(doc, "Anexo — puntos fuera del comparativo")
    tob = next((f for f in fuera if f.node_id == TOBALABA[0]), None)
    solo = [f for f in fuera if f.node_id != TOBALABA[0]]
    if tob:
        icco._p_justify(
            doc,
            f"Tobalaba ({tob.node_id}): medidor de pulso. Consumo 07–13/09 = "
            f"{format_number_chilean(tob.m3_sin, 1)} m³; 14–20/09 = "
            f"{format_number_chilean(tob.m3_con, 1)} m³. No se incluye en el ahorro atribuible "
            "al control por no ser comparable con el resto de la cohorte.",
        )
    if solo:
        partes = [
            f"{f.nombre} ({f.node_id}): 07–13 = {format_number_chilean(f.m3_sin, 1)} m³; "
            f"14–20 = {format_number_chilean(f.m3_con, 1)} m³"
            for f in solo
        ]
        icco._p_justify(
            doc,
            "Solo monitoreo (sin válvula de control): "
            + "; ".join(partes)
            + ". Se informan por transparencia; no forman parte del cálculo de ahorro WES.",
        )


def generar_informe(
    out_dir: Path,
    *,
    max_workers: int = 2,
) -> tuple[Path, Optional[Path], dict]:
    out_dir.mkdir(parents=True, exist_ok=True)
    charts = out_dir / "graficos"
    charts.mkdir(parents=True, exist_ok=True)

    precio = float(precio_referencia_clp())
    icco.CLP_POR_M3_REF = precio

    print("1) Consumos por colegio (API)…")
    filas = evaluar_colegios(max_workers=max_workers)
    fuera = evaluar_fuera_comparativo(max_workers=max_workers)
    tot_sin = float(sum(f.m3_sin for f in filas))
    tot_con = float(sum(f.m3_con for f in filas))
    ahorro = max(0.0, tot_sin - tot_con)
    pct = (100.0 * ahorro / tot_sin) if tot_sin > 1e-9 else 0.0
    clp_ahorro = ahorro * precio

    print(
        f"   Σ Sin WES={tot_sin:.1f} m³ | Σ Con WES={tot_con:.1f} m³ | "
        f"Ahorro={ahorro:.1f} m³ ({pct:.1f} %) ≈ {clp_ahorro:,.0f} CLP @ {precio:.0f}"
    )

    print("2) Gráficos…")
    png_barras = _png_barras_totales(tot_con, tot_sin, charts / "01_barras_con_sin.png")
    png_colegios = grafico_barras_colegios(filas, charts / "02_barras_por_colegio.png")
    png_diario = grafico_totales_diarios(filas, charts / "03_totales_diarios.png")
    png_clp = grafico_ahorro_clp(filas, precio, charts / "04_ahorro_clp.png")

    # Portada ICCO
    stem = "Auditoria_CORMUP_Vacaciones_14_20_Sep_2026"
    graf.NOMBRE_PUNTO = "CORMUP Peñalolén"
    icco.NOMBRE_PUNTO = "CORMUP Peñalolén"
    icco.PORTADA_TITULO = "Informe de Auditoría"
    icco.PORTADA_REFERENCIA_BORRADOR = stem
    icco.PORTADA_ESTABLECIMIENTO_LINEA1 = "CORMUP Peñalolén"
    icco.PORTADA_ESTABLECIMIENTO_LINEA2 = "Vacaciones Fiestas Patrias 2026"
    icco._ETIQUETA_COLEGIO_CUADRO_RESUMEN = "CORMUP"

    print("3) Word estilo ICCO…")
    doc = icco._open_icco_document(ROOT, None)
    icco._portada_y_indice(doc, ROOT, None, portada_completa=None)
    doc.add_paragraph("")
    _metodologia_cormup(doc)

    p_br = doc.add_paragraph()
    p_br.add_run().add_break(WD_BREAK.PAGE)

    _registros_cormup(
        doc,
        tot_con=tot_con,
        tot_sin=tot_sin,
        filas=filas,
        png_barras=png_barras,
        png_colegios=png_colegios,
        png_diario=png_diario,
        png_clp=png_clp,
    )

    p_br2 = doc.add_paragraph()
    p_br2.add_run().add_break(WD_BREAK.PAGE)

    # Resultados: total_ref = Con WES; total_aud = Sin WES (convención ICCO).
    ahorro_ref_menos_aud = tot_con - tot_sin
    rend_sobre_con = (100.0 * ahorro_ref_menos_aud / tot_con) if tot_con > 0 else 0.0
    icco._resultados_y_conclusiones(
        doc,
        tot_con,
        tot_sin,
        ahorro_ref_menos_aud,
        rend_sobre_con,
        len(PERIODO_REF.dias),
        ref=PERIODO_REF,
        aud=PERIODO_AUD,
    )

    icco._p_justify(
        doc,
        f"Valorización económica de la cohorte: volumen evitado "
        f"{format_number_chilean(ahorro, 1)} m³ × {format_number_chilean(precio, 0)} CLP/m³ ≈ "
        f"{format_currency_chilean(clp_ahorro)}. "
        "Los consumos residuales en 14–20/09 en colegios con obras o patinaje corresponden a "
        "habilitaciones autorizadas, no a fallas de corte.",
    )

    _anexo_fuera(doc, fuera)

    pie = doc.add_paragraph()
    pie.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    pr = pie.add_run(
        f"Documento generado automáticamente — "
        f"{datetime.now(timezone.utc).strftime('%d-%m-%Y %H:%M')} UTC"
    )
    pr.font.name = icco._BODY_FONT
    pr.font.size = Pt(9)
    pr.font.color.rgb = RGBColor(100, 100, 100)

    docx_path = out_dir / f"{stem}.docx"
    doc.save(docx_path)
    print(f"   Word: {docx_path}")

    print("4) PDF…")
    pdf_path = convertir_word_a_pdf(docx_path)

    meta = {
        "docx": str(docx_path.resolve()),
        "pdf": str(pdf_path.resolve()) if pdf_path else None,
        "precio_clp_m3": precio,
        "total_sin_m3": tot_sin,
        "total_con_m3": tot_con,
        "ahorro_m3": ahorro,
        "ahorro_pct": pct,
        "ahorro_clp": clp_ahorro,
        "n_colegios": len(filas),
        "periodo_con": f"{CON_CONTROL_INI:%Y-%m-%d}/{CON_CONTROL_FIN:%Y-%m-%d}",
        "periodo_sin": f"{SIN_CONTROL_INI:%Y-%m-%d}/{SIN_CONTROL_FIN:%Y-%m-%d}",
        "colegios": [
            {
                "node_id": f.node_id,
                "nombre": f.nombre,
                "m3_sin": f.m3_sin,
                "m3_con": f.m3_con,
                "ahorro_m3": f.ahorro_m3,
                "ahorro_pct": f.ahorro_pct,
            }
            for f in filas
        ],
    }
    (out_dir / "meta.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return docx_path, pdf_path, meta


def main() -> int:
    if sys.platform == "win32":
        try:
            sys.stdout.reconfigure(encoding="utf-8")
        except Exception:
            pass

    ts = datetime.now().strftime("%Y%m%d_%H%M")
    out_dir = OUT_BASE / f"AUDITORIA_{ts}"
    print(f"Salida: {out_dir}")
    try:
        docx, pdf, meta = generar_informe(out_dir, max_workers=2)
    except Exception as e:
        print(f"[ERROR] {e}")
        import traceback

        traceback.print_exc()
        return 1

    print(f"Word: {docx}")
    if pdf:
        print(f"PDF:  {pdf}")
    else:
        print("PDF:  (no generado)")
    print(
        f"Ahorro: {meta['ahorro_m3']:.1f} m³ ({meta['ahorro_pct']:.1f} %) ≈ "
        f"{meta['ahorro_clp']:,.0f} CLP"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
