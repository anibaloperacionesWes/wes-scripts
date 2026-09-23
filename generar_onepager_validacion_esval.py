"""
One-pager ejecutivo — Validación Matriz ESVAL (Fundo Zapallar).
Para entregar a audiencia no técnica.

Uso:
  python generar_onepager_validacion_esval.py
"""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

ROOT = Path(__file__).resolve().parent
_HEADING = RGBColor(31, 71, 136)
_MUTED = RGBColor(89, 89, 89)

NODE_ID = "000027-01"
NODO = "Matriz ESVAL"
EMPRESA = "Fundo Zapallar"
FECHA = "22-09-2026"
HORA_INI, HORA_FIN = "11:50", "15:00"

ITRON_DELTA = 21.74
US_DELTA = 22.39
ERROR_PCT = 2.9
TOL_US, TOL_ITRON = 1, 5
DIAMETRO_FINAL = 127
MORTERO = 3
M45_MAX = 1.5


def _fmt(x: float, dec: int = 2) -> str:
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def _font(run, *, size=11, bold=False, color=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _h(doc, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = _HEADING
        r.font.name = "Calibri"
        r.font.size = Pt(13 if level == 1 else 12)


def _p(doc, text: str, *, size=11) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.space_before = Pt(0)
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = para.add_run(text)
    _font(r, size=size)


def _bullet(doc, text: str) -> None:
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.space_after = Pt(1)
    para.paragraph_format.space_before = Pt(0)
    para.clear()
    r = para.add_run(text)
    _font(r, size=10)


def _row(tbl, i: int, vals: list[str], *, bold=False, header=False) -> None:
    for j, v in enumerate(vals):
        cell = tbl.rows[i].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(v)
        _font(r, size=10, bold=bold or header, color=_HEADING if header else None)


def main() -> int:
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out_dir = ROOT / "reports" / "Fundo_Zapallar" / "Informes_Tecnicos" / f"onepager_esval_{stamp}"
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"OnePager_Validacion_Matriz_ESVAL_{stamp}.docx"

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.4)
    sec.bottom_margin = Cm(1.2)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    try:
        from generar_reporte_word import add_logo_to_header

        add_logo_to_header(doc)
    except Exception as exc:
        print(f"[WARN] Logo: {exc}")

    title = doc.add_heading("Validación Matriz ESVAL — Fundo Zapallar", level=0)
    for r in title.runs:
        r.font.color.rgb = _HEADING
        r.font.name = "Calibri"
        r.font.size = Pt(16)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(6)
    r = meta.add_run(
        f"Punto {NODO} ({NODE_ID})  ·  {FECHA}  ·  {HORA_INI} → {HORA_FIN}  ·  Uso interno WES"
    )
    _font(r, size=9, color=_MUTED)

    _h(doc, "Qué se hizo")
    _bullet(doc, "Se revisó la calidad de señal del medidor ultrasónico (dentro de norma).")
    _bullet(doc, "Se verificaron cables, se limpió la tubería y se renovó la silicona de los sensores.")
    _bullet(
        doc,
        f"Se corrigió la configuración de cañería (diámetro y mortero interno) y el factor de escala "
        f"(tope {_fmt(M45_MAX, 1)}). Configuración final: Ø {DIAMETRO_FINAL} mm · mortero {MORTERO} mm.",
    )

    _h(doc, "Resultado de la validación")
    _p(
        doc,
        "Se comparó el consumo del medidor Itron (referencia ESVAL) con el ultrasónico WES "
        "en la misma ventana de tiempo.",
    )

    tbl = doc.add_table(rows=4, cols=2)
    tbl.style = "Table Grid"
    _row(tbl, 0, ["Medidor", "Consumo (m³)"], header=True)
    _row(tbl, 1, ["Itron (ESVAL)", _fmt(ITRON_DELTA)])
    _row(tbl, 2, ["Ultrasónico WES", _fmt(US_DELTA)])
    _row(tbl, 3, ["% Error", f"{_fmt(ERROR_PCT, 1)} %"], bold=True)
    doc.add_paragraph()

    _p(
        doc,
        f"Error = 1 − ({_fmt(ITRON_DELTA)} / {_fmt(US_DELTA)}) = {_fmt(ERROR_PCT, 1)} %. "
        f"No se espera igualdad exacta: el fabricante del ultrasónico declara ±{TOL_US} % "
        f"y el de Itron ±{TOL_ITRON} %. Un {_fmt(ERROR_PCT, 1)} % está dentro de lo aceptable.",
    )

    _h(doc, "Conclusión")
    _p(
        doc,
        f"Tras los ajustes en terreno, ambos medidores quedaron alineados. "
        f"La validación del {FECHA} confirma un error de {_fmt(ERROR_PCT, 1)} %, "
        "aceptable según las tolerancias de fabricante. El punto queda operativo para monitoreo WES.",
    )

    pie = doc.add_paragraph()
    pie.paragraph_format.space_before = Pt(8)
    r = pie.add_run(f"WES · generado {datetime.now().strftime('%d-%m-%Y %H:%M')}")
    _font(r, size=8, color=_MUTED)

    doc.save(out)
    print(f"[OK] {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
