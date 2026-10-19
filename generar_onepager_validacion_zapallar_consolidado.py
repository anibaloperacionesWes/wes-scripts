"""
One-pager consolidado — todas las validaciones relojería vs App WES
de Fundo Zapallar (Matriz ESVAL + Etapa N°5).

Uso:
  python generar_onepager_validacion_zapallar_consolidado.py
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "reports" / "Fundo_Zapallar" / "VALIDACION"
NAVY = HexColor("#1F4788")
MUTED = HexColor("#595959")
LIGHT = HexColor("#F5F8FB")
GREEN = HexColor("#E8F5E9")
GREEN_TXT = HexColor("#1B5E20")
AMBER = HexColor("#FFF8E1")
AMBER_TXT = HexColor("#E65100")
RED_BG = HexColor("#FFEBEE")
RED_TXT = HexColor("#B71C1C")

# Campañas desde la 3ª validación en adelante (orden cronológico).
# error = 1 − reloj/app.
CAMPANIAS = [
    {
        "id": "1",
        "ventana": "23-09 → 02-10",
        "detalle": "Relojería post-intervención",
        "matriz": {"reloj": 1281.9, "app": 1309.56, "error": 2.1, "lectura": "Aceptable"},
        "etapa5": {"reloj": 195.0, "app": 195.6, "error": 0.3, "lectura": "Excelente"},
        "nota": "Cierre con fotos del 02-10.",
    },
    {
        "id": "2",
        "ventana": "02-10 → 14-10",
        "detalle": "Fotos tarde 14-10 + Excel app",
        "matriz": {"reloj": 1391.5, "app": 1425.87, "error": 2.4, "lectura": "Aceptable"},
        "etapa5": {"reloj": 242.1, "app": 243.4, "error": 0.5, "lectura": "Excelente"},
        "nota": "Horas completas; excluye hora de cada foto.",
    },
    {
        "id": "3",
        "ventana": "02-10 → 15-10",
        "detalle": "Fotos mañana 15-10 + mismo Excel",
        "matriz": {"reloj": 1458.9, "app": 1496.32, "error": 2.5, "lectura": "Aceptable"},
        "etapa5": {"reloj": 258.0, "app": 261.7, "error": 1.4, "lectura": "Aceptable"},
        "nota": "Última campaña (vigente).",
    },
]


def _fmt(x: float, dec: int = 1) -> str:
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def _font_docx(run, *, size=10, bold=False, color=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _row(tbl, i: int, vals: list[str], *, bold=False, header=False) -> None:
    heading = RGBColor(31, 71, 136)
    for j, v in enumerate(vals):
        cell = tbl.rows[i].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(v)
        _font_docx(r, size=8, bold=bold or header, color=heading if header else None)


def build_docx(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.2)
    sec.bottom_margin = Cm(1.0)
    sec.left_margin = Cm(1.4)
    sec.right_margin = Cm(1.4)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(9)

    try:
        from generar_reporte_word import add_logo_to_header

        add_logo_to_header(doc)
    except Exception as exc:
        print(f"[WARN] Logo: {exc}")

    title = doc.add_heading("Validaciones Fundo Zapallar — consolidado", level=0)
    for r in title.runs:
        r.font.color.rgb = RGBColor(31, 71, 136)
        r.font.name = "Calibri"
        r.font.size = Pt(14)

    meta = doc.add_paragraph()
    r = meta.add_run(
        "Relojería vs App WES · Matriz ESVAL (000027-01) y Etapa N°5 (000027-03) · "
        "Error = 1 − (relojería / app) · Uso interno WES"
    )
    _font_docx(r, size=8, color=RGBColor(89, 89, 89))

    h = doc.add_heading("Historial de campañas", level=1)
    for r in h.runs:
        r.font.color.rgb = RGBColor(31, 71, 136)

    # Tabla ancha: ventana | Matriz reloj | Matriz app | Matriz err | E5 reloj | E5 app | E5 err | Lectura
    tbl = doc.add_table(rows=1 + len(CAMPANIAS), cols=8)
    tbl.style = "Table Grid"
    _row(
        tbl,
        0,
        [
            "Ventana",
            "Matriz reloj",
            "Matriz app",
            "Err %",
            "E5 reloj",
            "E5 app",
            "Err %",
            "Estado",
        ],
        header=True,
    )
    for i, c in enumerate(CAMPANIAS, start=1):
        m, e = c["matriz"], c["etapa5"]
        estado_parts = []
        if m:
            estado_parts.append(f"M:{m['lectura'][:3]}")
        if e:
            estado_parts.append(f"E5:{e['lectura'][:3]}")
        _row(
            tbl,
            i,
            [
                c["ventana"],
                f"{_fmt(m['reloj'], 2 if m['reloj'] < 100 else 1)} m³" if m else "—",
                f"{_fmt(m['app'], 2 if m['app'] < 100 else 2)} m³" if m else "—",
                f"{_fmt(m['error'], 1)} %" if m else "—",
                f"{_fmt(e['reloj'], 1)} m³" if e else "—",
                f"{_fmt(e['app'], 1)} m³" if e else "—",
                f"{_fmt(e['error'], 1)} %" if e else "—",
                " / ".join(estado_parts) if estado_parts else "—",
            ],
        )
    doc.add_paragraph()

    h = doc.add_heading("Notas por campaña", level=1)
    for r in h.runs:
        r.font.color.rgb = RGBColor(31, 71, 136)
    for c in CAMPANIAS:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.clear()
        r = p.add_run(f"{c['ventana']} — {c['detalle']}. {c['nota']}")
        _font_docx(r, size=8)

    h = doc.add_heading("Lectura consolidada", level=1)
    for r in h.runs:
        r.font.color.rgb = RGBColor(31, 71, 136)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(
        "Desde la relojería del 23-09 → 02-10, Matriz ESVAL se mantiene entre 2,1 % y 2,5 % "
        "(app por encima del Itron). Etapa N°5 cierra excelente/aceptable (0,3 % → 1,4 %). "
        "Última campaña vigente: 02-10 → 15-10 (Matriz 2,5 %, Etapa 5 1,4 %)."
    )
    _font_docx(r, size=9)

    pie = doc.add_paragraph()
    pie.paragraph_format.space_before = Pt(6)
    r = pie.add_run(f"WES · generado {datetime.now().strftime('%d-%m-%Y %H:%M')}")
    _font_docx(r, size=7, color=RGBColor(89, 89, 89))

    doc.save(path)
    print(f"[OK] DOCX {path}")


def _card(c: canvas.Canvas, x, y, w, h, fill) -> None:
    c.setFillColor(fill)
    c.setStrokeColor(HexColor("#D0D7E2"))
    c.setLineWidth(0.5)
    c.roundRect(x, y, w, h, 3, fill=1, stroke=1)


def _err_colors(err: float | None):
    if err is None:
        return MUTED, LIGHT
    if err < 1.0:
        return GREEN_TXT, GREEN
    if err <= 5.0:
        return AMBER_TXT, AMBER
    return RED_TXT, RED_BG


def build_pdf(path: Path, logo: Path | None) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    W, H = A4
    ml, mr, mt, mb = 1.3 * cm, 1.3 * cm, 1.15 * cm, 1.0 * cm
    usable = W - ml - mr

    y = H - mt
    if logo and logo.exists():
        c.drawImage(str(logo), ml, y - 0.9 * cm, width=3.8 * cm, height=0.85 * cm, mask="auto")
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 12.5)
    c.drawRightString(W - mr, y - 0.3 * cm, "Validaciones Fundo Zapallar")
    c.setFont("Helvetica", 8.5)
    c.setFillColor(MUTED)
    c.drawRightString(W - mr, y - 0.6 * cm, "Consolidado relojería vs App WES")
    y -= 1.1 * cm

    c.setStrokeColor(NAVY)
    c.setLineWidth(1.1)
    c.line(ml, y, W - mr, y)
    y -= 0.35 * cm

    c.setFillColor(MUTED)
    c.setFont("Helvetica", 7.5)
    c.drawString(
        ml,
        y,
        "Matriz ESVAL (000027-01) · Etapa N°5 (000027-03) · Error = 1 − (relojería / app) · Uso interno WES",
    )
    y -= 0.45 * cm

    # KPIs última campaña
    last = CAMPANIAS[-1]
    card_w = (usable - 0.3 * cm) / 2
    card_h = 1.55 * cm
    for i, (label, pto) in enumerate(
        (("Matriz ESVAL — última", last["matriz"]), ("Etapa N°5 — última", last["etapa5"]))
    ):
        tc, bg = _err_colors(pto["error"] if pto else None)
        x = ml + i * (card_w + 0.3 * cm)
        _card(c, x, y - card_h, card_w, card_h, bg)
        c.setFillColor(MUTED)
        c.setFont("Helvetica", 7)
        c.drawString(x + 5, y - 12, f"{label} · {last['ventana']}")
        c.setFillColor(tc)
        c.setFont("Helvetica-Bold", 15)
        c.drawString(x + 5, y - card_h + 10, f"{_fmt(pto['error'], 1)} %" if pto else "—")
    y -= card_h + 0.4 * cm

    # Tabla historial
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(ml, y, "Historial de campañas")
    y -= 0.28 * cm

    headers = ["#", "Ventana", "Matriz reloj", "App", "Err", "E5 reloj", "App", "Err", "Estado"]
    # relative weights
    weights = [0.35, 2.0, 1.35, 1.25, 0.7, 1.15, 1.15, 0.7, 1.5]
    col_w = [usable * w / sum(weights) for w in weights]
    row_h = 0.58 * cm

    def draw_row(yy, vals, header=False, alt=False, err_m=None, err_e=None):
        bg = NAVY if header else (LIGHT if alt else white)
        c.setFillColor(bg)
        c.rect(ml, yy - row_h, usable, row_h, fill=1, stroke=0)
        c.setStrokeColor(HexColor("#C5CEDA"))
        c.setLineWidth(0.35)
        c.rect(ml, yy - row_h, usable, row_h, fill=0, stroke=1)
        x = ml
        for i, (val, w) in enumerate(zip(vals, col_w)):
            if header:
                c.setFillColor(white)
                c.setFont("Helvetica-Bold", 6.5)
            else:
                c.setFillColor(HexColor("#222222"))
                c.setFont("Helvetica-Bold" if i in (4, 7, 8) else "Helvetica", 6.8)
                if i == 4 and err_m is not None:
                    c.setFillColor(_err_colors(err_m)[0])
                if i == 7 and err_e is not None:
                    c.setFillColor(_err_colors(err_e)[0])
            c.drawString(x + 2.5, yy - row_h + 0.18 * cm, val)
            x += w
        return yy - row_h

    y = draw_row(y, headers, header=True)
    for i, camp in enumerate(CAMPANIAS):
        m, e = camp["matriz"], camp["etapa5"]
        estado = []
        if m:
            estado.append(m["lectura"])
        if e:
            estado.append(e["lectura"])
        # compact estado: show worst or both short
        if e and e["error"] > 5:
            estado_txt = "M OK / E5 Revisar"
        elif all(p and p["lectura"] == "Excelente" for p in (m, e) if p):
            estado_txt = "Excelente"
        else:
            estado_txt = "Aceptable"

        vals = [
            camp["id"],
            camp["ventana"],
            f"{_fmt(m['reloj'], 2 if m['reloj'] < 100 else 1)}" if m else "—",
            f"{_fmt(m['app'], 2)}" if m else "—",
            f"{_fmt(m['error'], 1)}%" if m else "—",
            f"{_fmt(e['reloj'], 1)}" if e else "—",
            f"{_fmt(e['app'], 1)}" if e else "—",
            f"{_fmt(e['error'], 1)}%" if e else "—",
            estado_txt,
        ]
        y = draw_row(
            y,
            vals,
            alt=(i % 2 == 1),
            err_m=m["error"] if m else None,
            err_e=e["error"] if e else None,
        )
    y -= 0.35 * cm

    # Notas
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(ml, y, "Notas")
    y -= 0.28 * cm
    c.setFillColor(HexColor("#333333"))
    c.setFont("Helvetica", 7.5)
    for camp in CAMPANIAS:
        line = f"{camp['id']}. {camp['ventana']} — {camp['detalle']}. {camp['nota']}"
        # wrap lightly
        if len(line) > 118:
            # split at ~110
            cut = line.rfind(" ", 0, 115)
            c.drawString(ml, y, line[:cut])
            y -= 0.32 * cm
            c.drawString(ml + 10, y, line[cut + 1 :])
        else:
            c.drawString(ml, y, line)
        y -= 0.34 * cm

    y -= 0.15 * cm
    conc_h = 2.15 * cm
    _card(c, ml, y - conc_h, usable, conc_h, GREEN)
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 9)
    c.drawString(ml + 6, y - 0.35 * cm, "Lectura consolidada")
    c.setFillColor(HexColor("#1B3A1D"))
    c.setFont("Helvetica", 7.8)
    lines = [
        "Desde la relojería del 23-09 → 02-10, Matriz ESVAL se mantiene entre 2,1 % y 2,5 %",
        "(app por encima del Itron). Etapa N°5 cierra excelente/aceptable (0,3 % → 1,4 %).",
        "Última campaña vigente: 02-10 → 15-10 · Matriz 2,5 % · Etapa 5 1,4 %.",
    ]
    ty = y - 0.7 * cm
    for line in lines:
        c.drawString(ml + 6, ty, line)
        ty -= 0.32 * cm

    c.setFillColor(MUTED)
    c.setFont("Helvetica", 7)
    c.drawString(ml, mb - 0.05 * cm, f"WES · generado {datetime.now().strftime('%d-%m-%Y %H:%M')}")
    c.drawRightString(W - mr, mb - 0.05 * cm, "Fundo Zapallar · VALIDACION consolidado")

    c.showPage()
    c.save()
    print(f"[OK] PDF  {path}")


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    base = f"OnePager_Validacion_Fundo_Zapallar_CONSOLIDADO_{stamp}"
    docx_path = OUT_DIR / f"{base}.docx"
    pdf_path = OUT_DIR / f"{base}.pdf"
    logo = Path("/tmp/logo_wes.png")
    if not logo.exists():
        try:
            from PIL import Image

            Image.open(ROOT / "logo wes.bmp").convert("RGBA").save(logo)
        except Exception:
            logo = None

    build_docx(docx_path)
    build_pdf(pdf_path, logo if logo and logo.exists() else None)

    stable_pdf = OUT_DIR / "OnePager_Validacion_Fundo_Zapallar_CONSOLIDADO.pdf"
    stable_docx = OUT_DIR / "OnePager_Validacion_Fundo_Zapallar_CONSOLIDADO.docx"
    stable_pdf.write_bytes(pdf_path.read_bytes())
    stable_docx.write_bytes(docx_path.read_bytes())
    # cleanup stamped
    docx_path.unlink(missing_ok=True)
    pdf_path.unlink(missing_ok=True)
    print(f"[OK] estable {stable_pdf}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
