"""
Plan de alineación ultrasónico WES vs turbina Aguas Andinas — CORMUP Peñalolén.

Misma lógica de campo que Fundo Zapallar: la turbina de la sanitaria es la
referencia de consumo real; el ultrasónico WES se ajusta para coincidir.

Usa boletas con lectura real (no promedio) y WES + proyección de huecos
del comparativo ``generar_comparativo_facturaciones_cormup_penalolen``.

Uso:
  python generar_plan_alineacion_ultrasonico_turbina_cormup.py --skip-download
"""

from __future__ import annotations

import argparse
import statistics
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import List, Optional, Tuple

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Cm, Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from generar_comparativo_facturaciones_cormup_penalolen import (
    COMPANY_ID,
    FilaComparacion,
    NODOS_CORMUP,
    OUT_DIR,
    PDF_CACHE,
    Sitio,
    _cargar_sitios,
    _es_ciclo_septiembre,
    _filas_septiembre,
    convertir_a_pdf,
    cruzar_wes,
    descargar_facturaciones,
)
from generar_reporte_word import (
    add_formatted_heading,
    add_logo_to_header,
    add_picture_with_pagination,
    format_number_chilean,
    get_node_name,
)
from generar_comparativo_facturaciones_cormup_penalolen import (
    HEADING_RGB,
    _add_table_rows,
)

BAND_OK = 0.10  # |k-1| <= 10 % → alineado
BAND_MED = 0.30


@dataclass
class AlineacionSitio:
    sitio: Sitio
    validas: List[FilaComparacion]
    sep: Optional[FilaComparacion]
    k_sep: Optional[float]
    k_mediana: Optional[float]
    k_p25: Optional[float]
    k_p75: Optional[float]
    n_validas: int
    n_estimadas: int

    @property
    def k_recomendada(self) -> Optional[float]:
        if self.k_sep is not None:
            return self.k_sep
        return self.k_mediana

    @property
    def prioridad(self) -> str:
        k = self.k_recomendada
        if k is None:
            return "Sin lectura real"
        desv = abs(k - 1.0)
        if desv <= BAND_OK:
            return "Alineado"
        if desv <= BAND_MED:
            return "Ajuste menor"
        return "Prioridad terreno"

    @property
    def accion(self) -> str:
        k = self.k_recomendada
        if k is None:
            return "Falta boleta con lectura real para fijar factor"
        if abs(k - 1.0) <= BAND_OK:
            return "No tocar: ultrasónico ≈ turbina"
        if k > 1.0:
            return (
                f"WES lee de menos: subir factor/diámetro del ultrasónico "
                f"(× {format_number_chilean(k, 3)}) para igualar la turbina"
            )
        return (
            f"WES lee de más: bajar factor/diámetro del ultrasónico "
            f"(× {format_number_chilean(k, 3)}) para igualar la turbina"
        )


def _k(f: FilaComparacion) -> Optional[float]:
    if f.m3_wes <= 0 or not f.m3_cuenta:
        return None
    return float(f.m3_cuenta) / float(f.m3_wes)


def _percentil(vals: List[float], p: float) -> Optional[float]:
    if not vals:
        return None
    s = sorted(vals)
    if len(s) == 1:
        return s[0]
    idx = (len(s) - 1) * p
    lo = int(idx)
    hi = min(lo + 1, len(s) - 1)
    frac = idx - lo
    return s[lo] * (1 - frac) + s[hi] * frac


def _analizar(sitios: List[Sitio]) -> List[AlineacionSitio]:
    out: List[AlineacionSitio] = []
    for s in sitios:
        validas = [f for f in s.filas if f.valido and _k(f) is not None]
        estimadas = [f for f in s.filas if f.estimado]
        sep_cands = [f for f in validas if _es_ciclo_septiembre(f)]
        sep = sep_cands[-1] if sep_cands else None
        ks = [_k(f) for f in validas]
        ks_f = [x for x in ks if x is not None]
        out.append(
            AlineacionSitio(
                sitio=s,
                validas=validas,
                sep=sep,
                k_sep=_k(sep) if sep else None,
                k_mediana=statistics.median(ks_f) if ks_f else None,
                k_p25=_percentil(ks_f, 0.25),
                k_p75=_percentil(ks_f, 0.75),
                n_validas=len(validas),
                n_estimadas=len(estimadas),
            )
        )
    orden = {"Prioridad terreno": 0, "Ajuste menor": 1, "Alineado": 2, "Sin lectura real": 3}
    out.sort(key=lambda a: (orden.get(a.prioridad, 9), a.sitio.node_id))
    return out


def _write_excel(alineaciones: List[AlineacionSitio], out_xlsx: Path, generado: datetime) -> None:
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="003366")
    header_font = Font(color="FFFFFF", bold=True)
    fill_prio = PatternFill("solid", fgColor="FFE6E6")
    fill_ok = PatternFill("solid", fgColor="E2EFDA")
    fill_med = PatternFill("solid", fgColor="FFF2CC")

    ws = wb.active
    ws.title = "Plan_alineacion"
    headers = [
        "Prioridad",
        "Establecimiento",
        "Nodo WES",
        "k septiembre (turbina/WES)",
        "k mediana históricos válidos",
        "k p25–p75",
        "N° lecturas reales",
        "N° promedios (no sirven)",
        "Acción de campo",
    ]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    for a in alineaciones:
        rango = ""
        if a.k_p25 is not None and a.k_p75 is not None:
            rango = f"{a.k_p25:.3f} – {a.k_p75:.3f}"
        ws.append(
            [
                a.prioridad,
                a.sitio.node_name,
                a.sitio.node_id,
                None if a.k_sep is None else round(a.k_sep, 3),
                None if a.k_mediana is None else round(a.k_mediana, 3),
                rango,
                a.n_validas,
                a.n_estimadas,
                a.accion,
            ]
        )
        fill = {"Prioridad terreno": fill_prio, "Alineado": fill_ok, "Ajuste menor": fill_med}.get(a.prioridad)
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[3].number_format = "0.000"
        row[4].number_format = "0.000"
    ws.freeze_panes = "A2"
    for i, w in enumerate([18, 26, 12, 22, 22, 16, 16, 18, 70], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wd = wb.create_sheet("Factores_por_periodo")
    wd.append(
        [
            "Establecimiento",
            "Nodo",
            "Válida",
            "Emisión",
            "Período",
            "m³ turbina (cuenta)",
            "m³ ultrasónico (WES+proy)",
            "k = turbina / WES",
            "Ciclo septiembre",
            "Observación",
        ]
    )
    for col in range(1, 11):
        c = wd.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
    for a in alineaciones:
        for f in a.sitio.filas:
            kk = _k(f) if f.valido else None
            wd.append(
                [
                    a.sitio.node_name,
                    a.sitio.node_id,
                    "Sí" if f.valido else "No",
                    f.emision.strftime("%d-%m-%Y"),
                    f.periodo_txt,
                    f.m3_cuenta,
                    round(f.m3_wes, 1),
                    None if kk is None else round(kk, 3),
                    "Sí" if _es_ciclo_septiembre(f) else "",
                    f.observacion,
                ]
            )
    for row in wd.iter_rows(min_row=2, max_row=wd.max_row):
        row[5].number_format = "#,##0"
        row[6].number_format = "#,##0.0"
        row[7].number_format = "0.000"
    wd.freeze_panes = "A2"
    wd.auto_filter.ref = f"A1:{get_column_letter(wd.max_column)}{wd.max_row}"
    for i, w in enumerate([26, 12, 10, 12, 28, 16, 22, 16, 14, 50], start=1):
        wd.column_dimensions[get_column_letter(i)].width = w

    wn = wb.create_sheet("Metodo_Zapallar")
    wn["A1"] = "Método (como Fundo Zapallar / Matriz ESVAL)"
    wn["A2"] = (
        "La turbina de la sanitaria es la referencia de m³ reales. "
        "El ultrasónico WES se alinea para que k = m³_turbina / m³_WES ≈ 1,00."
    )
    wn["A3"] = (
        "k > 1: el ultrasónico registra menos que la turbina → subir factor/diámetro WES. "
        "k < 1: el ultrasónico registra más → bajar factor/diámetro WES."
    )
    wn["A4"] = (
        "Solo se usa boleta con lectura real. Promedio / estimado / medidor detenido / cerrado "
        "no sirven para fijar k."
    )
    wn["A5"] = (
        "La k de septiembre es la de campo (ciclo vigente). La mediana histórica avisa si el "
        "desvío es estable (calibración) o cambia de mes (hidráulica / bypass)."
    )
    wn["A6"] = f"Generado: {generado.strftime('%d-%m-%Y %H:%M')}"
    wn.column_dimensions["A"].width = 120
    wb.save(out_xlsx)


def _write_word(alineaciones: List[AlineacionSitio], out_docx: Path, generado: datetime) -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.4)
    add_logo_to_header(doc)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Plan de alineación ultrasónico WES vs turbina", level=1)
    if title.runs:
        title.runs[0].font.color.rgb = HEADING_RGB
    sub = doc.add_paragraph()
    r = sub.add_run("CORMUP / Colegios Peñalolén — mismo criterio que Fundo Zapallar (Matriz ESVAL)")
    r.italic = True
    r.font.color.rgb = HEADING_RGB
    p = doc.add_paragraph()
    p.add_run("Generado: ").bold = True
    p.add_run(generado.strftime("%d-%m-%Y %H:%M"))

    add_formatted_heading(doc, "1. Por qué hay que revisar todos", level=1)
    doc.add_paragraph(
        "En Zapallar la turbina de la sanitaria (Matriz ESVAL) es la referencia de consumo "
        "real y el ultrasónico WES de esa matriz se alineó a esa lectura. En Peñalolén cada "
        "colegio tiene el mismo par: turbina Aguas Andinas (boleta) y ultrasónico WES (API). "
        "El comparativo de cuentas mostró que, aun descartando promedios y rellenando huecos, "
        "casi todos los puntos desvían: no es un tema de boleta estimada, es calibración / "
        "instalación del ultrasónico respecto de la turbina."
    )
    doc.add_paragraph(
        "k = m³ turbina (cuenta, lectura real) ÷ m³ ultrasónico (WES medido + proyección de huecos). "
        "Objetivo de campo: k ≈ 1,00 (±10 %). "
        "k > 1 → WES lee de menos (subir factor). k < 1 → WES lee de más (bajar factor)."
    )

    n_prio = sum(1 for a in alineaciones if a.prioridad == "Prioridad terreno")
    n_med = sum(1 for a in alineaciones if a.prioridad == "Ajuste menor")
    n_ok = sum(1 for a in alineaciones if a.prioridad == "Alineado")
    doc.add_paragraph(
        f"Resultado con k de septiembre (o mediana si no hay ciclo sep. válido): "
        f"{n_prio} prioridad terreno, {n_med} ajuste menor, {n_ok} ya alineados. "
        f"Eduardo de la Barra (000008-02) no tiene facturas en Drive: hay que incluirlo en la ronda."
    )

    add_formatted_heading(doc, "2. Orden de visita (septiembre)", level=1)
    headers = [
        "Prioridad",
        "Establecimiento",
        "Nodo",
        "k sep.",
        "k mediana",
        "N° reales",
        "Acción",
    ]
    rows: List[List[str]] = []
    hi: List[int] = []
    for i, a in enumerate(alineaciones, start=1):
        rows.append(
            [
                a.prioridad,
                a.sitio.node_name,
                a.sitio.node_id,
                format_number_chilean(a.k_sep, 3) if a.k_sep is not None else "—",
                format_number_chilean(a.k_mediana, 3) if a.k_mediana is not None else "—",
                str(a.n_validas),
                a.accion,
            ]
        )
        if a.prioridad == "Prioridad terreno":
            hi.append(i)
    _add_table_rows(doc, headers, rows, highlight=hi, has_total=False)

    add_formatted_heading(doc, "3. Detalle por establecimiento (k de cada boleta real)", level=1)
    headers_d = [
        "Establecimiento",
        "Emisión",
        "Período",
        "m³ turbina",
        "m³ WES+proy",
        "k",
        "Sep.",
    ]
    for a in alineaciones:
        add_formatted_heading(doc, f"{a.sitio.node_name} ({a.sitio.node_id}) — {a.prioridad}", level=2)
        doc.add_paragraph(a.accion)
        if not a.validas:
            doc.add_paragraph("Sin períodos con lectura real.")
            continue
        rows_d: List[List[str]] = []
        hi_d: List[int] = []
        for i, f in enumerate(a.validas, start=1):
            kk = _k(f)
            rows_d.append(
                [
                    a.sitio.node_name,
                    f.emision.strftime("%d-%m-%Y"),
                    f.periodo_txt,
                    format_number_chilean(f.m3_cuenta, 0),
                    format_number_chilean(f.m3_wes, 1),
                    format_number_chilean(kk, 3) if kk is not None else "—",
                    "Sí" if _es_ciclo_septiembre(f) else "",
                ]
            )
            if _es_ciclo_septiembre(f):
                hi_d.append(i)
        _add_table_rows(doc, headers_d, rows_d, highlight=hi_d, has_total=False)
        if a.n_estimadas:
            doc.add_paragraph(
                f"{a.n_estimadas} boleta(s) a promedio/estimado se omiten para fijar k."
            )

    add_formatted_heading(doc, "4. Criterio de terreno (Zapallar)", level=1, page_break_before=True)
    doc.add_paragraph(
        "1) Confirmar que el ultrasónico está en la misma cañería que la turbina de Aguas Andinas "
        "(sin bypass, estanque ni ramal intermedio)."
    )
    doc.add_paragraph(
        "2) Verificar diámetro programado, sentido de flujo y tramo recto. Ajustar el factor "
        "del punto WES con la k de septiembre (ciclo vigente)."
    )
    doc.add_paragraph(
        "3) Si la k histórica salta de un mes a otro (p25–p75 muy abierto), no es solo "
        "calibración: revisar hidráulica (otra acometida, estanque, riego o medidor distinto)."
    )
    doc.add_paragraph(
        "4) Tras el ajuste, contrastar el siguiente ciclo con lectura real (no promedio). "
        "Objetivo: |k − 1| ≤ 0,10, como Carlos Fernández Peña y Matilde en septiembre."
    )
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)


def generar(skip_download: bool = False) -> Tuple[Path, Optional[Path], Path]:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    generado = datetime.now(timezone.utc).astimezone()
    ts = generado.strftime("%Y%m%d_%H%M")

    if skip_download and any(PDF_CACHE.glob("*/*.pdf")):
        locales = {p.name: p for p in PDF_CACHE.iterdir() if p.is_dir()}
        print(f"[INFO] Reuso cache {PDF_CACHE}", flush=True)
    else:
        locales = descargar_facturaciones(PDF_CACHE)

    sitios = _cargar_sitios(locales)
    cruzar_wes(sitios)
    alineaciones = _analizar(sitios)
    for a in alineaciones:
        print(
            f"  [{a.prioridad}] {a.sitio.node_id} {a.sitio.node_name} "
            f"k_sep={a.k_sep} k_med={a.k_mediana}",
            flush=True,
        )

    stem = f"Plan_alineacion_ultrasonico_vs_turbina_CORMUP_{ts}"
    out_xlsx = OUT_DIR / f"{stem}.xlsx"
    out_docx = OUT_DIR / f"{stem}.docx"
    _write_excel(alineaciones, out_xlsx, generado)
    _write_word(alineaciones, out_docx, generado)
    out_pdf = convertir_a_pdf(out_docx)
    print(f"[OK] Excel: {out_xlsx}", flush=True)
    print(f"[OK] Word:  {out_docx}", flush=True)
    if out_pdf:
        print(f"[OK] PDF:   {out_pdf}", flush=True)
    return out_docx, out_pdf, out_xlsx


def main() -> int:
    if sys.platform == "win32":
        for s in (sys.stdout, sys.stderr):
            try:
                s.reconfigure(encoding="utf-8", line_buffering=True)
            except Exception:
                pass
    ap = argparse.ArgumentParser(description="Plan alineación ultrasónico vs turbina CORMUP")
    ap.add_argument("--skip-download", action="store_true")
    args = ap.parse_args()
    generar(skip_download=args.skip_download)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
