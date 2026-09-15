"""Calendario enero 2026: qué días SÍ operó el control nocturno en Pizza Hut."""

from __future__ import annotations

import calendar
import csv
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Patch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.shared import Cm, Inches, Pt, RGBColor
from matplotlib.backends.backend_pdf import PdfPages

from analizar_pizza_hut_cero_nocturno import (
    NODE_ID,
    _es_cero,
    _fetch_dia,
    _fmt_fecha,
)
from wes_estilo_graficos_app import COLOR_BARRA_WES, COLOR_NOCHE, guardar_grafico_horario_24h_app

DOW = ["lun", "mar", "mié", "jue", "vie", "sáb", "dom"]
DOW_LARGO = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
COLOR_WES = "1F4E79"
EPS = 1e-9


def _tira_cero(horas: Dict[int, float]) -> Tuple[Optional[int], Optional[int]]:
    """Primera y última hora en cero de la tira continua más larga entre 00 y 07."""
    best = (None, None)
    best_len = 0
    i = 0
    while i <= 7:
        if i in horas and _es_cero(horas[i]):
            j = i
            while j <= 7 and j in horas and _es_cero(horas[j]):
                j += 1
            if j - i > best_len:
                best_len = j - i
                best = (i, j - 1)
            i = j
        else:
            i += 1
    return best


def clasificar(dia: date, horas: Dict[int, float]) -> dict:
    core = [1, 2, 3, 4]
    if not horas or not all(h in horas for h in core):
        return {
            "fecha": dia,
            "estado": "sin_datos",
            "n_cero": 0,
            "inicio": "",
            "fin": "",
            "horas": horas,
        }
    n_cero = sum(1 for h in core if _es_cero(horas[h]))
    ini, fin_h = _tira_cero(horas)
    inicio = f"{ini:02d}:00" if ini is not None else ""
    fin = f"{fin_h + 1:02d}:00" if fin_h is not None else ""
    # Completo: 01–04 en cero (toda la ventana 00:30–05:00).
    # Parcial / tardío: el corte se mantuvo hasta la mañana (04:00 y 05:00 en cero)
    # aunque a las 00:30 / 01:00 todavía hubiera agua.
    # Si el agua volvió a las 03:00, NO es control nocturno.
    se_mantuvo_manana = (
        4 in horas
        and 5 in horas
        and _es_cero(horas[4])
        and _es_cero(horas[5])
    )
    if n_cero == 4:
        estado = "completo"
    elif se_mantuvo_manana:
        estado = "parcial"
    else:
        estado = "no"
    return {
        "fecha": dia,
        "estado": estado,
        "n_cero": n_cero,
        "inicio": inicio,
        "fin": fin,
        "horas": horas,
    }


def calendario_enero(dias: List[dict], out: Path) -> Path:
    fig, ax = plt.subplots(figsize=(10.2, 7.2))
    fig.patch.set_facecolor("white")
    ax.set_xlim(0, 7)
    ax.set_ylim(0, 7)
    ax.axis("off")
    ax.set_title(
        "Pizza Hut — enero 2026\n¿Qué días operó el control de 00:30 a 05:00?",
        fontsize=14,
        fontweight="bold",
        color="#1F4E79",
        pad=8,
    )
    colores = {
        "completo": "#2E7D32",
        "parcial": "#EF6C00",
        "no": "#90A4AE",
        "sin_datos": "#EEEEEE",
    }
    by = {d["fecha"]: d for d in dias}
    for i, name in enumerate(DOW):
        ax.text(i + 0.5, 6.55, name, ha="center", va="center", fontsize=10, fontweight="bold")

    first = date(2026, 1, 1)
    # Monday=0
    offset = first.weekday()
    for day in range(1, 32):
        d = date(2026, 1, day)
        cell = offset + day - 1
        col = cell % 7
        row = 5 - cell // 7
        info = by.get(d, {"estado": "sin_datos", "inicio": "", "fin": ""})
        color = colores[info["estado"]]
        rect = FancyBboxPatch(
            (col + 0.06, row + 0.08),
            0.88,
            0.84,
            boxstyle="round,pad=0.02,rounding_size=0.08",
            facecolor=color,
            edgecolor="white",
            linewidth=1.5,
        )
        ax.add_patch(rect)
        txt = "white" if info["estado"] in {"completo", "parcial"} else "#333333"
        ax.text(col + 0.5, row + 0.58, str(day), ha="center", va="center", fontsize=13, fontweight="bold", color=txt)
        if info["estado"] == "completo":
            ax.text(col + 0.5, row + 0.28, "SÍ cortó", ha="center", va="center", fontsize=7, color="white")
        elif info["estado"] == "parcial":
            ax.text(
                col + 0.5,
                row + 0.28,
                f"{info['inicio']}–{info['fin']}",
                ha="center",
                va="center",
                fontsize=6.5,
                color="white",
            )

    ax.legend(
        handles=[
            Patch(facecolor="#2E7D32", label="SÍ — cortó 00:30 a 05:00 completo"),
            Patch(facecolor="#EF6C00", label="A medias — cortó más tarde (parte de la madrugada)"),
            Patch(facecolor="#90A4AE", label="NO — de madrugada siguió pasando agua"),
        ],
        loc="lower center",
        bbox_to_anchor=(0.5, -0.08),
        ncol=1,
        frameon=False,
        fontsize=9,
    )
    fig.tight_layout()
    out.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return out


def barras_enero(dias: List[dict], out: Path) -> Path:
    fig, ax = plt.subplots(figsize=(11.2, 3.8))
    xs = [d["fecha"].day for d in dias]
    colors = []
    for d in dias:
        colors.append({"completo": "#2E7D32", "parcial": "#EF6C00", "no": "#90A4AE", "sin_datos": "#EEEEEE"}[d["estado"]])
    ys = [1 if d["estado"] in {"completo", "parcial"} else 0 for d in dias]
    # show 3 levels
    ys = [{"completo": 2, "parcial": 1, "no": 0, "sin_datos": 0}[d["estado"]] for d in dias]
    ax.bar(xs, ys, color=colors, width=0.85)
    ax.set_xticks(xs)
    ax.set_xticklabels(xs, fontsize=8)
    ax.set_yticks([0, 1, 2])
    ax.set_yticklabels(["No cortó", "Cortó a medias", "Cortó 00:30–05:00"])
    ax.set_xlabel("Día de enero 2026")
    ax.set_title("Pizza Hut — enero 2026: ¿operó el control esa noche?", fontweight="bold", color="#1F4E79")
    ax.set_ylim(0, 2.4)
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(out, dpi=150, bbox_inches="tight")
    plt.close(fig)
    return out


def _shade(cell, hex_color: str) -> None:
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'w:val="clear" w:fill="{hex_color}"/>'
        )
    )


def _cell(cell, text, *, fill=None, bold=False, size=10, color=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Calibri"
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if fill:
        _shade(cell, fill)


def armar_word(
    dias: List[dict],
    cal: Path,
    barras: Path,
    perfiles: List[Path],
    out: Path,
) -> Path:
    comp = [d for d in dias if d["estado"] == "completo"]
    parc = [d for d in dias if d["estado"] == "parcial"]
    no = [d for d in dias if d["estado"] == "no"]

    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.5)
        s.bottom_margin = Cm(1.5)
        s.left_margin = Cm(1.7)
        s.right_margin = Cm(1.7)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("PIZZA HUT — ¿qué días cortó de noche?")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string(COLOR_WES)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("Enero 2026  ·  000025-07  ·  Parque Arauco Estación")
    r.font.size = Pt(12)

    p = doc.add_paragraph()
    r = p.add_run("En una frase: ")
    r.bold = True
    p.add_run(
        "en enero el control casi no cubrió toda la madrugada 00:30–05:00. "
        "Solo lo hizo 2 noches (domingo 25 y viernes 30). "
        "Varias noches sí cortó, pero más tarde (desde las 02:00 o 03:00). "
        "Eso es lo que suele verse en la app y por eso se recuerda enero."
    )

    h = doc.add_heading("Los 2 días que SÍ cortaron 00:30 a 05:00", 1)
    for run in h.runs:
        run.font.color.rgb = RGBColor.from_string(COLOR_WES)

    tbl = doc.add_table(rows=1 + len(comp), cols=4)
    tbl.style = "Table Grid"
    for i, x in enumerate(["Día", "¿Cortó?", "Desde", "Hasta (volvió el agua)"]):
        _cell(tbl.rows[0].cells[i], x, fill=COLOR_WES, bold=True, size=11, color="FFFFFF")
    for i, d in enumerate(comp, start=1):
        _cell(tbl.rows[i].cells[0], f"{DOW_LARGO[d['fecha'].weekday()]} {_fmt_fecha(d['fecha'])}", fill="C8E6C9", bold=True)
        _cell(tbl.rows[i].cells[1], "SÍ, completo", fill="C8E6C9")
        _cell(tbl.rows[i].cells[2], d["inicio"], fill="C8E6C9")
        _cell(tbl.rows[i].cells[3], d["fin"], fill="C8E6C9")

    doc.add_paragraph()
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(cal), width=Inches(6.3))

    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(barras), width=Inches(6.3))

    h = doc.add_heading("Días que cortaron a medias (más tarde)", 1)
    for run in h.runs:
        run.font.color.rgb = RGBColor.from_string(COLOR_WES)
    doc.add_paragraph(
        "Acá el agua todavía pasaba cerca de las 00:30, y el corte partió después. "
        "No cuenta como “ventana 00:30–05:00 en cero”, pero el control sí se movió."
    )
    tbl = doc.add_table(rows=1 + len(parc), cols=4)
    tbl.style = "Table Grid"
    for i, x in enumerate(["Día", "Horas en cero (01–04)", "Corte desde", "Volvió a las"]):
        _cell(tbl.rows[0].cells[i], x, fill="EF6C00", bold=True, size=10, color="FFFFFF")
    for i, d in enumerate(parc, start=1):
        fill = "FFE0B2"
        _cell(tbl.rows[i].cells[0], f"{DOW_LARGO[d['fecha'].weekday()]} {_fmt_fecha(d['fecha'])}", fill=fill)
        _cell(tbl.rows[i].cells[1], f"{d['n_cero']} de 4", fill=fill)
        _cell(tbl.rows[i].cells[2], d["inicio"], fill=fill)
        _cell(tbl.rows[i].cells[3], d["fin"], fill=fill)

    h = doc.add_heading("Cómo se ve una noche que SÍ cortó vs una a medias", 1)
    for run in h.runs:
        run.font.color.rgb = RGBColor.from_string(COLOR_WES)
    for pth in perfiles:
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(pth), width=Inches(6.2))

    doc.add_paragraph(
        f"Resumen enero 2026: {len(comp)} noches con corte completo 00:30–05:00, "
        f"{len(parc)} noches con corte a medias, {len(no)} noches sin corte. "
        f"Generado {datetime.now():%d-%m-%Y %H:%M}."
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def armar_pdf(dias: List[dict], cal: Path, barras: Path, perfiles: List[Path], out: Path) -> Path:
    comp = [d for d in dias if d["estado"] == "completo"]
    parc = [d for d in dias if d["estado"] == "parcial"]
    with PdfPages(out) as pdf:
        fig = plt.figure(figsize=(8.27, 11.69))
        ax = fig.add_subplot(111)
        ax.axis("off")
        ax.text(0.5, 0.95, "PIZZA HUT — enero 2026", ha="center", fontsize=16, color="#1F4E79", fontweight="bold")
        ax.text(0.5, 0.91, "Días en que SÍ operó el control nocturno", ha="center", fontsize=12)
        ax.text(
            0.08,
            0.84,
            "Solo 2 noches cortaron toda la ventana 00:30–05:00:\n"
            "  • domingo 25-01-2026  →  desde 00:00  hasta 08:00\n"
            "  • viernes 30-01-2026  →  desde 00:00  hasta 08:00\n\n"
            "El resto de enero: o no cortó, o cortó más tarde (02:00 / 03:00).\n"
            "Eso se ve en naranja en el calendario.",
            ha="left",
            va="top",
            fontsize=11,
            family="DejaVu Sans",
        )
        img = plt.imread(str(cal))
        ax2 = fig.add_axes([0.08, 0.08, 0.84, 0.58])
        ax2.imshow(img)
        ax2.axis("off")
        pdf.savefig(fig)
        plt.close(fig)

        fig = plt.figure(figsize=(8.27, 11.69))
        ax = fig.add_subplot(111)
        ax.axis("off")
        ax.set_title("Enero día por día", color="#1F4E79", fontsize=13, pad=12)
        img = plt.imread(str(barras))
        ax.imshow(img)
        ax.axis("off")
        pdf.savefig(fig)
        plt.close(fig)

        # table of complete + partial
        rows = [["Día", "Estado", "Desde", "Hasta"]]
        for d in dias:
            if d["estado"] not in {"completo", "parcial"}:
                continue
            est = "SÍ, 00:30–05:00" if d["estado"] == "completo" else "a medias"
            rows.append(
                [
                    f"{d['fecha'].strftime('%d-%m')} {DOW[d['fecha'].weekday()]}",
                    est,
                    d["inicio"],
                    d["fin"],
                ]
            )
        fig = plt.figure(figsize=(8.27, 11.69))
        ax = fig.add_subplot(111)
        ax.axis("off")
        ax.set_title("Lista de noches que sí se movió el control", color="#1F4E79", fontsize=13, pad=16)
        tbl = ax.table(cellText=rows, loc="upper center", cellLoc="center")
        tbl.auto_set_font_size(False)
        tbl.set_fontsize(8)
        tbl.scale(1.2, 1.45)
        for j in range(4):
            tbl[0, j].set_facecolor("#1F4E79")
            tbl[0, j].set_text_props(color="white", fontweight="bold")
        for i, d in enumerate([x for x in dias if x["estado"] in {"completo", "parcial"}], start=1):
            c = "#C8E6C9" if d["estado"] == "completo" else "#FFE0B2"
            for j in range(4):
                tbl[i, j].set_facecolor(c)
        pdf.savefig(fig)
        plt.close(fig)

        for pth in perfiles:
            fig = plt.figure(figsize=(8.27, 11.69))
            ax = fig.add_subplot(111)
            ax.imshow(plt.imread(str(pth)))
            ax.axis("off")
            fig.tight_layout()
            pdf.savefig(fig)
            plt.close(fig)
    return out


def main() -> int:
    out_dir = Path("reports/Puntos_En_Cero/Pizza_Hut/Enero_2026")
    out_dir.mkdir(parents=True, exist_ok=True)
    gra = out_dir / "graficos"
    gra.mkdir(parents=True, exist_ok=True)

    dias: List[dict] = []
    d = date(2026, 1, 1)
    print("[INFO] Descargando enero 2026...")
    while d <= date(2026, 1, 31):
        _, horas = _fetch_dia(d)
        dias.append(clasificar(d, horas))
        d += timedelta(days=1)

    comp = [x for x in dias if x["estado"] == "completo"]
    parc = [x for x in dias if x["estado"] == "parcial"]
    print(f"[INFO] Completos: {len(comp)}  Parciales: {len(parc)}")
    for x in comp:
        print(f"  SÍ  {_fmt_fecha(x['fecha'])}  {x['inicio']} → {x['fin']}")

    cal = calendario_enero(dias, gra / "calendario_enero_2026.png")
    barras = barras_enero(dias, gra / "barras_enero_2026.png")

    perfiles: List[Path] = []
    ejemplos = [
        (date(2026, 1, 25), "SÍ cortó 00:30–05:00 — domingo 25-01"),
        (date(2026, 1, 30), "SÍ cortó 00:30–05:00 — viernes 30-01"),
        (date(2026, 1, 23), "A medias — viernes 23-01 (cortó desde las 02:00)"),
        (date(2026, 1, 1), "NO cortó — jueves 01-01 (madrugada con agua)"),
    ]
    by = {x["fecha"]: x for x in dias}
    for dia, titulo in ejemplos:
        horas = by[dia]["horas"]
        p = gra / f"perfil_{dia.isoformat()}.png"
        guardar_grafico_horario_24h_app(horas, p, titulo=f"Pizza Hut — {titulo}", subtitulo="m³/h, hora Chile. Rojo = 00:00–06:00.")
        perfiles.append(p)

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    csv_path = out_dir / "enero_2026_dias_control.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["fecha", "estado", "horas_cero_01_04", "corte_desde", "volvio_a"])
        for x in dias:
            w.writerow([x["fecha"].isoformat(), x["estado"], x["n_cero"], x["inicio"], x["fin"]])

    docx_path = out_dir / f"Pizza_Hut_control_enero_2026_{stamp}.docx"
    pdf_path = out_dir / f"Pizza_Hut_control_enero_2026_{stamp}.pdf"
    armar_word(dias, cal, barras, perfiles, docx_path)
    armar_pdf(dias, cal, barras, perfiles, pdf_path)
    print(f"DOCX={docx_path}")
    print(f"PDF={pdf_path}")
    print(f"CAL={cal}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
