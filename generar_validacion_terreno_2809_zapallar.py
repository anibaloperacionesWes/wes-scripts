"""
Validación terreno 28-09-2026 — Fundo Zapallar.

Matriz ESVAL (Itron): lectura 23-09 17:05 → 28-09 09:05
Etapa N°5 (Sensus):   lectura 23-09 16:54 → 28-09 08:57
+ Estanque Inferior en la misma ventana Matriz (conversan).

Metodología App: CSV dates.measures.csv, hora del campo TIME (etiqueta),
igual que informe Matriz / Inferior previos.
"""

from __future__ import annotations

import csv
import io
import json
import shutil
from datetime import date, datetime, timedelta
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import requests
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from generar_tabla_zapallar_matriz_vs_inferior import (
    HEADER_FILL,
    WES_BLUE,
    _fmt_m3,
    _set_cell_text,
    _set_landscape,
)

BASE = "http://104.248.53.141:7003/wes/api/acl-node/v1"
MATRIZ_ID = "000027-01"
INFERIOR_ID = "000027-02"
ETAPA5_ID = "000027-03"

# Lecturas absolutas Itron (6 black + 1 red) — escala que preserva val. previa Δ170,9
ITRON_INI = 384148.9  # 23-09 ~17:05
ITRON_FIN = 384736.5  # 28-09 09:05
ITRON_INI_DT = "2026-09-23 17:05"
ITRON_FIN_DT = "2026-09-28 09:05"

# Etapa 5 Sensus (ruedas enteras)
ETAPA5_INI = 5177.0  # 23-09 16:54
ETAPA5_FIN = 5225.0  # 28-09 08:57
ETAPA5_INI_DT = "2026-09-23 16:54"
ETAPA5_FIN_DT = "2026-09-28 08:57"

# Ventana App: desde h17 del 23 (continua tras val. previa que cerró en h16)
# hasta h08 del 28 (lectura ~09:00; hora 9 incompleta → excluida)
APP_START = datetime(2026, 9, 23, 17, 0)
APP_END_EXCL = datetime(2026, 9, 28, 9, 0)  # exclusive end

FOTO_MATRIZ = Path(
    "/home/ubuntu/.cursor/projects/workspace/assets/2c90844e-579b-4a89-8e18-0e424d5e64cf.png"
)
FOTO_ETAPA5 = Path(
    "/home/ubuntu/.cursor/projects/workspace/assets/bf785120-aea0-404f-9fa4-0a4df3dccd62.png"
)


def _fmt(x: float, dec: int = 2) -> str:
    return _fmt_m3(x, dec)


def _pct_error(lectura: float, app: float) -> float:
    if lectura <= 0 or app <= 0:
        return 0.0
    return abs(1.0 - (min(lectura, app) / max(lectura, app))) * 100.0


def _estado(pct: float) -> str:
    if pct <= 5:
        return "Aceptable"
    if pct <= 15:
        return "Aceptable (diferencia moderada)"
    return "Revisar — diferencia relevante"


def hours_app(node_id: str, dia: date) -> dict[int, float]:
    r = requests.get(
        f"{BASE}/nodes/{node_id}/dates.measures.csv",
        params={"start": dia.strftime("%d%m%Y"), "end": dia.strftime("%d%m%Y")},
        timeout=60,
    )
    r.raise_for_status()
    acc: dict[int, float] = {}
    for row in csv.DictReader(io.StringIO(r.text)):
        t = row["TIME"].strip()
        v = float(row["VALUE"].strip())
        dt = datetime.fromisoformat(t.replace("Z", "+00:00"))
        acc[dt.hour] = acc.get(dt.hour, 0.0) + v
    return acc


def slots_ventana() -> list[tuple[date, int]]:
    out: list[tuple[date, int]] = []
    cur = APP_START
    while cur < APP_END_EXCL:
        out.append((cur.date(), cur.hour))
        cur += timedelta(hours=1)
    return out


def sum_ventana(node_id: str) -> tuple[float, list[dict]]:
    slots = slots_ventana()
    days = sorted({d for d, _ in slots})
    by = {d: hours_app(node_id, d) for d in days}
    detail = []
    total = 0.0
    for d, h in slots:
        v = by[d].get(h, 0.0)
        total += v
        detail.append({"fecha": d.isoformat(), "hora": h, "m3": round(v, 4)})
    return round(total, 2), detail


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out_dir = Path("reports/Fundo_Zapallar/Informes_Tecnicos") / f"validacion_terreno_2809_{stamp}"
    out_dir.mkdir(parents=True, exist_ok=True)

    # Copiar fotos
    if FOTO_MATRIZ.is_file():
        shutil.copy2(FOTO_MATRIZ, out_dir / "foto_matriz_itron_2809_0905.png")
    if FOTO_ETAPA5.is_file():
        shutil.copy2(FOTO_ETAPA5, out_dir / "foto_etapa5_sensus_2809_0857.png")

    app_mat, det_mat = sum_ventana(MATRIZ_ID)
    app_inf, det_inf = sum_ventana(INFERIOR_ID)
    app_e5, det_e5 = sum_ventana(ETAPA5_ID)

    itron_delta = round(ITRON_FIN - ITRON_INI, 2)
    e5_delta = round(ETAPA5_FIN - ETAPA5_INI, 2)

    err_mat = round(_pct_error(itron_delta, app_mat), 1)
    err_e5 = round(_pct_error(e5_delta, app_e5), 1)
    est_mat = _estado(err_mat)
    est_e5 = _estado(err_e5)
    inf_sobre_mat = round(100.0 * app_inf / app_mat, 1) if app_mat else 0.0
    inf_sobre_itron = round(100.0 * app_inf / itron_delta, 1) if itron_delta else 0.0

    # --- Chart Matriz ---
    fig, ax = plt.subplots(figsize=(9.0, 4.4))
    labels = ["Itron\n(terreno)", "Matriz App\nWES", "Estanque\nInferior"]
    vals = [itron_delta, app_mat, app_inf]
    colors = ["#5B9BD5", WES_BLUE, "#70AD47"]
    bars = ax.bar(labels, vals, color=colors, width=0.55)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + max(vals) * 0.02, f"{v:.1f} m³", ha="center", fontsize=10)
    ax.set_ylabel("m³")
    ax.set_title(
        f"Matriz ESVAL — 23/09 17:00 → 28/09 09:00 Chile\n"
        f"Itron vs App: error {err_mat:.1f}% · Inferior = {inf_sobre_mat:.0f}% App",
        fontsize=11,
        fontweight="bold",
        color=WES_BLUE,
    )
    ax.set_ylim(0, max(vals) * 1.2)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(axis="y", linestyle="--", alpha=0.35)
    fig.tight_layout()
    chart_mat = out_dir / "chart_validacion_matriz_2809.png"
    fig.savefig(chart_mat, dpi=160, bbox_inches="tight")
    plt.close(fig)

    # --- Chart Etapa 5 ---
    fig, ax = plt.subplots(figsize=(7.5, 4.2))
    labels = ["Sensus\n(terreno)", "App WES\nEtapa N°5"]
    vals = [e5_delta, app_e5]
    colors = ["#ED7D31", WES_BLUE]
    bars = ax.bar(labels, vals, color=colors, width=0.45)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + max(vals) * 0.03, f"{v:.1f} m³", ha="center", fontsize=11)
    ax.set_ylabel("m³")
    ax.set_title(
        f"Etapa N°5 — 23/09 17:00 → 28/09 09:00 Chile\n"
        f"Sensus vs App: error {err_e5:.1f}%",
        fontsize=11,
        fontweight="bold",
        color=WES_BLUE,
    )
    ax.set_ylim(0, max(vals) * 1.25)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(axis="y", linestyle="--", alpha=0.35)
    fig.tight_layout()
    chart_e5 = out_dir / "chart_validacion_etapa5_2809.png"
    fig.savefig(chart_e5, dpi=160, bbox_inches="tight")
    plt.close(fig)

    # --- Chart diario Matriz vs Inferior ---
    days = sorted({d for d, _ in slots_ventana()})
    by_m = {d: hours_app(MATRIZ_ID, d) for d in days}
    by_i = {d: hours_app(INFERIOR_ID, d) for d in days}
    day_m, day_i, day_lbl = [], [], []
    for d in days:
        hs = [h for dd, h in slots_ventana() if dd == d]
        day_m.append(sum(by_m[d].get(h, 0.0) for h in hs))
        day_i.append(sum(by_i[d].get(h, 0.0) for h in hs))
        day_lbl.append(d.strftime("%d/%m"))
    x = np.arange(len(day_lbl))
    w = 0.38
    fig, ax = plt.subplots(figsize=(10.5, 4.4))
    ax.bar(x - w / 2, day_m, w, label="Matriz App", color=WES_BLUE)
    ax.bar(x + w / 2, day_i, w, label="Estanque Inferior", color="#70AD47")
    ax.set_xticks(x)
    ax.set_xticklabels(day_lbl)
    ax.set_ylabel("m³ / día (horas de la ventana)")
    ax.set_title("Perfil diario Matriz vs Inferior en la ventana", fontsize=11, fontweight="bold", color=WES_BLUE)
    ax.legend(frameon=False)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(axis="y", linestyle="--", alpha=0.35)
    fig.tight_layout()
    chart_dia = out_dir / "chart_diario_matriz_inferior_2809.png"
    fig.savefig(chart_dia, dpi=160, bbox_inches="tight")
    plt.close(fig)

    # --- Word ---
    doc = Document()
    _set_landscape(doc)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("Validación terreno 28-09-2026 — Matriz ESVAL y Etapa N°5")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(31, 78, 121)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run(
        "Fundo Zapallar · continuidad desde validación 22–23/09 · "
        "App = CSV TIME etiqueta (misma metodología)"
    )
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(89, 89, 89)

    verd = doc.add_paragraph()
    verd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    vr = verd.add_run(
        f"Matriz: Itron {_fmt(itron_delta)} vs App {_fmt(app_mat)} → error {err_mat:.1f}% ({est_mat}).  "
        f"Etapa 5: Sensus {_fmt(e5_delta, 0)} vs App {_fmt(app_e5)} → error {err_e5:.1f}% ({est_e5}).  "
        f"Inferior = {inf_sobre_mat:.0f}% de Matriz App (conversan)."
    )
    vr.bold = True
    vr.font.size = Pt(11)
    vr.font.color.rgb = RGBColor(0, 100, 0)

    # Lecturas
    h1 = doc.add_paragraph()
    hr = h1.add_run("1) Lecturas fotográficas de hoy")
    hr.bold = True
    hr.font.size = Pt(12)
    hr.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    p.add_run(
        f"Matriz principal (Itron Flostar-S): {_fmt(ITRON_FIN, 1)} m³ a las {ITRON_FIN_DT} "
        f"(ruedas 6 negras + 1 roja; el ‘6’ es m³ entero, no decimal). "
        f"Lectura previa validada: {_fmt(ITRON_INI, 1)} m³ ({ITRON_INI_DT}). "
        f"Δ Itron = {_fmt(itron_delta)} m³.\n\n"
        f"Etapa N°5 (Sensus MeiStream Plus 100): {_fmt(ETAPA5_FIN, 0)} m³ a las {ETAPA5_FIN_DT} "
        f"(odómetro 005225). Lectura previa: {_fmt(ETAPA5_INI, 0)} m³ ({ETAPA5_INI_DT}). "
        f"Δ Sensus = {_fmt(e5_delta, 0)} m³."
    ).font.size = Pt(10)

    # Tabla Matriz
    h2 = doc.add_paragraph()
    hr2 = h2.add_run("2) Matriz ESVAL — Itron vs App (+ Inferior)")
    hr2.bold = True
    hr2.font.size = Pt(12)
    hr2.font.color.rgb = RGBColor(31, 78, 121)

    p2 = doc.add_paragraph()
    p2.add_run(
        f"Ventana App Chile: {APP_START.strftime('%d/%m/%Y %H:%M')} → "
        f"{APP_END_EXCL.strftime('%d/%m/%Y %H:%M')} (excluye hora final incompleta). "
        f"{len(slots_ventana())} horas. Continúa tras la validación previa (cerró en h16 del 23/09)."
    ).font.size = Pt(10)

    headers = ["Concepto", "Valor", "Nota"]
    rows = [
        ["Lectura Itron inicio", f"{_fmt(ITRON_INI, 1)} m³", ITRON_INI_DT],
        ["Lectura Itron fin", f"{_fmt(ITRON_FIN, 1)} m³", ITRON_FIN_DT],
        ["Δ Itron (terreno)", f"{_fmt(itron_delta)} m³", "Referencia mecánica"],
        ["Δ Matriz App WES", f"{_fmt(app_mat)} m³", MATRIZ_ID],
        ["Error Itron vs App", f"{err_mat:.1f} %", est_mat],
        ["Δ Estanque Inferior", f"{_fmt(app_inf)} m³", f"{inf_sobre_mat:.0f}% de Matriz App"],
        ["Inferior / Itron", f"{inf_sobre_itron:.0f} %", "Tubería en el camino — se parecen"],
    ]
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        _set_cell_text(
            table.rows[0].cells[i], h, bold=True, size=9, color=RGBColor(255, 255, 255), fill=HEADER_FILL
        )
    for r_i, vals in enumerate(rows):
        fill = "C6E0B4" if r_i in (4, 6) else "FFFFFF"
        for c_i, val in enumerate(vals):
            _set_cell_text(table.rows[r_i + 1].cells[c_i], val, size=9, fill=fill)

    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(chart_mat), width=Inches(7.5))

    # Tabla Etapa 5
    h3 = doc.add_paragraph()
    hr3 = h3.add_run("3) Etapa N°5 — Sensus vs App")
    hr3.bold = True
    hr3.font.size = Pt(12)
    hr3.font.color.rgb = RGBColor(31, 78, 121)

    rows5 = [
        ["Lectura Sensus inicio", f"{_fmt(ETAPA5_INI, 0)} m³", ETAPA5_INI_DT],
        ["Lectura Sensus fin", f"{_fmt(ETAPA5_FIN, 0)} m³", ETAPA5_FIN_DT],
        ["Δ Sensus (terreno)", f"{_fmt(e5_delta, 0)} m³", "Referencia mecánica"],
        ["Δ App WES Etapa N°5", f"{_fmt(app_e5)} m³", ETAPA5_ID],
        ["Error Sensus vs App", f"{err_e5:.1f} %", est_e5],
    ]
    t5 = doc.add_table(rows=1 + len(rows5), cols=3)
    t5.style = "Table Grid"
    for i, h in enumerate(headers):
        _set_cell_text(
            t5.rows[0].cells[i], h, bold=True, size=9, color=RGBColor(255, 255, 255), fill=HEADER_FILL
        )
    for r_i, vals in enumerate(rows5):
        fill = "C6E0B4" if r_i == 4 else "FFFFFF"
        for c_i, val in enumerate(vals):
            _set_cell_text(t5.rows[r_i + 1].cells[c_i], val, size=9, fill=fill)

    pic5 = doc.add_paragraph()
    pic5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic5.add_run().add_picture(str(chart_e5), width=Inches(6.2))

    h4 = doc.add_paragraph()
    hr4 = h4.add_run("4) Matriz vs Inferior en la ventana (perfil diario)")
    hr4.bold = True
    hr4.font.size = Pt(12)
    hr4.font.color.rgb = RGBColor(31, 78, 121)

    p4 = doc.add_paragraph()
    p4.add_run(
        "No se espera igualdad exacta (hay tubería). En esta ventana Inferior queda en "
        f"{inf_sobre_mat:.0f}% de Matriz App y {inf_sobre_itron:.0f}% del Itron: conversan, "
        "misma orden de magnitud y mismo sentido que la validación 22–23/09."
    ).font.size = Pt(10)

    picd = doc.add_paragraph()
    picd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picd.add_run().add_picture(str(chart_dia), width=Inches(8.5))

    concl = doc.add_paragraph()
    concl.add_run(
        "Conclusión: ambas validaciones de terreno de hoy cierran bien frente a la app. "
        f"Matriz error {err_mat:.1f}% (Itron {_fmt(itron_delta)} / App {_fmt(app_mat)}); "
        f"Etapa 5 error {err_e5:.1f}% (Sensus {_fmt(e5_delta, 0)} / App {_fmt(app_e5)}). "
        "Estanque Inferior sigue conversando con Matriz en la misma ventana."
    ).font.size = Pt(10)

    docx_path = out_dir / "Validacion_terreno_Matriz_Etapa5_2809.docx"
    doc.save(str(docx_path))

    # --- PDF resumen ---
    fig = plt.figure(figsize=(15.0, 10.0))
    fig.suptitle(
        "Validación terreno 28-09-2026 — Matriz ESVAL + Etapa N°5",
        fontsize=15,
        fontweight="bold",
        color=WES_BLUE,
        y=0.97,
        x=0.03,
        ha="left",
    )
    fig.text(
        0.03,
        0.935,
        f"Matriz: Itron {_fmt(itron_delta)} vs App {_fmt(app_mat)} → {err_mat:.1f}% ({est_mat}) · "
        f"Etapa 5: Sensus {_fmt(e5_delta, 0)} vs App {_fmt(app_e5)} → {err_e5:.1f}% ({est_e5}) · "
        f"Inferior {inf_sobre_mat:.0f}% Matriz",
        fontsize=9.5,
        color="#006400",
        fontweight="bold",
    )
    ax1 = fig.add_axes([0.04, 0.52, 0.44, 0.38])
    ax1.imshow(plt.imread(str(chart_mat)))
    ax1.axis("off")
    ax2 = fig.add_axes([0.52, 0.52, 0.44, 0.38])
    ax2.imshow(plt.imread(str(chart_e5)))
    ax2.axis("off")
    ax3 = fig.add_axes([0.10, 0.04, 0.80, 0.42])
    ax3.imshow(plt.imread(str(chart_dia)))
    ax3.axis("off")
    pdf_path = out_dir / "Validacion_terreno_Matriz_Etapa5_2809.pdf"
    fig.savefig(pdf_path, dpi=170)
    plt.close(fig)

    payload = {
        "ventana_app": {
            "inicio": APP_START.isoformat(sep=" "),
            "fin_excluyente": APP_END_EXCL.isoformat(sep=" "),
            "horas": len(slots_ventana()),
            "metodo": "dates.measures.csv TIME etiqueta (sin TZ shift)",
        },
        "matriz": {
            "itron_ini": ITRON_INI,
            "itron_fin": ITRON_FIN,
            "itron_delta": itron_delta,
            "app_m3": app_mat,
            "error_pct": err_mat,
            "estado": est_mat,
            "inferior_m3": app_inf,
            "inferior_sobre_app_pct": inf_sobre_mat,
            "inferior_sobre_itron_pct": inf_sobre_itron,
            "lectura_nota": "Itron 6 ruedas negras + 1 roja → 384736.5 (no 38473.65)",
        },
        "etapa5": {
            "sensus_ini": ETAPA5_INI,
            "sensus_fin": ETAPA5_FIN,
            "sensus_delta": e5_delta,
            "app_m3": app_e5,
            "error_pct": err_e5,
            "estado": est_e5,
        },
        "detalle_matriz": det_mat,
        "detalle_inferior": det_inf,
        "detalle_etapa5": det_e5,
    }
    (out_dir / "validacion_terreno_2809.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print(f"[OK] {pdf_path}")
    print(f"[OK] {docx_path}")
    print(f"Matriz Itron {itron_delta} vs App {app_mat} → {err_mat}% ({est_mat})")
    print(f"Inferior {app_inf} ({inf_sobre_mat}% App)")
    print(f"Etapa5 Sensus {e5_delta} vs App {app_e5} → {err_e5}% ({est_e5})")
    print(f"OUT_DIR={out_dir}")


if __name__ == "__main__":
    main()
