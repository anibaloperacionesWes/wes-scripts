"""
Comparativo Likankura vs Juan Pablo II — cuentas de
G:\\Mi unidad\\Colegios\\Peñalolén\\Facturaciones que sí corresponden.

Uso:
  python generar_comparativo_likankura_juanpablo.py --skip-download
"""

from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from facturacion_aguas_andinas_pdf import extraer_texto_pdf
from generar_comparativo_facturaciones_cormup_penalolen import (
    HEADING_RGB,
    MESES_CORTOS,
    OUT_DIR,
    PDF_CACHE,
    Sitio,
    _cargar_sitios,
    _extraer_cuenta,
    _extraer_total_pagar,
    _fmt_fecha,
    _medidas_wes,
    _set_cell,
    _tbl_full_width,
    convertir_a_pdf,
    cruzar_wes,
    descargar_facturaciones,
)
from generar_reporte_word import (
    add_formatted_heading,
    add_logo_to_header,
    estilizar_tabla_wes,
    format_number_chilean,
)

CARPETAS = ("Likankura", "Juan Pablo II")
NODOS = {"Likankura": "000008-13", "Juan Pablo II": "000008-14"}


def _meta_pdf(txt: str) -> dict:
    cuenta = _extraer_cuenta(txt, None)
    if not cuenta:
        for c in re.findall(r"(?<![\d.])(\d{5,7}-\d)(?!\d)", txt):
            if len(c.split("-")[0]) >= 6:
                cuenta = c
                break
    titular = ""
    direccion = ""
    lines = [ln.strip() for ln in txt.splitlines() if ln.strip()]
    for i, ln in enumerate(lines):
        if re.fullmatch(r"\d{5,7}-\d", ln) and i + 1 < len(lines):
            nxt = lines[i + 1]
            if nxt and not nxt.upper().startswith("AV. PRESIDENTE"):
                titular = nxt
                if i + 2 < len(lines) and not lines[i + 2].startswith("RUTA"):
                    direccion = lines[i + 2]
                break
    arranque = ""
    m = re.search(r"Arranque\s+individual-(\d+)", txt, re.I)
    if m:
        arranque = f"{m.group(1)} mm"
    return {"cuenta": cuenta or "", "titular": titular, "direccion": direccion, "arranque": arranque}


def _wes_periodo(meas, a, b) -> Tuple[float, int]:
    por_dia = {m.date.date(): float(m.total_m3) for m in meas if a <= m.date.date() <= b}
    return float(sum(por_dia.values())), len(por_dia)


def _cargar_par(skip_download: bool) -> List[Sitio]:
    if skip_download and any(PDF_CACHE.glob("*/*.pdf")):
        locales = {p.name: p for p in PDF_CACHE.iterdir() if p.is_dir() and p.name in CARPETAS}
        print(f"[INFO] Reuso cache {PDF_CACHE}", flush=True)
    else:
        todos = descargar_facturaciones(PDF_CACHE)
        locales = {k: v for k, v in todos.items() if k in CARPETAS}
    sitios = [s for s in _cargar_sitios(locales) if s.carpeta in CARPETAS]
    cruzar_wes(sitios)
    for s in sitios:
        for f in s.filas:
            txt = extraer_texto_pdf(PDF_CACHE / s.carpeta / f.pdf_name)
            meta = _meta_pdf(txt)
            f.cuenta = f.cuenta or meta["cuenta"]
            setattr(f, "titular", meta["titular"])
            setattr(f, "direccion", meta["direccion"])
            setattr(f, "arranque", meta["arranque"])
            setattr(f, "total_pagar", f.total_pagar or _extraer_total_pagar(txt))
    return sorted(sitios, key=lambda x: x.node_id)


def _cruzado(sitios: List[Sitio]) -> Dict[str, List[float]]:
    """WES del otro nodo en el mismo período de cada boleta."""
    out: Dict[str, List[float]] = {}
    by_id = {s.node_id: s for s in sitios}
    for s in sitios:
        otro = "000008-14" if s.node_id == "000008-13" else "000008-13"
        d0 = min(f.lectura_anterior.date() for f in s.filas)
        d1 = max(f.lectura_actual.date() for f in s.filas)
        meas = _medidas_wes(otro, d0, d1)
        vals = []
        for f in s.filas:
            m3, _n = _wes_periodo(meas, f.lectura_anterior.date(), f.lectura_actual.date())
            vals.append(m3)
        out[s.node_id] = vals
        print(f"[WES cruzado] boletas {s.node_name} vs nodo {by_id[otro].node_name}", flush=True)
    return out


def _write_excel(sitios: List[Sitio], cruz: Dict[str, List[float]], out_xlsx: Path, generado) -> None:
    wb = Workbook()
    azul = PatternFill("solid", fgColor="003366")
    blanco = Font(color="FFFFFF", bold=True, size=9)
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    ama = PatternFill("solid", fgColor="FFF2CC")
    verd = PatternFill("solid", fgColor="C6EFCE")
    naran = PatternFill("solid", fgColor="F8CBAD")

    wc = wb.active
    wc.title = "Cuentas"
    wc.append(["Dato", "Likankura (carpeta)", "Juan Pablo II (carpeta)"])
    lik = next(s for s in sitios if "Likan" in s.carpeta)
    jp = next(s for s in sitios if "Pablo" in s.carpeta)
    f_l = next((f for f in lik.filas if not f.estimado), lik.filas[0])
    f_j = next((f for f in jp.filas if not f.estimado), jp.filas[0])
    datos = [
        ("Nodo WES", lik.node_id, jp.node_id),
        ("N° cuenta Aguas Andinas", f_l.cuenta, f_j.cuenta),
        ("Titular en la boleta", getattr(f_l, "titular", ""), getattr(f_j, "titular", "")),
        ("Dirección de servicio", getattr(f_l, "direccion", ""), getattr(f_j, "direccion", "")),
        ("Medidor (turbina)", f_l.medidor, f_j.medidor),
        ("Arranque", getattr(f_l, "arranque", ""), getattr(f_j, "arranque", "")),
        ("Boletas en carpeta", str(len(lik.filas)), str(len(jp.filas))),
        ("Boletas a promedio", str(sum(1 for f in lik.filas if f.estimado)), str(sum(1 for f in jp.filas if f.estimado))),
    ]
    for row in datos:
        wc.append(list(row))
    for col in range(1, 4):
        c = wc.cell(1, col)
        c.fill = azul
        c.font = blanco
    for row in wc.iter_rows(min_row=1, max_row=wc.max_row, max_col=3):
        for cell in row:
            cell.border = thin
    wc.column_dimensions["A"].width = 28
    wc.column_dimensions["B"].width = 42
    wc.column_dimensions["C"].width = 42

    wm = wb.create_sheet("Mes_a_mes")
    headers = [
        "Mes emisión",
        "Lik cuenta m³",
        "Lik WES propio m³",
        "Lik % error",
        "Lik WES de JP2 m³",
        "JP2 cuenta m³",
        "JP2 WES propio m³",
        "JP2 % error",
        "JP2 WES de Likankura m³",
        "Promedio boleta",
    ]
    wm.append(headers)
    for col in range(1, len(headers) + 1):
        c = wm.cell(1, col)
        c.fill = azul
        c.font = blanco
        c.alignment = Alignment(wrap_text=True, horizontal="center")

    def _por_mes(sitio: Sitio):
        d = {}
        for f in sitio.filas:
            k = (f.emision.year, f.emision.month)
            d[k] = f
        return d

    ml, mj = _por_mes(lik), _por_mes(jp)
    cruz_l = {id(f): v for f, v in zip(lik.filas, cruz[lik.node_id])}
    cruz_j = {id(f): v for f, v in zip(jp.filas, cruz[jp.node_id])}
    meses = sorted(set(ml) | set(mj))
    tot = {k: 0.0 for k in ("lc", "lw", "lx", "jc", "jw", "jx")}
    for y, m in meses:
        fl, fj = ml.get((y, m)), mj.get((y, m))
        lc = fl.m3_cuenta if fl else None
        lw = fl.m3_wes if fl else None
        lx = cruz_l.get(id(fl)) if fl else None
        jc = fj.m3_cuenta if fj else None
        jw = fj.m3_wes if fj else None
        jx = cruz_j.get(id(fj)) if fj else None
        pl = (100.0 * (lw - lc) / lc) if fl and lc else None
        pj = (100.0 * (jw - jc) / jc) if fj and jc else None
        prom = ""
        if fl and fl.estimado:
            prom += "Lik "
        if fj and fj.estimado:
            prom += "JP2"
        wm.append(
            [
                f"{MESES_CORTOS[m]}-{y}",
                lc,
                None if lw is None else round(lw, 1),
                None if pl is None else round(pl, 1),
                None if lx is None else round(lx, 1),
                jc,
                None if jw is None else round(jw, 1),
                None if pj is None else round(pj, 1),
                None if jx is None else round(jx, 1),
                prom.strip() or "—",
            ]
        )
        row = wm.max_row
        if fl and fl.estimado:
            wm.cell(row, 2).fill = ama
        if fj and fj.estimado:
            wm.cell(row, 6).fill = ama
        for col, pct in ((4, pl), (8, pj)):
            if pct is None:
                continue
            wm.cell(row, col).fill = verd if abs(pct) < 15 else naran
        if fl:
            tot["lc"] += fl.m3_cuenta
            tot["lw"] += fl.m3_wes
            tot["lx"] += cruz_l[id(fl)]
        if fj:
            tot["jc"] += fj.m3_cuenta
            tot["jw"] += fj.m3_wes
            tot["jx"] += cruz_j[id(fj)]

    def _pct(c, w):
        return round(100.0 * (w - c) / c, 1) if c else None

    wm.append(
        [
            "TOTAL",
            tot["lc"],
            round(tot["lw"], 1),
            _pct(tot["lc"], tot["lw"]),
            round(tot["lx"], 1),
            tot["jc"],
            round(tot["jw"], 1),
            _pct(tot["jc"], tot["jw"]),
            round(tot["jx"], 1),
            "",
        ]
    )
    for row in wm.iter_rows(min_row=2, max_row=wm.max_row):
        for i in (1, 2, 4, 5, 6, 8):
            row[i].number_format = "#,##0.0" if i in (2, 4, 6, 8) else "#,##0"
        row[3].number_format = "0.0"
        row[7].number_format = "0.0"
        for cell in row:
            cell.border = thin
            cell.alignment = Alignment(horizontal="center")
    for cell in wm[wm.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    wm.freeze_panes = "B2"
    wm.row_dimensions[1].height = 32
    for i, w in enumerate([14, 14, 16, 12, 16, 14, 16, 12, 18, 14], start=1):
        wm.column_dimensions[get_column_letter(i)].width = w

    wn = wb.create_sheet("Notas")
    wn["A1"] = (
        "Comparativo de las carpetas Likankura y Juan Pablo II en "
        "Colegios/Peñalolén/Facturaciones. Estas son las cuentas que corresponden "
        "a esos dos establecimientos (boletas reales de cada carpeta)."
    )
    wn["A2"] = (
        "% error = (WES − cuenta) / cuenta × 100. "
        "Columna «WES de JP2 / de Likankura» = mismo período de la boleta cruzado "
        "con el otro nodo, para ver si las carpetas estuvieran invertidas."
    )
    wn["A3"] = (
        "Amarillo = boleta a promedio/estimado (no es lectura real de turbina). "
        f"Generado {generado.strftime('%d-%m-%Y %H:%M')}."
    )
    wn.column_dimensions["A"].width = 120
    wb.save(out_xlsx)


def _write_word(sitios: List[Sitio], cruz: Dict[str, List[float]], out_docx: Path, generado) -> None:
    lik = next(s for s in sitios if "Likan" in s.carpeta)
    jp = next(s for s in sitios if "Pablo" in s.carpeta)
    f_l = next((f for f in lik.filas if not f.estimado), lik.filas[0])
    f_j = next((f for f in jp.filas if not f.estimado), jp.filas[0])

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.2)
    add_logo_to_header(doc)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    title = doc.add_heading("Likankura vs Juan Pablo II — cuentas que corresponden", level=1)
    if title.runs:
        title.runs[0].font.color.rgb = HEADING_RGB
    p = doc.add_paragraph()
    p.add_run("CORMUP / Peñalolén / Facturaciones. ").bold = True
    p.add_run(
        "Comparativo de las dos carpetas cuyos números de cuenta sí corresponden "
        f"a esos colegios. Generado {generado.strftime('%d-%m-%Y %H:%M')}."
    )

    add_formatted_heading(doc, "1. Identificación de cada cuenta", level=1)
    headers = ["Dato", "Likankura", "Juan Pablo II"]
    rows = [
        ["Nodo WES", lik.node_id, jp.node_id],
        ["N° cuenta", f_l.cuenta or "—", f_j.cuenta or "—"],
        ["Titular", getattr(f_l, "titular", "—") or "—", getattr(f_j, "titular", "—") or "—"],
        ["Dirección", getattr(f_l, "direccion", "—") or "—", getattr(f_j, "direccion", "—") or "—"],
        ["Medidor", f_l.medidor or "—", f_j.medidor or "—"],
        ["Arranque", getattr(f_l, "arranque", "—") or "—", getattr(f_j, "arranque", "—") or "—"],
        ["Boletas", str(len(lik.filas)), str(len(jp.filas))],
    ]
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _tbl_full_width(table)
    for j, h in enumerate(headers):
        _set_cell(table.rows[0].cells[j], h, bold=True, size=10)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            _set_cell(table.rows[i].cells[j], val, bold=(j == 0), size=9)
    estilizar_tabla_wes(table)

    add_formatted_heading(doc, "2. Mes a mes: cuenta vs WES propio y vs el otro nodo", level=1)
    headers2 = [
        "Mes",
        "Lik cta",
        "Lik WES",
        "Lik %",
        "Lik vs JP2 WES",
        "JP2 cta",
        "JP2 WES",
        "JP2 %",
        "JP2 vs Lik WES",
    ]
    ml = {(f.emision.year, f.emision.month): f for f in lik.filas}
    mj = {(f.emision.year, f.emision.month): f for f in jp.filas}
    cruz_l = {id(f): v for f, v in zip(lik.filas, cruz[lik.node_id])}
    cruz_j = {id(f): v for f, v in zip(jp.filas, cruz[jp.node_id])}
    rows2 = []
    hi = []
    tot = {k: 0.0 for k in ("lc", "lw", "lx", "jc", "jw", "jx")}
    for i, k in enumerate(sorted(set(ml) | set(mj)), start=1):
        fl, fj = ml.get(k), mj.get(k)
        y, m = k

        def _pct(f):
            if not f or not f.m3_cuenta:
                return None
            return 100.0 * (f.m3_wes - f.m3_cuenta) / f.m3_cuenta

        pl, pj = _pct(fl), _pct(fj)
        rows2.append(
            [
                f"{MESES_CORTOS[m]}-{y}",
                format_number_chilean(fl.m3_cuenta, 0) if fl else "—",
                format_number_chilean(fl.m3_wes, 0) if fl else "—",
                format_number_chilean(pl, 1) if pl is not None else "—",
                format_number_chilean(cruz_l[id(fl)], 0) if fl else "—",
                format_number_chilean(fj.m3_cuenta, 0) if fj else "—",
                format_number_chilean(fj.m3_wes, 0) if fj else "—",
                format_number_chilean(pj, 1) if pj is not None else "—",
                format_number_chilean(cruz_j[id(fj)], 0) if fj else "—",
            ]
        )
        if (pl is not None and abs(pl) >= 15) or (pj is not None and abs(pj) >= 15):
            hi.append(i)
        if fl:
            tot["lc"] += fl.m3_cuenta
            tot["lw"] += fl.m3_wes
            tot["lx"] += cruz_l[id(fl)]
        if fj:
            tot["jc"] += fj.m3_cuenta
            tot["jw"] += fj.m3_wes
            tot["jx"] += cruz_j[id(fj)]
    rows2.append(
        [
            "TOTAL",
            format_number_chilean(tot["lc"], 0),
            format_number_chilean(tot["lw"], 0),
            format_number_chilean(100.0 * (tot["lw"] - tot["lc"]) / tot["lc"], 1) if tot["lc"] else "—",
            format_number_chilean(tot["lx"], 0),
            format_number_chilean(tot["jc"], 0),
            format_number_chilean(tot["jw"], 0),
            format_number_chilean(100.0 * (tot["jw"] - tot["jc"]) / tot["jc"], 1) if tot["jc"] else "—",
            format_number_chilean(tot["jx"], 0),
        ]
    )
    table2 = doc.add_table(rows=1 + len(rows2), cols=len(headers2))
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    _tbl_full_width(table2)
    for j, h in enumerate(headers2):
        _set_cell(table2.rows[0].cells[j], h, bold=True, size=8)
    for i, row in enumerate(rows2, start=1):
        is_tot = i == len(rows2)
        for j, val in enumerate(row):
            _set_cell(table2.rows[i].cells[j], val, bold=is_tot, size=8)
    estilizar_tabla_wes(table2, highlight_rows=hi, has_total_row=True)

    doc.add_paragraph(
        "Lik WES / JP2 WES = nodo propio. «Lik vs JP2 WES» y «JP2 vs Lik WES» = "
        "el mismo período de la boleta medido en el otro nodo, por si las carpetas "
        "estuvieran cruzadas. % = (WES − cuenta) / cuenta. "
        "Juan Pablo II septiembre es promedio (medidor cerrado): no es lectura real."
    )
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)


def generar(skip_download: bool = False) -> Tuple[Path, Optional[Path], Path]:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    generado = datetime.now(timezone.utc).astimezone()
    ts = generado.strftime("%Y%m%d_%H%M")
    sitios = _cargar_par(skip_download)
    if len(sitios) != 2:
        raise SystemExit(f"Se esperaban 2 sitios, hay {len(sitios)}: {[s.carpeta for s in sitios]}")
    cruz = _cruzado(sitios)
    stem = f"Comparativo_Likankura_vs_JuanPabloII_{ts}"
    out_xlsx = OUT_DIR / f"{stem}.xlsx"
    out_docx = OUT_DIR / f"{stem}.docx"
    _write_excel(sitios, cruz, out_xlsx, generado)
    _write_word(sitios, cruz, out_docx, generado)
    out_pdf = convertir_a_pdf(out_docx)
    print(f"[OK] Excel: {out_xlsx}", flush=True)
    print(f"[OK] Word:  {out_docx}", flush=True)
    if out_pdf:
        print(f"[OK] PDF:   {out_pdf}", flush=True)
    return out_docx, out_pdf, out_xlsx


def main() -> int:
    ap = argparse.ArgumentParser(description="Comparativo Likankura vs Juan Pablo II")
    ap.add_argument("--skip-download", action="store_true")
    args = ap.parse_args()
    generar(skip_download=args.skip_download)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
