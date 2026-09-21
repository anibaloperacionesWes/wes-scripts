# -*- coding: utf-8 -*-
"""
Listado de usuarios con cuenta en la aplicación WES (dashboard.wes.cl).

La API acl-entities no expone un GET masivo de usuarios ni último acceso/login.
Solo permite:
  - GET /users?email=correo_exacto
  - GET /users/{userId}

Este script consulta el universo de correos conocidos en el repositorio
(más correos extra) y reporta quiénes tienen cuenta, a qué compañía pertenecen
y cuántos nodos tienen habilitados.

Uso:
  python listar_usuarios_wes.py
  python listar_usuarios_wes.py --emails extra.txt
"""

from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Set, Tuple

import requests
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

ENTITY_BASE = "http://104.248.53.141:7001/wes/api/acl-entities/v1"
ROOT = Path(__file__).resolve().parent
EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9._\-]+\.[a-zA-Z]{2,}")

PLACEHOLDER_NEEDLES = (
    "example.com",
    "empresa.com",
    "ejemplo.com",
    "tu.correo",
    "tu@email",
    "destinatario@",
    "usuario@ejemplo",
    "nuevo_email@",
    "nuevo@",
    "cliente@empresa",
    "maria.garcia@",
    "carlos.rodriguez@",
)

# Correos truncados o con typo en archivos locales → forma canónica a consultar.
EMAIL_ALIASES = {
    "katiuska.varas@externos.para": "katiuska.varas@externos.parauco.com",
    "javiera.herrera@externos,parauco.com": "javiera.herrera@externos.parauco.com",
    "maurcioorellana@wes.cl": "mauricioorellana@wes.cl",
    "mauricioorellna@wes.cl": "mauricioorellana@wes.cl",
    "silvanaaraya.roja@gmail.com": "silvanaaraya.rojas@gmail.com",
}

EXTRA_EMAILS = {
    "katiuska.varas@externos.parauco.com",
    "javiera.herrera@externos.parauco.com",
    "moloya@externos.parauco.com",
    "medioambiente.dcl@parauco.com",
    "ignacio.gonzalez@parauco.com",
    "mauricioorellana@wes.cl",
    "operaciones@wes.cl",
    "jose.otarola@wildstream.cl",
    "joseluisricardo@hotmail.com",
    "felipecuevas.mancilla@gmail.com",
    "fcuevas@eragroup.com",
    "silvanaaraya.rojas@gmail.com",
    "go.salass@gmail.com",
}

SKIP_DIRS = {".git", "node_modules", "__pycache__", ".venv", "venv"}
HARVEST_EXTS = {".py", ".txt", ".md", ".csv", ".ps1"}

WES_BLUE = RGBColor(0, 51, 102)


def _empresas() -> Dict[str, str]:
    try:
        r = requests.get(f"{ENTITY_BASE}/configuration/companies", timeout=60)
        r.raise_for_status()
        data = r.json()
    except requests.RequestException:
        return {}
    out: Dict[str, str] = {}
    if not isinstance(data, list):
        return out
    for c in data:
        if not isinstance(c, dict):
            continue
        cid = str(c.get("companyId") or "").strip()
        name = str(c.get("name") or "").strip()
        if cid:
            out[cid] = name or cid
    return out


def _nombre_empresa(cid: str, catalogo: Dict[str, str]) -> str:
    cid = (cid or "").strip()
    if not cid:
        return ""
    return catalogo.get(cid, cid)


def _shading(cell, fill_hex: str) -> None:
    try:
        shading = parse_xml(
            f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'w:val="clear" w:fill="{fill_hex}"/>'
        )
        tc_pr = cell._element.get_or_add_tcPr()
        existing = tc_pr.find(qn("w:shd"))
        if existing is not None:
            tc_pr.remove(existing)
        tc_pr.append(shading)
    except Exception:
        pass


def _heading(doc: Document, text: str) -> None:
    h = doc.add_heading(text.upper(), level=1)
    for run in h.runs:
        run.bold = True
        run.font.color.rgb = WES_BLUE
    h.paragraph_format.space_after = Pt(6)


def _add_table(doc: Document, title: str, rows: List[List[str]]) -> None:
    cap = doc.add_paragraph()
    run = cap.add_run(title.upper())
    run.bold = True
    run.font.color.rgb = WES_BLUE
    cap.paragraph_format.space_after = Pt(4)
    if not rows:
        doc.add_paragraph("Sin datos disponibles.")
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = str(value)
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(1)
            para.paragraph_format.space_after = Pt(1)
            if para.runs:
                para.runs[0].font.size = Pt(8 if i else 9)
                if i == 0:
                    para.runs[0].bold = True
                    para.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                    _shading(cell, "003366")
                elif i % 2 == 0:
                    _shading(cell, "F2F6FC")
    doc.add_paragraph("")


def _add_logo(doc: Document) -> None:
    for name in ("logo wes.png", "logo_wes.png", "logo.png", "logo wes.bmp"):
        p = ROOT / name
        if p.is_file():
            try:
                section = doc.sections[0]
                section.header_distance = Inches(0.4)
                header = section.header
                para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
                para.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
                run = para.add_run()
                run.add_picture(str(p), width=Inches(1.2))
            except Exception:
                pass
            return


def _es_email_placeholder(email: str) -> bool:
    e = email.strip().lower()
    return any(n in e for n in PLACEHOLDER_NEEDLES)


def _normalizar_email(raw: str) -> Optional[str]:
    e = raw.strip().lower().rstrip(".,;")
    if not e or "@" not in e:
        return None
    e = EMAIL_ALIASES.get(e, e)
    if _es_email_placeholder(e):
        return None
    if not EMAIL_RE.fullmatch(e):
        return None
    return e


def _emails_desde_archivo(path: Path) -> Set[str]:
    if not path.is_file():
        return set()
    found: Set[str] = set()
    try:
        text = path.read_text(encoding="utf-8", errors="ignore")
    except OSError:
        return set()
    for m in EMAIL_RE.finditer(text):
        e = _normalizar_email(m.group(0))
        if e:
            found.add(e)
    # Comas en vez de punto (typo frecuente en listados PA).
    for line in text.splitlines():
        if "@externos,parauco.com" in line.lower():
            e = _normalizar_email(line.strip())
            if e:
                found.add(e)
    return found


def _recolectar_emails(extra: Iterable[Path]) -> List[str]:
    pool: Set[str] = set(EXTRA_EMAILS)
    try:
        from lista_contactos_reportes import CONTACTOS_REPORTES, CORREOS_AUTORIZADOS

        for c in CONTACTOS_REPORTES.values():
            e = _normalizar_email(c.get("email") or "")
            if e:
                pool.add(e)
        for row in CORREOS_AUTORIZADOS:
            e = _normalizar_email(row.get("email") or "")
            if e:
                pool.add(e)
    except ImportError:
        pass
    try:
        from config_correos_equipo import CORREOS_EQUIPO_WES

        for info in CORREOS_EQUIPO_WES.values():
            e = _normalizar_email(info.get("email") or "")
            if e:
                pool.add(e)
    except ImportError:
        pass

    for p in ROOT.rglob("*"):
        if any(part in SKIP_DIRS for part in p.parts):
            continue
        if not p.is_file() or p.suffix.lower() not in HARVEST_EXTS:
            continue
        if p.stat().st_size > 2_000_000:
            continue
        pool.update(_emails_desde_archivo(p))

    for fp in extra:
        pool.update(_emails_desde_archivo(fp))
    return sorted(pool)


def _fetch_user(email: str) -> Tuple[str, Optional[dict], str]:
    try:
        r = requests.get(
            f"{ENTITY_BASE}/users",
            params={"email": email},
            timeout=30,
        )
    except requests.RequestException as exc:
        return email, None, str(exc)
    if r.status_code == 404:
        return email, None, "no existe en API"
    if r.status_code != 200:
        return email, None, f"HTTP {r.status_code}"
    data = r.json()
    if not isinstance(data, dict):
        return email, None, "respuesta inválida"
    data.pop("password", None)
    return email, data, ""


def _nombre_usuario(user: dict) -> str:
    return f"{user.get('name', '')} {user.get('lastName', '')}".strip()


def _nodos_count(user: dict) -> int:
    allowed = user.get("allowedNodes") or []
    return len(allowed) if isinstance(allowed, list) else 0


def _escribir_csv_usuarios(path: Path, filas: List[dict]) -> None:
    cols = [
        "user_id",
        "email",
        "nombre",
        "company_id",
        "empresa",
        "nodos_count",
        "switch_enabled",
    ]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols, delimiter=";")
        w.writeheader()
        for row in filas:
            w.writerow({k: row.get(k, "") for k in cols})


def _escribir_csv_empresas(path: Path, filas_empresa: List[dict]) -> None:
    cols = ["company_id", "empresa", "usuarios", "emails"]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols, delimiter=";")
        w.writeheader()
        for row in filas_empresa:
            w.writerow(row)


def _escribir_xlsx(
    path: Path,
    filas: List[dict],
    filas_empresa: List[dict],
    sin_usuario: List[str],
) -> Optional[Path]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    header_fill = PatternFill("solid", fgColor="003366")
    header_font = Font(bold=True, color="FFFFFF")
    alt_fill = PatternFill("solid", fgColor="F2F6FC")

    def _sheet(ws, headers, rows, widths):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.fill = header_fill
            cell.font = header_font
        for r, row in enumerate(rows, 2):
            for c, val in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=val)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                if r % 2 == 0:
                    cell.fill = alt_fill
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Por_usuario"
    _sheet(
        ws1,
        [
            "user_id",
            "email",
            "nombre",
            "company_id",
            "empresa",
            "nodos_count",
            "switch_enabled",
        ],
        [
            [
                r["user_id"],
                r["email"],
                r["nombre"],
                r["company_id"],
                r["empresa"],
                r["nodos_count"],
                r["switch_enabled"],
            ]
            for r in filas
        ],
        (26, 36, 28, 12, 28, 14, 16),
    )

    ws2 = wb.create_sheet("Por_empresa")
    _sheet(
        ws2,
        ["company_id", "empresa", "usuarios", "emails"],
        [
            [r["company_id"], r["empresa"], r["usuarios"], r["emails"]]
            for r in filas_empresa
        ],
        (12, 28, 12, 80),
    )

    ws3 = wb.create_sheet("Sin_cuenta")
    _sheet(
        ws3,
        ["email_consultado"],
        [[e] for e in sin_usuario],
        (40,),
    )
    wb.save(path)
    return path


def _escribir_docx(
    path: Path,
    generado: datetime,
    emails_consultados: int,
    filas: List[dict],
    filas_empresa: List[dict],
    sin_usuario: List[str],
) -> Path:
    doc = Document()
    _add_logo(doc)
    title = doc.add_paragraph("Usuarios con acceso a la aplicación WES")
    title.style = "Title"
    for run in title.runs:
        run.font.size = Pt(22)
    sub = doc.add_paragraph(
        "MONITOREO WES\n"
        "Cuentas habilitadas en dashboard.wes.cl\n"
        f"Generado: {generado.strftime('%d-%m-%Y %H:%M')}"
    )
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    _heading(doc, "Alcance y limitación")
    doc.add_paragraph(
        "La API pública de WES no entrega un listado masivo de usuarios ni registra "
        "último inicio de sesión. Cada cuenta se confirma con GET /users?email=… "
        "sobre correos conocidos en este repositorio (equipo WES, contactos de "
        "reportes y listados de clientes). El resultado indica quién tiene cuenta "
        "y a qué compañía/nodos está habilitado; no quién entró a la app en una fecha."
    )
    doc.add_paragraph(
        "Para un universo exhaustivo hay que pedir un export al backoffice WES "
        "(base interna de usuarios)."
    )

    _heading(doc, "Resumen")
    _add_table(
        doc,
        "Cifras",
        [
            ["Indicador", "Valor"],
            ["Correos consultados", str(emails_consultados)],
            ["Cuentas encontradas", str(len(filas))],
            ["Correos sin cuenta en la API", str(len(sin_usuario))],
            ["Empresas con al menos un usuario", str(len(filas_empresa))],
        ],
    )

    _heading(doc, "Usuarios con cuenta")
    rows_u = [["Email", "Nombre", "Empresa", "Nodos", "Switch nodos"]]
    for r in filas:
        rows_u.append(
            [
                r["email"],
                r["nombre"] or "—",
                f"{r['company_id']} {r['empresa']}".strip(),
                str(r["nodos_count"]),
                "Sí" if r["switch_enabled"] == "True" else "No",
            ]
        )
    _add_table(
        doc,
        "Cuentas confirmadas en la API (switch = ON/OFF de nodos, no último acceso)",
        rows_u,
    )

    _heading(doc, "Usuarios por empresa")
    rows_e = [["ID", "Empresa", "Usuarios", "Correos"]]
    for r in filas_empresa:
        rows_e.append(
            [r["company_id"], r["empresa"], str(r["usuarios"]), r["emails"]]
        )
    _add_table(doc, "Agrupación por companyId", rows_e)

    if sin_usuario:
        _heading(doc, "Correos consultados sin cuenta")
        doc.add_paragraph(
            "Estos correos aparecen en el proyecto (destinatarios de informes, "
            "listados de contacto, etc.) pero la API no devolvió usuario WES."
        )
        rows_s = [["Email"]] + [[e] for e in sin_usuario]
        _add_table(doc, "Sin usuario en API", rows_s)

    doc.save(path)
    return path


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Listado de usuarios WES con cuenta (consulta por email)"
    )
    parser.add_argument(
        "--emails",
        type=Path,
        action="append",
        default=[],
        help="Archivo adicional con correos (uno por línea)",
    )
    parser.add_argument(
        "--salida",
        type=Path,
        default=ROOT / "reports" / "WES" / "Usuarios",
        help="Carpeta de salida (por defecto reports/WES/Usuarios)",
    )
    args = parser.parse_args()

    generado = datetime.now()
    out_dir = args.salida
    out_dir.mkdir(parents=True, exist_ok=True)

    catalogo = _empresas()
    emails = _recolectar_emails(args.emails)
    print(f"[INFO] Empresas en API: {len(catalogo)}")
    print(f"[INFO] Consultando {len(emails)} correo(s) conocidos ...")

    filas: List[dict] = []
    sin_usuario: List[str] = []
    por_empresa: Dict[str, Set[str]] = defaultdict(set)
    nombres_empresa: Dict[str, str] = {}

    with ThreadPoolExecutor(max_workers=8) as ex:
        futs = {ex.submit(_fetch_user, e): e for e in emails}
        for fut in as_completed(futs):
            email, user, obs = fut.result()
            if not user:
                sin_usuario.append(email if not obs or obs == "no existe en API" else f"{email} ({obs})")
                print(f"[--] {email} — {obs or 'sin cuenta'}")
                continue
            cid = str(user.get("companyId") or "").strip()
            empresa = _nombre_empresa(cid, catalogo)
            email_api = str(user.get("username") or email).strip()
            row = {
                "user_id": str(user.get("userId") or "").strip(),
                "email": email_api,
                "nombre": _nombre_usuario(user),
                "company_id": cid,
                "empresa": empresa,
                "nodos_count": _nodos_count(user),
                "switch_enabled": str(bool(user.get("switchEnabled"))),
            }
            filas.append(row)
            por_empresa[cid].add(email_api)
            nombres_empresa[cid] = empresa
            print(
                f"[OK] {email_api} — {empresa or cid} — {row['nodos_count']} nodo(s)"
            )

    filas.sort(key=lambda r: (r["empresa"].casefold(), r["email"].casefold()))
    sin_usuario = sorted(sin_usuario, key=str.casefold)

    filas_empresa: List[dict] = []
    for cid in sorted(por_empresa.keys()):
        mails = sorted(por_empresa[cid], key=str.casefold)
        filas_empresa.append(
            {
                "company_id": cid,
                "empresa": nombres_empresa.get(cid, cid),
                "usuarios": len(mails),
                "emails": ", ".join(mails),
            }
        )
    filas_empresa.sort(key=lambda r: (-r["usuarios"], r["empresa"].casefold()))

    csv_u = out_dir / "usuarios_wes_por_usuario.csv"
    csv_e = out_dir / "usuarios_wes_por_empresa.csv"
    xlsx = out_dir / "usuarios_wes.xlsx"
    docx = out_dir / "usuarios_wes.docx"
    resumen = out_dir / "RESUMEN.txt"

    _escribir_csv_usuarios(csv_u, filas)
    _escribir_csv_empresas(csv_e, filas_empresa)
    xlsx_ok = _escribir_xlsx(xlsx, filas, filas_empresa, sin_usuario)
    _escribir_docx(docx, generado, len(emails), filas, filas_empresa, sin_usuario)

    resumen.write_text(
        "\n".join(
            [
                "Aplicación: dashboard.wes.cl (API acl-entities)",
                f"Generado: {generado.isoformat(timespec='minutes')}",
                f"Correos consultados: {len(emails)}",
                f"Cuentas encontradas: {len(filas)}",
                f"Sin cuenta en API: {len(sin_usuario)}",
                f"Empresas con usuarios: {len(filas_empresa)}",
                "",
                "LIMITACIÓN: la API no lista todos los usuarios ni el último acceso.",
                "Este reporte cubre el universo de correos conocidos en el proyecto.",
                "Para listado exhaustivo solicitar export al backoffice WES.",
            ]
        ),
        encoding="utf-8",
    )

    print(f"\n[OK] Cuentas encontradas: {len(filas)}")
    print(f"[OK] CSV usuarios: {csv_u.resolve()}")
    print(f"[OK] CSV empresas: {csv_e.resolve()}")
    if xlsx_ok:
        print(f"[OK] XLSX:         {xlsx.resolve()}")
    print(f"[OK] DOCX:         {docx.resolve()}")
    print(f"[INFO] Sin cuenta: {len(sin_usuario)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
