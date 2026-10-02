"""
Horario por colegio CORMUP desde Drive:

  G:\\Mi unidad\\Colegios\\Peñalolén\\Facturaciones\\data de placas.xlsx

Una hoja H.* por establecimiento (F/H/M de la placa). Fila 28 = data placa
(suma de horas). Fila 29 = data app (totales diarios de la app). Las últimas
3 hojas de esa planilla (Erasmo Escala, Matilde Huici, CE Valle Hermoso)
están vacías: quedan como pendiente de descargar.

Uso:
  python generar_horario_placas_cormup.py
"""

from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from generar_horario_jp2_hora_consumo import (
    AZUL_HDR,
    CELESTE,
    CENTER,
    FONT_OK,
    HDR,
    NARANJA,
    NUM_FMT,
    ROJO,
    SUB,
    THIN,
    VERDE,
    construir_horario,
)
from wes_google_drive import credenciales_configuradas, obtener_servicio_drive, subir_a_drive

OUT_DIR = Path("reports/CORMUP/Facturaciones_vs_WES")
DRIVE_SUB = "CORMUP/Facturaciones_vs_WES"
DRIVE_NOMBRE = "Horario_placas_CORMUP.xlsx"
LOCAL_XLSX = OUT_DIR / "_pdfs" / "data_de_placas.xlsx"
PLACAS_NOMBRE = "data de placas.xlsx"
APP_DIR = OUT_DIR / "app_data"

# Nombre en data de placas.xlsx → (nombre, nodo, hoja destino)
COLEGIOS = [
    ("171", "Liceo Antonio Hermida", "000008-01", "H. 171"),
    ("eduardo de la barra", "Eduardo de la Barra", "000008-02", "H. E. Barra"),
    ("Carlos Fernandez Peña", "Carlos Fernández Peña", "000008-03", "H. C.F. Peña"),
    ("tobalaba", "Tobalaba", "000008-04", "H. Tobalaba"),
    ("santa mari", "Santa María", "000008-05", "H. Santa María"),
    ("Luis Arrieta Cañas", "Luis Arrieta Cañas", "000008-06", "H. Arrieta"),
    ("erasmo escala", "Erasmo Escala", "000008-07", "H. Erasmo"),
    ("alicura", "Alicura", "000008-08", "H. Alicura"),
    ("juan bta pastenes", "Juan Bautista Pastenes", "000008-09", "H. J.B. Pastenes"),
    ("matilde ", "Matilde Huici Navas", "000008-10", "H. Matilde"),
    ("ce valle", "CE Valle Hermoso", "000008-11", "H. Valle Hermoso"),
    ("Union Nacional ", "Unión Nacional Árabe", "000008-12", "H. Unión Árabe"),
    ("Likankura", "Likancura", "000008-13", "H. Likancura"),
    ("JP II", "Juan Pablo II", "000008-14", "H. JPII"),
]
PENDIENTES = {"000008-07", "000008-10", "000008-11"}

PAT_F = re.compile(r"^F:\s*(\d{2}/\d{2}/\d{4})")
PAT_H = re.compile(r"^H:\s*(\d{1,2}):00")
PAT_M = re.compile(r"^M:\s*([\d.,]+)")
PAT_R = re.compile(r"^R:")


def _cargar_data_app(node: str) -> dict[date, float]:
    """Totales diarios de la app (suma de horas del dump), por nodo."""
    path = APP_DIR / f"{node}_diario.json"
    if not path.is_file():
        return {}
    raw = json.loads(path.read_text(encoding="utf-8"))
    return {date.fromisoformat(str(k)): float(v) for k, v in raw.items()}


def _meta_data_app(node: str) -> dict[str, Any]:
    path = APP_DIR / f"{node}_meta.json"
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def descargar_data_de_placas() -> Path:
    from generar_comparativo_facturaciones_cormup_penalolen import (
        _buscar_carpeta_facturaciones,
        _descargar_archivo,
        _listar_hijos,
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    LOCAL_XLSX.parent.mkdir(parents=True, exist_ok=True)
    if not credenciales_configuradas():
        if LOCAL_XLSX.is_file():
            print(f"[INFO] sin Drive; uso {LOCAL_XLSX}", flush=True)
            return LOCAL_XLSX
        raise FileNotFoundError(LOCAL_XLSX)
    service = obtener_servicio_drive()
    folder_id = _buscar_carpeta_facturaciones(service)
    hijos = _listar_hijos(service, folder_id)
    target = next(
        (
            h
            for h in hijos
            if h["name"].lower().replace(" ", "") == PLACAS_NOMBRE.lower().replace(" ", "")
            or "data de placas" in h["name"].lower()
        ),
        None,
    )
    if not target:
        raise FileNotFoundError("No está data de placas.xlsx en Colegios/Peñalolén/Facturaciones")
    print(f"[Drive] {target['name']} id={target['id']}", flush=True)
    _descargar_archivo(service, target["id"], LOCAL_XLSX)
    print(f"[Drive] bajado {LOCAL_XLSX} ({LOCAL_XLSX.stat().st_size} bytes)", flush=True)
    return LOCAL_XLSX


def _parse_fhm_sheet(
    ws,
) -> tuple[dict[date, dict[int, float]], dict, set[tuple[date, int]]]:
    """Lee F:/H:/M: por columna. Si un (fecha, hora) se repite, se queda la primera."""
    col_tokens: dict[int, list[str]] = {}
    for col in range(1, ws.max_column + 1):
        toks: list[str] = []
        for row in range(1, ws.max_row + 1):
            raw = ws.cell(row, col).value
            if raw is None:
                continue
            s = str(raw).strip()
            if s:
                toks.append(s)
        if toks:
            col_tokens[col] = toks

    horas: dict[date, dict[int, float]] = {}
    dups_pares: set[tuple[date, int]] = set()
    n_triples = 0
    n_dup = 0
    n_conf = 0
    for tokens in col_tokens.values():
        pending_f: date | None = None
        pending_h: int | None = None
        for tok in tokens:
            if PAT_R.match(tok):
                pending_f = pending_h = None
                continue
            mf, mh, mm = PAT_F.match(tok), PAT_H.match(tok), PAT_M.match(tok)
            if mf:
                pending_f = datetime.strptime(mf.group(1), "%d/%m/%Y").date()
                pending_h = None
                continue
            if mh and pending_f is not None:
                pending_h = int(mh.group(1))
                continue
            if mm and pending_f is not None and pending_h is not None:
                val = float(mm.group(1).replace(",", "."))
                n_triples += 1
                prev = (horas.get(pending_f) or {}).get(pending_h)
                if prev is not None:
                    n_dup += 1
                    dups_pares.add((pending_f, pending_h))
                    if abs(prev - val) > 0.001:
                        n_conf += 1
                else:
                    horas.setdefault(pending_f, {})[pending_h] = val
                pending_f = pending_h = None
    dias = sorted(horas)
    stats = {
        "triples": n_triples,
        "unicos": sum(len(h) for h in horas.values()),
        "duplicados": n_dup,
        "dias_dup": len({d for d, _h in dups_pares}),
        "conflictos": n_conf,
        "dias": len(dias),
        "d0": dias[0] if dias else None,
        "d1": dias[-1] if dias else None,
        "m3": round(sum(sum(h.values()) for h in horas.values()), 2),
    }
    return horas, stats, dups_pares


def _hoja_pendiente(wb: Workbook, sheet_name: str, nombre: str, node: str) -> None:
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    ws["A1"] = (
        f"{nombre} ({node}) — data de placas pendiente de descargar. "
        "En la planilla de Drive esta hoja está vacía; cuando esté el pegado F/H/M "
        "se arma el horario igual que en los otros colegios."
    )
    ws["A1"].font = Font(bold=True, size=12, color="003366")
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 48
    ws["A3"] = "Estado"
    ws["B3"] = "Pendiente de descargar"
    ws["A3"].fill = NARANJA
    ws["B3"].fill = NARANJA
    ws["A3"].font = FONT_OK
    ws["B3"].font = FONT_OK
    ws["A3"].border = THIN
    ws["B3"].border = THIN
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 36


def _escribir_resumen(wb: Workbook, filas: list[dict]) -> None:
    if "Resumen" in wb.sheetnames:
        del wb["Resumen"]
    ws = wb.create_sheet("Resumen", 0)
    headers = [
        "Colegio",
        "Nodo",
        "Desde",
        "Hasta",
        "Días",
        "Horas únicas",
        "Duplicados",
        "m³ placa",
        "Días app",
        "m³ app",
        "Estado",
    ]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(1, col)
        c.fill = AZUL_HDR
        c.font = Font(color="FFFFFF", bold=True, size=9)
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        c.border = THIN
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws["A2"] = (
        "Fuente placa: Colegios / Peñalolén / Facturaciones / data de placas.xlsx. "
        "Una hoja horaria por colegio. Fila data placa = suma de horas de la placa "
        "(si un horario se repetía, queda una sola vez). Fila data app = totales "
        "diarios de la app. En cada hoja, encabezado y hora en rojo = ese día/"
        "horario se repetía en el pegado de la placa. Fila data app en rojo = no "
        "cuadra con data placa. Estado: Placa cargada = llegó el pegado F/H/M; "
        "Pendiente de placa = hay data app pero la hoja de la placa sigue vacía. "
        "Erasmo Escala, Matilde Huici y CE Valle Hermoso están pendientes de descargar."
    )
    ws["A2"].font = Font(bold=True, size=11, color="003366")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 52

    for info in filas:
        pendiente = info["pendiente"]
        n_app = int(info.get("dias_app") or 0)
        m3_app = info.get("m3_app")
        if pendiente and n_app:
            estado = "Pendiente de placa"
        elif pendiente:
            estado = "Pendiente de descargar"
        else:
            estado = "Placa cargada"
        ws.append(
            [
                info["nombre"],
                info["node"],
                info["d0"].strftime("%d/%m/%Y") if info["d0"] else "—",
                info["d1"].strftime("%d/%m/%Y") if info["d1"] else "—",
                info["dias"] or None,
                info["unicos"] or None,
                info["duplicados"] if not pendiente else None,
                info["m3"] if not pendiente else None,
                n_app or None,
                m3_app if n_app else None,
                estado,
            ]
        )
        row = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row, col)
            cell.border = THIN
            cell.alignment = CENTER
        ws.cell(row, 1).alignment = Alignment(horizontal="left", vertical="center")
        if pendiente:
            for col in range(1, len(headers) + 1):
                ws.cell(row, col).fill = NARANJA
            if n_app:
                ws.cell(row, 10).number_format = NUM_FMT
                ws.cell(row, 10).fill = VERDE
        else:
            ws.cell(row, 8).number_format = NUM_FMT
            ws.cell(row, 8).fill = CELESTE
            if info["duplicados"]:
                ws.cell(row, 7).fill = PatternFill("solid", fgColor="FFFF99")
            if n_app:
                ws.cell(row, 10).number_format = NUM_FMT
                ws.cell(row, 10).fill = VERDE
            ws.cell(row, 11).fill = VERDE
            ws.cell(row, 11).font = FONT_OK

    anchos = [28, 14, 14, 14, 10, 14, 14, 14, 12, 14, 26]
    for i, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    ws.page_setup.orientation = "landscape"
    ws.sheet_view.showGridLines = False


def main() -> None:
    path = descargar_data_de_placas()
    src = load_workbook(path, data_only=False)
    print("[INFO] hojas fuente", src.sheetnames, flush=True)

    wb = Workbook()
    tmp = wb.active
    tmp.title = "_tmp"
    filas_resumen: list[dict] = []

    for src_name, nombre, node, dest in COLEGIOS:
        if src_name not in src.sheetnames:
            print(f"[WARN] no está la hoja {src_name!r} en data de placas", flush=True)
            _hoja_pendiente(wb, dest, nombre, node)
            filas_resumen.append(
                {
                    "nombre": nombre,
                    "node": node,
                    "d0": None,
                    "d1": None,
                    "dias": 0,
                    "unicos": 0,
                    "duplicados": 0,
                    "m3": 0.0,
                    "dias_app": 0,
                    "m3_app": None,
                    "pendiente": True,
                }
            )
            continue
        horas, st, dups_pares = _parse_fhm_sheet(src[src_name])
        data_app = _cargar_data_app(node)
        meta_app = _meta_data_app(node)
        pendiente = node in PENDIENTES or not horas
        print(
            f"  {dest:18} {nombre:24} {st['triples']:5} triples "
            f"únicos={st['unicos']} dups={st['duplicados']} "
            f"días_dup={st['dias_dup']} conf={st['conflictos']} "
            f"días={st['dias']} m³={st['m3']} app={len(data_app)}"
            f"{' PENDIENTE' if pendiente else ''}",
            flush=True,
        )
        d0_res = st["d0"]
        d1_res = st["d1"]
        if data_app:
            d0_app, d1_app = min(data_app), max(data_app)
            d0_res = min(d0_res, d0_app) if d0_res else d0_app
            d1_res = max(d1_res, d1_app) if d1_res else d1_app
        filas_resumen.append(
            {
                "nombre": nombre,
                "node": node,
                "d0": d0_res,
                "d1": d1_res,
                "dias": st["dias"],
                "unicos": st["unicos"],
                "duplicados": st["duplicados"],
                "m3": st["m3"],
                "dias_app": len(data_app),
                "m3_app": round(sum(data_app.values()), 2) if data_app else None,
                "pendiente": pendiente,
            }
        )
        if pendiente and not data_app:
            _hoja_pendiente(wb, dest, nombre, node)
            continue
        dias = sorted(horas)
        if data_app:
            dias = sorted(set(dias) | set(data_app))
        nota = (
            "Fuente placa: data de placas.xlsx (pegado F/H/M). "
            "Fila data placa = celeste (suma de horas de la placa; "
            "si un horario se repetía, queda una sola vez). "
            "Fila data app = verde (totales diarios de la app). "
            "Encabezado y hora en rojo = ese día/horario se repetía en la placa. "
            "Fila data app en rojo = no coincide con data placa. "
            f"Duplicados placa={st['duplicados']} "
            f"(días={st['dias_dup']}, conflictos de valor={st['conflictos']}). "
        )
        if pendiente:
            nota += "Placa pendiente de descargar; esta hoja muestra solo data app. "
        nota += "Solo días con placa" + (" o con data app." if data_app else ".")
        if data_app:
            extra_app = meta_app.get("nota") or (
                f"Dump app {min(data_app).strftime('%d/%m/%Y')}–"
                f"{max(data_app).strftime('%d/%m/%Y')} ({len(data_app)} días)."
            )
            nota += " " + extra_app
        construir_horario(
            dias,
            horas,
            wb=wb,
            sheet_name=dest,
            nombre=nombre,
            node=node,
            d0=dias[0],
            d1=dias[-1],
            listado=data_app,
            nota=nota,
            resaltar=dups_pares or None,
            fill_resaltar=ROJO,
            etiqueta_total="data placa",
            etiqueta_listado="data app",
            mostrar_fila_listado=True,
        )

    if "_tmp" in wb.sheetnames:
        del wb["_tmp"]
    _escribir_resumen(wb, filas_resumen)

    ts = datetime.now().strftime("%Y%m%d_%H%M")
    out = OUT_DIR / f"Horario_placas_CORMUP_{ts}.xlsx"
    wb.save(out)
    estable = OUT_DIR / DRIVE_NOMBRE
    if out.resolve() != estable.resolve():
        estable.write_bytes(out.read_bytes())
    print("XLSX", out, "sheets", wb.sheetnames)
    if credenciales_configuradas():
        info = subir_a_drive(estable, subcarpeta=DRIVE_SUB, nombre=DRIVE_NOMBRE)
        print("DRIVE", info["id"], info["web_view_link"])


if __name__ == "__main__":
    main()
