"""
Totalizados mensuales por colegio: m³ cuenta (Aguas Andinas) vs m³ WES.

Fila = establecimiento. Columnas = ene-2026 en adelante (emisión de la boleta):
por cada mes, m³ que marcó la cuenta y m³ que marcó WES (medido + proyección de huecos).

Uso:
  python generar_totalizados_mensuales_cormup_penalolen.py --skip-download
"""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from generar_comparativo_facturaciones_cormup_penalolen import (
    COMPANY_ID,
    HEADING_RGB,
    MESES_CORTOS,
    OUT_DIR,
    PDF_CACHE,
    Sitio,
    _cargar_sitios,
    _set_cell,
    _tbl_full_width,
    convertir_a_pdf,
    cruzar_wes,
    descargar_facturaciones,
)
from generar_reporte_word import (
    add_formatted_heading,
    add_logo_to_header,
    estilizar_tabla_wes,
    format_number_chilean,
)

MES_DESDE = date(2026, 1, 1)


def _mes_key(d: date) -> Tuple[int, int]:
    return (d.year, d.month)


def _mes_label(y: int, m: int) -> str:
    return f"{MESES_CORTOS[m]}-{y}"


def _meses_rango(sitios: List[Sitio]) -> List[Tuple[int, int]]:
    keys = set()
    for s in sitios:
        for f in s.filas:
            em = f.emision.date()
            if em >= MES_DESDE:
                keys.add(_mes_key(em))
    if not keys:
        return []
    y0, m0 = min(keys)
    y1, m1 = max(keys)
    out: List[Tuple[int, int]] = []
    y, m = y0, m0
    while (y, m) <= (y1, m1):
        out.append((y, m))
        m += 1
        if m == 13:
            y += 1
            m = 1
    return out


def _matriz(sitios: List[Sitio], meses: List[Tuple[int, int]]):
    """sitio.node_id -> mes -> {cuenta, wes, n, n_est}."""
    data: Dict[str, Dict[Tuple[int, int], dict]] = {}
    for s in sitios:
        acc: Dict[Tuple[int, int], dict] = {
            k: {"cuenta": 0.0, "wes": 0.0, "n": 0, "n_est": 0} for k in meses
        }
        for f in s.filas:
            em = f.emision.date()
            if em < MES_DESDE:
                continue
            k = _mes_key(em)
            if k not in acc:
                continue
            acc[k]["cuenta"] += float(f.m3_cuenta)
            acc[k]["wes"] += float(f.m3_wes)
            acc[k]["n"] += 1
            if f.estimado:
                acc[k]["n_est"] += 1
        data[s.node_id] = acc
    return data


def _write_excel(
    sitios: List[Sitio],
    meses: List[Tuple[int, int]],
    data,
    out_xlsx: Path,
    generado: datetime,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Mensual_cuenta_vs_WES"

    azul = PatternFill("solid", fgColor="003366")
    azul2 = PatternFill("solid", fgColor="1F4E79")
    blanco = Font(color="FFFFFF", bold=True, size=9)
    tot_fill = PatternFill("solid", fgColor="D9E1F2")
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )

    # Fila 1: colegio | ene-2026 (merge 2) | feb-2026 ... | TOTAL
    n_meses = len(meses)
    ws.cell(1, 1, "Colegio")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    ws.cell(1, 2, "Nodo")
    ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
    col = 3
    for y, m in meses:
        ws.cell(1, col, _mes_label(y, m).upper())
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
        ws.cell(2, col, "m³ cuenta")
        ws.cell(2, col + 1, "m³ WES")
        col += 2
    ws.cell(1, col, "TOTAL ene→")
    ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
    ws.cell(2, col, "m³ cuenta")
    ws.cell(2, col + 1, "m³ WES")

    last_col = 2 + n_meses * 2 + 2
    for r in (1, 2):
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            cell.fill = azul if r == 1 else azul2
            cell.font = blanco
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 18

    sitios_ord = sorted(sitios, key=lambda s: s.node_id)
    tot_mes = {k: {"cuenta": 0.0, "wes": 0.0} for k in meses}
    tot_c = tot_w = 0.0
    row_i = 3
    for s in sitios_ord:
        ws.cell(row_i, 1, s.node_name)
        ws.cell(row_i, 2, s.node_id)
        col = 3
        sc = sw = 0.0
        for k in meses:
            cel = data[s.node_id][k]
            cta, wes = cel["cuenta"], cel["wes"]
            c1 = ws.cell(row_i, col, cta if cel["n"] else None)
            c2 = ws.cell(row_i, col + 1, wes if cel["n"] else None)
            c1.number_format = "#,##0"
            c2.number_format = "#,##0.0"
            if cel["n_est"] and cel["n"]:
                c1.fill = PatternFill("solid", fgColor="FFF2CC")
            sc += cta
            sw += wes
            tot_mes[k]["cuenta"] += cta
            tot_mes[k]["wes"] += wes
            col += 2
        ws.cell(row_i, col, sc).number_format = "#,##0"
        ws.cell(row_i, col + 1, sw).number_format = "#,##0.0"
        tot_c += sc
        tot_w += sw
        for c in range(1, last_col + 1):
            ws.cell(row_i, c).border = thin
            ws.cell(row_i, c).alignment = Alignment(horizontal="center")
        ws.cell(row_i, 1).alignment = Alignment(horizontal="left")
        row_i += 1

    ws.cell(row_i, 1, "TOTAL colegios")
    ws.cell(row_i, 2, COMPANY_ID)
    col = 3
    for k in meses:
        ws.cell(row_i, col, tot_mes[k]["cuenta"]).number_format = "#,##0"
        ws.cell(row_i, col + 1, tot_mes[k]["wes"]).number_format = "#,##0.0"
        col += 2
    ws.cell(row_i, col, tot_c).number_format = "#,##0"
    ws.cell(row_i, col + 1, tot_w).number_format = "#,##0.0"
    for c in range(1, last_col + 1):
        cell = ws.cell(row_i, c)
        cell.fill = tot_fill
        cell.font = Font(bold=True)
        cell.border = thin

    ws.freeze_panes = "C3"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 12
    for c in range(3, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 11

    wn = wb.create_sheet("Notas")
    wn["A1"] = "Cómo se arma el mes"
    wn["A2"] = (
        "Cada boleta se asigna al mes de FECHA DE EMISIÓN (ene-2026 en adelante). "
        "Si un colegio tiene más de una boleta en el mismo mes, se suman."
    )
    wn["A3"] = (
        "m³ cuenta = consumo total de la boleta Aguas Andinas. "
        "m³ WES = registro API del mismo período de lecturas + proyección de días sin dato."
    )
    wn["A4"] = (
        "Celda amarilla en «m³ cuenta»: esa boleta es promedio/estimado (no es lectura real de turbina). "
        "Igual se totaliza porque es lo que marcó la cuenta ese mes."
    )
    wn["A5"] = f"Generado: {generado.strftime('%d-%m-%Y %H:%M')}"
    wn.column_dimensions["A"].width = 120
    wb.save(out_xlsx)


def _write_word(
    sitios: List[Sitio],
    meses: List[Tuple[int, int]],
    data,
    out_docx: Path,
    generado: datetime,
) -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.2)
    add_logo_to_header(doc)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    title = doc.add_heading("Totalizados mensuales — cuenta vs WES", level=1)
    if title.runs:
        title.runs[0].font.color.rgb = HEADING_RGB
    p = doc.add_paragraph()
    p.add_run("CORMUP / Peñalolén. ").bold = True
    p.add_run(
        "Fila = colegio. Por cada mes (emisión de la boleta, ene-2026 en adelante): "
        "m³ que marcó la cuenta y m³ que marcó WES. "
        f"Generado {generado.strftime('%d-%m-%Y %H:%M')}."
    )

    # Encabezado de dos niveles no cabe fácil: una fila de headers compactos
    # "ene-26 Cta" / "ene-26 WES"
    headers = ["Colegio"]
    for y, m in meses:
        lab = f"{MESES_CORTOS[m]}-{str(y)[2:]}"
        headers.append(f"{lab} cta")
        headers.append(f"{lab} WES")
    headers += ["Total cta", "Total WES"]

    sitios_ord = sorted(sitios, key=lambda s: s.node_id)
    rows: List[List[str]] = []
    tot_mes = {k: {"cuenta": 0.0, "wes": 0.0} for k in meses}
    gt_c = gt_w = 0.0
    for s in sitios_ord:
        line = [s.node_name]
        sc = sw = 0.0
        for k in meses:
            cel = data[s.node_id][k]
            if cel["n"]:
                line.append(format_number_chilean(cel["cuenta"], 0))
                line.append(format_number_chilean(cel["wes"], 0))
                sc += cel["cuenta"]
                sw += cel["wes"]
                tot_mes[k]["cuenta"] += cel["cuenta"]
                tot_mes[k]["wes"] += cel["wes"]
            else:
                line.append("—")
                line.append("—")
        line.append(format_number_chilean(sc, 0))
        line.append(format_number_chilean(sw, 0))
        gt_c += sc
        gt_w += sw
        rows.append(line)
    tot_line = ["TOTAL"]
    for k in meses:
        tot_line.append(format_number_chilean(tot_mes[k]["cuenta"], 0))
        tot_line.append(format_number_chilean(tot_mes[k]["wes"], 0))
    tot_line += [format_number_chilean(gt_c, 0), format_number_chilean(gt_w, 0)]
    rows.append(tot_line)

    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _tbl_full_width(table)
    for j, h in enumerate(headers):
        _set_cell(table.rows[0].cells[j], h, bold=True, size=7)
    for i, row in enumerate(rows, start=1):
        is_tot = i == len(rows)
        for j, val in enumerate(row):
            _set_cell(table.rows[i].cells[j], val, bold=is_tot, size=7)
    estilizar_tabla_wes(table, has_total_row=True)

    doc.add_paragraph(
        "Mes = fecha de emisión de la boleta. WES = consumo API en el período de lecturas "
        "de esa misma boleta (más proyección si hubo huecos). "
        "En el Excel, la cuenta en amarillo es boleta a promedio (no lectura real)."
    )
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)


def generar(skip_download: bool = False) -> Tuple[Path, Optional[Path], Path]:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    generado = datetime.now(timezone.utc).astimezone()
    ts = generado.strftime("%Y%m%d_%H%M")

    if skip_download and any(PDF_CACHE.glob("*/*.pdf")):
        locales = {p.name: p for p in PDF_CACHE.iterdir() if p.is_dir()}
        print(f"[INFO] Reuso cache {PDF_CACHE}", flush=True)
    else:
        locales = descargar_facturaciones(PDF_CACHE)

    sitios = _cargar_sitios(locales)
    cruzar_wes(sitios)
    meses = _meses_rango(sitios)
    if not meses:
        raise SystemExit("No hay boletas con emisión desde enero 2026.")
    data = _matriz(sitios, meses)

    stem = f"Totalizados_mensuales_cuenta_vs_WES_CORMUP_{ts}"
    out_xlsx = OUT_DIR / f"{stem}.xlsx"
    out_docx = OUT_DIR / f"{stem}.docx"
    _write_excel(sitios, meses, data, out_xlsx, generado)
    _write_word(sitios, meses, data, out_docx, generado)
    out_pdf = convertir_a_pdf(out_docx)
    print(f"[OK] Excel: {out_xlsx}", flush=True)
    print(f"[OK] Word:  {out_docx}", flush=True)
    if out_pdf:
        print(f"[OK] PDF:   {out_pdf}", flush=True)
    return out_docx, out_pdf, out_xlsx


def main() -> int:
    if sys.platform == "win32":
        for s in (sys.stdout, sys.stderr):
            try:
                s.reconfigure(encoding="utf-8", line_buffering=True)
            except Exception:
                pass
    ap = argparse.ArgumentParser(description="Totalizados mensuales cuenta vs WES CORMUP")
    ap.add_argument("--skip-download", action="store_true")
    args = ap.parse_args()
    generar(skip_download=args.skip_download)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
