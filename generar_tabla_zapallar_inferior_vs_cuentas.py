"""
Tabla simple: mismo mes/ciclo de cuenta ESVAL vs Estanque Inferior WES.
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from generar_reporte_word import format_number_chilean
from generar_tabla_zapallar_matriz_vs_inferior import (
    HEADER_FILL,
    ROW_ALERTA_FILL,
    TOTAL_FILL,
    WES_BLUE,
    _fmt_m3,
    _pct,
    _set_cell_text,
    _set_landscape,
)
from generar_tabla_zapallar_mensual_esval_inferior import _fmt_pct_s

SRC = Path(
    "reports/Fundo_Zapallar/ANALISIS/tabla_mensual_esval_inferior_20260922_1510/"
    "mensual_esval_matriz_inferior.json"
)

HEADERS = [
    "Mes",
    "Ciclo de la cuenta",
    "N° boleta",
    "Cuenta ESVAL (m³)",
    "Estanque Inf. (m³)",
    "Dif. (m³)",
    "Dif. %",
    "¿Cuadra?",
]


def _chk(pct: float) -> str:
    a = abs(pct)
    if a <= 5:
        return "SI"
    if a <= 15:
        return "cerca"
    return "NO"


def _fill(chk: str) -> str:
    if chk == "SI":
        return "C6E0B4"
    if chk == "cerca":
        return "FFF2CC"
    return ROW_ALERTA_FILL


def main() -> None:
    data = json.loads(SRC.read_text(encoding="utf-8"))
    ciclos = data["ciclos"]
    rows = []
    for c in ciclos:
        pct = c["pct_inf_esval"]
        chk = "—" if c.get("transicion") else _chk(pct)
        rows.append(
            {
                "mes": c["mes"],
                "ciclo": c["ciclo"],
                "boleta": c["boleta"],
                "esval": c["esval"],
                "inf": c["inferior"],
                "dif": c["dif_inf_esval"],
                "pct": pct,
                "chk": chk,
                "transicion": c.get("transicion", False),
                "fase": c["fase"],
            }
        )

    body = [
        [
            r["mes"],
            r["ciclo"],
            r["boleta"],
            _fmt_m3(r["esval"], 0),
            _fmt_m3(r["inf"], 1),
            _fmt_m3(r["dif"], 1),
            _fmt_pct_s(r["pct"]),
            r["chk"],
        ]
        for r in rows
    ]
    tot_e = sum(r["esval"] for r in rows)
    tot_i = sum(r["inf"] for r in rows)
    tot = [
        "TOTAL 12",
        "",
        "",
        _fmt_m3(tot_e, 0),
        _fmt_m3(tot_i, 1),
        _fmt_m3(tot_i - tot_e, 1),
        _fmt_pct_s(_pct(tot_i - tot_e, tot_e)),
        "",
    ]
    sin = [r for r in rows if not r["transicion"]]
    se, si = sum(r["esval"] for r in sin), sum(r["inf"] for r in sin)
    tot_sin = [
        "TOTAL sin transición",
        "10 cuentas",
        "",
        _fmt_m3(se, 0),
        _fmt_m3(si, 1),
        _fmt_m3(si - se, 1),
        _fmt_pct_s(_pct(si - se, se)),
        "",
    ]

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out_dir = Path("reports/Fundo_Zapallar/ANALISIS") / f"inferior_vs_cuentas_{stamp}"
    out_dir.mkdir(parents=True, exist_ok=True)

    labels = [r["mes"] for r in rows]
    x = np.arange(len(labels))
    w = 0.38
    fig, ax = plt.subplots(figsize=(12.8, 4.4))
    ax.bar(x - w / 2, [r["esval"] for r in rows], w, label="Cuenta ESVAL (delta medidor)", color="#7F7F7F")
    ax.bar(x + w / 2, [r["inf"] for r in rows], w, label="Estanque Inferior WES", color="#70AD47")
    ax.set_ylabel("m³")
    ax.set_title("Mismo mes: cuenta ESVAL vs Estanque Inferior", fontsize=12, fontweight="bold", color=WES_BLUE)
    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontsize=9)
    ax.legend(frameon=False, fontsize=9)
    ax.grid(axis="y", linestyle="--", alpha=0.35)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    chart = out_dir / "chart_inferior_vs_cuentas.png"
    fig.savefig(chart, dpi=160, bbox_inches="tight")
    plt.close(fig)

    # Word
    doc = Document()
    _set_landscape(doc)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("Estanque Inferior vs cuentas ESVAL — mismo mes")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(31, 78, 121)
    n = doc.add_paragraph()
    n.add_run(
        "Cada fila usa el mismo ciclo de la boleta. Cuenta ESVAL = delta del medidor de turbina "
        "(no el ítem “facturado”). Inferior = WES día completo 00:00–24:00 Chile. "
        "±5 % cuadra, ±15 % cerca. Marzo–abril = transición, no calibrar."
    ).font.size = Pt(10)

    table = doc.add_table(rows=1 + len(body) + 2, cols=len(HEADERS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(HEADERS):
        fill = "C6E0B4" if i == 4 else HEADER_FILL
        color = RGBColor(31, 78, 121) if i == 4 else RGBColor(255, 255, 255)
        _set_cell_text(table.rows[0].cells[i], h, bold=True, size=9, color=color, fill=fill)
    for r_i, vals in enumerate(body):
        row = rows[r_i]
        for c_i, val in enumerate(vals):
            if c_i == 7:
                fill = _fill(row["chk"]) if row["chk"] != "—" else "F4C7C3"
            elif c_i == 4:
                fill = "C6E0B4" if row["chk"] in ("SI", "cerca") else ("F4C7C3" if row["transicion"] else "F8CBAD")
            else:
                fill = ROW_ALERTA_FILL if row["transicion"] else "FFFFFF"
            _set_cell_text(table.rows[r_i + 1].cells[c_i], val, size=9, fill=fill)
    for extra, vals in ((1 + len(body), tot), (2 + len(body), tot_sin)):
        for c_i, val in enumerate(vals):
            _set_cell_text(table.rows[extra].cells[c_i], val, bold=True, size=9, fill=TOTAL_FILL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(chart), width=Inches(10.0))
    docx_path = out_dir / "Inferior_vs_cuentas_ESVAL_mismo_mes.docx"
    doc.save(str(docx_path))

    # PDF
    headers_pdf = [
        "Mes",
        "Ciclo de la cuenta",
        "N° boleta",
        "Cuenta ESVAL\n(m³)",
        "Estanque Inf.\n(m³)",
        "Dif. (m³)",
        "Dif. %",
        "¿Cuadra?",
    ]
    cell = [headers_pdf] + body + [tot, tot_sin]
    fig = plt.figure(figsize=(15.4, 9.2))
    fig.suptitle(
        "Estanque Inferior vs cuentas ESVAL — mismo mes",
        fontsize=16,
        fontweight="bold",
        color=WES_BLUE,
        y=0.97,
        x=0.03,
        ha="left",
    )
    fig.text(
        0.03,
        0.925,
        "Mismo ciclo de lectura de cada boleta. Cuenta = delta del medidor turbina ESVAL. "
        "Inferior = WES 00:00–24:00 Chile. ±5 % cuadra · ±15 % cerca. Filas rosadas = transición (no calibrar).",
        fontsize=8.5,
        color="#444444",
    )
    ax = fig.add_axes([0.03, 0.40, 0.94, 0.50])
    ax.axis("off")
    col_w = [0.08, 0.16, 0.12, 0.14, 0.14, 0.13, 0.10, 0.10]
    tb = ax.table(cellText=cell, loc="center", cellLoc="center", colWidths=col_w)
    tb.auto_set_font_size(False)
    tb.set_fontsize(8)
    tb.scale(1, 1.55)
    n_rows, n_cols = len(cell), len(headers_pdf)
    for c in range(n_cols):
        h = tb[0, c]
        h.set_facecolor("#C6E0B4" if c == 4 else WES_BLUE)
        h.get_text().set_color(WES_BLUE if c == 4 else "white")
        h.get_text().set_fontweight("bold")
    for r in range(1, n_rows - 2):
        row = rows[r - 1]
        for c in range(n_cols):
            cell_i = tb[r, c]
            if c == 7:
                cell_i.set_facecolor(
                    "#C6E0B4" if row["chk"] == "SI" else ("#FFF2CC" if row["chk"] == "cerca" else "#F4C7C3")
                )
            elif c == 4:
                cell_i.set_facecolor(
                    "#C6E0B4" if row["chk"] in ("SI", "cerca") else ("#F4C7C3" if row["transicion"] else "#F8CBAD")
                )
            else:
                cell_i.set_facecolor("#F4C7C3" if row["transicion"] else "white")
    for r in (n_rows - 2, n_rows - 1):
        for c in range(n_cols):
            tb[r, c].set_facecolor("#D6E3F0")
            tb[r, c].get_text().set_fontweight("bold")
    axc = fig.add_axes([0.06, 0.05, 0.88, 0.32])
    axc.imshow(plt.imread(str(chart)))
    axc.axis("off")
    pdf_path = out_dir / "Inferior_vs_cuentas_ESVAL_mismo_mes.pdf"
    fig.savefig(pdf_path, dpi=170)
    plt.close(fig)

    payload = {
        "fuente": str(SRC),
        "filas": rows,
        "totales": {"esval": tot_e, "inferior": tot_i},
        "sin_transicion": {"esval": se, "inferior": si},
    }
    (out_dir / "inferior_vs_cuentas.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    print(f"[OK] {pdf_path}")
    print(f"[OK] {docx_path}")
    for r in rows:
        print(f"  {r['mes']} boleta {r['boleta']}: cuenta {r['esval']:.0f}  inf {r['inf']:.1f}  {r['pct']:+.1f}%  {r['chk']}")
    print(f"OUT_DIR={out_dir}")


if __name__ == "__main__":
    main()
