"""
Comparativo cuentas Aguas Andinas vs registro API WES — colegios CORMUP / Peñalolén.

Lee las facturaciones de Google Drive:
  G:\\Mi unidad\\Colegios\\Peñalolén\\Facturaciones
(una subcarpeta por establecimiento) y cruza cada período de lectura real
con el consumo diario del nodo WES. Las boletas a promedio no se comparan.
Si hay huecos de API se proyecta el m³ faltante y se suma al WES medido.

Salida (Word + PDF + Excel) en:
  reports/CORMUP/Facturaciones_vs_WES/

Uso:
  python generar_comparativo_facturaciones_cormup_penalolen.py
  python generar_comparativo_facturaciones_cormup_penalolen.py --skip-download
"""

from __future__ import annotations

import argparse
import io
import re
import shutil
import subprocess
import sys
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from googleapiclient.http import MediaIoBaseDownload
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from facturacion_aguas_andinas_pdf import extraer_texto_pdf, listar_periodos_desde_pdf
from generar_reporte_word import (
    acl_node_base_url,
    add_formatted_heading,
    add_logo_to_header,
    add_picture_with_pagination,
    estilizar_tabla_wes,
    fetch_json,
    flatten_measures,
    format_number_chilean,
    get_node_name,
    normalize_measures_payload,
    summarize_consumption,
)
from wes_estilo_graficos_app import COLOR_BARRA_FACT, COLOR_BARRA_WES
from wes_google_drive import obtener_servicio_drive

ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "reports" / "CORMUP" / "Facturaciones_vs_WES"
PDF_CACHE = OUT_DIR / "_pdfs"
COMPANY_ID = "000008"
DRIVE_PATH_PARTS = ("Colegios", "Peñalolén", "Facturaciones")
# ID conocido de G:\Mi unidad\Colegios\Peñalolén\Facturaciones (fallback si falla la búsqueda).
DRIVE_FOLDER_ID_FALLBACK = "1v7uQVAAlNLM1KelJU7ELzgkawgGoc8ew"

# Carpeta Drive → nodo WES CORMUP (000008-01 … 000008-14)
FOLDER_ALIASES: Dict[str, str] = {
    "171": "000008-01",
    "d-171": "000008-01",
    "d 171": "000008-01",
    "antonio hermida": "000008-01",
    "antonio hermida fabres": "000008-01",
    "liceo antonio hermida": "000008-01",
    "eduardo de la barra": "000008-02",
    "carlos fernandez pena": "000008-03",
    "carlos fernandez peña": "000008-03",
    "tobalaba": "000008-04",
    "santa maria": "000008-05",
    "santa maría": "000008-05",
    "santamaria": "000008-05",
    "arrieta canas": "000008-06",
    "arrieta cañas": "000008-06",
    "luis arrieta cana": "000008-06",
    "luis arrieta caña": "000008-06",
    "erasmo escala": "000008-07",
    "erasmos escala": "000008-07",
    "alicura": "000008-08",
    "juan bautista pasten": "000008-09",
    "juan bautista pastenes": "000008-09",
    "matilde huici navas": "000008-10",
    "ce valle hermoso": "000008-11",
    "ce valler hermoso": "000008-11",
    "valle hermoso": "000008-11",
    "union nacional arabe": "000008-12",
    "unión nacional árabe": "000008-12",
    "union nacional árabe": "000008-12",
    "likankura": "000008-13",
    "juan pablo ii": "000008-14",
    "juan pablo 2": "000008-14",
}

NODOS_CORMUP = [f"000008-{i:02d}" for i in range(1, 15)]
MESES_CORTOS = {
    1: "ene",
    2: "feb",
    3: "mar",
    4: "abr",
    5: "may",
    6: "jun",
    7: "jul",
    8: "ago",
    9: "sep",
    10: "oct",
    11: "nov",
    12: "dic",
}
HEADING_RGB = RGBColor(0, 51, 102)
UMBRAL_PCT_DESTACAR = 15.0


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.lower().replace("_", " ").replace("-", " ")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def _fmt_fecha(dt: datetime | date) -> str:
    if isinstance(dt, datetime):
        d = dt.date()
    else:
        d = dt
    return f"{d.day:02d}-{MESES_CORTOS[d.month]}-{d.year}"


def _to_ddmmyyyy(dt: datetime | date) -> str:
    if isinstance(dt, datetime):
        d = dt.date()
    else:
        d = dt
    return d.strftime("%d%m%Y")


def _resolver_nodo(nombre_carpeta: str) -> Optional[str]:
    key = _norm(nombre_carpeta)
    if key in FOLDER_ALIASES:
        return FOLDER_ALIASES[key]
    for alias, node_id in FOLDER_ALIASES.items():
        if alias and (alias in key or key in alias):
            return node_id
    return None


def _listar_hijos(service, parent_id: str) -> list[dict]:
    items: list[dict] = []
    page_token = None
    while True:
        resp = (
            service.files()
            .list(
                q=f"'{parent_id}' in parents and trashed=false",
                fields="nextPageToken, files(id, name, mimeType, size)",
                pageSize=100,
                pageToken=page_token,
                includeItemsFromAllDrives=True,
                supportsAllDrives=True,
            )
            .execute()
        )
        items.extend(resp.get("files", []))
        page_token = resp.get("nextPageToken")
        if not page_token:
            break
    return items


def _buscar_carpeta_por_nombre(service, nombre: str, parent_id: Optional[str] = None) -> Optional[dict]:
    safe = nombre.replace("'", "\\'")
    parent_clause = f"'{parent_id}' in parents and " if parent_id else ""
    query = (
        f"name='{safe}' and mimeType='application/vnd.google-apps.folder' "
        f"and {parent_clause}trashed=false"
    )
    items = (
        service.files()
        .list(
            q=query,
            fields="files(id, name, parents)",
            pageSize=10,
            includeItemsFromAllDrives=True,
            supportsAllDrives=True,
        )
        .execute()
        .get("files", [])
    )
    return items[0] if items else None


def _buscar_carpeta_facturaciones(service) -> str:
    """Resuelve Colegios / Peñalolén / Facturaciones en Drive."""
    parent = None
    current = None
    for parte in DRIVE_PATH_PARTS:
        found = None
        # Variantes de tilde / mayúsculas
        candidatos = [parte, parte.replace("é", "e"), parte.replace("ó", "o")]
        vistos = []
        for nom in candidatos:
            if nom in vistos:
                continue
            vistos.append(nom)
            found = _buscar_carpeta_por_nombre(service, nom, parent_id=parent)
            if found:
                break
        if not found and parte == "Peñalolén":
            for nom in ("Peñalolen", "Penalolen", "Peñalolén"):
                found = _buscar_carpeta_por_nombre(service, nom, parent_id=parent)
                if found:
                    break
        if not found:
            break
        current = found
        parent = found["id"]
    if current and current.get("name", "").lower() == "facturaciones":
        print(f"[Drive] Carpeta: {current['id']} ({'/'.join(DRIVE_PATH_PARTS)})", flush=True)
        return current["id"]
    print(
        f"[Drive] Búsqueda por ruta falló; uso ID conocido {DRIVE_FOLDER_ID_FALLBACK}",
        flush=True,
    )
    return DRIVE_FOLDER_ID_FALLBACK


def _descargar_archivo(service, file_id: str, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    req = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    buf = io.BytesIO()
    downloader = MediaIoBaseDownload(buf, req)
    done = False
    while not done:
        _status, done = downloader.next_chunk()
    dest.write_bytes(buf.getvalue())


def descargar_facturaciones(cache_dir: Path) -> Dict[str, Path]:
    """Descarga PDFs agrupados por subcarpeta (establecimiento). Retorna carpeta local por establecimiento."""
    service = obtener_servicio_drive()
    folder_id = _buscar_carpeta_facturaciones(service)
    cache_dir.mkdir(parents=True, exist_ok=True)
    locales: Dict[str, Path] = {}
    hijos = _listar_hijos(service, folder_id)
    subdirs = [h for h in hijos if h["mimeType"] == "application/vnd.google-apps.folder"]
    pdfs_raiz = [
        h
        for h in hijos
        if h["mimeType"] != "application/vnd.google-apps.folder"
        and h["name"].lower().endswith(".pdf")
    ]
    print(f"[Drive] {len(subdirs)} establecimientos, {len(pdfs_raiz)} PDF en raíz", flush=True)

    for sub in sorted(subdirs, key=lambda x: x["name"].lower()):
        local = cache_dir / re.sub(r'[<>:"/\\|?*]', "_", sub["name"]).strip()
        local.mkdir(parents=True, exist_ok=True)
        archivos = [
            a
            for a in _listar_hijos(service, sub["id"])
            if a["name"].lower().endswith(".pdf")
        ]
        print(f"  [{sub['name']}] {len(archivos)} PDF", flush=True)
        for a in archivos:
            dest = local / a["name"]
            if dest.is_file() and dest.stat().st_size > 10_000:
                continue
            _descargar_archivo(service, a["id"], dest)
        locales[sub["name"]] = local

    if pdfs_raiz:
        local = cache_dir / "_raiz"
        local.mkdir(parents=True, exist_ok=True)
        for a in pdfs_raiz:
            dest = local / a["name"]
            if not dest.is_file() or dest.stat().st_size < 10_000:
                _descargar_archivo(service, a["id"], dest)
        locales["_raiz"] = local
    return locales


def _extraer_cuenta(txt: str, fallback: Optional[str]) -> Optional[str]:
    if fallback:
        return fallback
    m = re.search(
        r"Nro\.?\s*de\s*cuenta[\s\S]{0,120}?(\d{5,7}-\d)",
        txt,
        flags=re.IGNORECASE,
    )
    if m:
        return m.group(1)
    # Tras dirección de Aguas Andinas suele ir el n° de cuenta sin puntos (no es RUT).
    cands = re.findall(r"(?<![\d.])(\d{5,7}-\d)(?!\d)", txt)
    for c in cands:
        cuerpo = c.split("-")[0]
        if 5 <= len(cuerpo) <= 7:
            return c
    return None


def _extraer_claves(txt: str) -> Tuple[str, str]:
    fact = ""
    lect = ""
    m = re.search(r"Clave\s+Facturaci[oó]n\s+([^\n]+)", txt, flags=re.IGNORECASE)
    if m:
        fact = re.sub(r"\s+", " ", m.group(1)).strip()
    m = re.search(r"Clave\s+Lectura\s+([^\n]+)", txt, flags=re.IGNORECASE)
    if m:
        lect = re.sub(r"\s+", " ", m.group(1)).strip()
    return fact, lect


def _es_estimado(clave_fact: str, clave_lect: str) -> bool:
    """Boletas a promedio / sin lectura real: no sirven para el cruce válido."""
    u = f"{clave_fact} {clave_lect}".upper()
    return any(
        t in u
        for t in (
            "ESTIMAD",
            "CASA CERRADA",
            "PROMEDIO",
            "DESCONTABLE",
            "MEDIDOR DETENIDO",
            "CERRADO",
            "TERMINO MEDIO",
            "TÉRMINO MEDIO",
        )
    )


def _es_ciclo_septiembre(f: "FilaComparacion") -> bool:
    return f.emision.month == 9 or f.lectura_actual.month == 9


def _fechas_inclusive(a: date, b: date) -> list[date]:
    out: list[date] = []
    d = a
    while d <= b:
        out.append(d)
        d += timedelta(days=1)
    return out


def _colapsar_huecos(fechas: list[date]) -> str:
    if not fechas:
        return ""
    rangos: list[tuple[date, date]] = []
    ini = prev = fechas[0]
    for d in fechas[1:]:
        if d == prev + timedelta(days=1):
            prev = d
        else:
            rangos.append((ini, prev))
            ini = prev = d
    rangos.append((ini, prev))
    partes = []
    for a, b in rangos:
        partes.append(_fmt_fecha(a) if a == b else f"{_fmt_fecha(a)}→{_fmt_fecha(b)}")
    return ", ".join(partes)


def _extraer_total_pagar(txt: str) -> Optional[int]:
    m = re.search(r"TOTAL\s+A\s+PAGAR\s*\$\s*([\d\.]+)", txt, flags=re.IGNORECASE)
    if not m:
        return None
    try:
        return int(m.group(1).replace(".", ""))
    except ValueError:
        return None


@dataclass
class FilaComparacion:
    establecimiento: str
    carpeta: str
    node_id: str
    node_name: str
    pdf_name: str
    boleta: str
    cuenta: str
    medidor: str
    emision: datetime
    lectura_anterior: datetime
    lectura_actual: datetime
    m3_cuenta: int
    m3_wes_medido: float
    m3_proyeccion: float
    dias_wes: int
    dias_esperados: int
    dias_hueco: int
    dias_periodo: int
    huecos_txt: str
    clave_facturacion: str
    clave_lectura: str
    estimado: bool
    total_pagar: Optional[int] = None
    error_wes: str = ""

    @property
    def valido(self) -> bool:
        return not self.estimado

    @property
    def m3_wes(self) -> float:
        """WES medido + proyección de huecos (lo comparable)."""
        return float(self.m3_wes_medido) + float(self.m3_proyeccion)

    @property
    def periodo_txt(self) -> str:
        return f"{_fmt_fecha(self.lectura_anterior)} → {_fmt_fecha(self.lectura_actual)}"

    @property
    def diff_m3(self) -> float:
        return float(self.m3_cuenta) - float(self.m3_wes)

    @property
    def pct(self) -> Optional[float]:
        if not self.m3_cuenta:
            return None
        return 100.0 * self.diff_m3 / float(self.m3_cuenta)

    @property
    def observacion(self) -> str:
        notas: list[str] = []
        if self.estimado:
            notas.append("NO VÁLIDO — promedio / estimado (no se compara)")
        if self.error_wes:
            notas.append(self.error_wes)
        elif self.dias_wes == 0 and not self.estimado:
            notas.append("Sin registro WES en el período")
        if self.dias_hueco > 0:
            notas.append(
                f"Huecos {self.dias_hueco} d ({self.huecos_txt}); "
                f"proyección +{format_number_chilean(self.m3_proyeccion, 1)} m³"
            )
        if self.valido and self.pct is not None and abs(self.pct) <= 10:
            notas.append("Alineado (±10 %)")
        return "; ".join(notas) if notas else ""


@dataclass
class Sitio:
    carpeta: str
    node_id: str
    node_name: str
    filas: List[FilaComparacion] = field(default_factory=list)
    errores_pdf: List[str] = field(default_factory=list)


def _cargar_sitios(locales: Dict[str, Path]) -> List[Sitio]:
    sitios: List[Sitio] = []
    no_mapeados: List[str] = []
    for carpeta, path in sorted(locales.items(), key=lambda x: x[0].lower()):
        node_id = _resolver_nodo(carpeta)
        if not node_id:
            no_mapeados.append(carpeta)
            print(f"[WARN] Carpeta sin nodo WES: {carpeta}", flush=True)
            continue
        sitio = Sitio(
            carpeta=carpeta,
            node_id=node_id,
            node_name=get_node_name(node_id),
        )
        pdfs = sorted(path.glob("*.pdf"))
        vistos: set[tuple[str, str]] = set()
        for pdf in pdfs:
            # Duplicados tipo "archivo (1).pdf"
            if " (1)" in pdf.name:
                orig = path / pdf.name.replace(" (1)", "", 1)
                if orig.is_file():
                    continue
            try:
                txt = extraer_texto_pdf(pdf)
                periodos = listar_periodos_desde_pdf(pdf)
            except Exception as e:
                sitio.errores_pdf.append(f"{pdf.name}: {e}")
                print(f"  [ERR] {carpeta}/{pdf.name}: {e}", flush=True)
                continue
            clave_f, clave_l = _extraer_claves(txt)
            total_pagar = _extraer_total_pagar(txt)
            estimado = _es_estimado(clave_f, clave_l)
            for per in periodos:
                cuenta = _extraer_cuenta(txt, per.cuenta) or ""
                key = (per.boleta, per.lectura_actual.strftime("%Y%m%d"))
                if key in vistos:
                    continue
                vistos.add(key)
                dias = (per.lectura_actual.date() - per.lectura_anterior.date()).days
                sitio.filas.append(
                    FilaComparacion(
                        establecimiento=sitio.node_name,
                        carpeta=carpeta,
                        node_id=node_id,
                        node_name=sitio.node_name,
                        pdf_name=pdf.name,
                        boleta=str(per.boleta),
                        cuenta=cuenta,
                        medidor=per.medidor or "",
                        emision=per.emision,
                        lectura_anterior=per.lectura_anterior,
                        lectura_actual=per.lectura_actual,
                        m3_cuenta=int(per.m3_cuenta),
                        m3_wes_medido=0.0,
                        m3_proyeccion=0.0,
                        dias_wes=0,
                        dias_esperados=int(dias) + 1,
                        dias_hueco=0,
                        dias_periodo=int(dias),
                        huecos_txt="",
                        clave_facturacion=clave_f,
                        clave_lectura=clave_l,
                        estimado=estimado,
                        total_pagar=total_pagar,
                    )
                )
        sitio.filas.sort(key=lambda f: f.lectura_actual)
        sitios.append(sitio)
        print(
            f"  [OK] {sitio.node_name} ({node_id}): {len(sitio.filas)} períodos, "
            f"{len(sitio.errores_pdf)} PDF con error",
            flush=True,
        )
    if no_mapeados:
        print(f"[WARN] Carpetas sin mapear: {no_mapeados}", flush=True)
    sitios.sort(key=lambda s: s.node_id)
    return sitios


def _medidas_wes(node_id: str, d0: date, d1: date) -> List:
    start = _to_ddmmyyyy(d0)
    end_api = _to_ddmmyyyy(d1 + timedelta(days=1))
    raw = fetch_json(
        f"{acl_node_base_url()}/nodes/measures/dates",
        params=[("id", node_id), ("start", start), ("end", end_api)],
    )
    norm = normalize_measures_payload(raw, node_id)
    return flatten_measures(norm)


def cruzar_wes(sitios: List[Sitio]) -> None:
    for sitio in sitios:
        if not sitio.filas:
            continue
        d0 = min(f.lectura_anterior.date() for f in sitio.filas)
        d1 = max(f.lectura_actual.date() for f in sitio.filas)
        print(
            f"[WES] {sitio.node_id} {sitio.node_name} {_fmt_fecha(d0)} → {_fmt_fecha(d1)}",
            flush=True,
        )
        try:
            meas = _medidas_wes(sitio.node_id, d0, d1)
        except Exception as e:
            print(f"  [WARN] API {sitio.node_id}: {e}", flush=True)
            for f in sitio.filas:
                f.error_wes = f"Error API WES: {e}"
            continue
        # Promedio diario del nodo (días con registro) para proyectar si el período viene vacío.
        tot_nodo = sum(m.total_m3 for m in meas)
        dias_nodo = len({m.date.date() for m in meas})
        prom_nodo = (tot_nodo / dias_nodo) if dias_nodo else 0.0
        for f in sitio.filas:
            a, b = f.lectura_anterior.date(), f.lectura_actual.date()
            esperadas = _fechas_inclusive(a, b)
            f.dias_esperados = len(esperadas)
            por_dia = {m.date.date(): float(m.total_m3) for m in meas if a <= m.date.date() <= b}
            f.dias_wes = len(por_dia)
            f.m3_wes_medido = float(sum(por_dia.values()))
            huecos = [d for d in esperadas if d not in por_dia]
            f.dias_hueco = len(huecos)
            f.huecos_txt = _colapsar_huecos(huecos)
            if f.dias_wes >= 3:
                prom = f.m3_wes_medido / f.dias_wes
            else:
                prom = prom_nodo
            f.m3_proyeccion = float(prom * f.dias_hueco) if huecos else 0.0
            if huecos:
                print(
                    f"  [proy] {f.boleta} huecos={f.dias_hueco} "
                    f"+{f.m3_proyeccion:.1f} m³ (prom {prom:.2f})",
                    flush=True,
                )


def _set_cell(cell, text: str, *, bold: bool = False, size: int = 8) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(str(text))
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run.bold = bold


def _tbl_full_width(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    for tag_nm in ("tblW", "tblLayout"):
        tq = qn(f"w:{tag_nm}")
        for el in list(tbl_pr):
            if el.tag == tq:
                tbl_pr.remove(el)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), "5000")
    tbl_w.set(qn("w:type"), "pct")
    tbl_pr.append(tbl_w)
    tbl_layout = OxmlElement("w:tblLayout")
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_pr.append(tbl_layout)


def _add_table_rows(doc: Document, headers: List[str], rows: List[List[str]], *, highlight: Optional[Iterable[int]] = None, has_total: bool = True) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _tbl_full_width(table)
    for j, h in enumerate(headers):
        _set_cell(table.rows[0].cells[j], h, bold=True, size=8)
    for i, row in enumerate(rows, start=1):
        is_total = has_total and i == len(rows)
        for j, val in enumerate(row):
            _set_cell(table.rows[i].cells[j], val, bold=is_total, size=8)
    estilizar_tabla_wes(table, highlight_rows=highlight, has_total_row=has_total)
    doc.add_paragraph("")


def _filas_validas(sitios: List[Sitio]) -> List[FilaComparacion]:
    return [f for s in sitios for f in s.filas if f.valido]


def _filas_septiembre(sitios: List[Sitio]) -> List[FilaComparacion]:
    return [f for s in sitios for f in s.filas if _es_ciclo_septiembre(f)]


def _grafico_barras(
    filas: List[FilaComparacion],
    out_png: Path,
    *,
    titulo: str,
    solo_validas: bool = True,
) -> Path:
    data = [f for f in filas if (f.valido if solo_validas else True)]
    labels = [f.node_name for f in data]
    m3_fact = [f.m3_cuenta for f in data]
    m3_wes = [f.m3_wes for f in data]
    fig, ax = plt.subplots(figsize=(11.5, 5.2))
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    x = range(len(labels))
    w = 0.38
    ax.bar([i - w / 2 for i in x], m3_fact, width=w, color=COLOR_BARRA_FACT, label="m³ cuenta (lectura real)")
    ax.bar([i + w / 2 for i in x], m3_wes, width=w, color=COLOR_BARRA_WES, label="m³ WES + proyección huecos")
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels, rotation=35, ha="right", fontsize=8)
    ax.set_ylabel("m³")
    ax.set_title(titulo)
    ax.legend(frameon=False)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.yaxis.grid(True, linestyle="--", linewidth=0.6, color="#cbd5e1", alpha=0.9)
    ax.set_axisbelow(True)
    fig.tight_layout()
    fig.savefig(out_png, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return out_png


def _style_header(ws, n_cols: int, header_fill, header_font) -> None:
    for col in range(1, n_cols + 1):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def _fila_detalle_excel(f: FilaComparacion) -> list:
    pct = f.pct if f.valido else None
    dif = f.diff_m3 if f.valido else None
    return [
        f.establecimiento,
        f.node_id,
        "Sí" if f.valido else "No",
        f.cuenta,
        f.medidor,
        f.boleta,
        f.emision.strftime("%d-%m-%Y"),
        f.lectura_anterior.strftime("%d-%m-%Y"),
        f.lectura_actual.strftime("%d-%m-%Y"),
        f.dias_esperados,
        f.m3_cuenta,
        round(f.m3_wes_medido, 1),
        f.dias_wes,
        f.dias_hueco,
        f.huecos_txt,
        round(f.m3_proyeccion, 1),
        round(f.m3_wes, 1),
        None if dif is None else round(dif, 1),
        None if pct is None else round(pct, 1),
        f.clave_facturacion,
        f.clave_lectura,
        f.observacion,
        f.pdf_name,
        f.total_pagar if f.total_pagar is not None else "",
    ]


HEADERS_DETALLE = [
    "Establecimiento",
    "Nodo WES",
    "Válida (lectura real)",
    "Cuenta Aguas Andinas",
    "Medidor",
    "N° factura / boleta",
    "Emisión",
    "Lectura anterior",
    "Lectura actual",
    "Días esperados",
    "m³ cuenta",
    "m³ WES medido",
    "Días WES",
    "Días hueco",
    "Huecos (fechas)",
    "m³ proyección huecos",
    "m³ WES + proyección",
    "Dif m³ (cuenta − WES+proy)",
    "% dif vs cuenta",
    "Clave facturación",
    "Clave lectura",
    "Observación",
    "Archivo PDF",
    "Total a pagar (CLP)",
]


def _write_sheet_filas(ws, filas: List[FilaComparacion], header_fill, header_font) -> None:
    ws.append(HEADERS_DETALLE)
    _style_header(ws, len(HEADERS_DETALLE), header_fill, header_font)
    for f in filas:
        ws.append(_fila_detalle_excel(f))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[9].number_format = "0"
        row[10].number_format = "#,##0"
        row[11].number_format = "#,##0.0"
        row[12].number_format = "0"
        row[13].number_format = "0"
        row[15].number_format = "#,##0.0"
        row[16].number_format = "#,##0.0"
        row[17].number_format = "#,##0.0"
        row[18].number_format = "0.0"
        row[23].number_format = "#,##0"
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    widths = [26, 12, 16, 16, 14, 16, 12, 14, 14, 12, 12, 14, 10, 10, 28, 16, 18, 20, 12, 28, 18, 40, 28, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_excel(sitios: List[Sitio], out_xlsx: Path, generado: datetime) -> None:
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="003366")
    header_font = Font(color="FFFFFF", bold=True)
    total_fill = PatternFill("solid", fgColor="D9E1F2")

    wsep = wb.active
    wsep.title = "Septiembre"
    sep = _filas_septiembre(sitios)
    sep.sort(key=lambda f: (not f.valido, f.node_id))
    _write_sheet_filas(wsep, sep, header_fill, header_font)

    ws = wb.create_sheet("Resumen_valido")
    headers_res = [
        "Establecimiento",
        "Nodo WES",
        "N° períodos válidos",
        "m³ cuenta (reales)",
        "m³ WES medido",
        "m³ proyección huecos",
        "m³ WES + proyección",
        "Dif m³ (cuenta − WES+proy)",
        "% dif vs cuenta",
        "Períodos no válidos (promedio)",
        "PDF con error",
    ]
    ws.append(headers_res)
    _style_header(ws, len(headers_res), header_fill, header_font)
    tot_c = tot_med = tot_pr = tot_n = tot_est = 0
    for s in sitios:
        validas = [f for f in s.filas if f.valido]
        m3c = sum(f.m3_cuenta for f in validas)
        m3m = sum(f.m3_wes_medido for f in validas)
        m3p = sum(f.m3_proyeccion for f in validas)
        m3w = m3m + m3p
        n_est = sum(1 for f in s.filas if f.estimado)
        pct = (100.0 * (m3c - m3w) / m3c) if m3c else None
        ws.append(
            [
                s.node_name,
                s.node_id,
                len(validas),
                m3c,
                round(m3m, 1),
                round(m3p, 1),
                round(m3w, 1),
                round(m3c - m3w, 1),
                None if pct is None else round(pct, 1),
                n_est,
                len(s.errores_pdf),
            ]
        )
        tot_c += m3c
        tot_med += m3m
        tot_pr += m3p
        tot_n += len(validas)
        tot_est += n_est
    tot_w = tot_med + tot_pr
    tot_pct = (100.0 * (tot_c - tot_w) / tot_c) if tot_c else None
    ws.append(
        [
            "TOTAL (solo lecturas reales)",
            COMPANY_ID,
            tot_n,
            tot_c,
            round(tot_med, 1),
            round(tot_pr, 1),
            round(tot_w, 1),
            round(tot_c - tot_w, 1),
            None if tot_pct is None else round(tot_pct, 1),
            tot_est,
            sum(len(s.errores_pdf) for s in sitios),
        ]
    )
    for cell in ws[ws.max_row]:
        cell.fill = total_fill
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[2].number_format = "0"
        row[3].number_format = "#,##0"
        row[4].number_format = "#,##0.0"
        row[5].number_format = "#,##0.0"
        row[6].number_format = "#,##0.0"
        row[7].number_format = "#,##0.0"
        row[8].number_format = "0.0"
    ws.freeze_panes = "A2"
    for i, w in enumerate([28, 12, 16, 18, 16, 18, 20, 24, 14, 24, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wd = wb.create_sheet("Detalle")
    todas = [f for s in sitios for f in s.filas]
    _write_sheet_filas(wd, todas, header_fill, header_font)

    wn = wb.create_sheet("Notas")
    wn["A1"] = "Criterio"
    wn["A2"] = (
        "Fuente de facturas: Google Drive → Colegios / Peñalolén / Facturaciones "
        "(una subcarpeta por establecimiento)."
    )
    wn["A3"] = (
        "Comparativo VÁLIDO: solo boletas con lectura real (no promedio, no estimado SISS, "
        "no casa cerrada, no medidor detenido, no cerrado). Esas no se contrastan contra WES."
    )
    wn["A4"] = (
        "Huecos: días del período de lecturas sin registro diario en la API WES. "
        "Se proyectan con el promedio m³/día de los días con dato en el mismo período "
        "(si hay < 3 días, se usa el promedio del nodo en todo el rango descargado) "
        "y se SUMAN al m³ WES medido."
    )
    wn["A5"] = (
        "Diferencia = m³ cuenta − (m³ WES medido + proyección de huecos). "
        "Positivo: la boleta facturó más que WES+proyección."
    )
    wn["A6"] = (
        "Septiembre: boletas con emisión o lectura actual en septiembre (ciclo más reciente)."
    )
    wn["A7"] = f"Generado: {generado.strftime('%d-%m-%Y %H:%M')}"
    wn["A8"] = (
        "Eduardo de la Barra (000008-02) no tiene subcarpeta de facturaciones en Drive."
    )
    wn.column_dimensions["A"].width = 120
    wb.save(out_xlsx)


def _write_word(
    sitios: List[Sitio],
    out_docx: Path,
    chart_sep: Path,
    chart_png: Path,
    generado: datetime,
) -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.4)
    add_logo_to_header(doc)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Comparación cuentas vs registro API WES", level=1)
    if title.runs:
        title.runs[0].font.color.rgb = HEADING_RGB
    sub = doc.add_paragraph()
    r = sub.add_run("CORMUP / Colegios Peñalolén — facturaciones Aguas Andinas")
    r.italic = True
    r.font.color.rgb = HEADING_RGB

    p = doc.add_paragraph()
    p.add_run("Empresa: ").bold = True
    p.add_run("CORMUP (000008)")
    p2 = doc.add_paragraph()
    p2.add_run("Fuente facturas: ").bold = True
    p2.add_run(r"G:\Mi unidad\Colegios\Peñalolén\Facturaciones")
    p3 = doc.add_paragraph()
    p3.add_run("Generado: ").bold = True
    p3.add_run(generado.strftime("%d-%m-%Y %H:%M"))

    add_formatted_heading(doc, "1. Alcance y criterio", level=1)
    doc.add_paragraph(
        "Cada establecimiento se asocia al nodo WES CORMUP homónimo. El cruce válido "
        "usa solo boletas con lectura real. Las facturadas a promedio / estimado SISS / "
        "casa cerrada / medidor detenido / cerrado NO se comparan: el m³ de la cuenta no "
        "es una lectura del medidor."
    )
    doc.add_paragraph(
        "Si en el período de lecturas faltan días en la API WES (huecos), se proyecta el "
        "consumo de esos días con el promedio m³/día de los días con dato del mismo período "
        "y se SUMA al total WES. Con menos de 3 días de base se usa el promedio del nodo."
    )
    presentes = {s.node_id for s in sitios}
    faltantes = [n for n in NODOS_CORMUP if n not in presentes]
    if faltantes:
        txt_falt = ", ".join(f"{n} ({get_node_name(n)})" for n in faltantes)
        doc.add_paragraph(f"Sin carpeta de facturaciones en Drive para: {txt_falt}.")

    add_formatted_heading(doc, "2. Ciclo septiembre (el que más importa)", level=1)
    doc.add_paragraph(
        "Boletas con emisión o lectura actual en septiembre. Primero las válidas "
        "(lectura real); al final las de promedio, solo informativas."
    )
    headers_sep = [
        "Establecimiento",
        "Período",
        "Emisión",
        "Válida",
        "m³ cuenta",
        "m³ WES med.",
        "Huecos",
        "m³ proy.",
        "m³ WES+proy",
        "Dif m³",
        "%",
        "Observación",
    ]
    sep = _filas_septiembre(sitios)
    sep.sort(key=lambda f: (not f.valido, f.node_id))
    rows_sep: List[List[str]] = []
    hi_sep: list[int] = []
    tot_c = tot_med = tot_pr = 0
    n_val = 0
    for i, f in enumerate(sep, start=1):
        dif_txt = format_number_chilean(f.diff_m3, 1) if f.valido else "—"
        pct_txt = format_number_chilean(f.pct or 0.0, 1) if f.valido else "—"
        rows_sep.append(
            [
                f.node_name,
                f.periodo_txt,
                f.emision.strftime("%d-%m-%Y"),
                "Sí" if f.valido else "No",
                format_number_chilean(f.m3_cuenta, 0),
                format_number_chilean(f.m3_wes_medido, 1),
                str(f.dias_hueco),
                format_number_chilean(f.m3_proyeccion, 1),
                format_number_chilean(f.m3_wes, 1),
                dif_txt,
                pct_txt,
                f.observacion,
            ]
        )
        if f.valido:
            tot_c += f.m3_cuenta
            tot_med += f.m3_wes_medido
            tot_pr += f.m3_proyeccion
            n_val += 1
            if abs(f.pct or 0) >= UMBRAL_PCT_DESTACAR:
                hi_sep.append(i)
        else:
            hi_sep.append(i)
    tot_w = tot_med + tot_pr
    tot_pct = (100.0 * (tot_c - tot_w) / tot_c) if tot_c else 0.0
    rows_sep.append(
        [
            "TOTAL válidas septiembre",
            "",
            "",
            str(n_val),
            format_number_chilean(tot_c, 0),
            format_number_chilean(tot_med, 1),
            "",
            format_number_chilean(tot_pr, 1),
            format_number_chilean(tot_w, 1),
            format_number_chilean(tot_c - tot_w, 1),
            format_number_chilean(tot_pct, 1),
            "",
        ]
    )
    _add_table_rows(doc, headers_sep, rows_sep, highlight=hi_sep, has_total=True)
    if chart_sep.exists():
        add_picture_with_pagination(doc, str(chart_sep), Inches(9.4), keep_with_next=False)

    add_formatted_heading(doc, "3. Resumen válido por establecimiento (todas las lecturas reales)", level=1)
    headers_r = [
        "Establecimiento",
        "Nodo",
        "N° válidas",
        "m³ cuenta",
        "m³ WES med.",
        "m³ proy.",
        "m³ WES+proy",
        "Dif m³",
        "% dif",
        "No válidas",
    ]
    rows_r: List[List[str]] = []
    highlight: list[int] = []
    tot_c = tot_med = tot_pr = tot_n = tot_e = 0
    for idx, s in enumerate(sitios, start=1):
        validas = [f for f in s.filas if f.valido]
        m3c = sum(f.m3_cuenta for f in validas)
        m3m = sum(f.m3_wes_medido for f in validas)
        m3p = sum(f.m3_proyeccion for f in validas)
        m3w = m3m + m3p
        n_est = sum(1 for f in s.filas if f.estimado)
        pct = (100.0 * (m3c - m3w) / m3c) if m3c else 0.0
        rows_r.append(
            [
                s.node_name,
                s.node_id,
                str(len(validas)),
                format_number_chilean(m3c, 0),
                format_number_chilean(m3m, 1),
                format_number_chilean(m3p, 1),
                format_number_chilean(m3w, 1),
                format_number_chilean(m3c - m3w, 1),
                format_number_chilean(pct, 1),
                str(n_est),
            ]
        )
        if validas and abs(pct) >= UMBRAL_PCT_DESTACAR:
            highlight.append(idx)
        tot_c += m3c
        tot_med += m3m
        tot_pr += m3p
        tot_n += len(validas)
        tot_e += n_est
    tot_w = tot_med + tot_pr
    tot_pct = (100.0 * (tot_c - tot_w) / tot_c) if tot_c else 0.0
    rows_r.append(
        [
            "TOTAL lecturas reales",
            COMPANY_ID,
            str(tot_n),
            format_number_chilean(tot_c, 0),
            format_number_chilean(tot_med, 1),
            format_number_chilean(tot_pr, 1),
            format_number_chilean(tot_w, 1),
            format_number_chilean(tot_c - tot_w, 1),
            format_number_chilean(tot_pct, 1),
            str(tot_e),
        ]
    )
    _add_table_rows(doc, headers_r, rows_r, highlight=highlight, has_total=True)

    if chart_png.exists():
        add_picture_with_pagination(doc, str(chart_png), Inches(9.4), keep_with_next=False)

    add_formatted_heading(doc, "4. Tablas por establecimiento", level=1)
    headers_d = [
        "Período (lecturas)",
        "Emisión",
        "N° factura",
        "Válida",
        "m³ cuenta",
        "m³ WES med.",
        "Huecos",
        "m³ proy.",
        "m³ WES+proy",
        "Dif m³",
        "%",
        "Observación",
    ]
    for s in sitios:
        add_formatted_heading(
            doc,
            f"{s.node_name} ({s.node_id})",
            level=2,
            page_break_before=True,
        )
        meta = doc.add_paragraph()
        meta.add_run("Carpeta Drive: ").bold = True
        meta.add_run(s.carpeta)
        cuentas = sorted({f.cuenta for f in s.filas if f.cuenta})
        medidores = sorted({f.medidor for f in s.filas if f.medidor})
        if cuentas:
            pcu = doc.add_paragraph()
            pcu.add_run("Cuenta(s): ").bold = True
            pcu.add_run(", ".join(cuentas))
        if medidores:
            pme = doc.add_paragraph()
            pme.add_run("Medidor(es): ").bold = True
            pme.add_run(", ".join(medidores))
        if s.errores_pdf:
            doc.add_paragraph("PDF no parseados: " + "; ".join(s.errores_pdf))
        if not s.filas:
            doc.add_paragraph("Sin períodos extraídos.")
            continue
        rows: List[List[str]] = []
        hi: list[int] = []
        for i, f in enumerate(s.filas, start=1):
            dif_txt = format_number_chilean(f.diff_m3, 1) if f.valido else "—"
            pct_txt = format_number_chilean(f.pct or 0.0, 1) if f.valido else "—"
            rows.append(
                [
                    f.periodo_txt,
                    f.emision.strftime("%d-%m-%Y"),
                    f.boleta,
                    "Sí" if f.valido else "No",
                    format_number_chilean(f.m3_cuenta, 0),
                    format_number_chilean(f.m3_wes_medido, 1),
                    str(f.dias_hueco),
                    format_number_chilean(f.m3_proyeccion, 1),
                    format_number_chilean(f.m3_wes, 1),
                    dif_txt,
                    pct_txt,
                    f.observacion,
                ]
            )
            if f.valido and f.pct is not None and abs(f.pct) >= UMBRAL_PCT_DESTACAR:
                hi.append(i)
            elif not f.valido:
                hi.append(i)
        validas = [f for f in s.filas if f.valido]
        m3c = sum(f.m3_cuenta for f in validas)
        m3m = sum(f.m3_wes_medido for f in validas)
        m3p = sum(f.m3_proyeccion for f in validas)
        m3w = m3m + m3p
        pct_t = (100.0 * (m3c - m3w) / m3c) if m3c else 0.0
        rows.append(
            [
                "TOTAL válidas",
                "",
                str(len(validas)),
                "",
                format_number_chilean(m3c, 0),
                format_number_chilean(m3m, 1),
                "",
                format_number_chilean(m3p, 1),
                format_number_chilean(m3w, 1),
                format_number_chilean(m3c - m3w, 1),
                format_number_chilean(pct_t, 1),
                "",
            ]
        )
        _add_table_rows(doc, headers_d, rows, highlight=hi, has_total=True)

    add_formatted_heading(doc, "5. Notas", level=1, page_break_before=True)
    doc.add_paragraph(
        "Fuente de medidas WES: GET /wes/api/acl-node/v1/nodes/measures/dates "
        "(consumo diario totalM3). El período de cada fila es el intervalo de lecturas "
        "de la boleta, no el mes calendario de emisión."
    )
    doc.add_paragraph(
        "Las filas en rojo / «NO VÁLIDO» son promedio o sin lectura real: no entran al "
        "total ni al % de diferencia. El m³ WES de esas filas se muestra solo como referencia."
    )
    doc.add_paragraph(
        "Proyección de huecos: promedio diario WES del mismo período × días sin dato, "
        "sumado al m³ medido. Un día con 0 m³ registrado no es hueco (sí hubo telemetría)."
    )
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)


def convertir_a_pdf(docx_path: Path) -> Optional[Path]:
    pdf_path = docx_path.with_suffix(".pdf")
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("[WARN] No hay LibreOffice/soffice; se omite PDF.", flush=True)
        return None
    subprocess.run(
        [
            soffice,
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            "--convert-to",
            "pdf",
            "--outdir",
            str(docx_path.parent),
            str(docx_path),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    return pdf_path if pdf_path.exists() else None


def generar(skip_download: bool = False) -> Tuple[Path, Optional[Path], Path]:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    generado = datetime.now(timezone.utc).astimezone()
    ts = generado.strftime("%Y%m%d_%H%M")

    if skip_download and any(PDF_CACHE.glob("*/*.pdf")):
        locales: Dict[str, Path] = {p.name: p for p in PDF_CACHE.iterdir() if p.is_dir()}
        print(f"[INFO] Reuso cache local {PDF_CACHE}", flush=True)
    else:
        locales = descargar_facturaciones(PDF_CACHE)

    sitios = _cargar_sitios(locales)
    if not sitios:
        raise SystemExit("No se mapearon establecimientos a nodos WES.")
    cruzar_wes(sitios)

    stem = f"Comparacion_cuentas_vs_WES_CORMUP_Penalolen_{ts}"
    out_xlsx = OUT_DIR / f"{stem}.xlsx"
    out_docx = OUT_DIR / f"{stem}.docx"
    chart_png = OUT_DIR / f"{stem}_barras.png"
    chart_sep = OUT_DIR / f"{stem}_septiembre.png"

    sep_validas = [f for f in _filas_septiembre(sitios) if f.valido]
    _grafico_barras(
        sep_validas,
        chart_sep,
        titulo="Septiembre — lecturas reales: m³ cuenta vs WES + proyección de huecos",
    )
    # Un punto por establecimiento (suma de boletas válidas)
    resumen_filas: List[FilaComparacion] = []
    for s in sitios:
        validas = [f for f in s.filas if f.valido]
        if not validas:
            continue
        acc = FilaComparacion(
            establecimiento=s.node_name,
            carpeta=s.carpeta,
            node_id=s.node_id,
            node_name=s.node_name,
            pdf_name="",
            boleta="",
            cuenta="",
            medidor="",
            emision=validas[0].emision,
            lectura_anterior=validas[0].lectura_anterior,
            lectura_actual=validas[-1].lectura_actual,
            m3_cuenta=sum(f.m3_cuenta for f in validas),
            m3_wes_medido=sum(f.m3_wes_medido for f in validas),
            m3_proyeccion=sum(f.m3_proyeccion for f in validas),
            dias_wes=sum(f.dias_wes for f in validas),
            dias_esperados=sum(f.dias_esperados for f in validas),
            dias_hueco=sum(f.dias_hueco for f in validas),
            dias_periodo=sum(f.dias_periodo for f in validas),
            huecos_txt="",
            clave_facturacion="Consumo real",
            clave_lectura="LECTURA NORMAL",
            estimado=False,
        )
        resumen_filas.append(acc)
    _grafico_barras(
        resumen_filas,
        chart_png,
        titulo="CORMUP Peñalolén — solo lecturas reales: m³ cuenta vs WES + proyección",
    )
    _write_excel(sitios, out_xlsx, generado)
    _write_word(sitios, out_docx, chart_sep, chart_png, generado)
    out_pdf = convertir_a_pdf(out_docx)

    print(f"[OK] Excel: {out_xlsx}", flush=True)
    print(f"[OK] Word:  {out_docx}", flush=True)
    if out_pdf:
        print(f"[OK] PDF:   {out_pdf}", flush=True)
    return out_docx, out_pdf, out_xlsx


def main() -> int:
    if sys.platform == "win32":
        for s in (sys.stdout, sys.stderr):
            try:
                s.reconfigure(encoding="utf-8", line_buffering=True)
            except Exception:
                pass
    ap = argparse.ArgumentParser(description="Comparativo facturaciones Peñalolén vs API WES")
    ap.add_argument("--skip-download", action="store_true", help="Usa PDFs ya descargados en _pdfs/")
    args = ap.parse_args()
    generar(skip_download=args.skip_download)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
