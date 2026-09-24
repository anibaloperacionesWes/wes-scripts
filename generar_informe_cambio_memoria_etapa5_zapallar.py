"""
Informe corto — cambio de memoria placa / Etapa N°5 Fundo Zapallar.

Validación al estilo Informe Interno Matriz ESVAL:
  - Inicio: lectura 22/09/2026 → consumo WES con hora 14 completa.
  - Fin: foto 23/09/2026 → consumo WES hasta las 17:00.
  - Tablas resumen (lectura / consumo / % error); sin tabla de huecos horarios.
  - Fotos compactas (~5,6 cm) solo relojería, lado a lado.

Uso:
  python generar_informe_cambio_memoria_etapa5_zapallar.py
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from zoneinfo import ZoneInfo

import requests
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "reports" / "Fundo_Zapallar" / "Informes_Tecnicos"
FOTOS_DIR = OUT_DIR / "fotos_etapa5"

try:
    from generar_reporte_word import CHILE_TZ, acl_node_base_url
except Exception:
    CHILE_TZ = ZoneInfo("America/Santiago")

    def acl_node_base_url() -> str:
        return "http://104.248.53.141:7003/wes/api/acl-node/v1"

COMPANY = "Fundo Zapallar"
COMPANY_ID = "000027"
NODE_ID = "000027-03"
NODE_NAME = "Etapa N°5"
FECHA_INTERVENCION = date(2026, 9, 22)
DIAMETRO = "DN90 fierro dúctil"
CAUDAL_MAX_REF_M3H = 60.0
COLOR_TITULO = RGBColor(0x1F, 0x47, 0x88)
COLOR_META = RGBColor(0x59, 0x59, 0x59)

# Lecturas mecánicas (fotos terreno)
LECTURA_AYER_M3 = 5144.0
LECTURA_AYER_DT = datetime(2026, 9, 22, 14, 30, tzinfo=CHILE_TZ)
LECTURA_HOY_M3 = 5177.0
LECTURA_HOY_DT = datetime(2026, 9, 23, 16, 54, tzinfo=CHILE_TZ)
# Suma WES hasta las 17:00 (incluye hora 16)
WES_HASTA_DT = datetime(2026, 9, 23, 17, 0, tzinfo=CHILE_TZ)

# Hueco 22/09 leído directo de la placa (m³/h).
# Extracción listó H:15:00 dos veces (0,60 y 2,90): se interpreta 0,60 como hora 14
# (coherente con lectura a las 14:30) y 2,90 como hora 15.
PLACA_HUECO_22: Dict[int, float] = {
    14: 0.60,
    15: 2.90,
    16: 5.50,
    17: 2.60,
    18: 0.90,
    19: 0.00,
    20: 0.00,
    21: 0.00,
    22: 0.00,
    23: 0.00,
}

MEDIDOR_MARCA = "Sensus"
MEDIDOR_MODELO = "MeiStream Plus 100"
MEDIDOR_SERIE = "8 SEN01 2370 9030"
MEDIDOR_Q3_M3H = 100.0
HRI_MODELO = "HRI-Mei B4 500 ms"
HRI_SERIE = "31730485"
HRI_PULSO_DN40_125_L = 100

FOTO_AYER = FOTOS_DIR / "lectura_20260922_1430_sensus_5144.png"
FOTO_HOY = FOTOS_DIR / "lectura_20260923_1654_sensus_5177.png"
FOTO_AYER_RELOJ = FOTOS_DIR / "lectura_20260922_1430_sensus_5144_reloj.png"
FOTO_HOY_RELOJ = FOTOS_DIR / "lectura_20260923_1654_sensus_5177_reloj.png"
# Cara del medidor (relojería / totalizador visible, sin brida)
CROP_AYER = (210, 340, 770, 840)
CROP_HOY = (200, 300, 780, 780)
FOTO_ANCHO = Inches(2.21)  # mismo tamaño que informe ESVAL Zapallar


@dataclass
class Validacion:
    lectura_ayer: float
    lectura_hoy: float
    delta_mecanico: float
    wes_m3: float
    detalle: List[Tuple[str, float, str]]
    hueco_horas_api: List[str]
    diferencia_m3: float
    diferencia_pct: float
    estado: str


def _set_run_font(run, *, bold: bool = False, size: int = 11, color: RGBColor | None = None) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color is not None:
        run.font.color.rgb = color


def _add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Title")
    p.clear()
    r = p.add_run(text)
    _set_run_font(r, bold=True, size=18, color=COLOR_TITULO)


def _add_meta(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_run_font(r, size=10, color=COLOR_META)


def _add_h(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = COLOR_TITULO
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def _add_body(doc: Document, text: str, *, size: int = 11) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    _set_run_font(r, size=size)


def _add_line(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_run_font(r, size=11, color=RGBColor(0, 0, 0))


def _set_cell_text(cell, text: str, *, bold: bool = False, size: int = 9, color: RGBColor | None = None) -> None:
    p = cell.paragraphs[0]
    # limpiar runs previos (evita run vacío sin formato)
    for child in list(p._element):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            p._element.remove(child)
    r = p.add_run(text)
    _set_run_font(r, bold=bold, size=size, color=color)


def _set_table_borders(table) -> None:
    """Bordes finos negros (estilo informe ESVAL Zapallar)."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    # quitar estilo Table Grid si aporta bordes gruesos / sombreado
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is not None:
        tblPr.remove(borders)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def _add_tabla_simple(doc: Document, filas: List[Tuple[str, ...]], *, header_blue: bool = True) -> None:
    cols = len(filas[0])
    t = doc.add_table(rows=len(filas), cols=cols)
    _set_table_borders(t)
    for i, row in enumerate(filas):
        for j, val in enumerate(row):
            if i == 0 and header_blue:
                _set_cell_text(t.rows[i].cells[j], val, bold=True, size=10, color=COLOR_TITULO)
            else:
                _set_cell_text(t.rows[i].cells[j], val, bold=False, size=10)


def _crop_relojeria(src: Path, dst: Path, box: Tuple[int, int, int, int]) -> Path:
    """Recorta la cara/relojería del medidor, upscale y nítidez leve."""
    from PIL import ImageEnhance, ImageFilter

    FOTOS_DIR.mkdir(parents=True, exist_ok=True)
    if not src.is_file():
        return dst
    im = Image.open(src).convert("RGB").crop(box)
    target_w = 900
    if im.width < target_w:
        scale = target_w / im.width
        im = im.resize(
            (int(im.width * scale), int(im.height * scale)),
            Image.Resampling.LANCZOS,
        )
    im = im.filter(ImageFilter.UnsharpMask(radius=1.5, percent=120, threshold=2))
    im = ImageEnhance.Contrast(im).enhance(1.12)
    im = ImageEnhance.Sharpness(im).enhance(1.25)
    im.save(dst, optimize=True)
    return dst


def _add_fotos_lado_a_lado(
    doc: Document,
    path_izq: Path,
    caption_izq: str,
    path_der: Path,
    caption_der: str,
) -> None:
    """Fotos compactas en 2 columnas (mismo layout que informe ESVAL)."""
    tbl = doc.add_table(rows=2, cols=2)
    _set_table_borders(tbl)
    for col_i, (path, caption) in enumerate(
        ((path_izq, caption_izq), (path_der, caption_der))
    ):
        cell_img = tbl.rows[0].cells[col_i]
        cell_img.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if path.is_file():
            run = cell_img.paragraphs[0].add_run()
            run.add_picture(str(path), width=FOTO_ANCHO)
        cell_cap = tbl.rows[1].cells[col_i]
        cell_cap.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell_cap.text = ""
        r = cell_cap.paragraphs[0].add_run(caption)
        _set_run_font(r, size=9, color=COLOR_META)


def _add_tabla_validacion_7(
    doc: Document,
    titulo: str,
    analisis: str,
    fecha_ini: str,
    lectura_ini: str,
    fecha_fin: str,
    lectura_fin: str,
    consumo: str,
    horas: str,
) -> None:
    """Tabla 7 columnas estilo Informe Interno Matriz ESVAL (texto azul, sin relleno)."""
    t = doc.add_table(rows=3, cols=7)
    _set_table_borders(t)
    t.rows[0].cells[0].merge(t.rows[0].cells[6])
    _set_cell_text(t.rows[0].cells[0], titulo, bold=True, size=11, color=COLOR_TITULO)
    headers = [
        "ANÁLISIS",
        "FECHA INICIAL",
        "LECTURA (m³)",
        "FECHA FINAL",
        "LECTURA (m³)",
        "CONSUMO m³",
        "HORAS",
    ]
    for j, htxt in enumerate(headers):
        _set_cell_text(t.rows[1].cells[j], htxt, bold=True, size=9, color=COLOR_TITULO)
    vals = [analisis, fecha_ini, lectura_ini, fecha_fin, lectura_fin, consumo, horas]
    for j, v in enumerate(vals):
        _set_cell_text(t.rows[2].cells[j], v, bold=False, size=9)


def _fmt(n: float, dec: int = 1) -> str:
    return f"{n:,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def obtener_api_por_hora() -> Dict[datetime, float]:
    base = acl_node_base_url()
    out: Dict[datetime, float] = {}
    for d in (date(2026, 9, 21), date(2026, 9, 22), date(2026, 9, 23)):
        ds = d.strftime("%d%m%Y")
        url = f"{base}/nodes/{NODE_ID}/dates.measures.csv"
        r = requests.get(url, params=[("start", ds), ("end", ds)], timeout=45)
        r.raise_for_status()
        lines = [ln for ln in r.text.strip().splitlines() if ln.strip()]
        start = 1 if lines and "TIME" in lines[0].upper() else 0
        for line in lines[start:]:
            parts = line.split(",")
            if len(parts) < 2:
                continue
            raw = parts[0].strip().replace("Z", "+00:00")
            dt = datetime.fromisoformat(raw)
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            out[dt.astimezone(CHILE_TZ)] = float(parts[1])
    return out


def valor_hora(dt: datetime, api: Dict[datetime, float]) -> Tuple[float, str]:
    """Prioriza valores de placa en el hueco del 22/09; si no, API."""
    if dt.date() == date(2026, 9, 22) and dt.hour in PLACA_HUECO_22:
        return PLACA_HUECO_22[dt.hour], "placa (hueco)"
    if dt in api:
        return api[dt], "API WES"
    return 0.0, "sin dato (0)"


def calcular_validacion(lectura_ayer: float, lectura_hoy: float) -> Validacion:
    api = obtener_api_por_hora()
    detalle: List[Tuple[str, float, str]] = []
    total = 0.0
    hueco_api: List[str] = []

    inicio_hora = LECTURA_AYER_DT.replace(minute=0, second=0, microsecond=0)
    v14, src14 = valor_hora(inicio_hora, api)
    detalle.append(
        (
            inicio_hora.strftime("%d/%m/%Y %H:%M"),
            v14,
            f"hora 14 completa; {_fmt(v14, 2)} m³/h [{src14}]",
        )
    )
    total += v14
    if "API" not in src14 and inicio_hora not in api:
        hueco_api.append(inicio_hora.strftime("%d/%m/%Y %H:%M"))

    cur = inicio_hora + timedelta(hours=1)
    while cur < WES_HASTA_DT:
        v, src = valor_hora(cur, api)
        detalle.append((cur.strftime("%d/%m/%Y %H:%M"), v, src))
        total += v
        if src.startswith("placa") or src.startswith("sin"):
            if cur not in api:
                hueco_api.append(cur.strftime("%d/%m/%Y %H:%M"))
        cur += timedelta(hours=1)

    delta = lectura_hoy - lectura_ayer
    dif = delta - total
    # Estilo ESVAL: % Error = |1 − (lectura / serie)| o |1 − (serie / lectura)|
    # Se reporta el valor absoluto respecto al mayor de ambos.
    if delta and total:
        pct = abs(1.0 - (min(delta, total) / max(delta, total))) * 100.0
    else:
        pct = 0.0
    if pct <= 5:
        estado = "Aceptable"
    elif pct <= 15:
        estado = "Aceptable (diferencia moderada)"
    else:
        estado = "Revisar — diferencia relevante"

    return Validacion(
        lectura_ayer=lectura_ayer,
        lectura_hoy=lectura_hoy,
        delta_mecanico=round(delta, 2),
        wes_m3=round(total, 2),
        detalle=detalle,
        hueco_horas_api=hueco_api,
        diferencia_m3=round(dif, 2),
        diferencia_pct=round(pct, 1),
        estado=estado,
    )


def build_doc(lectura_ayer: float, lectura_hoy: float) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ahora = datetime.now(CHILE_TZ)
    stamp = ahora.strftime("%Y%m%d_%H%M")
    out = OUT_DIR / f"Informe_Cambio_Memoria_Etapa5_Zapallar_{stamp}.docx"
    val = calcular_validacion(lectura_ayer, lectura_hoy)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    _add_title(doc, f"Informe interno — Cambio de memoria y validación {NODE_NAME}")
    _add_meta(
        doc,
        f"Empresa: {COMPANY} ({COMPANY_ID})\n"
        f"Punto: {NODE_NAME} ({NODE_ID})\n"
        f"Validación: {LECTURA_AYER_DT.strftime('%d-%m-%Y')} 14:00 → "
        f"{WES_HASTA_DT.strftime('%d-%m-%Y %H:%M')}\n"
        f"Clasificación: Uso interno WES",
    )

    _add_h(doc, "1. Objetivo", 1)
    _add_body(
        doc,
        "Documentar la intervención en terreno sobre la placa de Etapa N°5 ante caudales "
        "anómalos incompatibles con la red DN90, y presentar el cálculo de validación "
        "entre lecturas mecánicas del medidor Sensus y el consumo de la app WES.",
    )

    _add_h(doc, "2. Resumen", 1)
    _add_body(
        doc,
        "Se detectaron pulsos del orden de 92 y 200 m³/h. El sensor inductivo, los pulsos "
        "de revisión y los voltajes estaban correctos; el error provenía de la memoria de "
        "la placa. Se reemplazó la memoria y se normalizó el punto. La validación posterior "
        f"entre lectura Sensus y app WES arroja un error de {_fmt(val.diferencia_pct, 1)} % "
        f"({val.estado.lower()}).",
    )

    _add_h(doc, "3. Actividades realizadas", 1)

    _add_h(doc, "3.1 Revisión en terreno", 2)
    _add_body(
        doc,
        "Se verificó sensor Census / inductivo (OK), pulsos de revisión (OK) y voltajes de "
        "alimentación (correctos). El diagnóstico apuntó a falla de memoria de la placa.",
    )
    _add_line(doc, "Acción: cambio de memoria de la placa; punto reparado/normalizado.")

    _add_h(doc, "3.2 Medidor instalado", 2)
    _add_body(
        doc,
        f"Medidor {MEDIDOR_MARCA} {MEDIDOR_MODELO}, serie {MEDIDOR_SERIE} (2023), "
        f"Q3 = {_fmt(MEDIDOR_Q3_M3H, 0)} m³/h. Módulo {HRI_MODELO} (serie {HRI_SERIE}), "
        f"peso de pulso DN 40–125 = {HRI_PULSO_DN40_125_L} L/pulso.",
    )
    _add_line(doc, f"Lectura inicial: {_fmt(lectura_ayer, 0)} m³ ({LECTURA_AYER_DT.strftime('%d-%m-%Y')}).")
    _add_line(doc, f"Lectura final: {_fmt(lectura_hoy, 0)} m³ ({LECTURA_HOY_DT.strftime('%d-%m-%Y')}).")
    _add_line(doc, f"Δ mecánico: {_fmt(val.delta_mecanico, 0)} m³.")

    _add_h(doc, "4. Criterio hidráulico (DN90)", 1)
    _add_body(
        doc,
        f"La red es {DIAMETRO}. Techo práctico de red ≈ {CAUDAL_MAX_REF_M3H:.0f} m³/h "
        f"(~2,5 m/s). El medidor admite Q3 = {_fmt(MEDIDOR_Q3_M3H, 0)} m³/h, pero la red "
        "no puede entregar de forma realista 92–200 m³/h (error de memoria).",
    )
    _add_tabla_simple(
        doc,
        [
            ("Velocidad de referencia", "Caudal teórico DN90"),
            ("1,5 m/s", "≈ 34 m³/h"),
            ("2,0 m/s", "≈ 46 m³/h"),
            ("2,5 m/s (techo práctico)", "≈ 57–60 m³/h"),
        ],
    )

    _add_h(doc, "5. Cálculo de validación", 1)
    _add_h(
        doc,
        f"5.1 Sensus vs app WES ({LECTURA_AYER_DT.strftime('%d-%m')} 14:00 → "
        f"{WES_HASTA_DT.strftime('%d-%m %H:%M')})",
        2,
    )
    _add_body(
        doc,
        "Validación con lecturas fotográficas del medidor Sensus y el consumo registrado "
        "en la app WES. Ventana: hora 14 completa del 22-09-2026 hasta las 17:00 del 23-09-2026.",
    )

    foto_ayer = _crop_relojeria(FOTO_AYER, FOTO_AYER_RELOJ, CROP_AYER)
    foto_hoy = _crop_relojeria(FOTO_HOY, FOTO_HOY_RELOJ, CROP_HOY)
    if foto_ayer.is_file() or foto_hoy.is_file():
        _add_fotos_lado_a_lado(
            doc,
            foto_ayer,
            f"Sensus — {LECTURA_AYER_DT.strftime('%d-%m-%Y')}",
            foto_hoy,
            f"Sensus — {LECTURA_HOY_DT.strftime('%d-%m-%Y')}",
        )

    inicio_wes = LECTURA_AYER_DT.replace(minute=0, second=0, microsecond=0)
    horas_ventana = (WES_HASTA_DT - inicio_wes).total_seconds() / 3600.0
    _add_tabla_validacion_7(
        doc,
        "Validación Etapa N°5 — medidor Sensus (lectura mecánica)",
        NODE_NAME,
        inicio_wes.strftime("%d-%m-%Y %H:%M"),
        _fmt(val.lectura_ayer, 0),
        WES_HASTA_DT.strftime("%d-%m-%Y %H:%M"),
        _fmt(val.lectura_hoy, 0),
        _fmt(val.delta_mecanico, 2),
        _fmt(horas_ventana, 0),
    )

    doc.add_paragraph()
    t_err = doc.add_table(rows=3, cols=2)
    _set_table_borders(t_err)
    err_rows = [
        ("Total App WES", f"{_fmt(val.wes_m3, 2)} m³"),
        ("Total Lectura Sensus", f"{_fmt(val.delta_mecanico, 2)} m³"),
        ("% Error", f"{_fmt(val.diferencia_pct, 1)} %"),
    ]
    for i, (a, b) in enumerate(err_rows):
        _set_cell_text(t_err.rows[i].cells[0], a, bold=(i == 2), size=10)
        _set_cell_text(t_err.rows[i].cells[1], b, bold=(i == 2), size=10)

    dens = max(val.delta_mecanico, val.wes_m3)
    num = min(val.delta_mecanico, val.wes_m3)
    _add_body(
        doc,
        f"En base a las lecturas, el rango de error entre la lectura del medidor Sensus "
        f"y la app WES es de un {_fmt(val.diferencia_pct, 1)} % "
        f"(1 − {_fmt(num, 2)}/{_fmt(dens, 2)}). {val.estado}.",
    )

    _add_h(doc, "6. Conclusión", 1)
    _add_body(
        doc,
        "Se confirma falla de memoria de placa (sensor Sensus/HRI y voltajes OK). "
        f"Validación post-cambio: lectura mecánica {_fmt(val.delta_mecanico, 2)} m³ vs "
        f"app WES {_fmt(val.wes_m3, 2)} m³ (% error {_fmt(val.diferencia_pct, 1)} %). "
        f"{val.estado}. Seguimiento diario matutino del nodo {NODE_ID} con umbral "
        f"{CAUDAL_MAX_REF_M3H:.0f} m³/h (techo DN90).",
    )

    _add_meta(
        doc,
        f"Documento interno WES · {COMPANY} · generado {ahora.strftime('%d-%m-%Y %H:%M')}.",
    )

    doc.save(out)
    print(
        f"[VALIDACION] Δ mec={val.delta_mecanico} | serie={val.wes_m3} | "
        f"dif={val.diferencia_m3} ({val.diferencia_pct}%) | {val.estado}"
    )
    return out


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--lectura-ayer", type=float, default=LECTURA_AYER_M3)
    parser.add_argument("--lectura-hoy", type=float, default=LECTURA_HOY_M3)
    args = parser.parse_args()
    path = build_doc(args.lectura_ayer, args.lectura_hoy)
    print(f"[OK] Informe generado: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
