# -*- coding: utf-8 -*-
"""
Fichas one-page Parque Arauco — un caso por lámina.

Estética del consolidado (navy / gold / logo / cards). Cada slide: estado,
tres cifras y el gráfico hora a hora.

  python3 generar_fichas_onepage_pa.py
  python3 generar_fichas_onepage_pa.py --hasta 08/09/2026 --skip-refresh
"""

from __future__ import annotations

import argparse
import json
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Tuple

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptInches

from generar_consolidado_pa_ppt import JSON_DAILY, JSON_HOURS, JSON_HOURS_DIA, OUT_DIR, cargar_diario
from generar_informe_ahorro_pa_reunion import (
    AEB_PROP_DESDE,
    FALABELLA_PROP_DESDE,
    H_NOCHE_FIN,
    MATRIZ_MAQ,
    SI500,
    _perfil_hora,
    _serie_ventana,
    _stats_logrado,
    _stats_propuesta_noche,
    _stats_sur,
    chart_perfil_horario,
    clp,
    fn,
)
from generar_ppt_recorrido_ejecutivo_pa import (
    BAZAR,
    CTRL_NORTE,
    CTRL_SI500,
    DL_KENNEDY,
    FALABELLA,
    FONDO,
    GOLD,
    GRAY,
    LIGHT,
    LOGO,
    MAQ_ALZA,
    MATRIZ_AEB,
    NAVY,
    SUR_REPARACION,
    TARIFA_CLP_M3,
    TEAL,
    UMBRAL_FALABELLA_DIA,
    UMBRAL_MAQ_DIA,
    UMBRAL_MATRIZ_AEB_DIA,
    UMBRAL_PAK_BAZAR_DIA,
    UMBRAL_PAK_KEN_DIA,
    UMBRAL_SI500_DIA,
    WHITE,
    _caja,
    _header_bar,
    _tb,
)

CHARTS = OUT_DIR / "charts_fichas_onepage"


def _kpi(sl, x: float, y: float, w: float, etq: str, valor: str, nota: str, line) -> None:
    _caja(sl, x, y, w, 1.08, fill=WHITE, line=line)
    _tb(sl, x + 0.12, y + 0.08, w - 0.22, 0.22, [(etq, 11, True, GRAY)])
    _tb(sl, x + 0.12, y + 0.32, w - 0.22, 0.40, [(valor, 22, True, NAVY)])
    _tb(sl, x + 0.12, y + 0.74, w - 0.22, 0.26, [(nota, 11, False, GRAY)])


def _meta_cero(hoy: List[float]) -> List[float]:
    """Proyección: 01:00–05:00 a cero; 00:00–00:30 queda (no es noche)."""
    out = list(hoy)
    for i in range(1, min(6, len(out))):
        out[i] = 0.0
    return out


def _slide_ficha(
    prs,
    *,
    titulo: str,
    sub: str,
    estado: str,
    kpis: List[Tuple[str, str, str]],
    chart: Path,
    pie: str,
    logrado: bool,
) -> None:
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    _header_bar(sl, prs, titulo, sub)
    line = TEAL if logrado else GOLD
    fill = (232, 245, 233) if logrado else (255, 249, 235)
    _caja(sl, 0.22, 1.12, 2.35, 0.36, fill=fill, line=line)
    _tb(
        sl,
        0.30,
        1.16,
        2.20,
        0.28,
        [(estado, 12, True, TEAL if logrado else GOLD)],
        align=PP_ALIGN.CENTER,
    )
    gap = 0.14
    w = (12.88 - 2 * gap) / 3.0
    x = 0.22
    for etq, valor, nota in kpis:
        _kpi(sl, x, 1.58, w, etq, valor, nota, line)
        x += w + gap
    if chart.is_file():
        sl.shapes.add_picture(str(chart), PptInches(0.22), PptInches(2.82), width=PptInches(12.88))
    _tb(sl, 0.28, 7.12, 12.7, 0.28, [(pie, 11, False, GRAY)])


def _portada(prs, hasta: date) -> None:
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    if FONDO.is_file():
        pic = sl.shapes.add_picture(str(FONDO), 0, 0, width=prs.slide_width, height=prs.slide_height)
        spTree = sl.shapes._spTree
        spTree.remove(pic.element)
        spTree.insert(2, pic.element)
    from pptx.enum.shapes import MSO_SHAPE
    from generar_ppt_recorrido_ejecutivo_pa import _rgb

    veil = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, PptInches(3.40), prs.slide_width, PptInches(4.10))
    veil.fill.solid()
    veil.fill.fore_color.rgb = _rgb(NAVY)
    veil.line.fill.background()
    if LOGO.is_file():
        sl.shapes.add_picture(str(LOGO), PptInches(11.70), PptInches(3.58), width=PptInches(1.35))
    _tb(sl, 0.55, 3.62, 11.0, 0.36, [("WES  ·  Parque Arauco", 16, True, GOLD)])
    _tb(sl, 0.55, 4.05, 12.2, 0.70, [("Fichas one-page  ·  un caso por lámina", 28, True, WHITE)])
    _tb(
        sl,
        0.55,
        4.85,
        12.2,
        1.10,
        [
            (
                "Ya operativo: Estanque Sur, Estanque Norte, San Ignacio 500.  "
                "A copiar: Quilicura, Bazar Gourmet, DL Kennedy, El Bosque 1° piso, Falabella Maipú.",
                15,
                False,
                WHITE,
            ),
            (
                f"Datos al {hasta:%d/%m/%Y}  ·  tarifa ${fn(TARIFA_CLP_M3, 0)}/m³  ·  corte 00:30 → cero 01:00–05:00",
                14,
                False,
                GOLD,
            ),
        ],
    )


def build_ppt(casos: List[Dict[str, Any]], hasta: date) -> Path:
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    _portada(prs, hasta)
    for c in casos:
        _slide_ficha(prs, **c)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"Fichas_onepage_PA_{hasta.strftime('%Y%m%d')}.pptx"
    prs.save(str(path))
    print(f"[OK] PPT {path}", flush=True)
    return path


def main() -> int:
    parser = argparse.ArgumentParser(description="Fichas one-page PA por caso")
    parser.add_argument("--hasta", default="08/09/2026")
    parser.add_argument("--skip-refresh", action="store_true", default=True)
    args = parser.parse_args()
    hasta = datetime.strptime(args.hasta, "%d/%m/%Y").date()
    desde = date(2026, 5, 1)
    CHARTS.mkdir(parents=True, exist_ok=True)

    by = cargar_diario() if JSON_DAILY.is_file() else {}
    by_h = json.loads(JSON_HOURS.read_text(encoding="utf-8")).get("by_h") or {}
    by_h_dia = json.loads(JSON_HOURS_DIA.read_text(encoding="utf-8")).get("by_h") or {}

    sur = _stats_sur((by.get("000025-19") or {}).get("daily") or {}, hasta)
    norte = _stats_logrado(by_h, "000025-01", CTRL_NORTE, hasta, lookback_dias=14)
    bom = _stats_logrado(by_h, SI500, CTRL_SI500, hasta, lookback_dias=7)
    maq = _stats_propuesta_noche(_serie_ventana(by_h, MATRIZ_MAQ, 0, H_NOCHE_FIN), MAQ_ALZA, hasta)
    bazar = _stats_propuesta_noche(_serie_ventana(by_h, BAZAR, 0, H_NOCHE_FIN), date(2026, 7, 1), hasta)
    ken = _stats_propuesta_noche(_serie_ventana(by_h, DL_KENNEDY, 0, H_NOCHE_FIN), date(2026, 7, 1), hasta)
    aeb = _stats_propuesta_noche(_serie_ventana(by_h, MATRIZ_AEB, 0, H_NOCHE_FIN), AEB_PROP_DESDE, hasta)
    fala = _stats_propuesta_noche(
        _serie_ventana(by_h, FALABELLA, 0, H_NOCHE_FIN),
        FALABELLA_PROP_DESDE,
        hasta,
        min_noche=1.0,
    )

    sur_pre = _perfil_hora(by_h_dia, "000025-19", desde, SUR_REPARACION - timedelta(days=1), h_max=24)
    sur_post = _perfil_hora(by_h_dia, "000025-19", SUR_REPARACION + timedelta(days=1), hasta, h_max=24)
    norte_pre = _perfil_hora(
        by_h_dia, "000025-01", CTRL_NORTE - timedelta(days=14), CTRL_NORTE - timedelta(days=1), h_max=24
    )
    norte_post = _perfil_hora(
        by_h_dia, "000025-01", CTRL_NORTE, hasta, h_max=24, noche_cero=True
    )
    bom_pre = _perfil_hora(
        by_h_dia, SI500, CTRL_SI500 - timedelta(days=7), CTRL_SI500 - timedelta(days=1), h_max=24
    )
    bom_post = _perfil_hora(by_h_dia, SI500, CTRL_SI500, hasta, h_max=24, noche_cero=True)

    def _hoy6(nid: str, d0: date, *, min_noche: float = 0.0) -> List[float]:
        return _perfil_hora(by_h, nid, d0, hasta, h_max=6)

    maq_h = _hoy6(MATRIZ_MAQ, MAQ_ALZA)
    bazar_h = _hoy6(BAZAR, date(2026, 7, 1))
    ken_h = _hoy6(DL_KENNEDY, date(2026, 7, 1))
    aeb_h = _hoy6(MATRIZ_AEB, AEB_PROP_DESDE)
    fala_h = _perfil_hora(
        by_h, FALABELLA, FALABELLA_PROP_DESDE, hasta, h_max=6, min_noche=1.0
    )

    p_sur = CHARTS / "sur.png"
    chart_perfil_horario(
        p_sur,
        "m³/hora (mediana)",
        sur_pre,
        sur_post,
        "Antes 10/06",
        "Después 11/06",
        etiquetar=list(range(7)) + [10, 12, 15, 18],
        figsize=(10.8, 3.15),
    )
    p_norte = CHARTS / "norte.png"
    chart_perfil_horario(
        p_norte,
        "m³/hora (mediana)",
        norte_pre,
        norte_post,
        "Antes 05/08",
        "Con control (noche en cero)",
        corte_00_30=True,
        figsize=(10.8, 3.15),
    )
    p_bom = CHARTS / "si500.png"
    chart_perfil_horario(
        p_bom,
        "m³/hora (mediana)",
        bom_pre,
        bom_post,
        "Antes 17/07",
        "Con control (noche en cero)",
        corte_00_30=True,
        figsize=(10.8, 3.15),
    )
    p_maq = CHARTS / "maq.png"
    chart_perfil_horario(
        p_maq,
        "Noche 00:00–06:00  ·  meta a cero desde 00:30",
        maq_h,
        _meta_cero(maq_h),
        "Hoy",
        "Si se aprueba",
        corte_00_30=True,
        figsize=(10.8, 3.15),
    )
    p_bazar = CHARTS / "bazar.png"
    chart_perfil_horario(
        p_bazar,
        "Noche 00:00–06:00  ·  meta a cero desde 00:30",
        bazar_h,
        _meta_cero(bazar_h),
        "Hoy",
        "Si se aprueba",
        corte_00_30=True,
        figsize=(10.8, 3.15),
    )
    p_ken = CHARTS / "dl_kennedy.png"
    chart_perfil_horario(
        p_ken,
        "Noche 00:00–06:00  ·  meta a cero desde 00:30",
        ken_h,
        _meta_cero(ken_h),
        "Hoy",
        "Si se aprueba",
        corte_00_30=True,
        figsize=(10.8, 3.15),
    )
    p_aeb = CHARTS / "aeb.png"
    chart_perfil_horario(
        p_aeb,
        "Noche 00:00–06:00  ·  meta a cero desde 00:30",
        aeb_h,
        _meta_cero(aeb_h),
        "Hoy",
        "Si se aprueba",
        corte_00_30=True,
        figsize=(10.8, 3.15),
    )
    p_fala = CHARTS / "falabella.png"
    chart_perfil_horario(
        p_fala,
        "Noche 00:00–06:00  ·  meta a cero desde 00:30",
        fala_h,
        _meta_cero(fala_h),
        "Hoy",
        "Si se aprueba",
        corte_00_30=True,
        figsize=(10.8, 3.15),
    )

    casos = [
        {
            "titulo": "MAE  ·  Estanque Sur",
            "sub": f"Presostatos 10/06  ·  {SUR_REPARACION:%d/%m/%Y}  ·  000025-19  ·  {hasta:%d/%m/%Y}",
            "estado": "YA OPERATIVO",
            "kpis": [
                ("Día (mediana)", f"{fn(sur['pre'], 0)} → {fn(sur['post'], 0)}", "m³/día"),
                ("Ahorro mensual", fn(sur["ahorro_mes"], 0), "m³/mes"),
                ("Ahorro $", clp(sur["ahorro_mes"]), f"tarifa ${fn(TARIFA_CLP_M3, 0)}/m³"),
            ],
            "chart": p_sur,
            "pie": "No es on/off: se reparó la fuga. De madrugada el caudal casi desaparece; de día baja a menos de la mitad.",
            "logrado": True,
        },
        {
            "titulo": "MAE  ·  Estanque Norte",
            "sub": f"On/off desde el 05/08  ·  corte 00:30  ·  000025-01  ·  {hasta:%d/%m/%Y}",
            "estado": "YA OPERATIVO",
            "kpis": [
                ("Noche desde 00:30", f"{fn(norte['pre'], 1)} → 0", "m³/noche"),
                ("Ahorro mensual", fn(norte["ahorro_mes"], 0), "m³/mes"),
                ("Ahorro $", clp(norte["ahorro_mes"]), "01:00–05:00 en cero"),
            ],
            "chart": p_norte,
            "pie": "Dorado = noches con 01:00–05:00 en cero. Lo de 00:00–00:30 no es noche y no se resta del $.",
            "logrado": True,
        },
        {
            "titulo": "BOM  ·  San Ignacio 500",
            "sub": f"On/off desde el 17/07  ·  corte 00:30  ·  000025-18  ·  {hasta:%d/%m/%Y}",
            "estado": "YA OPERATIVO",
            "kpis": [
                ("Noche desde 00:30", f"{fn(bom['pre'], 1)} → 0", "m³/noche"),
                ("Ahorro mensual", fn(bom["ahorro_mes"], 0), "m³/mes"),
                ("Ahorro $", clp(bom["ahorro_mes"]), f"umbral 24 h {fn(UMBRAL_SI500_DIA, 0)} m³/día"),
            ],
            "chart": p_bom,
            "pie": "Antes ~5,8 m³/h toda la noche. Con control: 01:00–05:00 en cero. El tramo 00:00–00:30 no se resta.",
            "logrado": True,
        },
        {
            "titulo": "MAQ  ·  Matriz Principal",
            "sub": f"A copiar  ·  mismo corte 00:30  ·  000025-13  ·  umbral {fn(UMBRAL_MAQ_DIA, 0)} m³/día",
            "estado": "A PROPONER",
            "kpis": [
                ("Noche típica", f"{fn(maq['noche'], 1)} → 0", "m³ desde 00:30"),
                ("Proyección", fn(maq["ahorro_mes"], 0), "m³/mes"),
                ("Proyección $", clp(maq["ahorro_mes"]), "si se aprueba el on/off"),
            ],
            "chart": p_maq,
            "pie": "La Matriz concentra Quilicura. Desde el 22/06 el día se duplicó. Meta: esa noche a cero, igual que Norte y SI500.",
            "logrado": False,
        },
        {
            "titulo": "PAK  ·  Bazar Gourmet",
            "sub": f"A copiar  ·  corte 00:30 en Bazar (no las Sandías)  ·  000025-35  ·  umbral {fn(UMBRAL_PAK_BAZAR_DIA, 0)} m³/día",
            "estado": "A PROPONER",
            "kpis": [
                ("Noche típica", f"{fn(bazar['noche'], 1)} → 0", "m³ desde 00:30"),
                ("Proyección", fn(bazar["ahorro_mes"], 0), "m³/mes"),
                ("Proyección $", clp(bazar["ahorro_mes"]), "eslabón de mayor noche"),
            ],
            "chart": p_bazar,
            "pie": "Cortar en Bazar, no en las Sandías: las Sandías alimentan toda la cadena DL. Umbral DL 390 · Bazar 250 · DL Kennedy 20.",
            "logrado": False,
        },
        {
            "titulo": "PAK  ·  DL Kennedy",
            "sub": f"A copiar  ·  corte 00:30  ·  000025-36  ·  umbral {fn(UMBRAL_PAK_KEN_DIA, 0)} m³/día",
            "estado": "A PROPONER",
            "kpis": [
                ("Noche típica", f"{fn(ken['noche'], 1)} → 0", "m³ desde 00:30"),
                ("Proyección", fn(ken["ahorro_mes"], 0), "m³/mes"),
                ("Proyección $", clp(ken["ahorro_mes"]), "ramal que sale de DL"),
            ],
            "chart": p_ken,
            "pie": "Noche chica (0,9 m³) pero entra en la simulación. Mismo corte que Bazar; no se tocan las Sandías.",
            "logrado": False,
        },
        {
            "titulo": "AEB  ·  Matriz 1° piso",
            "sub": f"A copiar  ·  corte 00:30  ·  000025-11  ·  umbral {fn(UMBRAL_MATRIZ_AEB_DIA, 0)} m³/día",
            "estado": "A PROPONER",
            "kpis": [
                ("Noche típica", f"{fn(aeb['noche'], 1)} → 0", "m³ desde 00:30"),
                ("Proyección", fn(aeb["ahorro_mes"], 0), "m³/mes"),
                ("Proyección $", clp(aeb["ahorro_mes"]), "Anillo Plaza no entra en $"),
            ],
            "chart": p_aeb,
            "pie": "Matriz A.A. desactivada 15/05: el caudal de noche está en el primer piso. Anillo Plaza queda con umbral de día, sin $.",
            "logrado": False,
        },
        {
            "titulo": "MAM  ·  Falabella",
            "sub": f"A copiar  ·  corte 00:30  ·  000025-09  ·  umbral {fn(UMBRAL_FALABELLA_DIA, 0)} m³/día  ·  mall sale por Falabella desde 15/08",
            "estado": "A PROPONER",
            "kpis": [
                ("Noche típica", f"{fn(fala['noche'], 1)} → 0", "m³ desde 00:30"),
                ("Proyección", fn(fala["ahorro_mes"], 0), "m³/mes"),
                ("Proyección $", clp(fala["ahorro_mes"]), "solo noches con caudal"),
            ],
            "chart": p_fala,
            "pie": "Placa y Falabella son alimentación alternativa, no dos fallas. Arrow + punto 6 (Pasillo 2) sigue pendiente con Don Miguel.",
            "logrado": False,
        },
    ]
    ppt = build_ppt(casos, hasta)
    print("\n=== SALIDA ===")
    print(ppt)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
