"""
Tabla mensual Fundo Zapallar: medidor fiscal ESVAL vs Matriz WES vs Estanque Inferior.

Sirve para validar que, tras los ajustes del ultrasónico, los tres midan parecido:
la matriz principal llena el estanque inferior, y ambos deberían acercarse al
medidor de turbina de ESVAL.

WES = día completo Chile 00:00–24:00 (sin recorte 12:00), mismo criterio de
la última tabla de facturación.
"""

from __future__ import annotations

import json
import sys
from calendar import monthrange
from datetime import date, datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from generar_reporte_word import format_number_chilean
from generar_tabla_zapallar_matriz_vs_inferior import (
    CICLOS,
    HEADER_FILL,
    NODO_INFERIOR,
    NODO_MATRIZ,
    ROW_ALERTA_FILL,
    TOTAL_FILL,
    WES_BLUE,
    _fmt_m3,
    _pct,
    _set_cell_text,
    _set_landscape,
    load_hourly_range,
)
from generar_tabla_zapallar_wes_dia_completo import index_chile_hours, wes_dia_completo


def _fmt_pct_s(value: float, decimals: int = 1) -> str:
    if value > 0:
        sign = "+"
    elif value < 0:
        sign = "−"
    else:
        sign = ""
    return f"{sign}{format_number_chilean(abs(value), decimals)}%"

MES_ES = {
    1: "Ene",
    2: "Feb",
    3: "Mar",
    4: "Abr",
    5: "May",
    6: "Jun",
    7: "Jul",
    8: "Ago",
    9: "Sep",
    10: "Oct",
    11: "Nov",
    12: "Dic",
}

# Inf ≈ Matriz (la matriz llena el estanque): ±15 %.
LIM_INF_MATRIZ = 15.0
# Inf/Matriz ≈ medidor ESVAL: ±5 % "cuadra", ±15 % "cerca".
LIM_ESVAL_CUADRA = 5.0
LIM_ESVAL_CERCA = 15.0


def _check_inf_matriz(pct: float) -> str:
    return "SI" if abs(pct) <= LIM_INF_MATRIZ else "NO"


def _check_esval(pct: float) -> str:
    a = abs(pct)
    if a <= LIM_ESVAL_CUADRA:
        return "SI"
    if a <= LIM_ESVAL_CERCA:
        return "cerca"
    return "NO"


def _fill_check(label: str) -> str:
    if label == "SI":
        return "C6E0B4"
    if label == "cerca":
        return "FFF2CC"
    return ROW_ALERTA_FILL


def enriquecer(idx_mat, idx_inf) -> list[dict]:
    rows = []
    for c in CICLOS:
        mat, _, _ = wes_dia_completo(idx_mat, c["ini"], c["fin"])
        inf, _, _ = wes_dia_completo(idx_inf, c["ini"], c["fin"])
        esval = float(c["delta"])
        dif_inf_esval = inf - esval
        dif_mat_esval = mat - esval
        dif_inf_mat = inf - mat
        pct_inf_esval = _pct(dif_inf_esval, esval)
        pct_mat_esval = _pct(dif_mat_esval, esval)
        pct_inf_mat = _pct(dif_inf_mat, mat)
        mes = f"{MES_ES[c['fin'].month]}-{str(c['fin'].year)[2:]}"
        rows.append(
            {
                **c,
                "mes": mes,
                "esval": esval,
                "matriz": round(mat, 2),
                "inferior": round(inf, 2),
                "dif_inf_esval": round(dif_inf_esval, 2),
                "pct_inf_esval": round(pct_inf_esval, 2),
                "dif_mat_esval": round(dif_mat_esval, 2),
                "pct_mat_esval": round(pct_mat_esval, 2),
                "dif_inf_mat": round(dif_inf_mat, 2),
                "pct_inf_mat": round(pct_inf_mat, 2),
                "chk_inf_mat": _check_inf_matriz(pct_inf_mat),
                "chk_inf_esval": _check_esval(pct_inf_esval),
                "chk_mat_esval": _check_esval(pct_mat_esval),
                "transicion": "Transición" in c["fase"],
            }
        )
    return rows


def meses_calendario(idx_mat, idx_inf, start: date, end: date) -> list[dict]:
    out = []
    y, m = start.year, start.month
    while date(y, m, 1) <= end:
        last = date(y, m, monthrange(y, m)[1])
        ini = max(start, date(y, m, 1))
        fin = min(end, last)
        mat, _, _ = wes_dia_completo(idx_mat, ini, fin)
        inf, _, _ = wes_dia_completo(idx_inf, ini, fin)
        pct = _pct(inf - mat, mat) if mat else 0.0
        out.append(
            {
                "mes": f"{MES_ES[m]}-{y}",
                "ini": ini,
                "fin": fin,
                "matriz": round(mat, 2),
                "inferior": round(inf, 2),
                "dif": round(inf - mat, 2),
                "pct": round(pct, 2),
                "chk": _check_inf_matriz(pct),
            }
        )
        if m == 12:
            y, m = y + 1, 1
        else:
            m += 1
    return out


HEADERS_CICLO = [
    "Mes (ciclo ESVAL)",
    "Ciclo lectura",
    "Medidor ESVAL (m³)",
    "Matriz WES (m³)",
    "Estanque Inf. (m³)",
    "Inf vs ESVAL",
    "Matriz vs ESVAL",
    "Inf vs Matriz",
    "Inf ≈ Matriz",
    "Inf ≈ ESVAL",
]


def filas_ciclo(rows: list[dict]) -> tuple[list[list[str]], list[str], list[str], dict]:
    body = []
    for r in rows:
        body.append(
            [
                r["mes"],
                r["ciclo"],
                _fmt_m3(r["esval"], 0),
                _fmt_m3(r["matriz"], 1),
                _fmt_m3(r["inferior"], 1),
                f"{_fmt_m3(r['dif_inf_esval'], 1)} ({_fmt_pct_s(r['pct_inf_esval'])})",
                f"{_fmt_m3(r['dif_mat_esval'], 1)} ({_fmt_pct_s(r['pct_mat_esval'])})",
                f"{_fmt_m3(r['dif_inf_mat'], 1)} ({_fmt_pct_s(r['pct_inf_mat'])})",
                r["chk_inf_mat"],
                r["chk_inf_esval"],
            ]
        )
    tot_e = sum(r["esval"] for r in rows)
    tot_m = sum(r["matriz"] for r in rows)
    tot_i = sum(r["inferior"] for r in rows)
    tot = [
        "TOTAL 12 meses",
        "",
        _fmt_m3(tot_e, 0),
        _fmt_m3(tot_m, 1),
        _fmt_m3(tot_i, 1),
        f"{_fmt_m3(tot_i - tot_e, 1)} ({_fmt_pct_s(_pct(tot_i - tot_e, tot_e))})",
        f"{_fmt_m3(tot_m - tot_e, 1)} ({_fmt_pct_s(_pct(tot_m - tot_e, tot_e))})",
        f"{_fmt_m3(tot_i - tot_m, 1)} ({_fmt_pct_s(_pct(tot_i - tot_m, tot_m))})",
        "",
        "",
    ]
    sin = [r for r in rows if not r["transicion"]]
    se, sm, si = sum(r["esval"] for r in sin), sum(r["matriz"] for r in sin), sum(r["inferior"] for r in sin)
    tot_sin = [
        "TOTAL sin transición",
        "10 ciclos",
        _fmt_m3(se, 0),
        _fmt_m3(sm, 1),
        _fmt_m3(si, 1),
        f"{_fmt_m3(si - se, 1)} ({_fmt_pct_s(_pct(si - se, se))})",
        f"{_fmt_m3(sm - se, 1)} ({_fmt_pct_s(_pct(sm - se, se))})",
        f"{_fmt_m3(si - sm, 1)} ({_fmt_pct_s(_pct(si - sm, sm))})",
        "",
        "",
    ]
    summary = {"esval": tot_e, "matriz": tot_m, "inferior": tot_i, "esval_sin_tr": se, "matriz_sin_tr": sm, "inferior_sin_tr": si}
    return body, tot, tot_sin, summary


def crear_grafico_ciclo(rows: list[dict], out_path: Path) -> Path:
    labels = [r["mes"] for r in rows]
    x = np.arange(len(labels))
    w = 0.26
    fig, ax = plt.subplots(figsize=(13.8, 4.8))
    ax.bar(x - w, [r["esval"] for r in rows], w, label="Medidor ESVAL (turbina)", color="#7F7F7F")
    ax.bar(x, [r["matriz"] for r in rows], w, label="Matriz WES (ultrasónico)", color=WES_BLUE)
    ax.bar(x + w, [r["inferior"] for r in rows], w, label="Estanque Inferior WES", color="#70AD47")
    ax.set_ylabel("m³")
    ax.set_title(
        "m³ por ciclo ESVAL: medidor sanitario vs Matriz vs Estanque Inferior",
        fontsize=12,
        fontweight="bold",
        color=WES_BLUE,
        pad=10,
    )
    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontsize=9)
    ax.legend(frameon=False, fontsize=8.5)
    ax.grid(axis="y", linestyle="--", alpha=0.35)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    fig.savefig(out_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return out_path


def crear_grafico_calendario(meses: list[dict], out_path: Path) -> Path:
    labels = [m["mes"] for m in meses]
    x = np.arange(len(labels))
    w = 0.36
    fig, ax = plt.subplots(figsize=(13.8, 4.4))
    ax.bar(x - w / 2, [m["matriz"] for m in meses], w, label="Matriz WES", color=WES_BLUE)
    ax.bar(x + w / 2, [m["inferior"] for m in meses], w, label="Estanque Inferior", color="#70AD47")
    ax.set_ylabel("m³")
    ax.set_title(
        "m³ por mes calendario: Matriz ESVAL vs Estanque Inferior (la matriz llena el estanque)",
        fontsize=12,
        fontweight="bold",
        color=WES_BLUE,
        pad=10,
    )
    ax.set_xticks(x)
    ax.set_xticklabels(labels, rotation=30, ha="right", fontsize=8)
    ax.legend(frameon=False, fontsize=8.5)
    ax.grid(axis="y", linestyle="--", alpha=0.35)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    fig.savefig(out_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return out_path


def crear_pdf(
    rows: list[dict],
    meses: list[dict],
    chart_c: Path,
    chart_m: Path,
    out_path: Path,
) -> Path:
    body, total, tot_sin, _ = filas_ciclo(rows)
    headers = [
        "Mes\n(ciclo ESVAL)",
        "Ciclo lectura",
        "Medidor ESVAL\n(m³)",
        "Matriz WES\n(m³)",
        "Estanque Inf.\n(m³)",
        "Inf vs ESVAL",
        "Matriz vs ESVAL",
        "Inf vs Matriz",
        "Inf ≈\nMatriz",
        "Inf ≈\nESVAL",
    ]
    cell = [headers] + body + [total, tot_sin]
    n_rows, n_cols = len(cell), len(headers)
    fig = plt.figure(figsize=(16.6, 10.8))
    fig.suptitle(
        "Fundo Zapallar — validación mensual: medidor ESVAL, Matriz y Estanque Inferior",
        fontsize=14.5,
        fontweight="bold",
        color=WES_BLUE,
        y=0.985,
        x=0.02,
        ha="left",
    )
    fig.text(
        0.02,
        0.952,
        "WES día completo (00:00–24:00 Chile, ambos días de lectura inclusive). "
        "La matriz principal llena el estanque inferior: deberían medir casi igual (±15 %). "
        "Inf/Matriz vs medidor turbina ESVAL: ±5 % = cuadra, ±15 % = cerca.",
        fontsize=8,
        color="#444444",
    )
    ax = fig.add_axes([0.015, 0.50, 0.97, 0.43])
    ax.axis("off")
    col_w = [0.08, 0.12, 0.10, 0.10, 0.10, 0.13, 0.13, 0.13, 0.055, 0.055]
    table = ax.table(cellText=cell, loc="center", cellLoc="center", colWidths=col_w)
    table.auto_set_font_size(False)
    table.set_fontsize(6.9)
    table.scale(1, 1.48)
    for c in range(n_cols):
        h = table[0, c]
        fill = "#C6E0B4" if c in (4, 7, 8) else WES_BLUE
        h.set_facecolor(fill)
        h.get_text().set_color(WES_BLUE if c in (4, 7, 8) else "white")
        h.get_text().set_fontweight("bold")
        h.get_text().set_fontsize(6.6)
    for r in range(1, n_rows - 2):
        row = rows[r - 1]
        for c in range(n_cols):
            cell_i = table[r, c]
            if c == 8:
                cell_i.set_facecolor("#C6E0B4" if row["chk_inf_mat"] == "SI" else "#F4C7C3")
            elif c == 9:
                cell_i.set_facecolor(
                    "#C6E0B4"
                    if row["chk_inf_esval"] == "SI"
                    else ("#FFF2CC" if row["chk_inf_esval"] == "cerca" else "#F4C7C3")
                )
            elif c in (4, 7):
                cell_i.set_facecolor("#C6E0B4" if row["chk_inf_mat"] == "SI" else "#F8CBAD")
            else:
                cell_i.set_facecolor("#F4C7C3" if row["transicion"] else "white")
            if c in (0, 1):
                cell_i.get_text().set_horizontalalignment("left")
    for r_tot in (n_rows - 2, n_rows - 1):
        for c in range(n_cols):
            t = table[r_tot, c]
            t.set_facecolor("#C5D9F1" if c in (4, 7) else "#D6E3F0")
            t.get_text().set_fontweight("bold")

    axc = fig.add_axes([0.04, 0.27, 0.92, 0.21])
    axc.imshow(plt.imread(str(chart_c)))
    axc.axis("off")
    axm = fig.add_axes([0.04, 0.04, 0.92, 0.21])
    axm.imshow(plt.imread(str(chart_m)))
    axm.axis("off")
    fig.text(
        0.02,
        0.012,
        "Filas rosadas = ciclos de transición (desconexión 25/02/2026 e instalación ultrasónico 11/03/2026): "
        "ahí la matriz no es comparable. El gráfico inferior es mes calendario (solo WES).",
        fontsize=7.3,
        color="#444444",
    )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=170)
    plt.close(fig)
    return out_path


def crear_word(
    rows: list[dict],
    meses: list[dict],
    chart_c: Path,
    chart_m: Path,
    out_path: Path,
) -> Path:
    doc = Document()
    _set_landscape(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Fundo Zapallar — m³ mensuales: medidor ESVAL, Matriz y Estanque Inferior")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(31, 78, 121)

    note = doc.add_paragraph()
    note.add_run(
        "Validación tras los ajustes del estanque inferior y de la matriz ESVAL para que el ultrasónico "
        "cuente como el medidor de turbina de la sanitaria. La matriz principal llena el estanque inferior, "
        "así que ambos WES deberían coincidir (±15 %). Contra el medidor ESVAL: ±5 % cuadra, ±15 % cerca. "
        "WES día completo 00:00–24:00 Chile, sin recorte de 12:00."
    ).font.size = Pt(9)

    body, total, tot_sin, _ = filas_ciclo(rows)
    table = doc.add_table(rows=1 + len(body) + 2, cols=len(HEADERS_CICLO))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(HEADERS_CICLO):
        fill = "C6E0B4" if i in (4, 7, 8) else HEADER_FILL
        color = RGBColor(31, 78, 121) if i in (4, 7, 8) else RGBColor(255, 255, 255)
        _set_cell_text(table.rows[0].cells[i], h, bold=True, size=7.5, color=color, fill=fill)
    for r_i, vals in enumerate(body):
        row = rows[r_i]
        for c_i, val in enumerate(vals):
            if c_i == 8:
                fill = _fill_check(row["chk_inf_mat"])
            elif c_i == 9:
                fill = _fill_check(row["chk_inf_esval"])
            elif c_i in (4, 7):
                fill = "C6E0B4" if row["chk_inf_mat"] == "SI" else "F8CBAD"
            else:
                fill = ROW_ALERTA_FILL if row["transicion"] else "FFFFFF"
            _set_cell_text(
                table.rows[r_i + 1].cells[c_i],
                val,
                size=7.5,
                fill=fill,
                align="left" if c_i in (0, 1) else "center",
            )
    last = table.rows[1 + len(body)]
    for c_i, val in enumerate(total):
        fill = "C5D9F1" if c_i in (4, 7) else TOTAL_FILL
        _set_cell_text(last.cells[c_i], val, bold=True, size=7.5, fill=fill)
    last2 = table.rows[2 + len(body)]
    for c_i, val in enumerate(tot_sin):
        fill = "C5D9F1" if c_i in (4, 7) else TOTAL_FILL
        _set_cell_text(last2.cells[c_i], val, bold=True, size=7.5, fill=fill)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if chart_c.exists():
        p.add_run().add_picture(str(chart_c), width=Inches(10.3))

    doc.add_heading("Mes calendario (solo WES): Matriz vs Estanque Inferior", level=1)
    h2 = ["Mes", "Matriz WES (m³)", "Estanque Inf. (m³)", "Inf − Matriz (m³)", "Dif. %", "Inf ≈ Matriz"]
    t2 = doc.add_table(rows=1 + len(meses), cols=len(h2))
    t2.style = "Table Grid"
    for i, h in enumerate(h2):
        _set_cell_text(t2.rows[0].cells[i], h, bold=True, size=8, color=RGBColor(255, 255, 255), fill=HEADER_FILL)
    for i, m in enumerate(meses):
        vals = [
            m["mes"],
            _fmt_m3(m["matriz"], 1),
            _fmt_m3(m["inferior"], 1),
            _fmt_m3(m["dif"], 1),
            _fmt_pct_s(m["pct"]),
            m["chk"],
        ]
        for c, v in enumerate(vals):
            fill = _fill_check(m["chk"]) if c in (4, 5) else "FFFFFF"
            _set_cell_text(t2.rows[i + 1].cells[c], v, size=8, fill=fill)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if chart_m.exists():
        p2.add_run().add_picture(str(chart_m), width=Inches(10.3))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def main() -> int:
    if sys.platform == "win32":
        for s in (sys.stdout, sys.stderr):
            try:
                s.reconfigure(encoding="utf-8", line_buffering=True)
            except Exception:
                pass

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out_dir = Path("reports/Fundo_Zapallar/ANALISIS") / f"tabla_mensual_esval_inferior_{stamp}"
    out_dir.mkdir(parents=True, exist_ok=True)

    start_all = min(c["ini"] for c in CICLOS)
    end_all = max(c["fin"] for c in CICLOS)

    print("[INFO] Descargando Matriz ESVAL...", flush=True)
    mat_rows = load_hourly_range(NODO_MATRIZ, start_all, end_all)
    print("[INFO] Descargando Estanque Inferior...", flush=True)
    inf_rows = load_hourly_range(NODO_INFERIOR, start_all, end_all)
    if not mat_rows or not inf_rows:
        print("[ERROR] Faltan medidas WES.", flush=True)
        return 1

    idx_mat = index_chile_hours(mat_rows)
    idx_inf = index_chile_hours(inf_rows)
    ciclos = enriquecer(idx_mat, idx_inf)
    meses = meses_calendario(idx_mat, idx_inf, date(start_all.year, start_all.month, 1), end_all)
    _, _, _, summary = filas_ciclo(ciclos)

    json_path = out_dir / "mensual_esval_matriz_inferior.json"
    payload = {
        "generado": datetime.now().isoformat(timespec="seconds"),
        "metodo": "WES día completo Chile [ini 00:00, fin 24:00] inclusive",
        "limite_inf_matriz_pct": LIM_INF_MATRIZ,
        "limite_esval_cuadra_pct": LIM_ESVAL_CUADRA,
        "ciclos": [
            {k: (v.isoformat() if isinstance(v, date) else v) for k, v in r.items()}
            for r in ciclos
        ],
        "meses_calendario": [
            {k: (v.isoformat() if isinstance(v, date) else v) for k, v in m.items()}
            for m in meses
        ],
        "totales": summary,
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    chart_c = crear_grafico_ciclo(ciclos, out_dir / "chart_ciclo_esval_matriz_inf.png")
    chart_m = crear_grafico_calendario(meses, out_dir / "chart_mes_matriz_inf.png")
    docx_path = out_dir / "Tabla_mensual_ESVAL_Matriz_Estanque_Inferior.docx"
    crear_word(ciclos, meses, chart_c, chart_m, docx_path)
    pdf_path = out_dir / "Tabla_mensual_ESVAL_Matriz_Estanque_Inferior.pdf"
    crear_pdf(ciclos, meses, chart_c, chart_m, pdf_path)

    print(f"[OK] JSON {json_path}")
    print(f"[OK] Word {docx_path}")
    print(f"[OK] PDF  {pdf_path}")
    print(
        f"[OK] Totales ESVAL={summary['esval']:.0f}  "
        f"Matriz={summary['matriz']:.1f}  Inf={summary['inferior']:.1f}"
    )
    for r in ciclos:
        print(
            f"  {r['mes']} {r['ciclo']}: ESVAL={r['esval']:.0f}  "
            f"Matriz={r['matriz']:.1f}  Inf={r['inferior']:.1f}  "
            f"Inf≈Matriz {r['chk_inf_mat']} ({r['pct_inf_mat']:+.1f}%)  "
            f"Inf≈ESVAL {r['chk_inf_esval']} ({r['pct_inf_esval']:+.1f}%)"
        )
    print("Mes calendario Matriz vs Inf:")
    for m in meses:
        print(f"  {m['mes']}: M={m['matriz']:.1f} I={m['inferior']:.1f} {m['chk']} ({m['pct']:+.1f}%)")
    print(f"OUT_DIR={out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
