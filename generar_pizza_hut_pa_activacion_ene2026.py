"""Comparación Parque Arauco / Pizza Hut: activación 22-ene-2026.

Primera madrugada en cero = 23-ene-2026.
Se compara con el día anterior al 22 con mayor consumo nocturno (18-ene-2026).
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.shared import Cm, Inches, Pt, RGBColor
from matplotlib.backends.backend_pdf import PdfPages

from analizar_pizza_hut_cero_nocturno import _fetch_dia
from wes_estilo_graficos_app import COLOR_BARRA_WES, COLOR_CONSUMO, COLOR_NOCHE, guardar_grafico_horario_24h_app

COLOR_WES = "1F4E79"
DIA_ACTIVACION = date(2026, 1, 22)
DIA_PRIMERO_CERO = date(2026, 1, 23)
DIA_PEOR = date(2026, 1, 18)  # máximo 00–06 entre el 1 y el 21 de enero


def _nocturno(h: dict) -> float:
    return sum(float(h.get(i, 0.0)) for i in range(7))


def _diurno(h: dict) -> float:
    return sum(float(h.get(i, 0.0)) for i in range(7, 24))


def _perfil(ax, horas: dict, titulo: str, ymax: float) -> None:
    x = np.arange(24, dtype=float)
    y = np.array([float(horas.get(h, 0.0)) for h in range(24)])
    colors = [COLOR_NOCHE if h <= 6 else COLOR_BARRA_WES for h in range(24)]
    ax.axvspan(-0.5, 6.5, alpha=0.12, color=COLOR_NOCHE, zorder=0)
    ax.bar(x, y, width=0.78, color=colors, edgecolor="white", linewidth=0.5, zorder=2)
    ax.plot(x, y, color=COLOR_CONSUMO, linewidth=1.2, marker="o", markersize=3, zorder=4)
    ax.set_xlim(-0.5, 23.5)
    ax.set_ylim(0, ymax)
    ax.set_xticks(x)
    ax.set_xticklabels([f"{int(h):02d}" for h in x], fontsize=7)
    ax.set_ylabel("m³/h")
    ax.set_title(titulo, fontsize=11, fontweight="bold", color="#1F4E79")
    ax.grid(axis="y", alpha=0.35)
    ax.set_axisbelow(True)


def _shade(cell, hex_color: str) -> None:
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'w:val="clear" w:fill="{hex_color}"/>'
        )
    )


def _cell(cell, text, *, fill=None, bold=False, size=10, color=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Calibri"
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if fill:
        _shade(cell, fill)


def main() -> int:
    _, h_peor = _fetch_dia(DIA_PEOR)
    _, h_cero = _fetch_dia(DIA_PRIMERO_CERO)
    _, h_act = _fetch_dia(DIA_ACTIVACION)

    n_peor = _nocturno(h_peor)
    n_cero = _nocturno(h_cero)
    n_act = _nocturno(h_act)
    d_peor = _diurno(h_peor)
    d_cero = _diurno(h_cero)
    red = (n_peor - n_cero) / n_peor * 100.0 if n_peor else 0.0
    ymax = max(
        max(float(h_peor.get(i, 0.0)) for i in range(24)),
        max(float(h_cero.get(i, 0.0)) for i in range(24)),
        0.05,
    ) * 1.18

    out = Path("reports/Parque_Arauco/Estación/Pizza_Hut_activacion_ene2026")
    gra = out / "graficos"
    gra.mkdir(parents=True, exist_ok=True)

    guardar_grafico_horario_24h_app(
        h_peor,
        gra / "pizza_hut_20260118_sin_control_max_nocturno.png",
        titulo="Pizza Hut — domingo 18-01-2026 (antes de activar) — mayor consumo nocturno",
        subtitulo="Sin control. Madrugada 00–06 = 8,44 m³. Día elegido: el de más consumo nocturno antes del 22-01.",
    )
    guardar_grafico_horario_24h_app(
        h_cero,
        gra / "pizza_hut_20260123_primera_madrugada_cero.png",
        titulo="Pizza Hut — viernes 23-01-2026 (primera madrugada tras activación 22-01)",
        subtitulo="Activación 22-01. 00:00–01:00 residual; desde 02:00 en cero hasta las 08:00. Noche 00–06 = 0,90 m³.",
    )

    fig, axes = plt.subplots(2, 1, figsize=(10.4, 8.4), sharex=True)
    _perfil(
        axes[0],
        h_peor,
        "Antes — 18-01-2026 (máximo nocturno previo al 22-01)   00–06 = 8,44 m³",
        ymax,
    )
    _perfil(
        axes[1],
        h_cero,
        "Después — 23-01-2026 (primera madrugada con control)   00–06 = 0,90 m³",
        ymax,
    )
    axes[1].set_xlabel("Hora del día (Chile)")
    fig.suptitle(
        "Parque Arauco Estación · Pizza Hut (000025-07)\n"
        "Activación 22-01-2026  ·  primera madrugada en cero = 23-01-2026",
        fontsize=12,
        fontweight="bold",
        color="#1F4E79",
    )
    fig.tight_layout(rect=(0, 0.02, 1, 0.92))
    compar = gra / "comparacion_18ene_vs_23ene_misma_escala.png"
    fig.savefig(compar, dpi=160, bbox_inches="tight")
    plt.close(fig)

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    docx = out / f"Pizza_Hut_PA_activacion_22ene_vs_23ene_{stamp}.docx"
    pdf = out / f"Pizza_Hut_PA_activacion_22ene_vs_23ene_{stamp}.pdf"

    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.5)
        s.bottom_margin = Cm(1.5)
        s.left_margin = Cm(1.7)
        s.right_margin = Cm(1.7)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("PARQUE ARAUCO ESTACIÓN — PIZZA HUT")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(COLOR_WES)
    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("Activación del control 22-01-2026  ·  primera madrugada en cero 23-01-2026")
    r.font.size = Pt(11)

    doc.add_paragraph(
        "El sistema se activó el jueves 22 de enero de 2026. La primera madrugada que ya corre "
        "con control es la del viernes 23 de enero. Para el informe se compara esa noche con el día "
        "anterior al 22 que más consumió entre 00:00 y 06:00: domingo 18 de enero (8,44 m³)."
    )

    tbl = doc.add_table(rows=4, cols=4)
    tbl.style = "Table Grid"
    hdr = ["", "18-01 (antes, peor noche)", "23-01 (primera madrugada)", "Variación"]
    for i, h in enumerate(hdr):
        _cell(tbl.rows[0].cells[i], h, fill=COLOR_WES, bold=True, size=9, color="FFFFFF")
    filas = [
        ("Consumo 00:00–06:00", f"{n_peor:.2f} m³".replace(".", ","), f"{n_cero:.2f} m³".replace(".", ","), f"−{red:.0f} %"),
        ("Pico horario madrugada", f"{max(h_peor.get(i,0) for i in range(7)):.2f} m³/h".replace(".", ","), f"{max(h_cero.get(i,0) for i in range(7)):.2f} m³/h".replace(".", ","), "corte desde las 02:00"),
        ("Consumo diurno 07–23", f"{d_peor:.1f} m³".replace(".", ","), f"{d_cero:.1f} m³".replace(".", ","), "operación diurna se mantiene"),
    ]
    fills = ["FCE4D6", "C8E6C9", None]
    for i, row in enumerate(filas, start=1):
        for j, v in enumerate(row):
            fill = "D6E3F0" if j == 0 else (fills[i - 1] if j in (1, 2) else None)
            _cell(tbl.rows[i].cells[j], v, fill=fill, bold=(j == 0 or j == 3), size=9)

    doc.add_paragraph()
    doc.add_paragraph(
        "Nota para el informe: el 23-01 aún hay 0,69 m³/h a las 00:00 y 0,21 m³/h a la 01:00 "
        "(el corte tardó en cerrar esa primera noche). Desde las 02:00 hasta las 08:00 el punto queda en 0,00 m³/h. "
        "La madrugada del 22-01 (día de activación) todavía consume 7,85 m³: es la noche previa al primer ciclo."
    )
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(compar), width=Inches(6.3))
    for pth in [
        gra / "pizza_hut_20260118_sin_control_max_nocturno.png",
        gra / "pizza_hut_20260123_primera_madrugada_cero.png",
    ]:
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(pth), width=Inches(6.3))
    doc.save(str(docx))

    with PdfPages(pdf) as pdfp:
        fig = plt.figure(figsize=(8.27, 11.69))
        ax = fig.add_subplot(111)
        ax.axis("off")
        ax.text(0.5, 0.96, "Parque Arauco Estación — Pizza Hut", ha="center", fontsize=14, color="#1F4E79", fontweight="bold", transform=ax.transAxes)
        ax.text(
            0.08,
            0.90,
            "Activación: jueves 22-01-2026.\n"
            "Primera madrugada con control: viernes 23-01-2026.\n"
            "Día previo al 22 con más consumo nocturno: domingo 18-01-2026.\n\n"
            f"18-01  00–06 = {n_peor:.2f} m³   pico 2,20 m³/h\n"
            f"23-01  00–06 = {n_cero:.2f} m³   en cero desde las 02:00\n"
            f"Reducción primera noche: {red:.0f} %.\n\n"
            "Mismo eje Y en ambos gráficos (comparación justa).",
            ha="left",
            va="top",
            fontsize=11,
            transform=ax.transAxes,
            family="DejaVu Sans",
        )
        ax2 = fig.add_axes([0.06, 0.04, 0.88, 0.58])
        ax2.imshow(plt.imread(str(compar)))
        ax2.axis("off")
        pdfp.savefig(fig)
        plt.close(fig)
        for pth in [
            gra / "pizza_hut_20260118_sin_control_max_nocturno.png",
            gra / "pizza_hut_20260123_primera_madrugada_cero.png",
        ]:
            fig = plt.figure(figsize=(8.27, 11.69))
            ax = fig.add_subplot(111)
            ax.imshow(plt.imread(str(pth)))
            ax.axis("off")
            fig.tight_layout()
            pdfp.savefig(fig)
            plt.close(fig)

    print(f"ANTES={DIA_PEOR} noche={n_peor:.2f}")
    print(f"ACTIVACION={DIA_ACTIVACION} noche={n_act:.2f}")
    print(f"PRIMERO_CERO={DIA_PRIMERO_CERO} noche={n_cero:.2f} reduccion={red:.0f}%")
    print(f"DOCX={docx}")
    print(f"PDF={pdf}")
    print(f"PNG={compar}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
