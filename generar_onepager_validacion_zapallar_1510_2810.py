#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
One-pager — Validación relojería vs App WES (Fundo Zapallar).
Ventana 15-10-2026 → 28-10-2026 · Matriz ESVAL + Etapa N°5 · con fotos de terreno.

Uso:
  python generar_onepager_validacion_zapallar_1510_2810.py
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Inches
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "reports" / "Fundo_Zapallar" / "VALIDACION"
FOTOS = OUT_DIR / "fotos_20261028"
NAVY = HexColor("#1F4788")
MUTED = HexColor("#595959")
LIGHT = HexColor("#F5F8FB")
GREEN = HexColor("#E8F5E9")
GREEN_TXT = HexColor("#1B5E20")

VENTANA = "15-10-2026 → 28-10-2026"
MATRIZ = {
    "nombre": "Matriz ESVAL",
    "node": "000027-01",
    "medidor": "Itron FLOSTAR S · 15A1131586 B",
    "ini": "386.889,7 m³ · 15-10 08:19",
    "fin": "388.346,6 m³ · 28-10 17:48",
    "reloj": 1456.9,
    "app": 1467.28,
    "horas": 320,
    "suma": "09:00 del 15 → 16:00 del 28",
    "error": 0.70,
    "foto": FOTOS / "Matriz_ESVAL_Itron_20261028_1748.png",
}
ETAPA5 = {
    "nombre": "Etapa N°5",
    "node": "000027-03",
    "medidor": "Sensus MeiStream Plus 100 · 8 SEN01 2370 9030",
    "ini": "5.627 m³ · 15-10 08:36",
    "fin": "5.809 m³ · 28-10 17:41",
    "reloj": 182.0,
    "app": 181.6,
    "horas": 320,
    "suma": "09:00 del 15 → 16:00 del 28",
    "error": 0.22,  # |1 − reloj/app|; reloj 0,4 m³ sobre app (prorrateo bordes)
    "foto": FOTOS / "Etapa_N5_Sensus_20261028_1741.png",
}


def _fmt(x: float, dec: int = 2) -> str:
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def _font_docx(run, *, size=11, bold=False, color=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _row(tbl, i: int, vals: list[str], *, bold=False, header=False) -> None:
    heading = RGBColor(31, 71, 136)
    for j, v in enumerate(vals):
        cell = tbl.rows[i].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(v)
        _font_docx(r, size=9, bold=bold or header, color=heading if header else None)


def build_docx(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.2)
    sec.bottom_margin = Cm(1.0)
    sec.left_margin = Cm(1.4)
    sec.right_margin = Cm(1.4)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    try:
        from generar_reporte_word import add_logo_to_header

        add_logo_to_header(doc)
    except Exception as exc:
        print(f"[WARN] Logo: {exc}")

    title = doc.add_heading("Validación relojería vs App WES — Fundo Zapallar", level=0)
    for r in title.runs:
        r.font.color.rgb = RGBColor(31, 71, 136)
        r.font.size = Pt(14)

    meta = doc.add_paragraph()
    r = meta.add_run(
        f"Ventana {VENTANA}  ·  Horas completas (excluye hora de cada foto)  ·  Uso interno WES"
    )
    _font_docx(r, size=8, color=RGBColor(89, 89, 89))

    h = doc.add_heading("Resultado", level=1)
    for r in h.runs:
        r.font.color.rgb = RGBColor(31, 71, 136)

    tbl = doc.add_table(rows=3, cols=6)
    tbl.style = "Table Grid"
    _row(
        tbl,
        0,
        ["Punto", "Relojería (m³)", "App WES (m³)", "Diff", "Error", "Lectura"],
        header=True,
    )
    for i, pto in enumerate((MATRIZ, ETAPA5), start=1):
        diff = pto["app"] - pto["reloj"]
        _row(
            tbl,
            i,
            [
                f"{pto['nombre']} {pto['node']}",
                _fmt(pto["reloj"], 1),
                _fmt(pto["app"], 2),
                f"{'+' if diff >= 0 else ''}{_fmt(diff, 2)}",
                f"{_fmt(pto['error'], 1)} %",
                "Excelente",
            ],
        )

    for pto in (MATRIZ, ETAPA5):
        h = doc.add_heading(f"{pto['nombre']} ({pto['node']})", level=1)
        for r in h.runs:
            r.font.color.rgb = RGBColor(31, 71, 136)
        for line in (
            f"Medidor: {pto['medidor']}",
            f"Lectura inicial: {pto['ini']}",
            f"Lectura final: {pto['fin']}",
            f"Suma App: {pto['suma']} ({pto['horas']} horas)",
        ):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(line)
            _font_docx(r, size=9)
        if pto["foto"].exists():
            doc.add_picture(str(pto["foto"]), width=Inches(3.4))

    h = doc.add_heading("Conclusión", level=1)
    for r in h.runs:
        r.font.color.rgb = RGBColor(31, 71, 136)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(
        "Etapa N°5 cierra con error 0,2 % (Excelente). Matriz ESVAL cierra con 0,7 % "
        "(Excelente), mejor que las campañas 02-10→15-10 (2,5 %). Ambos puntos quedan "
        "validados para el período. App: totales diarios measures/dates; bordes del 15 y "
        "del 28 prorrateados por horas completas (CSV horario no disponible en cloud)."
    )
    _font_docx(r, size=9)

    pie = doc.add_paragraph()
    r = pie.add_run(f"WES · generado {datetime.now().strftime('%d-%m-%Y %H:%M')}")
    _font_docx(r, size=7, color=RGBColor(89, 89, 89))
    doc.save(path)
    print(f"[OK] DOCX {path}")


def _card(c: canvas.Canvas, x, y, w, h, fill) -> None:
    c.setFillColor(fill)
    c.setStrokeColor(HexColor("#D0D7E2"))
    c.setLineWidth(0.5)
    c.roundRect(x, y, w, h, 3, fill=1, stroke=1)


def build_pdf(path: Path, logo: Path | None) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    W, H = A4
    ml, mr, mt, mb = 1.2 * cm, 1.2 * cm, 1.1 * cm, 1.0 * cm
    usable = W - ml - mr
    y = H - mt

    if logo and logo.exists():
        c.drawImage(str(logo), ml, y - 0.85 * cm, width=3.6 * cm, height=0.8 * cm, mask="auto")
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 12)
    c.drawRightString(W - mr, y - 0.25 * cm, "Validación Fundo Zapallar")
    c.setFont("Helvetica", 8)
    c.setFillColor(MUTED)
    c.drawRightString(W - mr, y - 0.55 * cm, f"Relojería vs App · {VENTANA}")
    y -= 1.0 * cm
    c.setStrokeColor(NAVY)
    c.setLineWidth(1)
    c.line(ml, y, W - mr, y)
    y -= 0.4 * cm

    # KPI cards
    card_w = (usable - 0.25 * cm) / 2
    card_h = 1.7 * cm
    for i, pto in enumerate((MATRIZ, ETAPA5)):
        x = ml + i * (card_w + 0.25 * cm)
        _card(c, x, y - card_h, card_w, card_h, GREEN)
        c.setFillColor(MUTED)
        c.setFont("Helvetica", 7)
        c.drawString(x + 5, y - 12, f"{pto['nombre']} · {pto['node']}")
        c.setFillColor(GREEN_TXT)
        c.setFont("Helvetica-Bold", 16)
        c.drawString(x + 5, y - 0.85 * cm, f"{_fmt(pto['error'], 1)} %")
        c.setFillColor(MUTED)
        c.setFont("Helvetica", 7)
        c.drawString(
            x + 5,
            y - card_h + 8,
            f"reloj {_fmt(pto['reloj'], 1)} · app {_fmt(pto['app'], 2)}",
        )
    y -= card_h + 0.35 * cm

    # Table
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(ml, y, "Detalle de campaña")
    y -= 0.3 * cm
    headers = ["Punto", "Ini → Fin", "Reloj", "App", "Err", "Lectura"]
    weights = [1.6, 3.2, 1.1, 1.1, 0.7, 1.1]
    col_w = [usable * w / sum(weights) for w in weights]
    row_h = 0.55 * cm

    def draw_row(yy, vals, header=False, alt=False):
        bg = NAVY if header else (LIGHT if alt else white)
        c.setFillColor(bg)
        c.rect(ml, yy - row_h, usable, row_h, fill=1, stroke=0)
        c.setStrokeColor(HexColor("#C5CEDA"))
        c.rect(ml, yy - row_h, usable, row_h, fill=0, stroke=1)
        x = ml
        for i, (val, w) in enumerate(zip(vals, col_w)):
            c.setFillColor(white if header else HexColor("#222222"))
            c.setFont("Helvetica-Bold" if header or i in (4, 5) else "Helvetica", 6.5)
            if not header and i == 4:
                c.setFillColor(GREEN_TXT)
            c.drawString(x + 2, yy - row_h + 0.16 * cm, val)
            x += w
        return yy - row_h

    y = draw_row(y, headers, header=True)
    for i, pto in enumerate((MATRIZ, ETAPA5)):
        y = draw_row(
            y,
            [
                pto["nombre"],
                f"{pto['ini'].split('·')[0].strip()} → {pto['fin'].split('·')[0].strip()}",
                _fmt(pto["reloj"], 1),
                _fmt(pto["app"], 2),
                f"{_fmt(pto['error'], 1)}%",
                "Excelente",
            ],
            alt=i % 2 == 1,
        )
    y -= 0.35 * cm

    # Photos
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(ml, y, "Fotos de cierre · 28-10-2026")
    y -= 0.25 * cm
    photo_w = (usable - 0.3 * cm) / 2
    photo_h = 5.2 * cm
    for i, pto in enumerate((MATRIZ, ETAPA5)):
        x = ml + i * (photo_w + 0.3 * cm)
        if pto["foto"].exists():
            c.drawImage(
                ImageReader(str(pto["foto"])),
                x,
                y - photo_h,
                width=photo_w,
                height=photo_h,
                preserveAspectRatio=True,
                anchor="c",
                mask="auto",
            )
        c.setFillColor(MUTED)
        c.setFont("Helvetica", 6.5)
        c.drawCentredString(x + photo_w / 2, y - photo_h - 0.28 * cm, pto["fin"])
    y -= photo_h + 0.55 * cm

    # Conclusion
    conc_h = 1.7 * cm
    _card(c, ml, y - conc_h, usable, conc_h, GREEN)
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 9)
    c.drawString(ml + 6, y - 0.35 * cm, "Conclusión")
    c.setFillColor(HexColor("#1B3A1D"))
    c.setFont("Helvetica", 7.5)
    for j, line in enumerate(
        [
            "Matriz ESVAL 0,7 % y Etapa N°5 0,2 % — ambos Excelente. Mejor que 02-10→15-10 (2,5 % / 1,4 %).",
            "App: measures/dates diarios; bordes 15-10 (desde 09:00) y 28-10 (hasta 16:00) prorrateados.",
            "Medidores: Itron FLOSTAR S (Matriz) y Sensus MeiStream Plus 100 + HRI-Mei B4 (Etapa 5).",
        ]
    ):
        c.drawString(ml + 6, y - 0.7 * cm - j * 0.3 * cm, line)

    c.setFillColor(MUTED)
    c.setFont("Helvetica", 7)
    c.drawString(ml, mb, f"WES · generado {datetime.now().strftime('%d-%m-%Y %H:%M')}")
    c.drawRightString(W - mr, mb, "Fundo Zapallar · VALIDACION 15-10→28-10")
    c.showPage()
    c.save()
    print(f"[OK] PDF  {path}")


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    docx_path = OUT_DIR / "OnePager_Validacion_Fundo_Zapallar_1510_2810.docx"
    pdf_path = OUT_DIR / "OnePager_Validacion_Fundo_Zapallar_1510_2810.pdf"
    logo = Path("/tmp/logo_wes.png")
    if not logo.exists():
        try:
            from PIL import Image

            Image.open(ROOT / "logo wes.bmp").convert("RGBA").save(logo)
        except Exception:
            logo = None

    build_docx(docx_path)
    build_pdf(pdf_path, logo if logo and logo.exists() else None)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
