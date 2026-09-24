"""
Tabla Fundo Zapallar: boletas ESVAL vs WES Matriz (12h) + Estanque Inferior.

Replica la tabla de 12 ciclos de lectura y agrega:
- consumo WES 12h del Estanque Inferior (000027-02) en el mismo periodo
- diferencia WES Matriz principal (000027-01) − Estanque Inferior

La ventana WES 12h es la misma del informe de validación:
[día lectura inicial 12:00 Chile, día lectura final 12:00 Chile).
"""

from __future__ import annotations

import csv
import io
import json
import sys
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import requests
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.shared import Cm, Inches, Pt, RGBColor

from generar_reporte_word import format_number_chilean

CL = ZoneInfo("America/Santiago")
BASE = "http://104.248.53.141:7003/wes/api/acl-node/v1"
NODO_MATRIZ = "000027-01"
NODO_INFERIOR = "000027-02"
HEADER_FILL = "1F4E79"
ROW_ALERTA_FILL = "F4C7C3"
TOTAL_FILL = "D6E3F0"
NEW_COL_FILL = "E2EFDA"
WES_BLUE = "#1F4E79"

# Valores WES 12h de Matriz publicados en la tabla original (m³).
# Se mantienen para que la tabla nueva sea la misma + columnas extra.
CICLOS = [
    {
        "ciclo": "28/08–27/09/25",
        "ini": date(2025, 8, 28),
        "fin": date(2025, 9, 27),
        "boleta": "90097803",
        "delta": 3375,
        "facturado": 401,
        "wes_matriz": 3188.9,
        "fase": "Inductivo ESVAL",
    },
    {
        "ciclo": "27/09–29/10/25",
        "ini": date(2025, 9, 27),
        "fin": date(2025, 10, 29),
        "boleta": "90777916",
        "delta": 3973,
        "facturado": 383,
        "wes_matriz": 4086.3,
        "fase": "Inductivo ESVAL",
    },
    {
        "ciclo": "29/10–28/11/25",
        "ini": date(2025, 10, 29),
        "fin": date(2025, 11, 28),
        "boleta": "91459036",
        "delta": 5711,
        "facturado": 417,
        "wes_matriz": 5678.5,
        "fase": "Inductivo ESVAL",
    },
    {
        "ciclo": "28/11–30/12/25",
        "ini": date(2025, 11, 28),
        "fin": date(2025, 12, 30),
        "boleta": "92141513",
        "delta": 6319,
        "facturado": 676,
        "wes_matriz": 6269.4,
        "fase": "Inductivo ESVAL",
    },
    {
        "ciclo": "30/12/25–30/01/26",
        "ini": date(2025, 12, 30),
        "fin": date(2026, 1, 30),
        "boleta": "92825564",
        "delta": 7193,
        "facturado": 501,
        "wes_matriz": 7200.3,
        "fase": "Inductivo ESVAL",
    },
    {
        "ciclo": "30/01–27/02/26",
        "ini": date(2026, 1, 30),
        "fin": date(2026, 2, 27),
        "boleta": "93509803",
        "delta": 6478,
        "facturado": 311,
        "wes_matriz": 5716.3,
        "fase": "Transición (desconex./cambio)",
    },
    {
        "ciclo": "27/02–30/03/26",
        "ini": date(2026, 2, 27),
        "fin": date(2026, 3, 30),
        "boleta": "94194644",
        "delta": 6793,
        "facturado": 1028,
        "wes_matriz": 1867.1,
        "fase": "Transición (desconex./cambio)",
    },
    {
        "ciclo": "30/03–29/04/26",
        "ini": date(2026, 3, 30),
        "fin": date(2026, 4, 29),
        "boleta": "94879773",
        "delta": 6034,
        "facturado": 808,
        "wes_matriz": 2997.4,
        "fase": "Ultrasónico WES",
    },
    {
        "ciclo": "29/04–29/05/26",
        "ini": date(2026, 4, 29),
        "fin": date(2026, 5, 29),
        "boleta": "95565463",
        "delta": 4566,
        "facturado": 635,
        "wes_matriz": 3254.5,
        "fase": "Ultrasónico WES",
    },
    {
        "ciclo": "29/05–27/06/26",
        "ini": date(2026, 5, 29),
        "fin": date(2026, 6, 27),
        "boleta": "96251613",
        "delta": 3707,
        "facturado": 639,
        "wes_matriz": 2378.8,
        "fase": "Ultrasónico WES",
    },
    {
        "ciclo": "27/06–29/07/26",
        "ini": date(2026, 6, 27),
        "fin": date(2026, 7, 29),
        "boleta": "96937917",
        "delta": 5067,
        "facturado": 2328,
        "wes_matriz": 3239.2,
        "fase": "Ultrasónico WES",
    },
    {
        "ciclo": "29/07–28/08/26",
        "ini": date(2026, 7, 29),
        "fin": date(2026, 8, 28),
        "boleta": "97624985",
        "delta": 2079,
        "facturado": 331,
        "wes_matriz": 1753.7,
        "fase": "Ultrasónico WES",
    },
]


def _pct(num: float, den: float) -> float:
    if den == 0:
        return 0.0
    return 100.0 * num / den


def _fmt_pct(value: float, decimals: int = 1) -> str:
    sign = "+" if value > 0 else ""
    return f"{sign}{format_number_chilean(value, decimals)}%"


def _fmt_m3(value: float, decimals: int = 1) -> str:
    return format_number_chilean(value, decimals)


def _set_cell_shading(cell, hex_color: str) -> None:
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'w:fill="{hex_color}"/>'
        )
    )


def _set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    size: float = 8,
    color: RGBColor | None = None,
    align: str = "center",
    fill: str | None = None,
) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run.font.color.rgb = color or RGBColor(0, 0, 0)
    if fill:
        _set_cell_shading(cell, fill)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = parse_xml(
        '<w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:top w:w="40" w:type="dxa"/>'
        '<w:left w:w="40" w:type="dxa"/>'
        '<w:bottom w:w="40" w:type="dxa"/>'
        '<w:right w:w="40" w:type="dxa"/>'
        "</w:tcMar>"
    )
    tcPr.append(tcMar)


def fetch_day(nid: str, d: date) -> list[tuple[datetime, float]]:
    st = d.strftime("%d%m%Y")
    r = requests.get(
        f"{BASE}/nodes/{nid}/dates.measures.csv",
        params={"start": st, "end": st},
        timeout=60,
    )
    by: dict[datetime, float] = defaultdict(float)
    if not r.ok:
        return []
    reader = csv.DictReader(io.StringIO(r.text))
    for row in reader:
        raw_t = (row.get("TIME") or "").strip()
        raw_v = (row.get("VALUE") or "0").strip()
        if not raw_t:
            continue
        try:
            ts = datetime.fromisoformat(raw_t.replace("Z", "+00:00"))
            by[ts] += float(raw_v)
        except (ValueError, TypeError):
            continue
    return sorted(by.items())


def load_hourly_range(nid: str, start_d: date, end_d: date) -> dict[datetime, float]:
    days: list[date] = []
    d = start_d - timedelta(days=1)
    last = end_d + timedelta(days=1)
    while d <= last:
        days.append(d)
        d += timedelta(days=1)

    all_rows: dict[datetime, float] = {}
    print(f"[INFO] Descargando {len(days)} días horarios de {nid}...", flush=True)
    with ThreadPoolExecutor(max_workers=12) as ex:
        futs = {ex.submit(fetch_day, nid, day): day for day in days}
        done = 0
        for fut in as_completed(futs):
            day = futs[fut]
            try:
                rows = fut.result()
            except Exception as exc:
                print(f"[WARN] Falló {nid} {day}: {exc}", flush=True)
                rows = []
            for ts, v in rows:
                all_rows[ts] = all_rows.get(ts, 0.0) + v
            done += 1
            if done % 50 == 0 or done == len(days):
                print(f"  {nid}: {done}/{len(days)} días", flush=True)
    return all_rows


def wes_12h_from_rows(all_rows: dict[datetime, float], start_d: date, end_d: date) -> tuple[float, int]:
    start_local = datetime.combine(start_d, time(12, 0), tzinfo=CL)
    end_local = datetime.combine(end_d, time(12, 0), tzinfo=CL)
    tot = 0.0
    n = 0
    for ts, v in all_rows.items():
        if ts.tzinfo is None:
            ts = ts.replace(tzinfo=timezone.utc)
        local = ts.astimezone(CL)
        if start_local <= local < end_local:
            tot += v
            n += 1
    return tot, n


def enriquecer_ciclos(inf_rows: dict[datetime, float]) -> list[dict]:
    enriched = []
    for c in CICLOS:
        inf_m3, inf_h = wes_12h_from_rows(inf_rows, c["ini"], c["fin"])
        dif_boleta = c["wes_matriz"] - c["delta"]
        dif_boleta_pct = _pct(dif_boleta, c["delta"])
        dif_mi = c["wes_matriz"] - inf_m3
        dif_mi_pct = _pct(dif_mi, c["wes_matriz"])
        row = dict(c)
        row.update(
            {
                "wes_inferior": round(inf_m3, 2),
                "horas_inferior": inf_h,
                "dif_boleta": round(dif_boleta, 2),
                "dif_boleta_pct": round(dif_boleta_pct, 2),
                "dif_matriz_inf": round(dif_mi, 2),
                "dif_matriz_inf_pct": round(dif_mi_pct, 2),
                "alerta_boleta": abs(dif_boleta_pct) > 20,
            }
        )
        enriched.append(row)
    return enriched


def _set_landscape(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Cm(1.1)
    section.right_margin = Cm(1.1)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)


def _headers() -> list[str]:
    return [
        "Ciclo lectura",
        "Boleta",
        "Delta medidor (m³)",
        "Facturado",
        "WES Matriz 12h (m³)",
        "Dif. % boleta",
        "Estanque Inf. 12h (m³)",
        "Dif. Matriz−Inf (m³)",
        "Fase equipo",
    ]


def construir_filas(rows: list[dict]) -> tuple[list[list[str]], list[str], dict]:
    table_rows: list[list[str]] = []
    for r in rows:
        dif_txt = (
            f"{_fmt_m3(r['dif_matriz_inf'], 1)} "
            f"({_fmt_pct(r['dif_matriz_inf_pct'])})"
        )
        table_rows.append(
            [
                r["ciclo"],
                r["boleta"],
                _fmt_m3(r["delta"], 0),
                _fmt_m3(r["facturado"], 0),
                _fmt_m3(r["wes_matriz"], 1),
                _fmt_pct(r["dif_boleta_pct"]),
                _fmt_m3(r["wes_inferior"], 1),
                dif_txt,
                r["fase"],
            ]
        )
    tot_delta = sum(r["delta"] for r in rows)
    tot_fact = sum(r["facturado"] for r in rows)
    tot_mat = sum(r["wes_matriz"] for r in rows)
    tot_inf = sum(r["wes_inferior"] for r in rows)
    tot_dif_b_pct = _pct(tot_mat - tot_delta, tot_delta)
    tot_dif_mi = tot_mat - tot_inf
    tot_dif_mi_pct = _pct(tot_dif_mi, tot_mat)
    total = [
        "TOTAL 12 ciclos",
        "",
        _fmt_m3(tot_delta, 0),
        _fmt_m3(tot_fact, 0),
        _fmt_m3(tot_mat, 1),
        _fmt_pct(tot_dif_b_pct),
        _fmt_m3(tot_inf, 1),
        f"{_fmt_m3(tot_dif_mi, 1)} ({_fmt_pct(tot_dif_mi_pct)})",
        "",
    ]
    summary = {
        "delta": tot_delta,
        "facturado": tot_fact,
        "matriz": tot_mat,
        "inferior": tot_inf,
        "dif_boleta_pct": tot_dif_b_pct,
        "dif_matriz_inf": tot_dif_mi,
        "dif_matriz_inf_pct": tot_dif_mi_pct,
    }
    return table_rows, total, summary


def crear_word(rows: list[dict], chart_path: Path, out_path: Path) -> Path:
    doc = Document()
    _set_landscape(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "Fundo Zapallar — boletas ESVAL vs WES Matriz y Estanque Inferior"
    )
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(31, 78, 121)
    run.font.name = "Calibri"

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run(
        "Misma ventana 12h de la tabla original, con consumo del estanque inferior "
        "y diferencia respecto de la matriz principal (000027-01)."
    )
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(80, 80, 80)

    note = doc.add_paragraph()
    nrun = note.add_run(
        "Método WES 12h: suma horaria en hora Chile desde las 12:00 del día de lectura "
        "inicial hasta las 12:00 del día de lectura final (ventana semiabierta). "
        "Matriz principal = Matriz ESVAL (000027-01). Estanque Inferior = 000027-02. "
        "Dif. Matriz−Inf = WES Matriz 12h − WES Estanque Inferior 12h (positivo = la matriz midió más). "
        "Las filas en rojo son las mismas de la tabla original (|Dif. % boleta| > 20%). "
        "Las columnas nuevas van en verde."
    )
    nrun.font.size = Pt(9)

    headers = _headers()
    body, total, _ = construir_filas(rows)
    table = doc.add_table(rows=1 + len(body) + 1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for i, h in enumerate(headers):
        fill = NEW_COL_FILL if i in (6, 7) else HEADER_FILL
        color = RGBColor(31, 78, 121) if i in (6, 7) else RGBColor(255, 255, 255)
        _set_cell_text(table.rows[0].cells[i], h, bold=True, size=8, color=color, fill=fill)

    for r_i, vals in enumerate(body):
        alerta = rows[r_i]["alerta_boleta"]
        for c_i, val in enumerate(vals):
            if c_i in (6, 7):
                fill = "C6E0B4" if not alerta else "F8CBAD"
            else:
                fill = ROW_ALERTA_FILL if alerta else "FFFFFF"
            _set_cell_text(
                table.rows[r_i + 1].cells[c_i],
                val,
                bold=False,
                size=8,
                fill=fill,
                align="left" if c_i in (0, 8) else "center",
            )

    last = table.rows[1 + len(body)]
    for c_i, val in enumerate(total):
        fill = "C5D9F1" if c_i in (6, 7) else TOTAL_FILL
        _set_cell_text(last.cells[c_i], val, bold=True, size=8, fill=fill)

    doc.add_paragraph()
    pchart = doc.add_paragraph()
    pchart.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if chart_path.exists():
        pchart.add_run().add_picture(str(chart_path), width=Inches(10.4))

    pie = doc.add_paragraph()
    pie.add_run(
        "Nota: en los ciclos de transición (desconexión del sensor de matriz el 25/02/2026 "
        "e instalación del ultrasónico el 11/03/2026) la matriz no es comparable con el estanque inferior."
    ).font.size = Pt(8)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def crear_grafico(rows: list[dict], out_path: Path) -> Path:
    labels = [r["ciclo"] for r in rows]
    matriz = [r["wes_matriz"] for r in rows]
    inferior = [r["wes_inferior"] for r in rows]
    x = np.arange(len(labels))
    width = 0.38

    fig, ax = plt.subplots(figsize=(13.5, 4.8))
    ax.bar(x - width / 2, matriz, width, label="WES Matriz principal 12h", color=WES_BLUE)
    ax.bar(x + width / 2, inferior, width, label="Estanque Inferior 12h", color="#70AD47")
    ax.set_ylabel("m³", fontsize=10)
    ax.set_title(
        "Consumo WES 12h por ciclo: Matriz principal vs Estanque Inferior",
        fontsize=12,
        fontweight="bold",
        color=WES_BLUE,
        pad=10,
    )
    ax.set_xticks(x)
    ax.set_xticklabels(labels, rotation=35, ha="right", fontsize=8)
    ax.legend(frameon=False, fontsize=9)
    ax.grid(axis="y", linestyle="--", alpha=0.35)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return out_path


def crear_pdf_tabla(rows: list[dict], chart_path: Path, out_path: Path) -> Path:
    body, total, _ = construir_filas(rows)
    headers = [
        "Ciclo lectura",
        "Boleta",
        "Delta medidor\n(m³)",
        "Facturado",
        "WES Matriz\n12h (m³)",
        "Dif. %\nboleta",
        "Estanque Inf.\n12h (m³)",
        "Dif. Matriz−Inf\n(m³)",
        "Fase equipo",
    ]
    cell_text = [headers] + body + [total]
    n_rows, n_cols = len(cell_text), len(headers)

    fig = plt.figure(figsize=(16.4, 9.4))
    fig.suptitle(
        "Fundo Zapallar — boletas ESVAL vs WES Matriz y Estanque Inferior",
        fontsize=15,
        fontweight="bold",
        color=WES_BLUE,
        y=0.98,
        x=0.02,
        ha="left",
    )
    fig.text(
        0.02,
        0.942,
        "Ventana WES 12h: 12:00 Chile del día de lectura inicial → 12:00 del día de lectura final.  "
        "Dif. Matriz−Inf = WES Matriz − Estanque Inferior (positivo = la matriz midió más).  "
        "Columnas nuevas en verde.",
        fontsize=8,
        color="#444444",
        ha="left",
    )

    ax = fig.add_axes([0.02, 0.42, 0.96, 0.50])
    ax.axis("off")
    col_widths = [0.13, 0.08, 0.09, 0.07, 0.10, 0.08, 0.11, 0.14, 0.20]
    table = ax.table(
        cellText=cell_text,
        loc="center",
        cellLoc="center",
        colWidths=col_widths,
    )
    table.auto_set_font_size(False)
    table.set_fontsize(7.6)
    table.scale(1, 1.62)

    for c in range(n_cols):
        cell = table[0, c]
        if c in (6, 7):
            cell.set_facecolor("#C6E0B4")
            cell.get_text().set_color(WES_BLUE)
        else:
            cell.set_facecolor(WES_BLUE)
            cell.get_text().set_color("white")
        cell.get_text().set_fontweight("bold")
        cell.get_text().set_fontsize(7.3)

    for r in range(1, n_rows - 1):
        alerta = rows[r - 1]["alerta_boleta"]
        for c in range(n_cols):
            cell = table[r, c]
            if c in (6, 7):
                cell.set_facecolor("#C6E0B4" if not alerta else "#F8CBAD")
            else:
                cell.set_facecolor("#F4C7C3" if alerta else "white")
            if c in (0, 8):
                cell.get_text().set_horizontalalignment("left")

    for c in range(n_cols):
        cell = table[n_rows - 1, c]
        cell.set_facecolor("#C5D9F1" if c in (6, 7) else "#D6E3F0")
        cell.get_text().set_fontweight("bold")

    ax_c = fig.add_axes([0.05, 0.07, 0.90, 0.33])
    img = plt.imread(str(chart_path))
    ax_c.imshow(img)
    ax_c.axis("off")
    fig.text(
        0.02,
        0.02,
        "Nota: en transición (desconexión 25/02/2026 e instalación ultrasónico 11/03/2026) "
        "la matriz no es comparable con el estanque inferior. Las filas rojas son las de la tabla original "
        "(|Dif. % boleta| > 20%).",
        fontsize=7.5,
        color="#444444",
        ha="left",
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=170)
    plt.close(fig)
    return out_path


def main() -> int:
    if sys.platform == "win32":
        for s in (sys.stdout, sys.stderr):
            try:
                s.reconfigure(encoding="utf-8", line_buffering=True)
            except Exception:
                pass

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out_dir = Path("reports/Fundo_Zapallar/ANALISIS") / f"tabla_matriz_vs_inferior_{stamp}"
    out_dir.mkdir(parents=True, exist_ok=True)

    start_all = min(c["ini"] for c in CICLOS)
    end_all = max(c["fin"] for c in CICLOS)
    inf_rows = load_hourly_range(NODO_INFERIOR, start_all, end_all)
    if not inf_rows:
        print("[ERROR] No se obtuvieron medidas del Estanque Inferior.", flush=True)
        return 1

    rows = enriquecer_ciclos(inf_rows)
    _, _, summary = construir_filas(rows)

    json_path = out_dir / "matriz_vs_inferior_12_ciclos.json"
    payload = {
        "generado": datetime.now().isoformat(timespec="seconds"),
        "nodo_matriz": NODO_MATRIZ,
        "nodo_inferior": NODO_INFERIOR,
        "metodo": "WES 12h Chile [ini 12:00, fin 12:00)",
        "ciclos": [
            {
                **{k: (v.isoformat() if isinstance(v, date) else v) for k, v in r.items()},
            }
            for r in rows
        ],
        "totales": summary,
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    chart_path = out_dir / "chart_matriz_vs_inferior.png"
    crear_grafico(rows, chart_path)

    docx_path = out_dir / "Tabla_Matriz_WES_vs_Estanque_Inferior.docx"
    crear_word(rows, chart_path, docx_path)

    pdf_tabla = out_dir / "Tabla_Matriz_WES_vs_Estanque_Inferior.pdf"
    crear_pdf_tabla(rows, chart_path, pdf_tabla)

    print(f"[OK] JSON  {json_path}")
    print(f"[OK] Word  {docx_path}")
    print(f"[OK] PDF   {pdf_tabla}")
    print(
        f"[OK] Totales  Matriz={summary['matriz']:.1f}  "
        f"Inferior={summary['inferior']:.1f}  "
        f"Dif={summary['dif_matriz_inf']:.1f} ({summary['dif_matriz_inf_pct']:.1f}%)"
    )
    for r in rows:
        print(
            f"  {r['ciclo']}: Inf={r['wes_inferior']:.1f}  "
            f"Dif={r['dif_matriz_inf']:.1f} ({r['dif_matriz_inf_pct']:.1f}%)"
        )
    print(f"OUT_DIR={out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
