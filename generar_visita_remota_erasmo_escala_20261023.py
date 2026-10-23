"""
Formulario de visita remota — Colegio Erasmo Escala (CORMUP, Peñalolén).

Registro del llamado telefónico del viernes 23 de octubre de 2026, 09:00
(hora Chile): error de programación del viernes (corte en la mañana),
normalización de la programación, confirmación con la directora y
normalización del equipo que había activado el protocolo de emergencia bypass.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "reports" / "CORMUP" / "VISITA_REMOTA"
OUT_DOCX = OUT_DIR / "Formulario_Visita_Remota_Erasmo_Escala_Penalolen_20261023.docx"

AZUL = RGBColor(0, 51, 102)
AZUL_HEX = "003366"
GRIS_FILA = "F4F7FA"
VERDE_HEX = "1B7A3A"
LINEA = "D0D7E2"


def _set_run_font(run, *, name: str = "Calibri", size: int = 11, bold: bool = False, color: RGBColor | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _shade(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _margins(cell, top=40, bottom=40, left=80, right=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def _set_borders(cell, color: str = LINEA) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tc_borders.append(el)
    tc_pr.append(tc_borders)


def _paragraph(cell, text: str, *, bold: bool = False, size: int = 11, color: RGBColor | None = None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    _set_run_font(run, size=size, bold=bold, color=color)
    return p


def _add_logo(doc: Document) -> None:
    logo = ROOT / "logo wes.bmp"
    section = doc.sections[0]
    section.header_distance = Cm(0.4)
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if logo.is_file():
        run = p.add_run()
        run.add_picture(str(logo), width=Inches(1.35))


def _kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False
    table.allow_autofit = False
    table.columns[0].width = Cm(5.2)
    table.columns[1].width = Cm(12.0)
    for i, (label, value) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.width = Cm(5.2)
        c1.width = Cm(12.0)
        _shade(c0, AZUL_HEX)
        _shade(c1, "FFFFFF" if i % 2 == 0 else GRIS_FILA)
        _margins(c0)
        _margins(c1)
        _set_borders(c0, AZUL_HEX)
        _set_borders(c1)
        _paragraph(c0, label, bold=True, size=10, color=RGBColor(255, 255, 255))
        _paragraph(c1, value, size=11)


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    _set_run_font(run, size=13, bold=True, color=AZUL)


def _body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    _set_run_font(run, size=11)


def _actions(doc: Document, items: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1 + len(items), cols=2)
    table.autofit = False
    hdr = table.rows[0].cells
    for cell, text, width in (
        (hdr[0], "Acción", Cm(6.4)),
        (hdr[1], "Resultado", Cm(10.8)),
    ):
        cell.width = width
        _shade(cell, AZUL_HEX)
        _margins(cell)
        _set_borders(cell, AZUL_HEX)
        _paragraph(cell, text, bold=True, size=10, color=RGBColor(255, 255, 255))
    for i, (accion, resultado) in enumerate(items, start=1):
        c0, c1 = table.rows[i].cells
        c0.width = Cm(6.4)
        c1.width = Cm(10.8)
        _shade(c0, GRIS_FILA)
        _shade(c1, "FFFFFF")
        _margins(c0)
        _margins(c1)
        _set_borders(c0)
        _set_borders(c1)
        _paragraph(c0, accion, bold=True, size=10, color=AZUL)
        _paragraph(c1, resultado, size=10)


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.top_margin = Cm(1.15)
    section.bottom_margin = Cm(0.9)
    _add_logo(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(2)
    title.paragraph_format.space_after = Pt(0)
    run = title.add_run("Formulario de visita remota")
    _set_run_font(run, size=18, bold=True, color=AZUL)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(2)
    run = sub.add_run("Colegio Erasmo Escala  ·  CORMUP Peñalolén  ·  viernes 23 de octubre de 2026")
    _set_run_font(run, size=11, color=RGBColor(70, 80, 95))

    _heading(doc, "Identificación")
    _kv_table(
        doc,
        [
            ("Tipo de visita", "Remota (sin desplazamiento a terreno)"),
            ("Fecha", "Viernes 23 de octubre de 2026"),
            ("Hora del llamado", "09:00, hora de Chile (America/Santiago)"),
            ("Cliente", "CORMUP — Corporación Municipal de Peñalolén"),
            ("Establecimiento", "Centro Educacional Erasmo Escala Arriagada (Erasmo Escala)"),
            ("Comuna", "Peñalolén"),
            ("Punto de monitoreo", "000008-07"),
            ("Contacto del establecimiento", "Directora (confirmación telefónica)"),
            ("Estado del caso", "Cerrado en visita remota"),
        ],
    )

    _heading(doc, "Motivo de la atención")
    _body(
        doc,
        "Después del llamado telefónico realizado hoy viernes 23 de octubre de 2026 "
        "a las 09:00, se revisó la programación del equipo de control de agua del "
        "Colegio Erasmo Escala. La revisión detectó un error en la programación del "
        "día viernes: el corte estaba programado en la mañana.",
    )

    _heading(doc, "Trabajos realizados")
    _actions(
        doc,
        [
            (
                "Revisión de programación",
                "Se revisó la programación del viernes a partir del llamado de las 09:00.",
            ),
            (
                "Corrección del corte",
                "El viernes estaba programado con corte en la mañana. Se normalizó esa programación.",
            ),
            (
                "Confirmación con la directora",
                "Se habló por teléfono con la directora del colegio para confirmar la reparación.",
            ),
            (
                "Protocolo de emergencia",
                "El equipo había activado el protocolo de emergencia bypass. Se normalizó el equipo.",
            ),
        ],
    )

    _heading(doc, "Detalle")
    _body(
        doc,
        "El corte del viernes figuraba en la mañana y se normalizó en la atención remota. "
        "Horario de corte de referencia en el control nocturno de este punto (000008-07): "
        "01:00 a 05:59. La directora confirmó la reparación por teléfono y el equipo, "
        "que había activado el protocolo de emergencia bypass, quedó en operación regular.",
    )

    _heading(doc, "Cierre")
    _kv_table(
        doc,
        [
            ("Programación del viernes", "Normalizada. Se corrigió el corte que estaba programado en la mañana."),
            ("Protocolo de emergencia bypass", "Desactivado. Equipo normalizado."),
            ("Confirmación del establecimiento", "Directora informada por teléfono. Reparación confirmada."),
            ("Visita presencial", "No requerida. Atención cerrada en modalidad remota."),
        ],
    )

    cierre = doc.add_paragraph()
    cierre.paragraph_format.space_before = Pt(6)
    cierre.paragraph_format.space_after = Pt(0)
    run = cierre.add_run(
        "Registro del Agente WES — atención remota del 23 de octubre de 2026 — equipo de operaciones."
    )
    _set_run_font(run, size=9, color=RGBColor(90, 98, 110))

    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    path = build()
    print(path)
