"""
Fundo Zapallar: facturación ESVAL vs Matriz principal con día completo
(sin recorte 12:00) y revisión de huecos horarios.

WES día completo = suma de todas las horas Chile (00:00–24:00) desde el día
de lectura inicial hasta el día de lectura final, ambos inclusive.
"""

from __future__ import annotations

import csv
import json
import sys
from collections import defaultdict
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from matplotlib.backends.backend_pdf import PdfPages

from exclusiones_reportes import FUNDO_ZAPALLAR_NODE_IDS
from generar_tabla_zapallar_matriz_vs_inferior import (
    CICLOS,
    CL,
    HEADER_FILL,
    NODO_MATRIZ,
    ROW_ALERTA_FILL,
    TOTAL_FILL,
    WES_BLUE,
    _fmt_m3,
    _fmt_pct,
    _pct,
    _set_cell_text,
    _set_landscape,
    load_hourly_range,
    wes_12h_from_rows,
)

NOMBRES = {
    "000027-01": "Matriz ESVAL",
    "000027-02": "Estanque Inferior",
    "000027-03": "Etapa N°5",
    "000027-04": "Etapa N°1 al 4",
    "000027-06": "Etapa N°1",
    "000027-07": "Etapa N°2",
    "000027-08": "Etapa N°3",
    "000027-09": "Riego Llenado ESVAL",
}


def index_chile_hours(all_rows: dict[datetime, float]) -> dict[tuple[date, int], float]:
    acc: dict[tuple[date, int], float] = defaultdict(float)
    for ts, v in all_rows.items():
        if ts.tzinfo is None:
            ts = ts.replace(tzinfo=timezone.utc)
        local = ts.astimezone(CL)
        acc[(local.date(), local.hour)] += v
    return dict(acc)


def expected_hours_chile(d: date) -> set[int]:
    start = datetime.combine(d, time(0, 0), tzinfo=CL)
    end = datetime.combine(d + timedelta(days=1), time(0, 0), tzinfo=CL)
    hours: set[int] = set()
    t = start.astimezone(timezone.utc)
    end_u = end.astimezone(timezone.utc)
    while t < end_u:
        loc = t.astimezone(CL)
        if loc.date() == d:
            hours.add(loc.hour)
        t += timedelta(hours=1)
    return hours or set(range(24))


def wes_dia_completo(
    idx: dict[tuple[date, int], float], start_d: date, end_d: date
) -> tuple[float, int, int]:
    tot = 0.0
    presentes = 0
    esperadas = 0
    d = start_d
    while d <= end_d:
        expected = expected_hours_chile(d)
        esperadas += len(expected)
        for h in expected:
            key = (d, h)
            if key in idx:
                tot += idx[key]
                presentes += 1
        d += timedelta(days=1)
    return tot, presentes, esperadas


def huecos_rango(
    idx: dict[tuple[date, int], float], start_d: date, end_d: date
) -> list[dict]:
    out: list[dict] = []
    d = start_d
    while d <= end_d:
        expected = expected_hours_chile(d)
        present = {h for h in expected if (d, h) in idx}
        missing = sorted(expected - present)
        if missing:
            out.append(
                {
                    "date": d,
                    "missing": missing,
                    "n": len(missing),
                    "expected": len(expected),
                    "present": len(present),
                }
            )
        d += timedelta(days=1)
    return out


def colapsar_huecos(gaps: list[dict]) -> list[str]:
    """Agrupa días consecutivos sin ninguna hora en rangos; deja días parciales sueltos."""
    if not gaps:
        return []
    lines: list[str] = []
    i = 0
    while i < len(gaps):
        g = gaps[i]
        if g["n"] == g["expected"]:
            j = i
            while (
                j + 1 < len(gaps)
                and gaps[j + 1]["n"] == gaps[j + 1]["expected"]
                and gaps[j + 1]["date"] == gaps[j]["date"] + timedelta(days=1)
            ):
                j += 1
            if j == i:
                lines.append(f"{g['date'].strftime('%d/%m/%Y')}: día completo sin datos (24 h)")
            else:
                lines.append(
                    f"{gaps[i]['date'].strftime('%d/%m/%Y')}–{gaps[j]['date'].strftime('%d/%m/%Y')}: "
                    f"{(gaps[j]['date'] - gaps[i]['date']).days + 1} días sin datos"
                )
            i = j + 1
            continue
        horas = ", ".join(f"{h:02d}:00" for h in g["missing"][:12])
        extra = "" if len(g["missing"]) <= 12 else f" (+{len(g['missing']) - 12})"
        lines.append(f"{g['date'].strftime('%d/%m/%Y')}: {g['n']} h ({horas}{extra})")
        i += 1
    return lines


def enriquecer_ciclos(idx: dict[tuple[date, int], float], all_rows: dict[datetime, float]) -> list[dict]:
    enriched = []
    for c in CICLOS:
        full, presentes, esperadas = wes_dia_completo(idx, c["ini"], c["fin"])
        wes12, h12 = wes_12h_from_rows(all_rows, c["ini"], c["fin"])
        huecos = esperadas - presentes
        dif = full - c["delta"]
        dif_pct = _pct(dif, c["delta"])
        row = dict(c)
        row.update(
            {
                "wes_12h": round(wes12, 2),
                "horas_12h": h12,
                "wes_full": round(full, 2),
                "horas_presentes": presentes,
                "horas_esperadas": esperadas,
                "huecos": huecos,
                "cobertura_pct": round(_pct(presentes, esperadas), 1),
                "dif_full": round(dif, 2),
                "dif_full_pct": round(dif_pct, 2),
                "alerta": abs(dif_pct) > 20,
            }
        )
        enriched.append(row)
    return enriched


def construir_filas(rows: list[dict]) -> tuple[list[list[str]], list[str], dict]:
    body = []
    for r in rows:
        body.append(
            [
                r["ciclo"],
                r["boleta"],
                _fmt_m3(r["delta"], 0),
                _fmt_m3(r["facturado"], 0),
                _fmt_m3(r["wes_12h"], 1),
                _fmt_m3(r["wes_full"], 1),
                _fmt_pct(r["dif_full_pct"]),
                f"{r['huecos']} h ({_fmt_m3(r['cobertura_pct'], 1)}%)",
                r["fase"],
            ]
        )
    tot_delta = sum(r["delta"] for r in rows)
    tot_fact = sum(r["facturado"] for r in rows)
    tot_12 = sum(r["wes_12h"] for r in rows)
    tot_full = sum(r["wes_full"] for r in rows)
    tot_huecos = sum(r["huecos"] for r in rows)
    tot_esp = sum(r["horas_esperadas"] for r in rows)
    tot_pres = sum(r["horas_presentes"] for r in rows)
    tot_pct = _pct(tot_full - tot_delta, tot_delta)
    cob = _pct(tot_pres, tot_esp)
    total = [
        "TOTAL 12 ciclos",
        "",
        _fmt_m3(tot_delta, 0),
        _fmt_m3(tot_fact, 0),
        _fmt_m3(tot_12, 1),
        _fmt_m3(tot_full, 1),
        _fmt_pct(tot_pct),
        f"{tot_huecos} h ({_fmt_m3(cob, 1)}%)",
        "",
    ]
    summary = {
        "delta": tot_delta,
        "facturado": tot_fact,
        "wes_12h": tot_12,
        "wes_full": tot_full,
        "dif_full_pct": tot_pct,
        "huecos": tot_huecos,
        "cobertura_pct": cob,
    }
    return body, total, summary


HEADERS = [
    "Ciclo lectura",
    "Boleta",
    "Delta medidor (m³)",
    "Facturado",
    "WES 12h (m³)",
    "WES día completo (m³)",
    "Dif. % vs delta",
    "Huecos (cobertura)",
    "Fase equipo",
]


def crear_grafico_consumo(rows: list[dict], out_path: Path) -> Path:
    labels = [r["ciclo"] for r in rows]
    x = np.arange(len(labels))
    w = 0.27
    fig, ax = plt.subplots(figsize=(13.6, 4.6))
    ax.bar(x - w, [r["delta"] for r in rows], w, label="Delta medidor ESVAL", color="#7F7F7F")
    ax.bar(x, [r["wes_12h"] for r in rows], w, label="WES 12h (recorte 12:00)", color="#5B9BD5")
    ax.bar(x + w, [r["wes_full"] for r in rows], w, label="WES día completo", color=WES_BLUE)
    ax.set_ylabel("m³")
    ax.set_title(
        "Facturación (delta medidor) vs Matriz ESVAL: 12h vs día completo",
        fontsize=12,
        fontweight="bold",
        color=WES_BLUE,
        pad=10,
    )
    ax.set_xticks(x)
    ax.set_xticklabels(labels, rotation=35, ha="right", fontsize=8)
    ax.legend(frameon=False, fontsize=8.5)
    ax.grid(axis="y", linestyle="--", alpha=0.35)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    fig.savefig(out_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return out_path


def crear_grafico_huecos(
    gaps_por_nodo: dict[str, list[dict]],
    altas: dict[str, date | None],
    start_d: date,
    end_d: date,
    out_path: Path,
) -> Path:
    days: list[date] = []
    d = start_d
    while d <= end_d:
        days.append(d)
        d += timedelta(days=1)
    nids = [n for n in FUNDO_ZAPALLAR_NODE_IDS if n in gaps_por_nodo]
    mat = np.full((len(nids), len(days)), np.nan)
    day_ix = {day: i for i, day in enumerate(days)}
    for r, nid in enumerate(nids):
        alta = altas.get(nid)
        for i, day in enumerate(days):
            if alta and day >= alta:
                mat[r, i] = 0.0
        for g in gaps_por_nodo[nid]:
            i = day_ix.get(g["date"])
            if i is not None:
                mat[r, i] = g["n"]
    fig, ax = plt.subplots(figsize=(13.6, 4.8))
    cmap = plt.cm.YlOrRd.copy()
    cmap.set_bad("#D9D9D9")
    im = ax.imshow(mat, aspect="auto", cmap=cmap, interpolation="nearest", vmin=0, vmax=24)
    ax.set_yticks(range(len(nids)))
    ax.set_yticklabels([f"{NOMBRES.get(n, n)}" for n in nids], fontsize=8)
    step = max(1, len(days) // 14)
    xt = list(range(0, len(days), step))
    ax.set_xticks(xt)
    ax.set_xticklabels([days[i].strftime("%d/%m/%y") for i in xt], rotation=35, ha="right", fontsize=7.5)
    ax.set_title(
        "Huecos horarios por día (gris = antes del alta del punto; amarillo = serie completa; rojo = horas sin dato)",
        fontsize=10,
        fontweight="bold",
        color=WES_BLUE,
        pad=8,
    )
    fig.colorbar(im, ax=ax, fraction=0.02, pad=0.02, label="h faltantes")
    fig.tight_layout()
    fig.savefig(out_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return out_path


def _style_table(ax, cell_text, rows, headers_pdf):
    n_rows, n_cols = len(cell_text), len(headers_pdf)
    col_w = [0.12, 0.08, 0.09, 0.07, 0.09, 0.12, 0.09, 0.12, 0.12]
    table = ax.table(cellText=cell_text, loc="center", cellLoc="center", colWidths=col_w)
    table.auto_set_font_size(False)
    table.set_fontsize(7.3)
    table.scale(1, 1.55)
    for c in range(n_cols):
        cell = table[0, c]
        fill = "#C6E0B4" if c in (5, 7) else WES_BLUE
        cell.set_facecolor(fill)
        cell.get_text().set_color(WES_BLUE if c in (5, 7) else "white")
        cell.get_text().set_fontweight("bold")
        cell.get_text().set_fontsize(7.0)
    for r in range(1, n_rows - 1):
        alerta = rows[r - 1]["alerta"]
        huecos = rows[r - 1]["huecos"]
        for c in range(n_cols):
            cell = table[r, c]
            if c == 5:
                cell.set_facecolor("#C6E0B4" if not alerta else "#F8CBAD")
            elif c == 7:
                cell.set_facecolor("#F8CBAD" if huecos > 24 else "#C6E0B4")
            else:
                cell.set_facecolor("#F4C7C3" if alerta else "white")
            if c in (0, 8):
                cell.get_text().set_horizontalalignment("left")
    for c in range(n_cols):
        cell = table[n_rows - 1, c]
        cell.set_facecolor("#C5D9F1" if c in (5, 7) else "#D6E3F0")
        cell.get_text().set_fontweight("bold")
    return table


def crear_pdf(
    rows: list[dict],
    chart_cons: Path,
    chart_huecos: Path,
    resumen_nodos: list[dict],
    lineas_matriz: list[str],
    out_path: Path,
) -> Path:
    body, total, _ = construir_filas(rows)
    headers_pdf = [
        "Ciclo lectura",
        "Boleta",
        "Delta medidor\n(m³)",
        "Facturado",
        "WES 12h\n(m³)",
        "WES día\ncompleto (m³)",
        "Dif. %\nvs delta",
        "Huecos\n(cobertura)",
        "Fase equipo",
    ]
    cell_text = [headers_pdf] + body + [total]

    with PdfPages(out_path) as pdf:
        fig = plt.figure(figsize=(16.4, 9.4))
        fig.suptitle(
            "Fundo Zapallar — facturación ESVAL vs Matriz principal (día completo, sin corte 12h)",
            fontsize=14.5,
            fontweight="bold",
            color=WES_BLUE,
            y=0.98,
            x=0.02,
            ha="left",
        )
        fig.text(
            0.02,
            0.942,
            "WES día completo: todas las horas Chile 00:00–24:00 del día de lectura inicial al día de lectura final (ambos inclusive). "
            "WES 12h se muestra solo de referencia. Dif. % = (WES día completo − delta medidor) / delta.",
            fontsize=8,
            color="#444444",
        )
        ax = fig.add_axes([0.02, 0.42, 0.96, 0.50])
        ax.axis("off")
        _style_table(ax, cell_text, rows, headers_pdf)
        ax_c = fig.add_axes([0.05, 0.07, 0.90, 0.33])
        ax_c.imshow(plt.imread(str(chart_cons)))
        ax_c.axis("off")
        fig.text(
            0.02,
            0.018,
            "Las filas en rojo tienen |Dif. %| > 20 % respecto del delta del medidor fiscal. "
            "La columna verde es el WES sin recorte de 12:00. Huecos = horas Chile sin registro (el 0 m³ con señal no cuenta como hueco).",
            fontsize=7.4,
            color="#444444",
        )
        pdf.savefig(fig)
        plt.close(fig)

        fig2 = plt.figure(figsize=(16.4, 9.4))
        fig2.suptitle(
            "Fundo Zapallar — huecos horarios (serie WES, hora Chile)",
            fontsize=14.5,
            fontweight="bold",
            color=WES_BLUE,
            y=0.98,
            x=0.02,
            ha="left",
        )
        fig2.text(
            0.02,
            0.942,
            "Hueco = hora Chile esperada sin fila en dates.measures.csv (no es consumo cero). "
            "Revisión 28/08/2025–21/09/2026. Gris = antes del alta del punto. "
            "Etapa 1 y 3 entran en serie el 14/12/2025; Riego Llenado el 20/01/2026.",
            fontsize=8,
            color="#444444",
        )
        axh = fig2.add_axes([0.04, 0.48, 0.92, 0.43])
        axh.imshow(plt.imread(str(chart_huecos)))
        axh.axis("off")

        ax_t = fig2.add_axes([0.03, 0.06, 0.55, 0.38])
        ax_t.axis("off")
        nod_headers = ["Punto", "Alta serie", "Días con hueco", "Horas faltantes", "Peor día"]
        nod_rows = [nod_headers]
        for s in resumen_nodos:
            nod_rows.append(
                [
                    s["nombre"],
                    s.get("alta", "—"),
                    str(s["dias_con_hueco"]),
                    str(s["horas_faltantes"]),
                    s["peor_dia"],
                ]
            )
        t = ax_t.table(cellText=nod_rows, loc="upper left", cellLoc="left", colWidths=[0.28, 0.16, 0.16, 0.18, 0.22])
        t.auto_set_font_size(False)
        t.set_fontsize(8)
        t.scale(1, 1.45)
        for c in range(5):
            t[0, c].set_facecolor(WES_BLUE)
            t[0, c].get_text().set_color("white")
            t[0, c].get_text().set_fontweight("bold")
        for r in range(1, len(nod_rows)):
            for c in range(5):
                fill = "white"
                if resumen_nodos[r - 1]["horas_faltantes"] == 0:
                    fill = "#C6E0B4"
                elif resumen_nodos[r - 1]["horas_faltantes"] > 48:
                    fill = "#F4C7C3"
                t[r, c].set_facecolor(fill)

        ax_l = fig2.add_axes([0.60, 0.06, 0.38, 0.38])
        ax_l.axis("off")
        ax_l.set_title("Matriz ESVAL — detalle de huecos", loc="left", fontsize=10, color=WES_BLUE, fontweight="bold")
        if lineas_matriz:
            txt = "\n".join(lineas_matriz[:16])
            if len(lineas_matriz) > 16:
                txt += f"\n… +{len(lineas_matriz) - 16} tramos más (ver JSON/CSV)"
        else:
            txt = "Sin huecos horarios en la Matriz ESVAL en el periodo revisado."
        ax_l.text(0.0, 0.95, txt, va="top", ha="left", fontsize=7.6, family="DejaVu Sans", wrap=True)
        fig2.text(
            0.02,
            0.015,
            "Si un ciclo tiene muchos huecos, el WES día completo puede quedar bajo respecto del medidor fiscal sin que el caudalímetro esté descalibrado.",
            fontsize=7.4,
            color="#444444",
        )
        pdf.savefig(fig2)
        plt.close(fig2)
    return out_path


def crear_word(
    rows: list[dict],
    chart_cons: Path,
    chart_huecos: Path,
    resumen_nodos: list[dict],
    lineas_matriz: list[str],
    out_path: Path,
) -> Path:
    doc = Document()
    _set_landscape(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Fundo Zapallar — facturación vs Matriz ESVAL (día completo)")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(31, 78, 121)

    note = doc.add_paragraph()
    note.add_run(
        "Sin recorte de 12:00. WES día completo suma todas las horas Chile (00:00–24:00) "
        "desde el día de lectura inicial hasta el día de lectura final, ambos inclusive. "
        "La columna WES 12h queda de referencia. Hueco = hora Chile sin registro en la API "
        "(un 0 m³ con señal no es hueco)."
    ).font.size = Pt(9)

    body, total, _ = construir_filas(rows)
    table = doc.add_table(rows=1 + len(body) + 1, cols=len(HEADERS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(HEADERS):
        fill = "C6E0B4" if i in (5, 7) else HEADER_FILL
        color = RGBColor(31, 78, 121) if i in (5, 7) else RGBColor(255, 255, 255)
        _set_cell_text(table.rows[0].cells[i], h, bold=True, size=8, color=color, fill=fill)
    for r_i, vals in enumerate(body):
        alerta = rows[r_i]["alerta"]
        huecos = rows[r_i]["huecos"]
        for c_i, val in enumerate(vals):
            if c_i == 5:
                fill = "C6E0B4" if not alerta else "F8CBAD"
            elif c_i == 7:
                fill = "F8CBAD" if huecos > 24 else "C6E0B4"
            else:
                fill = ROW_ALERTA_FILL if alerta else "FFFFFF"
            _set_cell_text(
                table.rows[r_i + 1].cells[c_i],
                val,
                size=8,
                fill=fill,
                align="left" if c_i in (0, 8) else "center",
            )
    last = table.rows[1 + len(body)]
    for c_i, val in enumerate(total):
        fill = "C5D9F1" if c_i in (5, 7) else TOTAL_FILL
        _set_cell_text(last.cells[c_i], val, bold=True, size=8, fill=fill)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if chart_cons.exists():
        p.add_run().add_picture(str(chart_cons), width=Inches(10.2))

    doc.add_heading("Huecos horarios en Fundo Zapallar", level=1)
    t2 = doc.add_table(rows=1 + len(resumen_nodos), cols=5)
    t2.style = "Table Grid"
    for i, h in enumerate(["Punto", "Alta serie", "Días con hueco", "Horas faltantes", "Peor día"]):
        _set_cell_text(t2.rows[0].cells[i], h, bold=True, size=9, color=RGBColor(255, 255, 255), fill=HEADER_FILL)
    for i, s in enumerate(resumen_nodos):
        fill = "C6E0B4" if s["horas_faltantes"] == 0 else ("F4C7C3" if s["horas_faltantes"] > 48 else "FFFFFF")
        vals = [s["nombre"], s.get("alta", "—"), str(s["dias_con_hueco"]), str(s["horas_faltantes"]), s["peor_dia"]]
        for c, v in enumerate(vals):
            _set_cell_text(t2.rows[i + 1].cells[c], v, size=9, fill=fill, align="left" if c in (0, 4) else "center")

    doc.add_paragraph()
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if chart_huecos.exists():
        p2.add_run().add_picture(str(chart_huecos), width=Inches(10.2))

    doc.add_heading("Detalle Matriz ESVAL", level=2)
    if lineas_matriz:
        for line in lineas_matriz:
            doc.add_paragraph(line, style="List Bullet")
    else:
        doc.add_paragraph("Sin huecos horarios en la Matriz ESVAL en el periodo revisado.")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def primera_fecha_con_dato(idx: dict[tuple[date, int], float], start_d: date, end_d: date) -> date | None:
    d = start_d
    while d <= end_d:
        if any((d, h) in idx for h in range(24)):
            return d
        d += timedelta(days=1)
    return None


def resumen_nodo(nid: str, gaps: list[dict], alta: date | None) -> dict:
    nombre = NOMBRES.get(nid, nid)
    if not gaps:
        return {
            "node_id": nid,
            "nombre": nombre,
            "alta": alta.strftime("%d/%m/%Y") if alta else "—",
            "dias_con_hueco": 0,
            "horas_faltantes": 0,
            "peor_dia": "—",
        }
    peor = max(gaps, key=lambda g: (g["n"], g["date"]))
    return {
        "node_id": nid,
        "nombre": nombre,
        "alta": alta.strftime("%d/%m/%Y") if alta else "—",
        "dias_con_hueco": len(gaps),
        "horas_faltantes": int(sum(g["n"] for g in gaps)),
        "peor_dia": f"{peor['date'].strftime('%d/%m/%Y')} ({peor['n']} h)",
    }


def main() -> int:
    if sys.platform == "win32":
        for s in (sys.stdout, sys.stderr):
            try:
                s.reconfigure(encoding="utf-8", line_buffering=True)
            except Exception:
                pass

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out_dir = Path("reports/Fundo_Zapallar/ANALISIS") / f"tabla_matriz_dia_completo_{stamp}"
    out_dir.mkdir(parents=True, exist_ok=True)

    start_all = min(c["ini"] for c in CICLOS)
    # Incluye ayer (21/09/2026) para la revisión de huecos posterior a la última boleta.
    end_huecos = date(2026, 9, 21)

    print("[INFO] Descargando Matriz ESVAL (tabla + huecos)...", flush=True)
    mat_rows = load_hourly_range(NODO_MATRIZ, start_all, end_huecos)
    if not mat_rows:
        print("[ERROR] Sin medidas de la Matriz ESVAL.", flush=True)
        return 1
    mat_idx = index_chile_hours(mat_rows)
    ciclos = enriquecer_ciclos(mat_idx, mat_rows)
    altas: dict[str, date | None] = {}
    alta_mat = primera_fecha_con_dato(mat_idx, start_all, end_huecos)
    altas[NODO_MATRIZ] = alta_mat
    gaps_matriz = huecos_rango(mat_idx, alta_mat or start_all, end_huecos) if alta_mat else []
    lineas_matriz = colapsar_huecos(gaps_matriz)

    gaps_por_nodo: dict[str, list[dict]] = {NODO_MATRIZ: gaps_matriz}
    resumen_nodos = [resumen_nodo(NODO_MATRIZ, gaps_matriz, alta_mat)]

    otros = [n for n in FUNDO_ZAPALLAR_NODE_IDS if n != NODO_MATRIZ]
    for nid in otros:
        print(f"[INFO] Huecos {nid} {NOMBRES.get(nid, '')}...", flush=True)
        rows_n = load_hourly_range(nid, start_all, end_huecos)
        idx_n = index_chile_hours(rows_n)
        alta_n = primera_fecha_con_dato(idx_n, start_all, end_huecos)
        altas[nid] = alta_n
        gaps_n = huecos_rango(idx_n, alta_n or start_all, end_huecos) if alta_n else []
        gaps_por_nodo[nid] = gaps_n
        resumen_nodos.append(resumen_nodo(nid, gaps_n, alta_n))

    _, _, summary = construir_filas(ciclos)

    json_path = out_dir / "matriz_dia_completo_y_huecos.json"
    payload = {
        "generado": datetime.now().isoformat(timespec="seconds"),
        "metodo_wes": "día completo Chile [ini 00:00, fin 24:00] inclusive",
        "huecos_desde": start_all.isoformat(),
        "huecos_hasta": end_huecos.isoformat(),
        "alta_serie": {nid: (d.isoformat() if d else None) for nid, d in altas.items()},
        "ciclos": [
            {k: (v.isoformat() if isinstance(v, date) else v) for k, v in r.items()}
            for r in ciclos
        ],
        "totales": summary,
        "huecos_por_nodo": {
            nid: [
                {
                    "date": g["date"].isoformat(),
                    "n": g["n"],
                    "expected": g["expected"],
                    "missing": g["missing"],
                }
                for g in gaps
            ]
            for nid, gaps in gaps_por_nodo.items()
        },
        "resumen_nodos": resumen_nodos,
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    csv_path = out_dir / "huecos_horarios_fundo_zapallar.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["node_id", "nombre", "alta_serie", "fecha", "horas_faltantes", "horas_esperadas", "horas_missing"])
        for nid, gaps in gaps_por_nodo.items():
            alta_s = altas[nid].isoformat() if altas.get(nid) else ""
            for g in gaps:
                w.writerow(
                    [
                        nid,
                        NOMBRES.get(nid, nid),
                        alta_s,
                        g["date"].isoformat(),
                        g["n"],
                        g["expected"],
                        " ".join(f"{h:02d}" for h in g["missing"]),
                    ]
                )

    chart_cons = crear_grafico_consumo(ciclos, out_dir / "chart_delta_vs_wes_full.png")
    chart_huecos = crear_grafico_huecos(
        gaps_por_nodo, altas, start_all, end_huecos, out_dir / "chart_huecos_horarios.png"
    )

    docx_path = out_dir / "Tabla_Facturacion_vs_Matriz_dia_completo.docx"
    crear_word(ciclos, chart_cons, chart_huecos, resumen_nodos, lineas_matriz, docx_path)
    pdf_path = out_dir / "Tabla_Facturacion_vs_Matriz_dia_completo.pdf"
    crear_pdf(ciclos, chart_cons, chart_huecos, resumen_nodos, lineas_matriz, pdf_path)

    print(f"[OK] JSON {json_path}")
    print(f"[OK] CSV  {csv_path}")
    print(f"[OK] Word {docx_path}")
    print(f"[OK] PDF  {pdf_path}")
    print(
        f"[OK] Totales delta={summary['delta']:.0f}  "
        f"12h={summary['wes_12h']:.1f}  full={summary['wes_full']:.1f}  "
        f"dif={summary['dif_full_pct']:.1f}%  huecos_matriz={resumen_nodos[0]['horas_faltantes']}"
    )
    for r in ciclos:
        print(
            f"  {r['ciclo']}: full={r['wes_full']:.1f} vs delta={r['delta']} "
            f"({r['dif_full_pct']:+.1f}%)  huecos={r['huecos']} h  cob={r['cobertura_pct']}%"
        )
    print("Huecos por punto:")
    for s in resumen_nodos:
        print(f"  {s['nombre']}: alta {s.get('alta','—')} | {s['horas_faltantes']} h en {s['dias_con_hueco']} días | peor {s['peor_dia']}")
    if lineas_matriz:
        print("Detalle Matriz:")
        for line in lineas_matriz[:20]:
            print("   ", line)
    else:
        print("Matriz ESVAL sin huecos horarios.")
    print(f"OUT_DIR={out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
