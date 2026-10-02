"""
Perfil de consumo Weir / Vulco (San Bernardo) a partir de las boletas
Aguas Andinas de septiembre 2026.

El gráfico «Consumo últimos 13 meses» de cada boleta no trae tabla numérica.
Se mide la altura de cada barra y se escala para que el último mes coincida
con los m³ facturados en esa boleta.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent
BOLETAS = ROOT / "reports" / "Weir" / "boletas_202609"
OUT = ROOT / "reports" / "Weir" / "Perfil_consumo"
UPLOADS = Path("/home/ubuntu/.cursor/projects/workspace/uploads")

MESES = [
    "Sep-25",
    "Oct-25",
    "Nov-25",
    "Dic-25",
    "Ene-26",
    "Feb-26",
    "Mar-26",
    "Abr-26",
    "May-26",
    "Jun-26",
    "Jul-26",
    "Ago-26",
    "Sep-26",
]

pdfmetrics.registerFont(TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("DejaVuBold", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))

NAVY = (14 / 255, 58 / 255, 93 / 255)
TEAL = (0 / 255, 122 / 255, 140 / 255)
BAR = "#1F6AA5"
NAVY_HEX = RGBColor(14, 58, 93)


@dataclass
class Cuenta:
    src_name: str
    archivo: str
    titulo: str
    sitio: str
    direccion: str
    titular: str
    rut: str
    cuenta: str
    medidor: str
    documento: str
    lectura: str
    consumo_m3: int
    total_pagar: str
    clave: str
    nota_extra: str
    lectura_perfil: str


CUENTAS = [
    Cuenta(
        src_name="Baquedano_1215_9189.pdf",
        archivo="weir Baquedano 1215",
        titulo="Baquedano 1215",
        sitio="Estadio Vulco",
        direccion="Baquedano 1215, San Bernardo",
        titular="Vulco S.A.",
        rut="91.619.000-K",
        cuenta="295661-6",
        medidor="110176970",
        documento="Factura electrónica N° 9455958",
        lectura="31-07-2026 al 31-08-2026",
        consumo_m3=397,
        total_pagar="$ 550.552",
        clave="Consumo real · lectura ratificada",
        nota_extra="Despacho electrónico: weirbills01@weir.urjanet.com. Arranque 40 mm.",
        lectura_perfil=(
            "El consumo es bajo en primavera (cerca de 500–600 m³) y sube fuerte "
            "entre febrero y abril de 2026, cuando se acerca a 3.800–3.900 m³. "
            "Desde mayo baja y septiembre 2026 es el mes más bajo de la serie (397 m³, lectura real). "
            "El salto de verano respecto del mes actual es del orden de 10 veces."
        ),
    ),
    Cuenta(
        src_name="Almirante_Riveros_7621.pdf",
        archivo="weir Almirante Riveros",
        titulo="Almirante Riveros 1050",
        sitio="Vulco",
        direccion="Almirante Riveros 1050, San Bernardo",
        titular="Gómez González, Alejandro Javier",
        rut="No figura en la boleta (boleta, no factura)",
        cuenta="2713987-6",
        medidor="2.128",
        documento="Boleta electrónica N° 322122961",
        lectura="31-07-2026 al 31-08-2026 (sin lecturas de medidor)",
        consumo_m3=253,
        total_pagar="$ 351.198",
        clave="Consumo estimado · medidor cerrado",
        nota_extra=(
            "La boleta no trae lectura actual ni anterior. Informa consumo promedio "
            "descontable de 253 m³ y 12.144 m³ por abonar en la próxima facturación. "
            "Despacho: weirbills01@weir.urjanet.com."
        ),
        lectura_perfil=(
            "La serie del gráfico es plana: los 13 meses quedan en 253 m³. "
            "No es un consumo medido. Aguas Andinas facturó por estimación porque "
            "la clave de lectura es CERRADO. Hasta que haya una lectura real, "
            "este perfil no sirve para comparar meses ni para dimensionar un medidor WES "
            "contra la facturación."
        ),
    ),
    Cuenta(
        src_name="San_Jos__0895_850e.pdf",
        archivo="weir San José 0895",
        titulo="San José 0895",
        sitio="Matriz San José 895",
        direccion="San José 0895, San Bernardo",
        titular="Vulco S.A.",
        rut="91.619.000-K",
        cuenta="296813-4",
        medidor="123707699",
        documento="Factura electrónica N° 9455953",
        lectura="31-07-2026 al 31-08-2026",
        consumo_m3=4014,
        total_pagar="$ 5.466.709",
        clave="Consumo real · lectura normal",
        nota_extra=(
            "Arranque 50 mm. Indemnización por corte de agua potable el 01-09-2026 "
            "(2 h 20 min): −$ 91.255. Despacho: weirbills01@weir.urjanet.com."
        ),
        lectura_perfil=(
            "Es la cuenta de mayor volumen. Tras un septiembre 2025 más bajo "
            "(cerca de 2.600 m³), el resto del año se mueve entre unos 3.700 y 7.100 m³. "
            "El máximo está en marzo 2026. Los últimos seis meses se estabilizan "
            "alrededor de 3.700–4.400 m³. Septiembre 2026 facturó 4.014 m³ con lectura real."
        ),
    ),
    Cuenta(
        src_name="San_Jos__0815_b1f8.pdf",
        archivo="weir San José 0815",
        titulo="San José 0815",
        sitio="Medidor fundición",
        direccion="San José 0815, San Bernardo",
        titular="Vulco S.A.",
        rut="91.619.000-K",
        cuenta="296816-9",
        medidor="131163780",
        documento="Factura electrónica N° 9455955",
        lectura="30-05-2026 al 02-09-2026 (95 días)",
        consumo_m3=3646,
        total_pagar="$ 4.690.886",
        clave="Consumo real · lectura normal",
        nota_extra=(
            "Diferencia de lecturas 4.803 m³, abono de 1.157 m³ descontables, "
            "consumo facturado 3.646 m³. Indemnización por corte el 11-08-2026 "
            "(20 h 01 min): −$ 357.614. Arranque 40 mm. "
            "Despacho: hector.almendares@mail.weir."
        ),
        lectura_perfil=(
            "La barra de septiembre no es un mes calendario: la factura cubre del "
            "30-05-2026 al 02-09-2026 (95 días) y concentra 3.646 m³. "
            "Mayo y junio también aparecen altos (cerca de 3.100 y 2.300 m³). "
            "Agosto queda en torno a 280 m³. Conviene no usar septiembre como "
            "consumo mensual típico al comparar con las otras direcciones."
        ),
    ),
    Cuenta(
        src_name="San_Jos__0783_b9dd.pdf",
        archivo="weir San José 0783",
        titulo="San José 0783",
        sitio="Servicio San José 0783",
        direccion="San José 0783, San Bernardo",
        titular="Vulco S.A.",
        rut="91.619.000-K",
        cuenta="296817-7",
        medidor="122714081",
        documento="Factura electrónica N° 9457198",
        lectura="31-07-2026 al 31-08-2026",
        consumo_m3=1072,
        total_pagar="$ 1.485.313",
        clave="Consumo real · lectura normal",
        nota_extra=(
            "El membrete de la factura dice San José 0815 (domicilio de Vulco S.A.). "
            "En el cuerpo, Aguas Andinas imprime DIRECCIÓN: SAN JOSE 0783. "
            "Arranque 50 mm. Indemnización por corte el 11-08-2026 (4 h 55 min)."
        ),
        lectura_perfil=(
            "Consumo relativamente estable, entre unos 660 y 1.330 m³. "
            "El máximo es enero 2026 y el mínimo agosto 2026. "
            "Septiembre 2026 facturó 1.072 m³ con lectura real, en línea con el promedio del año."
        ),
    ),
]


def _boleta_path(cuenta: Cuenta) -> Path:
    local = BOLETAS / cuenta.src_name
    if local.is_file():
        return local
    return UPLOADS / cuenta.src_name


def medir_alturas(pdf: Path) -> list[int]:
    """Altura en píxeles de las 13 barras, de izquierda a derecha."""
    doc = pymupdf.open(pdf)
    page = doc[0]
    clip = pymupdf.Rect(18, 312, 172, 348)
    pix = page.get_pixmap(matrix=pymupdf.Matrix(6, 6), clip=clip, alpha=False)
    doc.close()
    arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
    if arr.shape[2] == 4:
        arr = arr[:, :, :3]
    r = arr[:, :, 0].astype(int)
    g = arr[:, :, 1].astype(int)
    b = arr[:, :, 2].astype(int)
    bar = (np.abs(r - g) < 12) & (np.abs(r - b) < 12) & (r >= 75) & (r <= 120)
    col = bar.sum(axis=0)
    spans: list[tuple[int, int]] = []
    inicio = None
    for i, activo in enumerate(col > 6):
        if activo and inicio is None:
            inicio = i
        elif not activo and inicio is not None:
            if 12 <= i - inicio <= 80:
                spans.append((inicio, i - 1))
            inicio = None
    barras: list[tuple[int, int]] = []
    bottoms: list[int] = []
    for x0, x1 in spans:
        densidad = bar[:, x0 : x1 + 1].sum(axis=1)
        rows = np.where(densidad > (x1 - x0) * 0.4)[0]
        if len(rows) < 3:
            continue
        barras.append((x0, x1))
        bottoms.append(int(rows.max()))
    if len(barras) != 13:
        raise RuntimeError(f"{pdf.name}: se esperaban 13 barras y hay {len(barras)}")
    base = int(np.median(bottoms))
    corregidas = []
    for x0, x1 in barras:
        densidad = bar[:, x0 : x1 + 1].sum(axis=1)
        rows = np.where(densidad > (x1 - x0) * 0.4)[0]
        top = int(rows.min())
        corregidas.append(max(base - top, 1))
    return corregidas


def serie_m3(alturas: list[int], consumo_ultimo: int, plana: bool) -> list[int]:
    if plana or max(alturas) - min(alturas) <= 2:
        return [consumo_ultimo] * 13
    escala = consumo_ultimo / alturas[-1]
    valores = [int(round(h * escala / 10.0) * 10) for h in alturas[:-1]]
    valores.append(consumo_ultimo)
    return valores


def fmt_m3(n: int) -> str:
    return f"{n:,}".replace(",", ".")


def stats(valores: list[int]) -> dict:
    ultimos_12 = valores[1:]
    return {
        "promedio_12": int(round(sum(ultimos_12) / 12)),
        "suma_12": sum(ultimos_12),
        "max": max(valores),
        "min": min(valores),
        "mes_max": MESES[valores.index(max(valores))],
        "mes_min": MESES[valores.index(min(valores))],
    }


def grafico(valores: list[int], titulo: str, destino: Path) -> None:
    fig, ax = plt.subplots(figsize=(10.2, 3.6), dpi=140)
    colores = [BAR] * 12 + ["#0E3A5D"]
    ax.bar(MESES, valores, color=colores, width=0.72, zorder=3)
    ax.set_ylabel("m³")
    ax.set_title(titulo, loc="left", fontsize=11, color="#0E3A5D", pad=8)
    ax.yaxis.grid(True, linestyle="--", linewidth=0.6, color="#D0D7DE", zorder=0)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(axis="x", labelrotation=40, labelsize=8)
    ax.tick_params(axis="y", labelsize=8)
    ymax = max(valores) * 1.18
    ax.set_ylim(0, ymax)
    for i, v in enumerate(valores):
        ax.text(i, v + ymax * 0.02, fmt_m3(v), ha="center", va="bottom", fontsize=6.5, color="#243140")
    fig.tight_layout()
    fig.savefig(destino, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def _wrap(c: canvas.Canvas, text: str, font: str, size: float, width: float) -> list[str]:
    words = text.split()
    lines: list[str] = []
    cur = ""
    for w in words:
        trial = w if not cur else cur + " " + w
        if c.stringWidth(trial, font, size) <= width:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def escribir_pdf(cuenta: Cuenta, valores: list[int], png: Path, destino: Path) -> None:
    s = stats(valores)
    c = canvas.Canvas(str(destino), pagesize=A4)
    w, h = A4
    c.setFillColorRGB(*NAVY)
    c.rect(0, h - 78, w, 78, fill=1, stroke=0)
    c.setFillColorRGB(1, 1, 1)
    c.setFont("DejaVuBold", 15)
    c.drawString(36, h - 34, f"Perfil de consumo · Weir · {cuenta.titulo}")
    c.setFont("DejaVu", 9)
    c.drawString(36, h - 52, "Vulco S.A. · San Bernardo · Aguas Andinas · últimos 13 meses de la boleta")
    c.setFont("DejaVu", 8)
    c.drawRightString(w - 36, h - 52, "Octubre 2026")

    y = h - 100
    c.setFillColorRGB(*NAVY)
    c.setFont("DejaVuBold", 11)
    c.drawString(36, y, "Cuenta")
    y -= 16
    c.setFillColorRGB(0.12, 0.16, 0.2)
    filas = [
        ("Dirección", cuenta.direccion),
        ("Sitio", cuenta.sitio),
        ("Titular", f"{cuenta.titular} · RUT {cuenta.rut}"),
        ("N° de cuenta", cuenta.cuenta),
        ("Medidor", cuenta.medidor),
        ("Documento", cuenta.documento),
        ("Período de lectura", cuenta.lectura),
        ("Consumo facturado", f"{fmt_m3(cuenta.consumo_m3)} m³"),
        ("Total a pagar", f"{cuenta.total_pagar} · vencimiento 26-09-2026"),
        ("Clave", cuenta.clave),
    ]
    c.setFont("DejaVu", 8.5)
    for etiqueta, valor in filas:
        c.setFont("DejaVuBold", 8.5)
        c.drawString(36, y, etiqueta)
        c.setFont("DejaVu", 8.5)
        c.drawString(150, y, valor)
        y -= 13

    y -= 6
    c.setFillColorRGB(*TEAL)
    c.setFont("DejaVuBold", 11)
    c.setFillColorRGB(*NAVY)
    c.drawString(36, y, "Últimos 13 meses")
    y -= 4
    img = ImageReader(str(png))
    iw, ih = Image.open(png).size
    draw_w = w - 72
    draw_h = draw_w * ih / iw
    y -= draw_h
    c.drawImage(img, 36, y, width=draw_w, height=draw_h, mask="auto")

    y -= 18
    c.setFont("DejaVu", 8.5)
    c.setFillColorRGB(0.12, 0.16, 0.2)
    resumen = (
        f"Últimos 12 meses (oct-25 a sep-26): {fmt_m3(s['suma_12'])} m³ en total, "
        f"promedio {fmt_m3(s['promedio_12'])} m³/mes. "
        f"Máximo {fmt_m3(s['max'])} m³ ({s['mes_max']}). "
        f"Mínimo {fmt_m3(s['min'])} m³ ({s['mes_min']})."
    )
    for line in _wrap(c, resumen, "DejaVu", 8.5, w - 72):
        c.drawString(36, y, line)
        y -= 12

    y -= 6
    c.setFont("DejaVuBold", 11)
    c.setFillColorRGB(*NAVY)
    c.drawString(36, y, "Cómo se lee")
    y -= 14
    c.setFillColorRGB(0.12, 0.16, 0.2)
    c.setFont("DejaVu", 8.5)
    for parrafo in (cuenta.lectura_perfil, cuenta.nota_extra):
        for line in _wrap(c, parrafo, "DejaVu", 8.5, w - 72):
            c.drawString(36, y, line)
            y -= 11
        y -= 4

    y -= 2
    c.setFont("DejaVu", 7.5)
    c.setFillColorRGB(0.35, 0.4, 0.45)
    nota = (
        "Fuente: boleta o factura Aguas Andinas emitida el 08-09-2026. "
        "Los meses previos al último se leyeron del gráfico «Consumo últimos 13 meses» "
        "(altura de la barra, ajustada para que septiembre 2026 coincida con los m³ facturados; "
        "precisión aproximada de ±1 barra). La barra de septiembre está en azul oscuro y es el dato de la boleta."
    )
    for line in _wrap(c, nota, "DejaVu", 7.5, w - 72):
        c.drawString(36, y, line)
        y -= 10
    c.setFont("DejaVu", 8)
    c.setFillColorRGB(*NAVY)
    c.drawString(36, 28, "WES · Perfil de consumo Weir Minerals · San Bernardo")
    c.showPage()
    c.save()


def escribir_docx(cuenta: Cuenta, valores: list[int], png: Path, destino: Path) -> None:
    s = stats(valores)
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)
    titulo = doc.add_heading(f"Perfil de consumo · Weir · {cuenta.titulo}", level=1)
    for run in titulo.runs:
        run.font.color.rgb = NAVY_HEX
    p = doc.add_paragraph("Vulco S.A. · San Bernardo · Aguas Andinas · boleta de septiembre 2026")
    p.runs[0].italic = True

    tabla = doc.add_table(rows=10, cols=2)
    tabla.style = "Table Grid"
    datos = [
        ("Dirección", cuenta.direccion),
        ("Sitio", cuenta.sitio),
        ("Titular", f"{cuenta.titular} · RUT {cuenta.rut}"),
        ("N° de cuenta", cuenta.cuenta),
        ("Medidor", cuenta.medidor),
        ("Documento", cuenta.documento),
        ("Período de lectura", cuenta.lectura),
        ("Consumo facturado", f"{fmt_m3(cuenta.consumo_m3)} m³"),
        ("Total a pagar", f"{cuenta.total_pagar} · vence 26-09-2026"),
        ("Clave", cuenta.clave),
    ]
    for i, (k, v) in enumerate(datos):
        tabla.rows[i].cells[0].text = k
        tabla.rows[i].cells[1].text = v

    doc.add_paragraph()
    doc.add_picture(str(png), width=Inches(6.4))
    cap = doc.add_paragraph(
        f"Últimos 12 meses (oct-25 a sep-26): {fmt_m3(s['suma_12'])} m³ · "
        f"promedio {fmt_m3(s['promedio_12'])} m³/mes · "
        f"máximo {fmt_m3(s['max'])} m³ ({s['mes_max']}) · "
        f"mínimo {fmt_m3(s['min'])} m³ ({s['mes_min']})."
    )
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_heading("Serie mensual (m³)", level=2)
    t2 = doc.add_table(rows=2, cols=13)
    t2.style = "Table Grid"
    for i, mes in enumerate(MESES):
        t2.rows[0].cells[i].text = mes
        t2.rows[1].cells[i].text = fmt_m3(valores[i])

    doc.add_heading("Cómo se lee", level=2)
    doc.add_paragraph(cuenta.lectura_perfil)
    doc.add_paragraph(cuenta.nota_extra)
    nota = doc.add_paragraph(
        "Los meses previos a septiembre 2026 se obtuvieron midiendo el gráfico "
        "«Consumo últimos 13 meses» de la boleta emitida el 08-09-2026 y ajustando "
        "la escala para que esa última barra coincida con los m³ facturados."
    )
    for run in nota.runs:
        run.font.size = Pt(9)
        run.italic = True
    pie = doc.add_paragraph("WES · Perfil de consumo Weir Minerals · San Bernardo · octubre 2026")
    for run in pie.runs:
        run.font.color.rgb = NAVY_HEX
        run.font.size = Pt(9)
    doc.save(destino)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "graficos").mkdir(exist_ok=True)
    for cuenta in CUENTAS:
        pdf_src = _boleta_path(cuenta)
        texto = "\n".join(page.get_text() for page in pymupdf.open(pdf_src))
        if cuenta.cuenta not in texto.replace(" ", ""):
            # el número puede venir con espacios; buscar dígitos
            digitos = re.sub(r"\D", "", cuenta.cuenta)
            if digitos not in re.sub(r"\D", "", texto):
                raise RuntimeError(f"No está la cuenta {cuenta.cuenta} en {pdf_src.name}")
        alturas = medir_alturas(pdf_src)
        plana = "estimado" in cuenta.clave.lower()
        valores = serie_m3(alturas, cuenta.consumo_m3, plana)
        png = OUT / "graficos" / f"{cuenta.archivo}.png"
        grafico(valores, f"{cuenta.titulo} · m³ por mes", png)
        pdf_out = OUT / f"{cuenta.archivo}.pdf"
        docx_out = OUT / f"{cuenta.archivo}.docx"
        escribir_pdf(cuenta, valores, png, pdf_out)
        escribir_docx(cuenta, valores, png, docx_out)
        print(cuenta.archivo)
        print(" ", list(zip(MESES, valores)))
        print(" ", stats(valores))
        print(" ", pdf_out)


if __name__ == "__main__":
    main()
