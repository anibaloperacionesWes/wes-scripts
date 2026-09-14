"""
Corte total 24 h — colegios CORMUP con control (Peñalolén).

Periodo operativo: desde el domingo indicado (Chile), las 24 horas del día.
No incluye Eduardo de la Barra, Alicura ni Likankura (sin control).

Uso:
  python cormup_corte_total.py              # programa Excel 24 h + verifica + informe
  python cormup_corte_total.py --restaurar  # vuelve horarios nocturnos originales
"""
from __future__ import annotations

import argparse
import json
import shutil
from dataclasses import dataclass, asdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from zoneinfo import ZoneInfo

import requests
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl import load_workbook

from control_nocturno import obtener_datos_horarios_dia, parse_horario_a_horas

CHILE = ZoneInfo("America/Santiago")
NODE_BASE = "http://104.248.53.141:7003/wes/api/acl-node/v1"
ROOT = Path(__file__).resolve().parent
EXCEL_HORARIOS = ROOT / "reports" / "HORARIOS CONTROL NOCTURNO.xlsx"
BACKUP_JSON = ROOT / "reports" / "CORMUP" / "horarios_cormup_antes_corte_total.json"
HORARIO_CORTE_TOTAL = "00:00 a 23:59"
COMPANY_ID = "000008"

# 11 colegios con control (lista acordada con Aníbal / Don David).
COLEGIOS_CONTROL: List[Tuple[str, str]] = [
    ("000008-01", "Lic. Antonio Hermida F"),
    ("000008-03", "Carlos Fernandez P."),
    ("000008-04", "Tobalaba"),
    ("000008-05", "Santa Maria"),
    ("000008-06", "Luis Arrieta Caña"),
    ("000008-07", "Erasmo Escala"),
    ("000008-09", "Juan Bautista Pasten"),
    ("000008-10", "Matilde Huici Navas"),
    ("000008-11", "CE Valle Hermoso"),
    ("000008-12", "Unión Nacional Árabe"),
    ("000008-14", "Juan Pablo II"),
]

SIN_CONTROL: List[Tuple[str, str]] = [
    ("000008-02", "Eduardo de la Barra"),
    ("000008-08", "Alicura"),
    ("000008-13", "Likankura"),
]

CONTROL_IDS = {nid for nid, _ in COLEGIOS_CONTROL}


@dataclass
class ResultadoNodo:
    node_id: str
    nombre: str
    con_control: bool
    wes_status: Optional[str]
    mch: Optional[str]
    lpm: Optional[str]
    last_update: Optional[str]
    is_controllable: Optional[bool]
    total_domingo_m3: float
    horas_consumo_domingo: List[int]
    max_hora_domingo: float
    total_lunes_m3: float
    horas_consumo_lunes: List[int]
    max_hora_lunes: float
    corte_24h_domingo: bool
    comando_off_http: Optional[int] = None


def _fmt_m3(v: float) -> str:
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _horas_txt(horas: List[int]) -> str:
    if not horas:
        return "—"
    if horas == list(range(24)):
        return "00:00-24:00 (todas)"
    # rangos compactos
    partes: List[str] = []
    start = prev = horas[0]
    for h in horas[1:]:
        if h == prev + 1:
            prev = h
            continue
        partes.append(f"{start:02d}:00-{prev + 1:02d}:00" if start != prev else f"{start:02d}:00")
        start = prev = h
    partes.append(f"{start:02d}:00-{prev + 1:02d}:00" if start != prev else f"{start:02d}:00")
    return ", ".join(partes)


def domingo_periodo(hoy: Optional[date] = None) -> date:
    hoy = hoy or datetime.now(CHILE).date()
    # Lunes → domingo anterior
    return hoy - __import__("datetime").timedelta(days=(hoy.weekday() + 1) % 7)


def leer_estado_nodo(node_id: str) -> Dict[str, Any]:
    out: Dict[str, Any] = {
        "wesStatus": None,
        "mch": None,
        "lpm": None,
        "lastUpdate": None,
        "isControllable": None,
        "ok": False,
    }
    try:
        r = requests.get(f"{NODE_BASE}/nodes/{node_id}", timeout=15)
        if r.status_code != 200:
            return out
        d = r.json() if r.content else {}
        out["ok"] = True
        out["wesStatus"] = d.get("wesStatus")
        out["mch"] = d.get("mch")
        out["lpm"] = d.get("lpm")
        out["lastUpdate"] = d.get("lastUpdate")
        out["isControllable"] = d.get("isControllable")
        return out
    except Exception:
        return out


def enviar_comando_off(node_id: str) -> int:
    """POST /nodes/status. En la API pública no cierra la válvula física; se registra el intento."""
    r = requests.post(
        f"{NODE_BASE}/nodes/status",
        json={"nodeId": node_id, "status": "OFF"},
        timeout=20,
    )
    return int(r.status_code)


def consumo_dia(node_id: str, dia: date) -> Tuple[float, List[int], float]:
    hourly = obtener_datos_horarios_dia(node_id, datetime.combine(dia, datetime.min.time()))
    total = 0.0
    horas: List[int] = []
    max_h = 0.0
    for h in range(24):
        v = float(hourly.get(h, 0.0) or 0.0)
        total += v
        if v > 0.001:
            horas.append(h)
            if v > max_h:
                max_h = v
    return total, horas, max_h


def backup_y_programar_excel(horario: str = HORARIO_CORTE_TOTAL) -> Dict[str, str]:
    if not EXCEL_HORARIOS.is_file():
        raise FileNotFoundError(EXCEL_HORARIOS)
    wb = load_workbook(EXCEL_HORARIOS)
    ws = wb[wb.sheetnames[0]]
    originales: Dict[str, str] = {}
    n = 0
    for row in ws.iter_rows(min_row=2):
        cells = list(row)
        if len(cells) < 5:
            continue
        cliente = str(cells[1].value or "").strip().upper()
        node_id = str(cells[3].value or "").strip()
        if cliente != "CORMUP" or node_id not in CONTROL_IDS:
            continue
        prev = "" if cells[4].value is None else str(cells[4].value).strip()
        originales[node_id] = prev
        cells[4].value = horario
        n += 1
    BACKUP_JSON.parent.mkdir(parents=True, exist_ok=True)
    if not BACKUP_JSON.exists():
        BACKUP_JSON.write_text(
            json.dumps(
                {
                    "excel": str(EXCEL_HORARIOS.name),
                    "guardado": datetime.now(CHILE).isoformat(timespec="seconds"),
                    "horarios_originales": originales,
                },
                indent=2,
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
    wb.save(EXCEL_HORARIOS)
    print(f"[OK] Excel: {n} colegios CORMUP con control → horario '{horario}'")
    return originales


def restaurar_excel() -> int:
    if not BACKUP_JSON.is_file():
        raise FileNotFoundError(f"No hay backup: {BACKUP_JSON}")
    data = json.loads(BACKUP_JSON.read_text(encoding="utf-8"))
    originales: Dict[str, str] = data.get("horarios_originales") or {}
    wb = load_workbook(EXCEL_HORARIOS)
    ws = wb[wb.sheetnames[0]]
    n = 0
    for row in ws.iter_rows(min_row=2):
        cells = list(row)
        if len(cells) < 5:
            continue
        node_id = str(cells[3].value or "").strip()
        if node_id in originales:
            cells[4].value = originales[node_id]
            n += 1
    wb.save(EXCEL_HORARIOS)
    print(f"[OK] Restaurados {n} horarios CORMUP desde {BACKUP_JSON.name}")
    return n


def evaluar_nodos(domingo: date, lunes: date, enviar_off: bool) -> List[ResultadoNodo]:
    out: List[ResultadoNodo] = []
    todos = [(nid, nom, True) for nid, nom in COLEGIOS_CONTROL] + [
        (nid, nom, False) for nid, nom in SIN_CONTROL
    ]
    for nid, nom, ctrl in todos:
        est = leer_estado_nodo(nid)
        tot_d, hrs_d, max_d = consumo_dia(nid, domingo)
        tot_l, hrs_l, max_l = consumo_dia(nid, lunes)
        code = None
        if ctrl and enviar_off:
            try:
                code = enviar_comando_off(nid)
            except Exception:
                code = -1
        corte = ctrl and tot_d <= 0.05 and len(hrs_d) <= 1
        out.append(
            ResultadoNodo(
                node_id=nid,
                nombre=nom,
                con_control=ctrl,
                wes_status=est.get("wesStatus"),
                mch=None if est.get("mch") is None else str(est.get("mch")),
                lpm=None if est.get("lpm") is None else str(est.get("lpm")),
                last_update=None if est.get("lastUpdate") is None else str(est.get("lastUpdate")),
                is_controllable=est.get("isControllable"),
                total_domingo_m3=tot_d,
                horas_consumo_domingo=hrs_d,
                max_hora_domingo=max_d,
                total_lunes_m3=tot_l,
                horas_consumo_lunes=hrs_l,
                max_hora_lunes=max_l,
                corte_24h_domingo=corte,
                comando_off_http=code,
            )
        )
        marca = "CTRL" if ctrl else "NO  "
        ok = "SI" if corte else "NO"
        print(
            f"  {nid} {marca} {nom[:22]:22} ves={est.get('wesStatus') or '-':3} "
            f"dom={tot_d:8.2f} m3 hrs={len(hrs_d):2}  24h={ok}  lun={tot_l:8.2f}"
        )
    return out


def _sombrear(cell, fill: str) -> None:
    try:
        shading = parse_xml(
            f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'w:val="clear" w:fill="{fill}"/>'
        )
        tc_pr = cell._element.get_or_add_tcPr()
        if tc_pr.find(qn("w:shd")) is None:
            tc_pr.append(shading)
    except Exception:
        pass


def _celda(cell, texto: str, *, bold: bool = False, center: bool = True, color: Optional[RGBColor] = None) -> None:
    cell.text = texto
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    if p.runs:
        p.runs[0].bold = bold
        p.runs[0].font.size = Pt(8)
        p.runs[0].font.color.rgb = color or RGBColor(0, 0, 0)


def crear_informe_word(
    filas: List[ResultadoNodo],
    domingo: date,
    lunes: date,
    out_docx: Path,
    horario_excel: str,
) -> Path:
    doc = Document()
    title = doc.add_heading("CORMUP — CORTE TOTAL 24 HORAS", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    title.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    gen = datetime.now(CHILE).strftime("%d-%m-%Y %H:%M")
    p = doc.add_paragraph(
        f"Peñalolén (company 000008)  |  Generado {gen} Chile\n"
        f"Ventana programada: {horario_excel}  |  Desde domingo {domingo:%d/%m/%Y} (24 h)"
    )
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p.runs[0].font.size = Pt(10)

    ctrl = [f for f in filas if f.con_control]
    ok = [f for f in ctrl if f.corte_24h_domingo]
    no = [f for f in ctrl if not f.corte_24h_domingo]
    valvula_on = [f for f in ctrl if (f.wes_status or "").upper() == "ON"]

    doc.add_heading("Resumen ejecutivo", 1)
    doc.add_paragraph(
        f"Se programó corte total (24 horas) en el Excel de horarios para los {len(ctrl)} colegios "
        f"CORMUP con control. El periodo parte el domingo {domingo:%d/%m/%Y} y sigue el lunes "
        f"{lunes:%d/%m/%Y} (día en curso).\n\n"
        f"Cumplen corte 24 h el domingo (consumo ≈ 0): {len(ok)} de {len(ctrl)}.\n"
        f"No cumplen (hubo consumo en el día): {len(no)}.\n"
        f"Válvula reportada ON ahora (wesStatus): {len(valvula_on)} de {len(ctrl)}.\n\n"
        "No se programan Eduardo de la Barra, Alicura ni Likankura (sin equipo de control)."
    )

    if no:
        doc.add_heading("Fuera de corte total (prioridad)", 1)
        for f in sorted(no, key=lambda x: x.total_domingo_m3, reverse=True):
            doc.add_paragraph(
                f"{f.nombre} ({f.node_id}): domingo {_fmt_m3(f.total_domingo_m3)} m³ "
                f"en {_horas_txt(f.horas_consumo_domingo)}; "
                f"lunes (parcial) {_fmt_m3(f.total_lunes_m3)} m³; "
                f"wesStatus={f.wes_status or 's/d'}, caudal {f.mch or 's/d'} m³/h, {f.lpm or 's/d'} L/min.",
                style="List Number",
            )

    doc.add_heading("Colegios con control — domingo y lunes", 1)
    headers = [
        "Colegio",
        "ID",
        "Válvula ahora",
        "Domingo m³",
        "Horas con consumo (dom)",
        "24 h OK",
        "Lunes m³ (parcial)",
    ]
    table = doc.add_table(rows=1 + len(ctrl), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        _celda(cell, h, bold=True, color=RGBColor(255, 255, 255))
        _sombrear(cell, "1F4E79")

    for r_i, f in enumerate(ctrl, start=1):
        vals = [
            f.nombre,
            f.node_id,
            f"{f.wes_status or 's/d'}" + (f"  {f.mch} m³/h" if f.mch else ""),
            _fmt_m3(f.total_domingo_m3),
            _horas_txt(f.horas_consumo_domingo),
            "Sí" if f.corte_24h_domingo else "No",
            _fmt_m3(f.total_lunes_m3),
        ]
        fill = "C6EFCE" if f.corte_24h_domingo else "FFC7CE"
        for c_i, v in enumerate(vals):
            cell = table.rows[r_i].cells[c_i]
            col = RGBColor(0, 97, 0) if f.corte_24h_domingo else RGBColor(156, 0, 6)
            _celda(cell, v, bold=(c_i == 5), color=col)
            _sombrear(cell, fill)

    doc.add_heading("Sin control (no programados)", 1)
    for f in [x for x in filas if not x.con_control]:
        nota_datos = "sin datos API" if f.last_update is None and f.wes_status is None else (
            f"wesStatus={f.wes_status or 's/d'}; domingo {_fmt_m3(f.total_domingo_m3)} m³"
        )
        doc.add_paragraph(f"{f.nombre} ({f.node_id}): {nota_datos}.", style="List Bullet")

    doc.add_heading("Qué se programó y qué no hace la API", 1)
    doc.add_paragraph(
        "1) Excel `HORARIOS CONTROL NOCTURNO.xlsx`: los 11 colegios con control quedaron en "
        f"'{horario_excel}' (24 horas). El control nocturno diario evaluará todo el día, no solo madrugada.\n"
        "2) POST /nodes/status con status=OFF: la API pública de WES acepta el llamado (HTTP 201) "
        "pero no cambia wesStatus ni cierra la válvula. El programa semanal de cada PIC se carga "
        "desde la app WES. Hay que dejar OFF las 24 h en la grilla de cada colegio con control.\n"
        "3) Para volver a horarios nocturnos habituales: python cormup_corte_total.py --restaurar"
    )

    horas_excel = parse_horario_a_horas(horario_excel)
    nota = doc.add_paragraph(
        f"Criterio '24 h OK': consumo domingo ≤ 0,05 m³ y a lo más 1 hora con registro > 0,001 m³. "
        f"Horas que cubre el Excel: {len(horas_excel)} (0–23)."
    )
    nota.runs[0].italic = True
    nota.runs[0].font.size = Pt(9)
    nota.runs[0].font.color.rgb = RGBColor(89, 89, 89)

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_docx))
    return out_docx


def guardar_pdf_simple(filas: List[ResultadoNodo], domingo: date, lunes: date, pdf_path: Path) -> Path:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    ctrl = [f for f in filas if f.con_control]
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    with PdfPages(pdf_path) as pdf:
        fig = plt.figure(figsize=(11.69, 8.27))
        ax = fig.add_subplot(111)
        ax.axis("off")
        ax.set_title("CORMUP — Corte total 24 horas (colegios con control)", fontsize=14, pad=12)
        headers = ["Colegio", "ID", "Válvula", "Domingo m³", "24 h OK", "Lunes m³"]
        data = [headers]
        for f in ctrl:
            data.append(
                [
                    f.nombre[:22],
                    f.node_id,
                    str(f.wes_status or "s/d"),
                    f"{f.total_domingo_m3:.2f}",
                    "Sí" if f.corte_24h_domingo else "No",
                    f"{f.total_lunes_m3:.2f}",
                ]
            )
        table = ax.table(cellText=data, loc="center", cellLoc="center")
        table.auto_set_font_size(False)
        table.set_fontsize(8)
        table.scale(1.2, 1.35)
        for j in range(len(headers)):
            table[0, j].set_facecolor("#1F4E79")
            table[0, j].set_text_props(color="white", weight="bold")
        for i, f in enumerate(ctrl, start=1):
            color = "#C6EFCE" if f.corte_24h_domingo else "#FFC7CE"
            for j in range(len(headers)):
                table[i, j].set_facecolor(color)
        ax.text(
            0.02,
            0.06,
            f"Periodo: domingo {domingo:%d/%m/%Y} (24 h) y lunes {lunes:%d/%m/%Y} parcial. "
            f"Excel programado 00:00 a 23:59. Cumplen {sum(1 for f in ctrl if f.corte_24h_domingo)}/{len(ctrl)}.",
            transform=ax.transAxes,
            fontsize=8,
        )
        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)
    return pdf_path


def convertir_pdf(docx_path: Path, filas: Optional[List[ResultadoNodo]] = None,
                  domingo: Optional[date] = None, lunes: Optional[date] = None) -> Optional[Path]:
    pdf = docx_path.with_suffix(".pdf")
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        import subprocess

        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
                check=True,
                timeout=120,
                capture_output=True,
            )
            if pdf.exists():
                return pdf
        except Exception:
            pass
    if filas is not None and domingo is not None and lunes is not None:
        return guardar_pdf_simple(filas, domingo, lunes, pdf)
    return None


def main() -> int:
    ap = argparse.ArgumentParser(description="Corte total 24 h CORMUP (colegios con control)")
    ap.add_argument("--restaurar", action="store_true", help="Restaurar horarios nocturnos del backup")
    ap.add_argument("--sin-comando-off", action="store_true", help="No llamar POST /nodes/status")
    ap.add_argument("--domingo", default="", help="YYYY-MM-DD (default: domingo anterior a hoy Chile)")
    args = ap.parse_args()

    if args.restaurar:
        restaurar_excel()
        return 0

    if args.domingo:
        domingo = date.fromisoformat(args.domingo)
    else:
        domingo = domingo_periodo()
    lunes = domingo + __import__("datetime").timedelta(days=1)

    print("=" * 72)
    print("CORMUP corte total 24 h")
    print(f"Domingo {domingo} → lunes {lunes} (parcial si es hoy)")
    print("=" * 72)

    originales = backup_y_programar_excel(HORARIO_CORTE_TOTAL)
    print(f"[INFO] Backup horarios: {BACKUP_JSON}")
    print("[INFO] Originales:", originales)

    print("\n[INFO] Consumo y estado de válvula")
    filas = evaluar_nodos(domingo, lunes, enviar_off=not args.sin_comando_off)

    ts = datetime.now(CHILE).strftime("%Y%m%d_%H%M")
    out_dir = ROOT / "reports" / "CORMUP" / "CORTE_TOTAL" / f"CORTE_TOTAL_{ts}"
    out_dir.mkdir(parents=True, exist_ok=True)
    json_path = out_dir / "resultado.json"
    json_path.write_text(
        json.dumps([asdict(f) for f in filas], indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    docx = crear_informe_word(
        filas,
        domingo,
        lunes,
        out_dir / f"Corte_Total_CORMUP_{domingo:%Y%m%d}_{lunes:%Y%m%d}.docx",
        HORARIO_CORTE_TOTAL,
    )
    print(f"[OK] Word: {docx}")
    pdf = convertir_pdf(docx, filas, domingo, lunes)
    if pdf:
        print(f"[OK] PDF: {pdf}")
    else:
        print("[AVISO] No se pudo generar PDF; se entrega Word.")
        pdf = None

    ctrl_ok = sum(1 for f in filas if f.con_control and f.corte_24h_domingo)
    ctrl_n = sum(1 for f in filas if f.con_control)
    print(f"\nCumplen 24 h domingo: {ctrl_ok}/{ctrl_n}")
    print(f"JSON: {json_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
