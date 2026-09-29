"""
Juan Pablo II (000008-14): 1ª columna Hora, consumo por fecha a la derecha.

Fila 28 = Total (00:00–23:59, suma horaria WES).
Fila 29 = Listado diario pegado por fecha (valores entregados por el usuario).

Uso:
  python generar_horario_jp2_hora_consumo.py
"""

from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from copy import copy

from wes_estilo_graficos_app import horas_api_chile
from wes_google_drive import credenciales_configuradas, subir_a_drive

NODE = "000008-14"
NOMBRE = "Juan Pablo II"
D0 = date(2026, 7, 29)
D1 = date(2026, 9, 28)
OUT_DIR = Path("reports/CORMUP/Facturaciones_vs_WES")
DRIVE_SUB = "CORMUP/Facturaciones_vs_WES"

# Valores diarios del listado (coma decimal chilena → float), alineados por fecha.
LISTADO = {
    date(2026, 7, 29): 4.17,
    date(2026, 7, 30): 5.16,
    date(2026, 7, 31): 4.87,
    date(2026, 8, 1): 5.2,
    date(2026, 8, 2): 0.02,
    date(2026, 8, 3): 7.29,
    date(2026, 8, 4): 6.22,
    date(2026, 8, 5): 5.49,
    date(2026, 8, 6): 4.81,
    date(2026, 8, 7): 5.02,
    date(2026, 8, 8): 0.41,
    date(2026, 8, 9): 0.03,
    date(2026, 8, 10): 4.94,
    date(2026, 8, 11): 4.14,
    date(2026, 8, 12): 3.3,
    date(2026, 8, 13): 3.15,
    date(2026, 8, 14): 4.01,
    date(2026, 8, 15): 0.57,
    date(2026, 8, 16): 0.04,
    date(2026, 8, 17): 2.98,
    date(2026, 8, 18): 3.02,
    date(2026, 8, 19): 3.67,
    date(2026, 8, 20): 5.14,
    date(2026, 8, 21): 3.43,
    date(2026, 8, 22): 0.4,
    date(2026, 8, 23): 0.03,
    date(2026, 8, 24): 4.62,
    date(2026, 8, 25): 4.2,
    date(2026, 8, 26): 5.26,
    date(2026, 8, 27): 3.7,
    date(2026, 8, 28): 3.99,
    date(2026, 8, 29): 0.03,
    date(2026, 8, 30): 0.03,
    date(2026, 8, 31): 0.03,
    date(2026, 9, 1): 4.33,
    date(2026, 9, 2): 3.86,
    date(2026, 9, 3): 4.26,
    date(2026, 9, 4): 4.09,
    date(2026, 9, 5): 0.03,
    date(2026, 9, 6): 0.03,
    date(2026, 9, 7): 4.54,
    date(2026, 9, 8): 3.65,
    date(2026, 9, 9): 4.16,
    date(2026, 9, 10): 3.88,
    date(2026, 9, 11): 4.16,
    date(2026, 9, 12): 0.03,
    date(2026, 9, 13): 0.03,
    date(2026, 9, 14): 3.48,
    date(2026, 9, 15): 4.24,
    date(2026, 9, 16): 4.09,
    date(2026, 9, 17): 4.59,
    date(2026, 9, 18): 4.14,
    date(2026, 9, 19): 0.03,
    date(2026, 9, 20): 0.03,
    date(2026, 9, 21): 4.16,
    date(2026, 9, 22): 4.4,
    date(2026, 9, 23): 3.91,
    date(2026, 9, 24): 4.07,
    date(2026, 9, 25): 3.99,
    date(2026, 9, 26): 0.03,
    date(2026, 9, 27): 0.03,
    date(2026, 9, 28): 4.5,
}

HDR = PatternFill("solid", fgColor="D9D9D9")
SUB = PatternFill("solid", fgColor="F2F2F2")
TOT = PatternFill("solid", fgColor="D9E1F2")
LIS = PatternFill("solid", fgColor="C6EFCE")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


NUM_FMT = "0.00"


def _num(v: float) -> float:
    """Número real con 2 decimales (punto) para que Excel/Sheets pueda sumar."""
    return round(float(v), 2)


def _dias_unicos(d0: date, d1: date) -> list[date]:
    dias: list[date] = []
    seen: set[date] = set()
    d = d0
    while d <= d1:
        if d not in seen:
            seen.add(d)
            dias.append(d)
        d += timedelta(days=1)
    return dias


def _cargar_horas_xlsx(path: Path) -> dict[date, dict[int, float]]:
    """Reusa horas ya exportadas para no volver a consultar días viejos."""
    out: dict[date, dict[int, float]] = {}
    if not path.is_file():
        return out
    wb = load_workbook(path, data_only=False)
    ws = wb.active
    for col in range(2, ws.max_column + 1):
        hdr = str(ws.cell(2, col).value or "")
        try:
            dia = datetime.strptime(hdr.replace("Fecha", "").strip(), "%d/%m/%Y").date()
        except ValueError:
            continue
        horas: dict[int, float] = {}
        for h in range(24):
            v = ws.cell(4 + h, col).value
            if isinstance(v, (int, float)):
                horas[h] = float(v)
            elif isinstance(v, str) and v and not v.startswith("="):
                horas[h] = float(v.replace(",", "."))
            else:
                horas[h] = 0.0
        out[dia] = horas
    return out


def _horas_rango(dias: list[date], cache: dict[date, dict[int, float]] | None = None) -> dict[date, dict[int, float]]:
    horas: dict[date, dict[int, float]] = dict(cache or {})
    for dia in dias:
        if dia in horas and len(horas[dia]) == 24:
            continue
        try:
            h = horas_api_chile(NODE, datetime.combine(dia, datetime.min.time()))
            horas[dia] = {i: float(h.get(i, 0.0)) for i in range(24)}
            print(f"  API {dia.isoformat()}", flush=True)
        except Exception as exc:
            print(f"  API FAIL {dia.isoformat()}: {exc}", flush=True)
            horas[dia] = {i: 0.0 for i in range(24)}
    return horas


def construir_horario(dias: list[date], horas: dict[date, dict[int, float]]) -> Workbook:

    wb = Workbook()
    ws = wb.active
    ws.title = "Horario"

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=1 + min(8, len(dias)))
    ws["A1"] = (
        f"{NOMBRE} ({NODE}) — 1ª columna Hora; consumo por fecha a la derecha "
        f"({D0.strftime('%d-%m')} a {D1.strftime('%d-%m')}). Día repetido no se suma. "
        "Fila bajo Total = listado diario según fecha."
    )
    ws["A1"].font = Font(bold=True, size=12, color="003366")
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)
    c = ws.cell(2, 1, "Hora")
    c.fill = HDR
    c.font = Font(bold=True, size=11)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = THIN
    ws.cell(3, 1).border = THIN
    ws.cell(3, 1).fill = HDR

    for i, dia in enumerate(dias):
        col = 2 + i
        top = ws.cell(2, col, f"Fecha  {dia.strftime('%d/%m/%Y')}")
        top.fill = HDR
        top.font = Font(bold=True, size=9)
        top.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        top.border = THIN
        sub = ws.cell(3, col, "Consumo")
        sub.fill = SUB
        sub.font = Font(bold=True, size=9)
        sub.alignment = CENTER
        sub.border = THIN
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.column_dimensions["A"].width = 12
    ws.row_dimensions[2].height = 22

    for h in range(24):
        r = 4 + h
        bg = PatternFill("solid", fgColor="FFFFFF" if h % 2 == 0 else "FCFCFC")
        a = ws.cell(r, 1, f"{h:02d}:00")
        a.alignment = CENTER
        a.border = THIN
        a.fill = bg
        a.font = Font(bold=True, size=9)
        for i, dia in enumerate(dias):
            b = ws.cell(r, 2 + i, _num(horas[dia][h]))
            b.number_format = NUM_FMT
            b.alignment = CENTER
            b.border = THIN
            b.fill = bg

    rt = 28
    t = ws.cell(rt, 1, "Total")
    t.font = Font(bold=True, size=9)
    t.alignment = CENTER
    t.border = THIN
    t.fill = TOT
    for i, dia in enumerate(dias):
        col = get_column_letter(2 + i)
        b = ws.cell(rt, 2 + i, f"=SUM({col}4:{col}27)")
        b.number_format = NUM_FMT
        b.font = Font(bold=True, size=9)
        b.alignment = CENTER
        b.border = THIN
        b.fill = TOT

    rl = 29
    lab = ws.cell(rl, 1, "Listado")
    lab.font = Font(bold=True, size=9)
    lab.alignment = CENTER
    lab.border = THIN
    lab.fill = LIS
    for i, dia in enumerate(dias):
        val = LISTADO.get(dia)
        cell = ws.cell(rl, 2 + i, _num(val) if val is not None else None)
        if val is not None:
            cell.number_format = NUM_FMT
        cell.font = Font(bold=True, size=9)
        cell.alignment = CENTER
        cell.border = THIN
        cell.fill = LIS

    ws.freeze_panes = "B4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.print_title_rows = "1:3"
    return wb


def _copiar_a_general(ws) -> Path | None:
    gen = OUT_DIR / "Lecturas_vs_WES_medio_dia_CORMUP_20260929_1418.xlsx"
    if not gen.is_file():
        return None
    gwb = load_workbook(gen)
    if "JP2 horario" in gwb.sheetnames:
        del gwb["JP2 horario"]
    gws = gwb.create_sheet("JP2 horario", 1)
    for row in ws.iter_rows():
        for c in row:
            dest = gws.cell(c.row, c.column, c.value)
            dest.font = copy(c.font)
            dest.fill = copy(c.fill)
            dest.alignment = copy(c.alignment)
            dest.border = copy(c.border)
            dest.number_format = c.number_format
    for m in ws.merged_cells.ranges:
        gws.merge_cells(str(m))
    gws.column_dimensions["A"].width = 12
    for i in range(2, ws.max_column + 1):
        gws.column_dimensions[get_column_letter(i)].width = 14
    gws.freeze_panes = "B4"
    gws.row_dimensions[1].height = 32
    gwb.save(gen)
    return gen


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    dias = _dias_unicos(D0, D1)
    prev = sorted(OUT_DIR.glob("Horario_JP2_hora_consumo_*.xlsx"))
    cache = _cargar_horas_xlsx(prev[-1]) if prev else {}
    print(f"[INFO] cache {len(cache)} días; faltan {sum(1 for d in dias if d not in cache)}", flush=True)
    horas = _horas_rango(dias, cache)
    wb = construir_horario(dias, horas)
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    out = OUT_DIR / f"Horario_JP2_hora_consumo_{ts}.xlsx"
    wb.save(out)
    print("XLSX", out, "cols", wb.active.max_column)
    gen = _copiar_a_general(wb.active)
    if credenciales_configuradas():
        info = subir_a_drive(out, subcarpeta=DRIVE_SUB)
        print("DRIVE", info["web_view_link"])
        if gen is not None:
            info2 = subir_a_drive(gen, subcarpeta=DRIVE_SUB)
            print("DRIVE general", info2["web_view_link"])


if __name__ == "__main__":
    main()
