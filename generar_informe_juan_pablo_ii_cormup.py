"""
Informe formal — Colegio Juan Pablo II (CORMUP), nodo 000008-14.

Cruza las facturaciones Aguas Andinas de lectura real (cuenta 364260-7,
medidor 130738981) con el consumo del nodo WES y grafica el histórico y
la diferencia mensual (m³ facturados − m³ WES).

Las boletas y sus m³ provienen del comparativo CORMUP / Peñalolén
(facturas en Drive, lectura real). El m³ WES se vuelve a leer de la API
para el mismo intervalo de lecturas. La boleta a promedio no entra al
cálculo.

Uso:
  python generar_informe_juan_pablo_ii_cormup.py
"""

from __future__ import annotations

import shutil
import subprocess
import sys
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import List, Optional

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor

from generar_reporte_word import (
    acl_node_base_url,
    add_formatted_heading,
    add_logo_to_header,
    add_picture_with_pagination,
    estilizar_tabla_wes,
    fetch_json,
    flatten_measures,
    format_number_chilean,
    normalize_measures_payload,
)
from wes_estilo_graficos_app import COLOR_BARRA_FACT, COLOR_BARRA_WES

ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "reports" / "CORMUP" / "Juan_Pablo_II"
NODE_ID = "000008-14"
NODE_NOMBRE = "Juan Pablo II"
CUENTA = "364260-7"
MEDIDOR = "130738981"
TARIFA_CLP_M3 = 1270  # KPI vigente del nodo desde 2024-12-09
MESES = (
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
)
MESES_CORTOS = {
    1: "ene",
    2: "feb",
    3: "mar",
    4: "abr",
    5: "may",
    6: "jun",
    7: "jul",
    8: "ago",
    9: "sep",
    10: "oct",
    11: "nov",
    12: "dic",
}


@dataclass
class Periodo:
    lectura_anterior: date
    lectura_actual: date
    emision: date
    boleta: str
    m3_cuenta: int
    total_pagar: int
    estimado: bool


# Lecturas reales extraídas de las boletas Aguas Andinas (comparativo 21-09-2026).
# La de septiembre 2026 está facturada a promedio SISS: se informa y no se compara.
PERIODOS: List[Periodo] = [
    Periodo(date(2025, 12, 4), date(2026, 1, 2), date(2026, 1, 10), "9004459", 10, 14250, False),
    Periodo(date(2026, 1, 2), date(2026, 2, 2), date(2026, 2, 9), "9062561", 27, 36940, False),
    Periodo(date(2026, 2, 2), date(2026, 3, 5), date(2026, 3, 10), "9119121", 28, 38670, False),
    Periodo(date(2026, 3, 5), date(2026, 4, 2), date(2026, 4, 10), "9174713", 18, 64460, False),
    Periodo(date(2026, 4, 2), date(2026, 5, 4), date(2026, 5, 11), "9231585", 19, 28120, False),
    Periodo(date(2026, 5, 4), date(2026, 6, 3), date(2026, 6, 9), "9289980", 36, 79240, False),
    Periodo(date(2026, 6, 3), date(2026, 7, 3), date(2026, 7, 9), "9345231", 23, 33090, False),
    Periodo(date(2026, 7, 3), date(2026, 8, 3), date(2026, 8, 10), "9403874", 35, 43560, False),
    Periodo(date(2026, 8, 3), date(2026, 9, 3), date(2026, 9, 9), "9458934", 26, 31780, True),
]


def _fmt_m3(v: float, dec: int = 1) -> str:
    return format_number_chilean(v, dec)


def _fmt_clp(v: float) -> str:
    return "$" + f"{v:,.0f}".replace(",", ".")


def _fmt_fecha(d: date) -> str:
    return f"{d.day:02d}-{MESES_CORTOS[d.month]}-{d.year}"


def _etiqueta(p: Periodo) -> str:
    return f"{MESES_CORTOS[p.emision.month]}-{p.emision.year % 100:02d}"


def _fechas(a: date, b: date) -> list[date]:
    out = []
    d = a
    while d <= b:
        out.append(d)
        d += timedelta(days=1)
    return out


def _medidas(node_id: str, d0: date, d1: date):
    raw = fetch_json(
        f"{acl_node_base_url()}/nodes/measures/dates",
        params=[
            ("id", node_id),
            ("start", d0.strftime("%d%m%Y")),
            ("end", (d1 + timedelta(days=1)).strftime("%d%m%Y")),
        ],
    )
    return flatten_measures(normalize_measures_payload(raw, node_id))


def _mes_wes(year: int, month: int) -> tuple[float, int]:
    import calendar

    last = calendar.monthrange(year, month)[1]
    d0, d1 = date(year, month, 1), date(year, month, last)
    if d1 > date.today():
        d1 = date.today()
    meas = _medidas(NODE_ID, d0, d1)
    por_dia = {}
    for m in meas:
        dia = m.date.date()
        if d0 <= dia <= d1:
            por_dia[dia] = float(m.total_m3)
    return float(sum(por_dia.values())), len(por_dia)


def cruzar(periodos: List[Periodo]) -> list[dict]:
    validos = [p for p in periodos if not p.estimado]
    d0 = min(p.lectura_anterior for p in periodos)
    d1 = max(p.lectura_actual for p in periodos)
    meas = _medidas(NODE_ID, d0, d1)
    filas = []
    for p in periodos:
        esperadas = _fechas(p.lectura_anterior, p.lectura_actual)
        por_dia = {
            m.date.date(): float(m.total_m3)
            for m in meas
            if p.lectura_anterior <= m.date.date() <= p.lectura_actual
        }
        m3_wes = float(sum(por_dia.values()))
        huecos = [d for d in esperadas if d not in por_dia]
        dias_wes = len(por_dia)
        prom = (m3_wes / dias_wes) if dias_wes >= 3 else 0.0
        proy = prom * len(huecos) if huecos else 0.0
        wes_comp = m3_wes + proy
        diff = None if p.estimado else float(p.m3_cuenta) - wes_comp
        pct = (100.0 * diff / p.m3_cuenta) if diff is not None and p.m3_cuenta else None
        filas.append(
            {
                "p": p,
                "dias": len(esperadas),
                "dias_wes": dias_wes,
                "dias_hueco": len(huecos),
                "m3_wes": m3_wes,
                "m3_proy": proy,
                "m3_wes_comp": wes_comp,
                "diff": diff,
                "pct": pct,
            }
        )
    return filas


def _estilo_ejes(ax) -> None:
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(axis="y", linestyle=":", alpha=0.45)
    ax.set_axisbelow(True)


def grafico_historico(filas: list[dict], out: Path) -> None:
    validas = [f for f in filas if not f["p"].estimado]
    labels = [_etiqueta(f["p"]) for f in validas]
    cuenta = [f["p"].m3_cuenta for f in validas]
    wes = [f["m3_wes_comp"] for f in validas]
    x = range(len(labels))
    w = 0.38
    fig, ax = plt.subplots(figsize=(10.2, 4.8), dpi=140)
    fig.patch.set_facecolor("white")
    b1 = ax.bar([i - w / 2 for i in x], cuenta, width=w, color=COLOR_BARRA_FACT, label="m³ facturados (lectura real)")
    b2 = ax.bar([i + w / 2 for i in x], wes, width=w, color=COLOR_BARRA_WES, label="m³ WES (medido + huecos)")
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels, fontsize=9)
    ax.set_ylabel("m³ del período de lectura")
    ax.set_title("Juan Pablo II — histórico facturación vs consumo WES\nNodo 000008-14 · cuenta 364260-7")
    ax.legend(frameon=False, fontsize=9)
    ax.bar_label(b1, labels=[_fmt_m3(v, 0) for v in cuenta], fontsize=7, padding=2)
    ax.bar_label(b2, labels=[_fmt_m3(v, 0) for v in wes], fontsize=7, padding=2)
    _estilo_ejes(ax)
    fig.tight_layout()
    fig.savefig(out, bbox_inches="tight")
    plt.close(fig)


def grafico_diferencia(filas: list[dict], out: Path) -> None:
    validas = [f for f in filas if f["diff"] is not None]
    labels = [_etiqueta(f["p"]) for f in validas]
    diffs = [f["diff"] for f in validas]
    colors = ["#1d7a46" if v >= 0 else "#b42318" for v in diffs]
    fig, ax = plt.subplots(figsize=(10.2, 4.8), dpi=140)
    fig.patch.set_facecolor("white")
    bars = ax.bar(labels, diffs, color=colors, width=0.62)
    ax.axhline(0, color="#334155", linewidth=0.8)
    ax.set_ylabel("m³  (facturados − WES)")
    ax.set_title(
        "Diferencia mensual: facturación − consumo WES\n"
        "Positivo = la cuenta supera a WES; negativo = WES supera a la cuenta"
    )
    ax.bar_label(bars, labels=[_fmt_m3(v, 0) for v in diffs], fontsize=8, padding=3)
    _estilo_ejes(ax)
    fig.tight_layout()
    fig.savefig(out, bbox_inches="tight")
    plt.close(fig)


def grafico_valorizacion(filas: list[dict], out: Path) -> None:
    validas = [f for f in filas if f["diff"] is not None]
    labels = [_etiqueta(f["p"]) for f in validas]
    clp = [f["diff"] * TARIFA_CLP_M3 for f in validas]
    colors = ["#1d7a46" if v >= 0 else "#b42318" for v in clp]
    fig, ax = plt.subplots(figsize=(10.2, 4.8), dpi=140)
    fig.patch.set_facecolor("white")
    bars = ax.bar(labels, [v / 1000 for v in clp], color=colors, width=0.62)
    ax.axhline(0, color="#334155", linewidth=0.8)
    ax.set_ylabel("Miles de CLP")
    ax.set_title(
        f"Valorización de la diferencia mensual\n"
        f"Tarifa de referencia del nodo: {_fmt_clp(TARIFA_CLP_M3)} / m³"
    )
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"{v:,.0f}".replace(",", ".")))
    ax.bar_label(
        bars,
        labels=[_fmt_clp(v) for v in clp],
        fontsize=7,
        padding=3,
    )
    _estilo_ejes(ax)
    fig.tight_layout()
    fig.savefig(out, bbox_inches="tight")
    plt.close(fig)


def grafico_serie_mensual(serie: list[tuple[str, float]], out: Path) -> None:
    labels = [s[0] for s in serie]
    vals = [s[1] for s in serie]
    fig, ax = plt.subplots(figsize=(10.2, 4.6), dpi=140)
    fig.patch.set_facecolor("white")
    ax.plot(labels, vals, color=COLOR_BARRA_WES, marker="o", linewidth=1.8)
    ax.fill_between(range(len(vals)), vals, color=COLOR_BARRA_WES, alpha=0.15)
    ax.set_xticks(range(len(labels)))
    ax.set_xticklabels(labels, rotation=55, ha="right", fontsize=7)
    ax.set_ylabel("m³ / mes")
    ax.set_title("Consumo mensual registrado por WES — Juan Pablo II (000008-14)")
    _estilo_ejes(ax)
    fig.tight_layout()
    fig.savefig(out, bbox_inches="tight")
    plt.close(fig)


def _para(doc: Document, text: str, *, bold: bool = False, size: int = 11) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(30, 41, 59)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)


def _tabla(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(8)
        run.font.name = "Calibri"
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            run.font.size = Pt(8)
            run.font.name = "Calibri"
            run.bold = i == len(rows)
    estilizar_tabla_wes(table, has_total_row=True)


def escribir_word(
    filas: list[dict],
    serie: list[tuple[str, float]],
    pngs: dict[str, Path],
    out_docx: Path,
) -> None:
    validas = [f for f in filas if f["diff"] is not None]
    estimada = next(f for f in filas if f["p"].estimado)
    sum_cuenta = sum(f["p"].m3_cuenta for f in validas)
    sum_wes = sum(f["m3_wes_comp"] for f in validas)
    sum_diff = sum_cuenta - sum_wes
    pct = 100.0 * sum_diff / sum_cuenta if sum_cuenta else 0.0
    generado = datetime.now()

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.6)
        section.right_margin = Cm(1.6)
    add_logo_to_header(doc)

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = titulo.add_run("INFORME DE AUDITORÍA Y RENDIMIENTO")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0, 51, 102)
    r.font.name = "Calibri"

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("Colegio Juan Pablo II  ·  CORMUP / Peñalolén\nNodo WES 000008-14")
    rs.font.size = Pt(12)
    rs.font.color.rgb = RGBColor(0, 51, 102)
    rs.font.name = "Calibri"

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rm = meta.add_run(
        f"Cuenta Aguas Andinas {CUENTA}  ·  Medidor {MEDIDOR}\n"
        f"Generado el {generado.strftime('%d-%m-%Y %H:%M')}"
    )
    rm.font.size = Pt(10)
    rm.font.name = "Calibri"
    rm.font.color.rgb = RGBColor(71, 85, 105)

    add_formatted_heading(doc, "1. Alcance", 1)
    _para(
        doc,
        "Este informe presenta el histórico de facturación de agua potable del Colegio Juan Pablo II "
        "frente al consumo registrado por el nodo WES 000008-14, y la diferencia mensual entre ambos. "
        "La comparación usa solo boletas con lectura real del medidor. La boleta emitida a promedio "
        "(consumo estimado SISS) se informa aparte y no entra al total ni a los gráficos.",
    )
    _para(
        doc,
        "La diferencia se calcula como m³ facturados menos m³ WES del mismo intervalo de lecturas. "
        "Si en ese intervalo faltan días en la API, esos días se proyectan con el promedio diario "
        "de los días con dato del mismo período (se exige al menos 3 días de base) y se suman al WES. "
        f"La valorización usa la tarifa vigente del nodo: {_fmt_clp(TARIFA_CLP_M3)} por m³ "
        "(configuración KPI desde el 09-12-2024). No es el total de la boleta: la cuenta incluye cargo fijo y otros ítems.",
    )

    add_formatted_heading(doc, "2. Resultado", 1)
    if sum_diff < 0:
        lectura = (
            f"En los {len(validas)} períodos con lectura real, la cuenta registra "
            f"{_fmt_m3(sum_cuenta)} m³ y WES {_fmt_m3(sum_wes)} m³. "
            f"La diferencia acumulada es {_fmt_m3(sum_diff)} m³ ({pct:.1f}% respecto de la cuenta): "
            "el nodo mide más agua que la que factura Aguas Andinas. "
            "No hay un ahorro de la boleta frente al registro WES; la brecha es de sobre-registro del nodo "
            "(o de sub-registro del medidor de la compañía) y se mantiene en todos los meses válidos."
        )
    else:
        lectura = (
            f"En los {len(validas)} períodos con lectura real, la cuenta registra "
            f"{_fmt_m3(sum_cuenta)} m³ y WES {_fmt_m3(sum_wes)} m³. "
            f"La diferencia acumulada es {_fmt_m3(sum_diff)} m³ ({pct:.1f}%)."
        )
    _para(doc, lectura)
    _para(
        doc,
        f"Valorización de esa diferencia a {_fmt_clp(TARIFA_CLP_M3)}/m³: {_fmt_clp(sum_diff * TARIFA_CLP_M3)}. "
        f"La boleta de {_fmt_fecha(estimada['p'].emision)} (N° {estimada['p'].boleta}, "
        f"{estimada['p'].m3_cuenta} m³, {_fmt_clp(estimada['p'].total_pagar)}) está facturada a promedio "
        "y queda fuera del comparativo.",
    )

    add_formatted_heading(doc, "3. Histórico por período de lectura", 1)
    headers = [
        "Período de lecturas",
        "Emisión",
        "Boleta",
        "m³ cuenta",
        "m³ WES",
        "Huecos (d)",
        "m³ WES+proy",
        "Dif. m³",
        "% vs cuenta",
        "Total boleta",
    ]
    rows = []
    for f in filas:
        p: Periodo = f["p"]
        if p.estimado:
            dif = "—"
            pct_txt = "—"
            nota_wes = _fmt_m3(f["m3_wes"])
        else:
            dif = _fmt_m3(f["diff"])
            pct_txt = f"{f['pct']:.1f}".replace(".", ",")
            nota_wes = _fmt_m3(f["m3_wes"])
        rows.append(
            [
                f"{_fmt_fecha(p.lectura_anterior)} → {_fmt_fecha(p.lectura_actual)}",
                _fmt_fecha(p.emision),
                p.boleta + (" *" if p.estimado else ""),
                str(p.m3_cuenta),
                nota_wes,
                str(f["dias_hueco"]),
                "—" if p.estimado else _fmt_m3(f["m3_wes_comp"]),
                dif,
                pct_txt,
                _fmt_clp(p.total_pagar),
            ]
        )
    rows.append(
        [
            "TOTAL lecturas reales",
            "",
            "",
            _fmt_m3(sum_cuenta, 0),
            _fmt_m3(sum(f["m3_wes"] for f in validas)),
            str(sum(f["dias_hueco"] for f in validas)),
            _fmt_m3(sum_wes),
            _fmt_m3(sum_diff),
            f"{pct:.1f}".replace(".", ","),
            _fmt_clp(sum(f["p"].total_pagar for f in validas)),
        ]
    )
    _tabla(doc, headers, rows)
    _para(
        doc,
        "* Boleta a promedio / estimado: no se suma al total ni se grafica. "
        "Los tres primeros períodos tienen días sin dato en la API; la columna m³ WES+proy incorpora esa proyección.",
        size=9,
    )

    add_formatted_heading(doc, "4. Gráficos", 1)
    add_formatted_heading(doc, "4.1 Facturación y consumo WES", 2)
    _para(
        doc,
        "Cada barra agrupa el mes de emisión de la boleta. En naranja, los m³ de la lectura real; "
        "en azul, el consumo WES del mismo intervalo, con la proyección de huecos cuando corresponde.",
    )
    add_picture_with_pagination(doc, str(pngs["historico"]), Inches(6.3))

    add_formatted_heading(doc, "4.2 Diferencia mensual", 2)
    _para(
        doc,
        "La diferencia mensual es m³ facturados menos m³ WES (con proyección de huecos). "
        "Un valor positivo sería un volumen facturado por sobre el registro WES. "
        "En Juan Pablo II todos los meses válidos son negativos: WES queda por encima de la cuenta. "
        "Esa brecha no debe leerse como ahorro de agua del establecimiento.",
    )
    add_picture_with_pagination(doc, str(pngs["diferencia"]), Inches(6.3))

    add_formatted_heading(doc, "4.3 Valorización de la diferencia", 2)
    _para(
        doc,
        f"La misma diferencia, en pesos, a {_fmt_clp(TARIFA_CLP_M3)} por m³. "
        "Sirve para dimensionar la brecha. No reemplaza el total a pagar de la boleta.",
    )
    add_picture_with_pagination(doc, str(pngs["clp"]), Inches(6.3))

    add_formatted_heading(doc, "4.4 Histórico mensual del nodo WES", 2)
    _para(
        doc,
        "Consumo mensual que devuelve la API del nodo, independiente del corte de la boleta. "
        "Los meses de vacaciones y los primeros meses de medición quedan bajos o en cero.",
    )
    add_picture_with_pagination(doc, str(pngs["serie"]), Inches(6.3))

    add_formatted_heading(doc, "5. Conclusión", 1)
    _para(
        doc,
        f"Entre {_fmt_fecha(validas[0]['p'].lectura_anterior)} y {_fmt_fecha(validas[-1]['p'].lectura_actual)}, "
        f"las lecturas reales suman {_fmt_m3(sum_cuenta)} m³ facturados contra {_fmt_m3(sum_wes)} m³ WES "
        f"(diferencia {_fmt_m3(sum_diff)} m³, {_fmt_clp(sum_diff * TARIFA_CLP_M3)} a tarifa de referencia). "
        "El desvío es estable y de gran magnitud porcentual, incluso en los meses sin huecos de API "
        "(abril a agosto 2026). Conviene revisar en terreno que el medidor de la compañía y el punto "
        "WES midan el mismo ramal, y contrastar la lectura física del medidor  "
        f"{MEDIDOR} con el acumulado del nodo.",
    )
    _para(
        doc,
        "Fuente de las boletas: facturaciones Aguas Andinas del establecimiento "
        "(carpeta Colegios / Peñalolén / Facturaciones). "
        "Fuente del consumo: API WES, nodo 000008-14. "
        "La auditoría de habilitación 2024 (lecturas con control apagado y encendido) no se mezcla "
        "con esta serie: ese ejercicio mide el efecto del control; este informe mide si la boleta "
        "y el nodo coinciden mes a mes.",
    )

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_docx))


def a_pdf(docx: Path) -> Optional[Path]:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
        check=False,
        timeout=180,
    )
    pdf = docx.with_suffix(".pdf")
    return pdf if pdf.is_file() else None


def main() -> int:
    if sys.platform == "win32":
        try:
            sys.stdout.reconfigure(encoding="utf-8")
        except Exception:
            pass
    print("[1] Cruce facturación vs WES…", flush=True)
    filas = cruzar(PERIODOS)
    for f in filas:
        p = f["p"]
        print(
            f"  {p.boleta} cuenta={p.m3_cuenta} wes={f['m3_wes']:.1f} "
            f"proy={f['m3_proy']:.1f} diff={f['diff']}",
            flush=True,
        )
    print("[2] Serie mensual API…", flush=True)
    serie: list[tuple[str, float]] = []
    y, m = 2024, 11
    hoy = date.today()
    while (y, m) <= (hoy.year, hoy.month):
        total, dias = _mes_wes(y, m)
        serie.append((f"{MESES_CORTOS[m]}-{y % 100:02d}", total))
        print(f"  {y}-{m:02d} {total:.1f} m³ ({dias} d)", flush=True)
        m += 1
        if m == 13:
            m = 1
            y += 1

    ts = datetime.now().strftime("%Y%m%d_%H%M")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    pngs = {
        "historico": OUT_DIR / f"jp2_historico_facturacion_vs_wes_{ts}.png",
        "diferencia": OUT_DIR / f"jp2_diferencia_mensual_{ts}.png",
        "clp": OUT_DIR / f"jp2_valorizacion_diferencia_{ts}.png",
        "serie": OUT_DIR / f"jp2_consumo_mensual_wes_{ts}.png",
    }
    print("[3] Gráficos…", flush=True)
    grafico_historico(filas, pngs["historico"])
    grafico_diferencia(filas, pngs["diferencia"])
    grafico_valorizacion(filas, pngs["clp"])
    grafico_serie_mensual(serie, pngs["serie"])

    docx = OUT_DIR / f"Informe_Auditoria_Rendimiento_Juan_Pablo_II_000008-14_{ts}.docx"
    print("[4] Word…", flush=True)
    escribir_word(filas, serie, pngs, docx)
    print(f"[OK] {docx}", flush=True)
    pdf = a_pdf(docx)
    if pdf:
        print(f"[OK] {pdf}", flush=True)
    else:
        print("[WARN] Sin LibreOffice: no se generó PDF.", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
