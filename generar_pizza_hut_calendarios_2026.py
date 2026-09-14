"""Calendarios 2026 Pizza Hut: verde = cerró a tiempo, naranjo = tardó en cerrar."""

from __future__ import annotations

import calendar
import csv
from datetime import date, datetime, timedelta
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import FancyBboxPatch, Patch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.shared import Cm, Inches, Pt, RGBColor

from analizar_pizza_hut_cero_nocturno import _fetch_dia
from generar_pizza_hut_enero_control import DOW, clasificar

COLOR_WES = "1F4E79"
MESES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre"]
VERDE = "#2E7D32"
NARANJO = "#EF6C00"
GRIS = "#90A4AE"


def calendario_mes(dias: list, year: int, month: int, out: Path) -> Path:
    nd = calendar.monthrange(year, month)[1]
    first = date(year, month, 1)
    offset = first.weekday()
    n_rows = (offset + nd - 1) // 7 + 1
    n_v = sum(1 for d in dias if d["estado"] == "completo")
    n_n = sum(1 for d in dias if d["estado"] == "parcial")
    n_g = sum(1 for d in dias if d["estado"] == "no")

    fig = plt.figure(figsize=(10.4, 2.6 + n_rows * 1.05))
    ax = fig.add_axes([0.04, 0.16, 0.92, 0.74])
    ax.set_xlim(0, 7)
    ax.set_ylim(0, n_rows + 0.7)
    ax.axis("off")
    fig.suptitle(
        f"Pizza Hut — {MESES[month]} {year}\n"
        f"Verde {n_v} a tiempo   ·   Naranjo {n_n} tardó en cerrar   ·   Gris {n_g} no cortó",
        fontsize=13,
        fontweight="bold",
        color="#1F4E79",
        y=0.98,
    )
    by = {d["fecha"]: d for d in dias}
    for i, name in enumerate(DOW):
        ax.text(i + 0.5, n_rows + 0.38, name, ha="center", va="center", fontsize=10, fontweight="bold")
    for day in range(1, nd + 1):
        d = date(year, month, day)
        cell = offset + day - 1
        col = cell % 7
        row = n_rows - 1 - cell // 7
        info = by.get(d, {"estado": "sin_datos", "inicio": "", "fin": ""})
        color = {"completo": VERDE, "parcial": NARANJO, "no": GRIS, "sin_datos": "#EEEEEE"}[info["estado"]]
        ax.add_patch(
            FancyBboxPatch(
                (col + 0.05, row + 0.06),
                0.90,
                0.88,
                boxstyle="round,pad=0.02,rounding_size=0.08",
                facecolor=color,
                edgecolor="white",
                linewidth=1.6,
            )
        )
        txt = "white" if info["estado"] in {"completo", "parcial"} else "#333333"
        ax.text(col + 0.5, row + 0.60, str(day), ha="center", va="center", fontsize=14, fontweight="bold", color=txt)
        if info["estado"] == "completo":
            ax.text(col + 0.5, row + 0.28, "a tiempo", ha="center", va="center", fontsize=7, color="white")
        elif info["estado"] == "parcial":
            ax.text(col + 0.5, row + 0.28, f"cerró {info['inicio']}", ha="center", va="center", fontsize=6.5, color="white")

    fig.legend(
        handles=[
            Patch(facecolor=VERDE, label="Verde — cerró a tiempo (00:30–05:00 en cero)"),
            Patch(facecolor=NARANJO, label="Naranjo — tardó en cerrar (partió después, se mantuvo hasta la mañana)"),
            Patch(facecolor=GRIS, label="Gris — no cortó"),
        ],
        loc="lower center",
        bbox_to_anchor=(0.5, 0.01),
        ncol=1,
        frameon=False,
        fontsize=8.5,
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out, dpi=160)
    plt.close(fig)
    return out


def barras_conteo(resumen: list, out: Path) -> Path:
    labels = [r["mes"][:3] for r in resumen]
    v = [r["verde"] for r in resumen]
    n = [r["naranjo"] for r in resumen]
    g = [r["gris"] for r in resumen]
    x = list(range(len(labels)))
    fig, ax = plt.subplots(figsize=(10.6, 4.2))
    ax.bar(x, v, color=VERDE, label="Cerró a tiempo")
    ax.bar(x, n, bottom=v, color=NARANJO, label="Tardó en cerrar")
    ax.bar(x, g, bottom=[a + b for a, b in zip(v, n)], color=GRIS, label="No cortó")
    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.set_ylabel("Noches")
    ax.set_title("Pizza Hut 2026 — conteo: a tiempo vs tardó en cerrar", fontweight="bold", color="#1F4E79")
    ax.legend(frameon=False, fontsize=9)
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    fig.savefig(out, dpi=150, bbox_inches="tight")
    plt.close(fig)
    return out


def _shade(cell, hex_color: str) -> None:
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'w:val="clear" w:fill="{hex_color}"/>'
        )
    )


def _cell(cell, text, *, fill=None, bold=False, size=10, color=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Calibri"
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if fill:
        _shade(cell, fill)


def main() -> int:
    hoy = date.today()
    hasta = min(hoy, date(2026, 9, 14))
    out_dir = Path("reports/Puntos_En_Cero/Pizza_Hut/2026/calendarios")
    gra = out_dir / "graficos"
    gra.mkdir(parents=True, exist_ok=True)

    print("[INFO] Descargando 2026...")
    dias_all = []
    d0 = date(2026, 1, 1)
    d = d0
    while d <= hasta:
        dias_all.append(d)
        d += timedelta(days=1)
    from concurrent.futures import ThreadPoolExecutor, as_completed

    por_mes: dict = {m: [] for m in range(1, 10)}
    tmp = {}
    with ThreadPoolExecutor(max_workers=16) as ex:
        futs = {ex.submit(_fetch_dia, x): x for x in dias_all}
        done = 0
        for fut in as_completed(futs):
            dia, horas = fut.result()
            tmp[dia] = horas
            done += 1
            if done % 40 == 0:
                print(f"  {done}/{len(dias_all)}...")
    for dia in sorted(tmp):
        por_mes[dia.month].append(clasificar(dia, tmp[dia]))
    print(f"[INFO] {len(tmp)} días")

    resumen = []
    cals = []
    for m in range(1, 10):
        dias = por_mes[m]
        if not dias:
            continue
        cal = calendario_mes(dias, 2026, m, gra / f"calendario_{m:02d}_2026.png")
        cals.append(cal)
        resumen.append(
            {
                "mes": MESES[m],
                "verde": sum(1 for x in dias if x["estado"] == "completo"),
                "naranjo": sum(1 for x in dias if x["estado"] == "parcial"),
                "gris": sum(1 for x in dias if x["estado"] == "no"),
                "dias": len(dias),
            }
        )
        print(f"  {MESES[m]}: verde {resumen[-1]['verde']}  naranjo {resumen[-1]['naranjo']}  gris {resumen[-1]['gris']}")

    bar = barras_conteo(resumen, gra / "conteo_2026_verde_naranjo.png")
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    csv_path = out_dir / "conteo_verde_naranjo_2026.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["mes", "verde_a_tiempo", "naranjo_tardo_cerrar", "gris_no_corto", "dias"])
        for r in resumen:
            w.writerow([r["mes"], r["verde"], r["naranjo"], r["gris"], r["dias"]])

    docx_path = out_dir / f"Pizza_Hut_calendarios_verde_naranjo_2026_{stamp}.docx"
    pdf_path = out_dir / f"Pizza_Hut_calendarios_verde_naranjo_2026_{stamp}.pdf"

    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.4)
        s.bottom_margin = Cm(1.4)
        s.left_margin = Cm(1.6)
        s.right_margin = Cm(1.6)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("PIZZA HUT 2026 — calendarios verde / naranjo")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor.from_string(COLOR_WES)
    doc.add_paragraph(
        "Verde = cerró a tiempo (00:30–05:00 en cero). "
        "Naranjo = el control sí cortó, solo tardó en cerrar (partió después de las 00:30 y se mantuvo hasta la mañana). "
        "Gris = esa noche no cortó."
    )

    tot_v = sum(r["verde"] for r in resumen)
    tot_n = sum(r["naranjo"] for r in resumen)
    tot_g = sum(r["gris"] for r in resumen)
    tbl = doc.add_table(rows=1 + len(resumen) + 1, cols=4)
    tbl.style = "Table Grid"
    hdr = ["Mes", "Verde (a tiempo)", "Naranjo (tardó en cerrar)", "Gris (no cortó)"]
    for i, h in enumerate(hdr):
        _cell(tbl.rows[0].cells[i], h, fill=COLOR_WES, bold=True, size=10, color="FFFFFF")
    for i, r in enumerate(resumen, start=1):
        _cell(tbl.rows[i].cells[0], r["mes"], size=10, bold=True)
        _cell(tbl.rows[i].cells[1], str(r["verde"]), fill="C8E6C9", size=11, bold=True)
        _cell(tbl.rows[i].cells[2], str(r["naranjo"]), fill="FFE0B2", size=11, bold=True)
        _cell(tbl.rows[i].cells[3], str(r["gris"]), fill="ECEFF1", size=11)
    last = tbl.rows[-1]
    _cell(last.cells[0], "TOTAL 2026", fill="D6E3F0", bold=True)
    _cell(last.cells[1], str(tot_v), fill="C8E6C9", bold=True, size=12)
    _cell(last.cells[2], str(tot_n), fill="FFE0B2", bold=True, size=12)
    _cell(last.cells[3], str(tot_g), fill="D6E3F0", bold=True)

    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(bar), width=Inches(6.3))
    for cal in cals:
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(cal), width=Inches(6.3))
    doc.save(str(docx_path))

    with PdfPages(pdf_path) as pdf:
        fig = plt.figure(figsize=(8.27, 11.69))
        ax = fig.add_subplot(111)
        ax.axis("off")
        ax.set_title("Conteo 2026 — a tiempo vs tardó en cerrar", color="#1F4E79", fontsize=13, pad=10)
        ax.imshow(plt.imread(str(bar)))
        fig.tight_layout()
        pdf.savefig(fig)
        plt.close(fig)
        for cal in cals:
            fig = plt.figure(figsize=(8.27, 11.69))
            ax = fig.add_subplot(111)
            ax.imshow(plt.imread(str(cal)))
            ax.axis("off")
            fig.tight_layout()
            pdf.savefig(fig)
            plt.close(fig)

    print(f"DOCX={docx_path}")
    print(f"PDF={pdf_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
