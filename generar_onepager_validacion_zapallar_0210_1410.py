"""
One-pager — Validación relojería vs App WES (Fundo Zapallar).
Ventana 02-10-2026 → 14-10-2026 · Matriz ESVAL + Etapa N°5.

Uso:
  python generar_onepager_validacion_zapallar_0210_1410.py
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm, mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "reports" / "Fundo_Zapallar" / "VALIDACION"
NAVY = HexColor("#1F4788")
MUTED = HexColor("#595959")
LIGHT = HexColor("#F5F8FB")
GREEN = HexColor("#E8F5E9")
GREEN_TXT = HexColor("#1B5E20")
AMBER = HexColor("#FFF8E1")
AMBER_TXT = HexColor("#E65100")

# --- Datos validados (02-10 → 14-10) ---
VENTANA = "02-10-2026 → 14-10-2026"
MATRIZ = {
    "nombre": "Matriz ESVAL",
    "node": "000027-01",
    "medidor": "Itron FLOSTAR S",
    "ini": "385.430,8 m³ · 02-10 13:22",
    "fin": "386.822,3 m³ · 14-10 17:08",
    "reloj": 1391.5,
    "app": 1425.87,
    "horas": 291,
    "suma": "14:00 del 02 → 16:00 del 14",
    "error": 2.41,
}
ETAPA5 = {
    "nombre": "Etapa N°5",
    "node": "000027-03",
    "medidor": "Sensus MeiStream Plus 100",
    "ini": "5.369 m³ · 02-10 08:44",
    "fin": "5.611,1 m³ · 14-10 17:00",
    "reloj": 242.1,
    "app": 243.4,
    "horas": 296,
    "suma": "09:00 del 02 → 16:00 del 14",
    "error": 0.53,
}


def _fmt(x: float, dec: int = 2) -> str:
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def _font_docx(run, *, size=11, bold=False, color=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _row(tbl, i: int, vals: list[str], *, bold=False, header=False) -> None:
    heading = RGBColor(31, 71, 136)
    for j, v in enumerate(vals):
        cell = tbl.rows[i].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(v)
        _font_docx(r, size=9, bold=bold or header, color=heading if header else None)


def build_docx(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.3)
    sec.bottom_margin = Cm(1.1)
    sec.left_margin = Cm(1.6)
    sec.right_margin = Cm(1.6)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    try:
        from generar_reporte_word import add_logo_to_header

        add_logo_to_header(doc)
    except Exception as exc:
        print(f"[WARN] Logo: {exc}")

    title = doc.add_heading("Validación relojería vs App WES — Fundo Zapallar", level=0)
    for r in title.runs:
        r.font.color.rgb = RGBColor(31, 71, 136)
        r.font.name = "Calibri"
        r.font.size = Pt(15)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(6)
    r = meta.add_run(
        f"Ventana {VENTANA}  ·  Horas completas (se excluye la hora de cada foto)  ·  Uso interno WES"
    )
    _font_docx(r, size=9, color=RGBColor(89, 89, 89))

    h = doc.add_heading("Resultado", level=1)
    for r in h.runs:
        r.font.color.rgb = RGBColor(31, 71, 136)

    tbl = doc.add_table(rows=3, cols=6)
    tbl.style = "Table Grid"
    _row(
        tbl,
        0,
        ["Punto", "Relojería (m³)", "App WES (m³)", "Diff", "Error", "Lectura"],
        header=True,
    )
    _row(
        tbl,
        1,
        [
            f"{MATRIZ['nombre']}\n{MATRIZ['node']}",
            _fmt(MATRIZ["reloj"], 1),
            _fmt(MATRIZ["app"], 2),
            f"+{_fmt(MATRIZ['app'] - MATRIZ['reloj'], 2)}",
            f"{_fmt(MATRIZ['error'], 1)} %",
            "Aceptable",
        ],
    )
    _row(
        tbl,
        2,
        [
            f"{ETAPA5['nombre']}\n{ETAPA5['node']}",
            _fmt(ETAPA5["reloj"], 1),
            _fmt(ETAPA5["app"], 1),
            f"+{_fmt(ETAPA5['app'] - ETAPA5['reloj'], 1)}",
            f"{_fmt(ETAPA5['error'], 1)} %",
            "Excelente",
        ],
    )
    doc.add_paragraph()

    for pto in (MATRIZ, ETAPA5):
        h = doc.add_heading(f"{pto['nombre']} ({pto['node']})", level=2)
        for r in h.runs:
            r.font.color.rgb = RGBColor(31, 71, 136)
            r.font.size = Pt(11)
        for line in (
            f"Medidor: {pto['medidor']}",
            f"Lectura inicial: {pto['ini']}",
            f"Lectura final: {pto['fin']}",
            f"Suma App: {pto['suma']} ({pto['horas']} horas)",
        ):
            para = doc.add_paragraph(style="List Bullet")
            para.paragraph_format.space_after = Pt(1)
            para.clear()
            r = para.add_run(line)
            _font_docx(r, size=9)

    h = doc.add_heading("Conclusión", level=1)
    for r in h.runs:
        r.font.color.rgb = RGBColor(31, 71, 136)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(
        "Etapa N°5 cierra con error 0,5 % (excelente). Matriz ESVAL cierra con 2,4 %, "
        "app por encima del Itron, del mismo orden que el cierre anterior (2,1 % el 02-10). "
        "Ambos puntos quedan validados para el período."
    )
    _font_docx(r, size=10)

    pie = doc.add_paragraph()
    pie.paragraph_format.space_before = Pt(8)
    r = pie.add_run(f"WES · generado {datetime.now().strftime('%d-%m-%Y %H:%M')}")
    _font_docx(r, size=8, color=RGBColor(89, 89, 89))

    doc.save(path)
    print(f"[OK] DOCX {path}")


def _card(c: canvas.Canvas, x, y, w, h, fill) -> None:
    c.setFillColor(fill)
    c.setStrokeColor(HexColor("#D0D7E2"))
    c.setLineWidth(0.6)
    c.roundRect(x, y, w, h, 4, fill=1, stroke=1)


def _kpi(c: canvas.Canvas, x, y, w, h, label: str, value: str, fill, value_color) -> None:
    _card(c, x, y, w, h, fill)
    c.setFillColor(MUTED)
    c.setFont("Helvetica", 8)
    c.drawString(x + 6, y + h - 12, label)
    c.setFillColor(value_color)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(x + 6, y + 10, value)


def build_pdf(path: Path, logo: Path | None) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    W, H = A4
    ml, mr, mt, mb = 1.5 * cm, 1.5 * cm, 1.3 * cm, 1.2 * cm
    usable_w = W - ml - mr

    y = H - mt
    if logo and logo.exists():
        c.drawImage(str(logo), ml, y - 1.0 * cm, width=4.2 * cm, height=0.95 * cm, mask="auto")
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 14)
    c.drawRightString(W - mr, y - 0.35 * cm, "Validación relojería vs App WES")
    c.setFont("Helvetica", 9)
    c.setFillColor(MUTED)
    c.drawRightString(W - mr, y - 0.7 * cm, "Fundo Zapallar")
    y -= 1.25 * cm

    c.setStrokeColor(NAVY)
    c.setLineWidth(1.2)
    c.line(ml, y, W - mr, y)
    y -= 0.45 * cm

    c.setFillColor(MUTED)
    c.setFont("Helvetica", 8.5)
    c.drawString(ml, y, f"Ventana {VENTANA}  ·  Horas completas (excluye hora de cada foto)  ·  Uso interno WES")
    y -= 0.55 * cm

    # KPIs
    card_w = (usable_w - 0.4 * cm) / 2
    card_h = 2.0 * cm
    _kpi(
        c,
        ml,
        y - card_h,
        card_w,
        card_h,
        "Matriz ESVAL — error",
        f"{_fmt(MATRIZ['error'], 1)} %",
        AMBER,
        AMBER_TXT,
    )
    _kpi(
        c,
        ml + card_w + 0.4 * cm,
        y - card_h,
        card_w,
        card_h,
        "Etapa N°5 — error",
        f"{_fmt(ETAPA5['error'], 1)} %",
        GREEN,
        GREEN_TXT,
    )
    y -= card_h + 0.55 * cm

    # Tabla
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 11)
    c.drawString(ml, y, "Resultado")
    y -= 0.35 * cm

    headers = ["Punto", "Relojería", "App WES", "Diff", "Error", "Lectura"]
    col_w = [3.6 * cm, 2.4 * cm, 2.4 * cm, 2.2 * cm, 1.8 * cm, 2.4 * cm]
    # normalize to usable_w
    scale = usable_w / sum(col_w)
    col_w = [w * scale for w in col_w]
    row_h = 0.72 * cm

    def draw_table_row(yy, vals, header=False, alt=False):
        x = ml
        bg = NAVY if header else (LIGHT if alt else white)
        c.setFillColor(bg)
        c.rect(ml, yy - row_h, usable_w, row_h, fill=1, stroke=0)
        c.setStrokeColor(HexColor("#C5CEDA"))
        c.setLineWidth(0.4)
        c.rect(ml, yy - row_h, usable_w, row_h, fill=0, stroke=1)
        for i, (val, w) in enumerate(zip(vals, col_w)):
            c.setFillColor(white if header else NAVY if i == 0 else HexColor("#222222"))
            c.setFont("Helvetica-Bold" if header or i in (4, 5) else "Helvetica", 8.5)
            c.drawString(x + 4, yy - row_h + 0.25 * cm, val)
            x += w
        return yy - row_h

    y = draw_table_row(y, headers, header=True)
    y = draw_table_row(
        y,
        [
            MATRIZ["nombre"],
            f"{_fmt(MATRIZ['reloj'], 1)} m³",
            f"{_fmt(MATRIZ['app'], 2)} m³",
            f"+{_fmt(MATRIZ['app'] - MATRIZ['reloj'], 2)}",
            f"{_fmt(MATRIZ['error'], 1)} %",
            "Aceptable",
        ],
        alt=False,
    )
    y = draw_table_row(
        y,
        [
            ETAPA5["nombre"],
            f"{_fmt(ETAPA5['reloj'], 1)} m³",
            f"{_fmt(ETAPA5['app'], 1)} m³",
            f"+{_fmt(ETAPA5['app'] - ETAPA5['reloj'], 1)}",
            f"{_fmt(ETAPA5['error'], 1)} %",
            "Excelente",
        ],
        alt=True,
    )
    y -= 0.55 * cm

    # Detalle por punto
    for pto, badge_fill, badge_txt in (
        (MATRIZ, AMBER, AMBER_TXT),
        (ETAPA5, GREEN, GREEN_TXT),
    ):
        box_h = 3.15 * cm
        _card(c, ml, y - box_h, usable_w, box_h, LIGHT)
        c.setFillColor(NAVY)
        c.setFont("Helvetica-Bold", 10)
        c.drawString(ml + 8, y - 0.4 * cm, f"{pto['nombre']}  ({pto['node']})")
        # badge
        bw = 2.6 * cm
        c.setFillColor(badge_fill)
        c.roundRect(W - mr - bw - 6, y - 0.55 * cm, bw, 0.45 * cm, 3, fill=1, stroke=0)
        c.setFillColor(badge_txt)
        c.setFont("Helvetica-Bold", 8)
        c.drawCentredString(W - mr - bw / 2 - 6, y - 0.42 * cm, f"Error {_fmt(pto['error'], 1)} %")

        c.setFillColor(HexColor("#333333"))
        c.setFont("Helvetica", 8.5)
        lines = [
            f"Medidor: {pto['medidor']}",
            f"Inicial: {pto['ini']}     Final: {pto['fin']}",
            f"Relojería: {_fmt(pto['reloj'], 1)} m³     App WES: {_fmt(pto['app'], 2 if pto is MATRIZ else 1)} m³     "
            f"Diff: +{_fmt(pto['app'] - pto['reloj'], 2 if pto is MATRIZ else 1)} m³",
            f"Suma App: {pto['suma']}  ·  {pto['horas']} horas completas",
        ]
        ly = y - 0.85 * cm
        for line in lines:
            c.drawString(ml + 8, ly, line)
            ly -= 0.42 * cm
        y -= box_h + 0.35 * cm

    # Conclusión
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 11)
    c.drawString(ml, y, "Conclusión")
    y -= 0.2 * cm
    conc_h = 2.0 * cm
    _card(c, ml, y - conc_h, usable_w, conc_h, GREEN)
    c.setFillColor(HexColor("#1B3A1D"))
    c.setFont("Helvetica", 9)
    text = (
        "Etapa N°5 cierra con error 0,5 % (excelente). Matriz ESVAL cierra con 2,4 %, "
        "app por encima del Itron, del mismo orden que el cierre anterior (2,1 % el 02-10). "
        "Ambos puntos quedan validados para el período."
    )
    # simple wrap
    words = text.split()
    lines, cur = [], ""
    max_chars = 105
    for w in words:
        trial = (cur + " " + w).strip()
        if len(trial) <= max_chars:
            cur = trial
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    ty = y - 0.45 * cm
    for line in lines:
        c.drawString(ml + 8, ty, line)
        ty -= 0.38 * cm

    c.setFillColor(MUTED)
    c.setFont("Helvetica", 7.5)
    c.drawString(ml, mb - 0.1 * cm, f"WES · generado {datetime.now().strftime('%d-%m-%Y %H:%M')}")
    c.drawRightString(W - mr, mb - 0.1 * cm, "Fundo Zapallar · VALIDACION")

    c.showPage()
    c.save()
    print(f"[OK] PDF  {path}")


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    base = f"OnePager_Validacion_Fundo_Zapallar_0210_1410_{stamp}"
    docx_path = OUT_DIR / f"{base}.docx"
    pdf_path = OUT_DIR / f"{base}.pdf"
    logo = Path("/tmp/logo_wes.png")
    if not logo.exists():
        try:
            from PIL import Image

            Image.open(ROOT / "logo wes.bmp").convert("RGBA").save(logo)
        except Exception:
            logo = None

    build_docx(docx_path)
    build_pdf(pdf_path, logo if logo and logo.exists() else None)

    # also copy stable names without stamp for easy find
    stable_pdf = OUT_DIR / "OnePager_Validacion_Fundo_Zapallar_0210_1410.pdf"
    stable_docx = OUT_DIR / "OnePager_Validacion_Fundo_Zapallar_0210_1410.docx"
    stable_pdf.write_bytes(pdf_path.read_bytes())
    stable_docx.write_bytes(docx_path.read_bytes())
    print(f"[OK] estable {stable_pdf}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
