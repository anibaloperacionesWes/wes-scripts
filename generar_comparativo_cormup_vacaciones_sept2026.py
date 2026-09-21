"""
Comparativo vacaciones CORMUP / Peñalolén — septiembre 2026.

Semana sin control:     07–13/09/2026
Semana con control especial por vacaciones: 14–20/09/2026

Cohorte: colegios con control, excluyendo Tobalaba (pulso) y los tres sin control
(Eduardo de la Barra, Alicura, Likankura).

Uso:
  python generar_comparativo_cormup_vacaciones_sept2026.py
  python generar_comparativo_cormup_vacaciones_sept2026.py --sin-agregados
"""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple
from zoneinfo import ZoneInfo

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from matplotlib.backends.backend_pdf import PdfPages

from auditoria_cpa_icco_renca_grafico import _vectores_m3h_por_dias
from generar_reporte_word import (
    add_logo_to_header,
    add_picture_with_pagination,
    estilizar_tabla_wes,
    format_number_chilean,
    generate_aggregated_report,
)
from wes_estilo_graficos_app import COLOR_BARRA_WES, COLOR_CONSUMO, COLOR_NOCHE

CHILE = ZoneInfo("America/Santiago")
ROOT = Path(__file__).resolve().parent
COMPANY_ID = "000008"

SIN_CONTROL_INI = date(2026, 9, 7)
SIN_CONTROL_FIN = date(2026, 9, 13)
CON_CONTROL_INI = date(2026, 9, 14)
CON_CONTROL_FIN = date(2026, 9, 20)

# 10 colegios con control (Tobalaba fuera por pulso).
COLEGIOS: List[Tuple[str, str]] = [
    ("000008-01", "Lic. Antonio Hermida F"),
    ("000008-03", "Carlos Fernandez P."),
    ("000008-05", "Santa Maria"),
    ("000008-06", "Luis Arrieta Caña"),
    ("000008-07", "Erasmo Escala"),
    ("000008-09", "Juan Bautista Pasten"),
    ("000008-10", "Matilde Huici Navas"),
    ("000008-11", "CE Valle Hermoso"),
    ("000008-12", "Unión Nacional Árabe"),
    ("000008-14", "Juan Pablo II"),
]

EXCLUIDOS_NOTA = [
    "Tobalaba — fuera por problemas de pulso",
    "Eduardo de la Barra — sin control",
    "Alicura — sin control",
    "Likankura — sin control",
]

HORARIOS_ESPECIALES = [
    "Lic. Antonio Hermida F: corte programado (24 h).",
    "Carlos Fernandez P.: obras — agua habilitada toda la semana (incluye sáb/dom) 09:00–18:00.",
    "Santa Maria: corte programado (24 h).",
    "Luis Arrieta Caña: patinaje lun–mar 17:30–21:00; martes 15/09 habilitación adicional por revisión del colegio.",
    "Erasmo Escala: corte programado (24 h).",
    "Juan Bautista Pasten: obras — agua habilitada toda la semana (incluye sáb/dom) 09:00–18:00.",
    "Matilde Huici Navas: corte programado (24 h).",
    "CE Valle Hermoso: patinaje lun y mié 17:30–21:00; miércoles 16/09 habilitación especial ≈10:00–tarde.",
    "Unión Nacional Árabe: corte programado (24 h).",
    "Juan Pablo II: corte programado (24 h).",
]


@dataclass
class ResultadoColegio:
    node_id: str
    nombre: str
    m3_sin: float
    m3_con: float
    por_dia_sin: List[float] = field(default_factory=list)
    por_dia_con: List[float] = field(default_factory=list)

    @property
    def ahorro_m3(self) -> float:
        return self.m3_sin - self.m3_con

    @property
    def ahorro_pct(self) -> Optional[float]:
        if self.m3_sin <= 1e-9:
            return None
        return 100.0 * (self.m3_sin - self.m3_con) / self.m3_sin


def _rango(ini: date, fin: date) -> List[date]:
    out: List[date] = []
    d = ini
    while d <= fin:
        out.append(d)
        d += timedelta(days=1)
    return out


def _fmt_pct(v: Optional[float]) -> str:
    if v is None:
        return "—"
    return f"{format_number_chilean(v, 1)} %"


def _sombrear(cell, hex_color: str) -> None:
    try:
        shading = parse_xml(
            f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'w:val="clear" w:fill="{hex_color}"/>'
        )
        tc_pr = cell._element.get_or_add_tcPr()
        if tc_pr.find(qn("w:shd")) is None:
            tc_pr.append(shading)
    except Exception:
        pass


def _consumo_periodo(node_id: str, dias: Sequence[date]) -> Tuple[float, List[float]]:
    vecs = _vectores_m3h_por_dias(node_id, dias)
    por_dia = [float(sum(v)) for v in vecs]
    return float(sum(por_dia)), por_dia


def evaluar_colegios(max_workers: int = 8) -> List[ResultadoColegio]:
    dias_sin = _rango(SIN_CONTROL_INI, SIN_CONTROL_FIN)
    dias_con = _rango(CON_CONTROL_INI, CON_CONTROL_FIN)
    resultados: Dict[str, ResultadoColegio] = {}

    def _uno(nid: str, nombre: str) -> ResultadoColegio:
        m3_sin, por_sin = _consumo_periodo(nid, dias_sin)
        m3_con, por_con = _consumo_periodo(nid, dias_con)
        return ResultadoColegio(
            node_id=nid,
            nombre=nombre,
            m3_sin=m3_sin,
            m3_con=m3_con,
            por_dia_sin=por_sin,
            por_dia_con=por_con,
        )

    with ThreadPoolExecutor(max_workers=max_workers) as ex:
        futs = {ex.submit(_uno, nid, nom): nid for nid, nom in COLEGIOS}
        for fut in as_completed(futs):
            r = fut.result()
            resultados[r.node_id] = r
            print(
                f"  [OK] {r.nombre}: sin={r.m3_sin:.2f} m³ | con={r.m3_con:.2f} m³ "
                f"| ahorro={r.ahorro_m3:.2f} m³ ({_fmt_pct(r.ahorro_pct)})"
            )

    return [resultados[nid] for nid, _ in COLEGIOS]


def grafico_barras_colegios(filas: List[ResultadoColegio], out_png: Path) -> Path:
    nombres = [f.nombre for f in filas]
    sin_vals = [f.m3_sin for f in filas]
    con_vals = [f.m3_con for f in filas]
    x = list(range(len(filas)))
    w = 0.38

    fig, ax = plt.subplots(figsize=(12, 6.2))
    ax.bar([i - w / 2 for i in x], sin_vals, width=w, color=COLOR_NOCHE, label="Sin control (7–13/09)")
    ax.bar([i + w / 2 for i in x], con_vals, width=w, color=COLOR_BARRA_WES, label="Con control vacaciones (14–20/09)")
    ax.set_xticks(x)
    ax.set_xticklabels(nombres, rotation=35, ha="right", fontsize=8)
    ax.set_ylabel("Consumo (m³)")
    ax.set_title("CORMUP Peñalolén — consumo semanal por colegio")
    ax.legend(frameon=False)
    ax.grid(axis="y", linestyle=":", alpha=0.4)
    fig.tight_layout()
    out_png.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_png, dpi=150)
    plt.close(fig)
    return out_png


def grafico_totales_diarios(filas: List[ResultadoColegio], out_png: Path) -> Path:
    dias_sin = _rango(SIN_CONTROL_INI, SIN_CONTROL_FIN)
    dias_con = _rango(CON_CONTROL_INI, CON_CONTROL_FIN)
    tot_sin = [sum(f.por_dia_sin[i] for f in filas) for i in range(len(dias_sin))]
    tot_con = [sum(f.por_dia_con[i] for f in filas) for i in range(len(dias_con))]

    fig, ax = plt.subplots(figsize=(11, 5.2))
    ax.plot(
        [d.strftime("%d/%m") for d in dias_sin],
        tot_sin,
        marker="o",
        color=COLOR_NOCHE,
        label="Sin control (7–13/09)",
    )
    ax.plot(
        [d.strftime("%d/%m") for d in dias_con],
        tot_con,
        marker="o",
        color=COLOR_CONSUMO,
        label="Con control vacaciones (14–20/09)",
    )
    ax.set_ylabel("Consumo agregado (m³/día)")
    ax.set_title("CORMUP Peñalolén — total diario de los 10 colegios")
    ax.legend(frameon=False)
    ax.grid(True, linestyle=":", alpha=0.4)
    fig.tight_layout()
    fig.savefig(out_png, dpi=150)
    plt.close(fig)
    return out_png


def crear_informe_word(
    filas: List[ResultadoColegio],
    chart_barras: Path,
    chart_diario: Path,
    out_docx: Path,
) -> Path:
    tot_sin = sum(f.m3_sin for f in filas)
    tot_con = sum(f.m3_con for f in filas)
    ahorro = tot_sin - tot_con
    pct = (100.0 * ahorro / tot_sin) if tot_sin > 1e-9 else None

    doc = Document()
    add_logo_to_header(doc)

    title = doc.add_heading("CORMUP Peñalolén — Comparativo vacaciones septiembre 2026", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(0, 51, 102)

    gen = datetime.now(CHILE).strftime("%d-%m-%Y %H:%M")
    sub = doc.add_paragraph(
        f"Sin control: {SIN_CONTROL_INI:%d/%m/%Y}–{SIN_CONTROL_FIN:%d/%m/%Y}  |  "
        f"Con control especial vacaciones: {CON_CONTROL_INI:%d/%m/%Y}–{CON_CONTROL_FIN:%d/%m/%Y}\n"
        f"Generado {gen} (hora Chile)"
    )
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    for run in sub.runs:
        run.font.size = Pt(10)

    doc.add_heading("Resumen ejecutivo", level=1)
    doc.add_paragraph(
        f"Se comparan dos semanas de 7 días sobre los {len(filas)} colegios CORMUP con control "
        f"operativo (sin Tobalaba ni los establecimientos sin válvula).\n\n"
        f"Consumo semana sin control (7–13/09): {format_number_chilean(tot_sin, 1)} m³.\n"
        f"Consumo semana con control vacaciones (14–20/09): {format_number_chilean(tot_con, 1)} m³.\n"
        f"Variación: {format_number_chilean(ahorro, 1)} m³"
        + (f" ({_fmt_pct(pct)})." if pct is not None else ".")
        + "\n\n"
        "La semana con control no fue corte 24 h uniforme: hubo habilitaciones por obras, "
        "patinaje y revisiones puntuales (detalle más abajo). Esos consumos esperados reducen "
        "el ahorro aparente frente a un corte total."
    )

    doc.add_heading("Alcance y exclusiones", level=1)
    doc.add_paragraph("Colegios incluidos:")
    for _, nom in COLEGIOS:
        doc.add_paragraph(nom, style="List Bullet")
    doc.add_paragraph("Excluidos de este informe:")
    for txt in EXCLUIDOS_NOTA:
        doc.add_paragraph(txt, style="List Bullet")

    doc.add_heading("Horarios de control especial (14–20/09)", level=1)
    for txt in HORARIOS_ESPECIALES:
        doc.add_paragraph(txt, style="List Bullet")

    doc.add_heading("Comparación por colegio", level=1)
    headers = [
        "Colegio",
        "Sin control m³\n(7–13/09)",
        "Con control m³\n(14–20/09)",
        "Ahorro m³",
        "Ahorro %",
    ]
    table = doc.add_table(rows=1 + len(filas) + 1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9)
        _sombrear(cell, "1F4E79")

    orden = sorted(filas, key=lambda f: f.ahorro_m3, reverse=True)
    for r_i, f in enumerate(orden, start=1):
        vals = [
            f.nombre,
            format_number_chilean(f.m3_sin, 1),
            format_number_chilean(f.m3_con, 1),
            format_number_chilean(f.ahorro_m3, 1),
            _fmt_pct(f.ahorro_pct),
        ]
        for c_i, v in enumerate(vals):
            cell = table.rows[r_i].cells[c_i]
            cell.text = v
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if c_i >= 3:
                        run.font.color.rgb = (
                            RGBColor(0, 97, 0) if f.ahorro_m3 >= 0 else RGBColor(156, 0, 6)
                        )

    total_row = table.rows[1 + len(filas)]
    totales = [
        "TOTAL (10 colegios)",
        format_number_chilean(tot_sin, 1),
        format_number_chilean(tot_con, 1),
        format_number_chilean(ahorro, 1),
        _fmt_pct(pct),
    ]
    for c_i, v in enumerate(totales):
        cell = total_row.cells[c_i]
        cell.text = v
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
        _sombrear(cell, "D6EAF8")

    estilizar_tabla_wes(table, has_total_row=True)

    doc.add_heading("Gráficos", level=1)
    add_picture_with_pagination(doc, str(chart_barras), width=Inches(6.3))
    doc.add_paragraph("")
    add_picture_with_pagination(doc, str(chart_diario), width=Inches(6.3))

    doc.add_heading("Lectura operativa", level=1)
    doc.add_paragraph(
        "• Hermida, Santa Maria, Erasmo Escala, Matilde Huici, Unión Nacional Árabe y Juan Pablo II "
        "tenían corte programado; residuales en esa semana deben revisarse como eventual fuga o "
        "habilitación no registrada.\n"
        "• Carlos Fernandez y Juan Bautista Pasten concentraron consumo diurno esperado por obras "
        "(09:00–18:00 todos los días).\n"
        "• Luis Arrieta y Valle Hermoso tuvieron ventanas de patinaje y habilitaciones puntuales "
        "(martes 15 y miércoles 16, respectivamente).\n"
        "• Tobalaba se excluyó del análisis por falla de pulso; no forma parte del total."
    )

    doc.add_heading("Conclusiones", level=1)
    if ahorro > 0:
        doc.add_paragraph(
            f"En el conjunto de los {len(filas)} colegios, la semana con control especial por "
            f"vacaciones registró {format_number_chilean(ahorro, 1)} m³ menos que la semana previa "
            f"sin control ({_fmt_pct(pct)}). El resultado incorpora las habilitaciones de obra y "
            f"patinaje acordadas con el cliente."
        )
    else:
        doc.add_paragraph(
            f"En el conjunto de los {len(filas)} colegios, la semana con control especial no muestra "
            f"ahorro neto frente a la semana sin control (variación "
            f"{format_number_chilean(ahorro, 1)} m³). Revisar habilitaciones de obra/patinaje y "
            f"posibles residuales fuera de ventana."
        )

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_docx))
    return out_docx


def convertir_pdf(docx_path: Path, filas: List[ResultadoColegio], chart_barras: Path, chart_diario: Path) -> Path:
    pdf = docx_path.with_suffix(".pdf")
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        try:
            subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(docx_path.parent),
                    str(docx_path),
                ],
                check=True,
                timeout=180,
                capture_output=True,
            )
            if pdf.exists():
                return pdf
        except Exception as exc:
            print(f"[AVISO] LibreOffice falló: {exc}")

    # Fallback: PDF de síntesis con tablas/gráficos.
    tot_sin = sum(f.m3_sin for f in filas)
    tot_con = sum(f.m3_con for f in filas)
    ahorro = tot_sin - tot_con
    pct = (100.0 * ahorro / tot_sin) if tot_sin > 1e-9 else None

    with PdfPages(pdf) as pages:
        fig = plt.figure(figsize=(11.69, 8.27))
        ax = fig.add_subplot(111)
        ax.axis("off")
        ax.set_title(
            "CORMUP Peñalolén — Comparativo vacaciones 7–13 vs 14–20/09/2026",
            fontsize=13,
            pad=16,
        )
        headers = ["Colegio", "Sin control m³", "Con control m³", "Ahorro m³", "Ahorro %"]
        data = [headers]
        for f in sorted(filas, key=lambda x: x.ahorro_m3, reverse=True):
            data.append(
                [
                    f.nombre[:26],
                    f"{f.m3_sin:.1f}",
                    f"{f.m3_con:.1f}",
                    f"{f.ahorro_m3:.1f}",
                    "—" if f.ahorro_pct is None else f"{f.ahorro_pct:.1f}%",
                ]
            )
        data.append(
            [
                "TOTAL",
                f"{tot_sin:.1f}",
                f"{tot_con:.1f}",
                f"{ahorro:.1f}",
                "—" if pct is None else f"{pct:.1f}%",
            ]
        )
        table = ax.table(cellText=data, loc="center", cellLoc="center")
        table.auto_set_font_size(False)
        table.set_fontsize(8)
        table.scale(1.15, 1.35)
        for j in range(len(headers)):
            table[0, j].set_facecolor("#1F4E79")
            table[0, j].set_text_props(color="white", weight="bold")
        for i in range(1, len(data)):
            color = "#D6EAF8" if i == len(data) - 1 else ("#C6EFCE" if float(data[i][3]) >= 0 else "#FFC7CE")
            for j in range(len(headers)):
                table[i, j].set_facecolor(color)
        ax.text(
            0.02,
            0.04,
            "Excluidos: Tobalaba (pulso), Eduardo de la Barra, Alicura, Likankura. "
            "Control 14–20 con obras/patinaje/habilitaciones puntuales.",
            transform=ax.transAxes,
            fontsize=8,
        )
        pages.savefig(fig, bbox_inches="tight")
        plt.close(fig)

        for img in (chart_barras, chart_diario):
            fig2 = plt.figure(figsize=(11.69, 8.27))
            ax2 = fig2.add_subplot(111)
            ax2.axis("off")
            ax2.imshow(plt.imread(img))
            pages.savefig(fig2, bbox_inches="tight")
            plt.close(fig2)
    return pdf


def generar_agregados_apoyo() -> List[Path]:
    nodes = [nid for nid, _ in COLEGIOS]
    nota_sin = (
        "Periodo de referencia sin control operativo de vacaciones (07–13/09/2026). "
        "Excluidos: Tobalaba (pulso), Eduardo de la Barra, Alicura y Likankura."
    )
    nota_con = (
        "Periodo con control especial por vacaciones (14–20/09/2026). "
        "Excluidos: Tobalaba (pulso), Eduardo de la Barra, Alicura y Likankura. "
        "Horarios: corte 24 h salvo Carlos Fernandez y Juan Bautista Pasten "
        "(obra agua 09:00–18:00 toda la semana); Luis Arrieta patinaje lun–mar 17:30–21:00 "
        "(+ habilitación mar 15 por revisión); Valle Hermoso patinaje lun y mié 17:30–21:00 "
        "(+ habilitación mié 16 ≈10:00–tarde)."
    )
    outs: List[Path] = []
    for start, end, nota in (
        ("07/09/2026", "13/09/2026", nota_sin),
        ("14/09/2026", "20/09/2026", nota_con),
    ):
        print(f"\n[INFO] Agregado apoyo {start}–{end}")
        path = generate_aggregated_report(
            company_id=COMPANY_ID,
            node_ids=list(nodes),
            start_date=start,
            end_date=end,
            apply_exclusions=False,
            generate_ppt=False,
            nota_contexto_periodo=nota,
            parallel_node_fetch=True,
            max_parallel_workers=8,
        )
        outs.append(Path(path))
        print(f"[OK] Agregado: {path}")
    return outs


def main() -> int:
    ap = argparse.ArgumentParser(description="Comparativo vacaciones CORMUP sept 2026")
    ap.add_argument("--sin-agregados", action="store_true", help="Solo el comparativo (sin agregados Word)")
    ap.add_argument("--workers", type=int, default=8)
    args = ap.parse_args()

    print("=" * 72)
    print("CORMUP Peñalolén — comparativo vacaciones 7–13 vs 14–20/09/2026")
    print("=" * 72)

    filas = evaluar_colegios(max_workers=args.workers)
    ts = datetime.now(CHILE).strftime("%Y%m%d_%H%M")
    out_dir = ROOT / "reports" / "CORMUP" / "VACACIONES" / f"COMPARATIVO_{ts}"
    out_dir.mkdir(parents=True, exist_ok=True)

    (out_dir / "resultado.json").write_text(
        json.dumps([asdict(f) for f in filas], indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    chart_barras = grafico_barras_colegios(filas, out_dir / "chart_barras_colegios.png")
    chart_diario = grafico_totales_diarios(filas, out_dir / "chart_totales_diarios.png")

    docx = crear_informe_word(
        filas,
        chart_barras,
        chart_diario,
        out_dir / "Comparativo_CORMUP_Vacaciones_20260907_20260920.docx",
    )
    print(f"[OK] Word: {docx}")

    pdf = convertir_pdf(docx, filas, chart_barras, chart_diario)
    print(f"[OK] PDF: {pdf}")

    agregados: List[Path] = []
    if not args.sin_agregados:
        try:
            agregados = generar_agregados_apoyo()
        except Exception as exc:
            print(f"[AVISO] No se generaron agregados de apoyo: {exc}")

    meta = {
        "out_dir": str(out_dir),
        "docx": str(docx),
        "pdf": str(pdf),
        "agregados": [str(p) for p in agregados],
        "total_sin_m3": sum(f.m3_sin for f in filas),
        "total_con_m3": sum(f.m3_con for f in filas),
    }
    (out_dir / "meta.json").write_text(json.dumps(meta, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(meta, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
